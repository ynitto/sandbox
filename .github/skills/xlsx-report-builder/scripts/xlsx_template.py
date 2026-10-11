#!/usr/bin/env python3
"""xlsx テンプレートへデータを流し込む（書式・不変部分はテンプレートのまま保つ）。

2 つのコマンド:

    analyze  テンプレートを解析し、表構造の定義ファイル（JSON）の下書きと確認用の要約を出す
    render   テンプレート + 定義ファイル + データ から xlsx を再構成する

openpyxl で読み書きせず、xlsx（zip）内のシート XML の「可変の行」だけを書き換える。
スタイル・テーマ・図形・画像・マクロ・他シートなど、触らない部品は元のバイト列のままコピーする。
定義ファイルとデータの書式は references/template.md を参照。
"""
from __future__ import annotations

import argparse
import base64
import datetime as dt
import io
import json
import os
import posixpath
import re
import sys
import zipfile
from copy import deepcopy
from typing import Any

from lxml import etree
from openpyxl.formula.tokenizer import Tokenizer, Token
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, get_column_letter

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"
NS_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
NS_C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

DEF_VERSION = 1
CELL_RE = re.compile(r"^([A-Za-z]{1,3})(\d+)$")
REF_PART_RE = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")
TOTAL_WORDS = ("合計", "小計", "総計", "計", "total", "sum", "subtotal")
TOTAL_FUNCS = re.compile(r"^(SUM|SUBTOTAL|AVERAGE|COUNT|COUNTA|MAX|MIN)\(", re.I)
BUILTIN_FORMATS = {1: "0", 2: "0.00", 3: "#,##0", 4: "#,##0.00", 9: "0%", 10: "0.00%", 11: "0.00E+00", 12: "# ?/?",
                   14: "m/d/yyyy", 15: "d-mmm-yy", 20: "h:mm", 21: "h:mm:ss", 22: "m/d/yyyy h:mm", 49: "@"}
BUILTIN_DATE_IDS = set(range(14, 23)) | set(range(27, 37)) | set(range(45, 48)) | set(range(50, 59))


class TemplateError(Exception):
    """定義ファイル・データ・テンプレートの不整合。"""


def q(tag: str) -> str:
    return f"{{{NS}}}{tag}"


# ---------------------------------------------------------------------------
# パッケージ（zip）の読み書き
# ---------------------------------------------------------------------------

class Package:
    def __init__(self, path: "str | bytes"):
        with zipfile.ZipFile(io.BytesIO(path) if isinstance(path, (bytes, bytearray)) else path) as z:
            self.infos = z.infolist()
            self.data = {i.filename: z.read(i.filename) for i in self.infos}
        self.removed: set[str] = set()

    def xml(self, name: str) -> etree._Element:
        return etree.fromstring(self.data[name])

    def put_xml(self, name: str, root: etree._Element) -> None:
        self.data[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def save(self, path: str) -> None:
        if os.path.dirname(path):
            os.makedirs(os.path.dirname(path), exist_ok=True)   # 出力先のフォルダが無ければ作る
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for info in self.infos:
                if info.filename in self.removed:
                    continue
                zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
                zi.compress_type = zipfile.ZIP_DEFLATED
                zi.external_attr = info.external_attr
                z.writestr(zi, self.data[info.filename])

    def rels_of(self, part: str) -> dict[str, tuple[str, str]]:
        """part の rels を {rId: (type, 解決済みターゲットパス)} で返す。"""
        rels_name = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
        if rels_name not in self.data:
            return {}
        out = {}
        for r in self.xml(rels_name):
            target = r.get("Target", "")
            if r.get("TargetMode") == "External":
                continue
            if target.startswith("/"):
                resolved = target.lstrip("/")
            else:
                resolved = posixpath.normpath(posixpath.join(posixpath.dirname(part), target))
            out[r.get("Id")] = (r.get("Type", ""), resolved)
        return out

    def sheets(self) -> list[tuple[str, str]]:
        """[(シート名, シート part パス)] をブック順で返す。"""
        wb = self.xml("xl/workbook.xml")
        rels = self.rels_of("xl/workbook.xml")
        out = []
        for s in wb.find(q("sheets")):
            rid = s.get(f"{{{NS_R}}}id")
            if rid in rels:
                out.append((s.get("name"), rels[rid][1]))
        return out


def read_shared_strings(pkg: Package) -> list[str]:
    name = "xl/sharedStrings.xml"
    if name not in pkg.data:
        return []
    out = []
    for si in pkg.xml(name):
        out.append("".join(t.text or "" for t in si.iter(q("t")) if t.getparent().tag != q("rPh")))
    return out


class Styles:
    """cellXfs の添字から、日付書式か・書式コード・太字/塗りの有無を引く。"""

    def __init__(self, pkg: Package):
        self.date: dict[int, bool] = {}
        self.code: dict[int, str] = {}
        self.emphasis: dict[int, bool] = {}
        self.desc: dict[int, str] = {}
        if "xl/styles.xml" not in pkg.data:
            return
        root = pkg.xml("xl/styles.xml")
        custom = {}
        nf = root.find(q("numFmts"))
        if nf is not None:
            custom = {int(n.get("numFmtId")): n.get("formatCode", "") for n in nf}
        fonts = root.find(q("fonts"))
        fills = root.find(q("fills"))
        xfs = root.find(q("cellXfs"))
        if xfs is None:
            return
        for i, xf in enumerate(xfs):
            fid = int(xf.get("numFmtId", "0"))
            code = custom.get(fid, "")
            self.code[i] = code or ("General" if fid == 0 else BUILTIN_FORMATS.get(fid, f"builtin:{fid}"))
            self.date[i] = fid in BUILTIN_DATE_IDS or bool(code and _is_date_code(code))
            bold = False
            if fonts is not None and int(xf.get("fontId", "0")) < len(fonts):
                bold = fonts[int(xf.get("fontId", "0"))].find(q("b")) is not None
            filled = False
            if fills is not None and int(xf.get("fillId", "0")) < len(fills):
                pf = fills[int(xf.get("fillId", "0"))].find(q("patternFill"))
                filled = pf is not None and pf.get("patternType") not in (None, "none")
            self.emphasis[i] = bold or filled
            parts = []
            if bold:
                parts.append("太字")
            if filled:
                fg = pf.find(q("fgColor")) if pf is not None else None
                color = (fg.get("rgb") or (f"theme{fg.get('theme')}" if fg.get("theme") else "")) if fg is not None else ""
                parts.append("塗り" + color[-6:] if color else "塗り")
            borders = root.find(q("borders"))
            bid = int(xf.get("borderId", "0"))
            if borders is not None and bid < len(borders) and any(
                    side.get("style") for side in borders[bid]):
                parts.append("罫線")
            if self.code[i] != "General":
                parts.append(self.code[i])
            self.desc[i] = "・".join(parts) or "標準"

    def describe(self, s: str | None) -> str:
        return self.desc.get(int(s or 0), "標準")

    def is_date(self, s: str | None) -> bool:
        return self.date.get(int(s or 0), False)

    def is_number(self, s: str | None) -> bool:
        """数値の書式（`#,##0`・`0.0%` など）のセルか。標準・文字列（@）・日付は含めない。"""
        code = self.code.get(int(s or 0), "General")
        return code not in ("General", "@") and not code.startswith("builtin:") and not self.is_date(s) \
            and bool(re.search(r"[0#]", re.sub(r"\[[^\]]*\]", "", _strip_literals(code))))


def _strip_literals(code: str) -> str:
    """書式コードから、書式の記号ではない文字（"文字列"・\\x・_x・*x）を除く。"""
    return re.sub(r'"[^"]*"|[\\_*].', "", code)


def _is_date_code(code: str) -> bool:
    # Excel の角括弧は色・条件・ロケール指定にも使われるが、
    # [h] / [m] / [s]（および繰り返し）は 24 時間を超える経過時間の書式。
    # それらだけは日付・時刻判定用に残し、他の角括弧指定は無視する。
    stripped = re.sub(
        r"\[([hms]+)\]",
        lambda m: m.group(1) if len(set(m.group(1).lower())) == 1 else "",
        _strip_literals(code),
        flags=re.I,
    )
    stripped = re.sub(r"\[[^\]]*\]", "", stripped)
    return bool(re.search(r"[ymdhs]", stripped, re.I))


# ---------------------------------------------------------------------------
# セル参照・数式の書き換え
# ---------------------------------------------------------------------------

def split_ref(ref: str) -> tuple[int, int]:
    m = CELL_RE.match(ref)
    if not m:
        raise TemplateError(f"セル参照が不正です: {ref!r}")
    return column_index_from_string(m.group(1).upper()), int(m.group(2))


def parse_row_spec(spec: Any) -> tuple[int, int]:
    """`20` / `"20"` / `"20:22"` / `"20-22"` / `[20, 22]` を (先頭, 末尾) にする。"""
    if isinstance(spec, (list, tuple)) and len(spec) == 2:
        a, b = int(spec[0]), int(spec[1])
    else:
        parts = re.split(r"[:\-]", str(spec))
        a, b = int(parts[0]), int(parts[-1])
    if a < 1 or b < a:
        raise TemplateError(f"行の指定が不正です: {spec!r}")
    return a, b


def parse_cell_spec(spec: str) -> list[tuple[int, int]]:
    """`B5` / `B5:D7` を (列, 行) の一覧にする。"""
    a, _, b = str(spec).partition(":")
    (c1, r1), (c2, r2) = split_ref(a), split_ref(b or a)
    return [(c, r) for r in range(min(r1, r2), max(r1, r2) + 1) for c in range(min(c1, c2), max(c1, c2) + 1)]


class RowMap:
    """可変表の行数の増減に応じた、元の行番号 → 新しい行番号 の対応。

    tables: dict(first, end, count, n) のリスト。first..end が元のサンプル行、n が出力行数。
    """

    def __init__(self, tables: list[dict]):
        self.tables = sorted(tables, key=lambda t: t["first"])

    def shift_before(self, r: int) -> int:
        return sum(t["n"] - t["count"] for t in self.tables if t["end"] < r)

    def new_last(self, t: dict) -> int:
        return t["first"] + self.shift_before(t["first"]) + t["n"] - 1

    def map(self, r: int, end_to_last: bool = True) -> int:
        if end_to_last:
            for t in self.tables:
                if r == t["end"]:
                    return self.new_last(t)
        return r + self.shift_before(r)

    def map_span(self, r1: int, r2: int) -> "tuple[int, int] | None":
        """範囲の行 r1..r2 を新しい行に写す。表の残らない行（縮めた・取り除いた分）は、表の新しい範囲に寄せる。
        範囲がすべて消えるなら None（縮めた表で、先頭と末尾が逆転した範囲を作らない）。"""
        def one(r: int, is_end: bool) -> int:
            for t in self.tables:
                if t["first"] <= r <= t["end"]:
                    new_first = t["first"] + self.shift_before(t["first"])
                    if is_end and (r == t["end"] or r - t["first"] >= t["n"]):
                        return new_first + t["n"] - 1
                    return new_first + min(r - t["first"], t["n"])
            return r + self.shift_before(r)
        a, b = one(r1, False), one(r2, True)
        return (a, b) if a <= b else None


class Rewriter:
    """数式・範囲文字列中の行参照を RowMap で書き換える。他シートの参照はそのシートの RowMap を使う。"""

    def __init__(self, this_sheet: str, maps: dict[str, RowMap]):
        self.this_sheet = this_sheet
        self.maps = maps

    def _split_sheet(self, operand: str) -> tuple[str | None, str, str]:
        if "!" not in operand:
            return None, "", operand
        prefix, _, rest = operand.rpartition("!")
        name = prefix
        if name.startswith("'") and name.endswith("'"):
            name = name[1:-1].replace("''", "'")
        return name, prefix + "!", rest

    def operand(self, operand: str, row_fn=None) -> str:
        """1 つの範囲オペランド（Sheet1!A1:B2 など）を書き換える。row_fn(rowmap, idx, last, abs, r) → 新行。"""
        sheet, prefix, rest = self._split_sheet(operand)
        target = sheet if sheet is not None else self.this_sheet
        rowmap = self.maps.get(target)
        if rowmap is None:
            return operand
        parts = rest.split(":")
        if not all(REF_PART_RE.match(p) for p in parts):
            return operand
        out = []
        for idx, p in enumerate(parts):
            m = REF_PART_RE.match(p)
            ca, col, ra, row = m.group(1), m.group(2), m.group(3), int(m.group(4))
            last = idx == len(parts) - 1
            if row_fn is None:
                new = rowmap.map(row, last)
            else:
                new = row_fn(rowmap, idx, last, bool(ra), row)
            out.append(f"{ca}{col}{ra}{new}")
        return prefix + ":".join(out)

    def formula(self, formula: str, row_fn=None) -> str:
        """`=` なしの数式文字列を書き換える。"""
        try:
            tok = Tokenizer("=" + formula)
        except Exception:
            return formula
        pieces = []
        for t in tok.items:
            if t.type == Token.OPERAND and t.subtype == Token.RANGE:
                pieces.append(self.operand(t.value, row_fn))
            else:
                pieces.append(t.value)
        return "".join(pieces)

    def sqref(self, sqref: str) -> str:
        """`A1:B2 D5` のような範囲リストを、範囲の先頭は据え置き・末尾はサンプル最終行 → 新最終行で書き換える。"""
        out = []
        for piece in sqref.split():
            parts = piece.split(":")
            if not all(REF_PART_RE.match(p.replace("$", "")) for p in parts):
                out.append(piece)
                continue
            rowmap = self.maps.get(self.this_sheet)
            if rowmap is None:
                out.append(piece)
                continue
            first = REF_PART_RE.match(parts[0].replace("$", ""))
            last = REF_PART_RE.match(parts[-1].replace("$", ""))
            span = rowmap.map_span(int(first.group(4)), int(last.group(4)))
            if span is None:
                continue   # 範囲の行がすべて消えた
            a, b = f"{first.group(2)}{span[0]}", f"{last.group(2)}{span[1]}"
            out.append(a if a == b else f"{a}:{b}")
        return " ".join(out)


# ---------------------------------------------------------------------------
# シート XML の読み取り補助
# ---------------------------------------------------------------------------

def cell_col(c) -> int:
    return split_ref(c.get("r"))[0]


def cell_value(c, sst: list[str]) -> Any:
    t = c.get("t")
    if t == "inlineStr":
        return "".join(x.text or "" for x in c.iter(q("t")))
    v = c.find(q("v"))
    if v is None or v.text is None:
        return None
    if t == "s":
        return sst[int(v.text)]
    if t in ("str", "e"):
        return v.text
    if t == "b":
        return v.text == "1"
    try:
        n = float(v.text)
        return int(n) if n.is_integer() else n
    except ValueError:
        return v.text


def expand_shared_formulas(root: etree._Element) -> None:
    """共有数式を、各セルが自分の数式を持つ形に展開する（行の複製・移動のため）。"""
    masters: dict[str, tuple[str, str]] = {}
    for c in root.iter(q("c")):
        f = c.find(q("f"))
        if f is not None and f.get("t") == "shared" and f.text:
            masters[f.get("si")] = (f.text, c.get("r"))
    for c in root.iter(q("c")):
        f = c.find(q("f"))
        if f is None or f.get("t") != "shared":
            continue
        text, origin = masters[f.get("si")]
        if not f.text:
            f.text = Translator("=" + text, origin=origin).translate_formula(c.get("r"))[1:]
        for a in ("t", "si", "ref"):
            f.attrib.pop(a, None)


def sheet_rows(root: etree._Element) -> dict[int, etree._Element]:
    return {int(r.get("r")): r for r in root.find(q("sheetData"))}


# ---------------------------------------------------------------------------
# analyze
# ---------------------------------------------------------------------------

def analyze(template: str) -> dict:
    pkg = Package(template)
    sst = read_shared_strings(pkg)
    styles = Styles(pkg)
    sheets_def = []
    # 前の文書の作成者・会社名。残す文字（タイトルなど）に含まれていたら、前の案件の値の疑い
    props = read_properties(pkg)
    words = {v2: k for k, v in props.items() if k in ("creator", "lastModifiedBy", "company", "manager")
             for v2 in {v.strip(), re.sub(r"(株式会社|有限会社|\(株\)|（株）|様|御中|\s)", "", v)} if len(v2) >= 2}
    doc = DocValues()   # 文書の値（年度・作成日・宛先など）の名前は、ブック全体で 1 つ
    for name, part in pkg.sheets():
        root = pkg.xml(part)
        expand_shared_formulas(root)
        sheets_def.append(_analyze_sheet(pkg, name, part, root, sst, styles, words, doc))
    # 表のあるシートが 2 つ以上なら、データのキーをシート名にする（同じ items だと、どのシートにも同じ明細が入る）
    with_tables = [s for s in sheets_def if s["tables"]]
    if len(with_tables) > 1:
        for sd in with_tables:
            base = re.sub(r"[.\s]+", "_", sd["name"]).strip("_") or "sheet"
            for n, t in enumerate(sd["tables"], start=1):
                t["key"] = base if len(sd["tables"]) == 1 else f"{base}_{n}"
    # タブの名前が同じ言葉で始まるもの（「受注_一覧」「受注_明細」）は、データを分けるときの 1 つのまとまり（group）にする
    heads = [re.split(r"[_\-－・\s（(]", sd["name"], maxsplit=1)[0] for sd in sheets_def]
    for sd, head in zip(sheets_def, heads):
        if head and head != sd["name"] and heads.count(head) >= 2:
            sd["group"] = head
            sd["_notes"].append(f"データを分けるときは、名前が「{head}」で始まるタブと 1 つのファイルにまとめる（group）")
    # 固定セルのデータのキーは、ブック全体で重ならないようにする（別のシートの同じラベルに、黙って同じ値が入らない）
    used = {t["key"] for sd in sheets_def for t in sd["tables"]}
    for sd in sheets_def:
        renamed: dict[str, str] = {}   # 年・月・日に分かれた欄は、同じキーのまま名前を変える
        for ref, spec in sd["cells"].items():
            if isinstance(spec, dict) and "text" in spec:
                continue
            key = spec["key"] if isinstance(spec, dict) else spec
            if key.startswith(DOC_PREFIX + "."):
                continue   # 文書の値は、どのタブでも同じキー（同じ値が入る）
            if key not in renamed:
                base, new, n = key, key, 2
                while new in used:
                    new, n = f"{base}{n}", n + 1
                used.add(new)
                renamed[key] = new
            sd["cells"][ref] = dict(spec, key=renamed[key]) if isinstance(spec, dict) else renamed[key]
    # テンプレートの値は、keep に挙げたもの（見出し・ラベル）以外を出力に残さない。来歴（作成者・コメントなど）も消す
    return {
        "version": DEF_VERSION,
        "template": posixpath.basename(template),
        "strict": True,
        "properties": {"scrub": True},
        "sheets": sheets_def,
    }


def _serial_to_iso(value) -> str:
    """日付の書式のセルの数値（Excel の通し番号）を、読める形（2025-01-01・2025-01-01 09:30）にする。"""
    d = dt.datetime(1899, 12, 30) + dt.timedelta(days=float(value))
    return d.date().isoformat() if d.time() == dt.time() else d.isoformat(sep=" ", timespec="minutes")


def _grid(root, sst, styles: "Styles | None" = None):
    grid: dict[int, dict[int, dict]] = {}
    for row in root.find(q("sheetData")):
        r = int(row.get("r"))
        cols = {}
        for c in row:
            f = c.find(q("f"))
            value = cell_value(c, sst)
            if styles is not None and isinstance(value, (int, float)) and not isinstance(value, bool) \
                    and styles.is_date(c.get("s")) and 0 < value < 2958466:
                value = _serial_to_iso(value)
            cols[cell_col(c)] = {
                "ref": c.get("r"),
                "s": c.get("s", "0"),
                "value": value,
                "formula": ("=" + f.text) if f is not None and f.text else None,
            }
        grid[r] = cols
    return grid


def _sig(cols: dict[int, dict]) -> tuple:
    return tuple(sorted((c, i["s"]) for c, i in cols.items()))


def _is_total_row(cols: dict[int, dict]) -> bool:
    for i in cols.values():
        if i["formula"] and TOTAL_FUNCS.match(i["formula"][1:]) and _sums_rows_above(i):
            return True
        v = i["value"]
        if isinstance(v, str) and v.strip().lower().replace(" ", "").replace("　", "") in TOTAL_WORDS:
            return True
    return False


def _sums_rows_above(cell: dict) -> bool:
    """SUM などが、上の行を縦に集めているか（合計の行）。同じ行を横に足す `=SUM(B2:E2)` は、明細の行の合計の列。"""
    own = split_ref(cell["ref"])[1]
    rows = [int(m.group(1)) for m in re.finditer(r"(?<![A-Za-z0-9_.])\$?[A-Za-z]{1,3}\$?(\d+)(?![\d(])", cell["formula"])]
    return not rows or min(rows) < own


def _find_tables(grid, styles) -> list[dict]:
    tables = []
    rows = sorted(grid)
    r_idx = 0
    while r_idx < len(rows):
        r = rows[r_idx]
        cols = grid[r]
        heads = {c: i for c, i in cols.items() if isinstance(i["value"], str) and i["value"].strip() and not i["formula"]}
        nxt = grid.get(r + 1)
        if len(heads) >= 2 and nxt is not None:
            span = sorted(heads)
            body_cols = [c for c in span if c in nxt]
            same_style = len({heads[c]["s"] for c in span}) == 1
            emph = all(styles.emphasis.get(int(heads[c]["s"]), False) for c in span)
            differs = any(nxt[c]["s"] != heads[c]["s"] for c in body_cols)
            # 2 段の見出し（上の段が「判定」「対象OS」のようなまとまり）なら、下の段を見出しにする
            nxt_heads = [c for c, i in nxt.items() if isinstance(i["value"], str) and i["value"].strip() and not i["formula"]]
            two_tier = len(nxt_heads) > len(span) and all(styles.emphasis.get(int(nxt[c]["s"]), False) for c in nxt_heads) \
                and (r + 2) in grid
            # 書式の強調が無い見出しは、隙間なく並ぶときだけ表と見る（「宛先 ○○ 承認」のような記入欄の行と取り違えない）
            no_gap = span[-1] - span[0] + 1 == len(span)
            if len(body_cols) * 2 >= len(span) and (emph or (same_style and len(span) >= 3 and differs and no_gap)) \
                    and not _is_total_row(nxt) and not two_tier:
                lo, hi = span[0], span[-1]
                body = []
                rr = r + 1
                while rr in grid:
                    g = grid[rr]
                    inside = [c for c in g if lo <= c <= hi]
                    if len(inside) * 2 < len(span) or _is_total_row(g):
                        break
                    body.append(rr)
                    rr += 1
                if body:
                    # 見出しの無い端の列でも、× や - の既定の印が並ぶなら表の列（余白として残す）。表の範囲に入れる
                    def marks_only(c):
                        cells = [grid[b][c]["value"] for b in body if c in grid[b]]
                        return not (cols.get(c) and cols[c]["value"] not in (None, "")) and len(cells) == len(body) \
                            and all(isinstance(v, str) and v.strip() in MARK_OFF for v in cells)
                    while marks_only(hi + 1):
                        hi += 1
                    while lo > 1 and marks_only(lo - 1):
                        lo -= 1
                    tables.append(_table_info(grid, r, body, lo, hi, rr, styles))
                    r_idx = next((k for k, x in enumerate(rows) if x >= rr), len(rows))
                    continue
        r_idx += 1
    return tables


def _table_info(grid, header_row, body, lo, hi, after, styles) -> dict:
    sigs = [_sig({c: i for c, i in grid[r].items() if lo <= c <= hi}) for r in body]
    period = next((p for p in (1, 2, 3) if p <= len(sigs) and all(sigs[k] == sigs[k % p] for k in range(len(sigs)))), None)
    total = after if after in grid and _is_total_row(grid[after]) else None
    return {
        "header_row": header_row, "body": body, "lo": lo, "hi": hi,
        "period": period, "total_row": total,
    }


def _analyze_sheet(pkg, name, part, root, sst, styles, words: "dict | None" = None,
                   doc: "DocValues | None" = None) -> dict:
    grid = _grid(root, sst, styles)
    found = _find_tables(grid, styles)
    notes: list[str] = []
    tables = []
    covered: set[int] = set()
    heads: set[int] = set()
    for n, t in enumerate(found, start=1):
        first, last = t["body"][0], t["body"][-1]
        confirm = ["データのキー（key）と各列の key が意図どおりか"]
        if t["period"] is None:
            pattern = [first]
            confirm.append("サンプル行の書式が揃っていない。複製元にする行（pattern）を選ぶ")
        else:
            pattern = t["body"][:t["period"]]
        if t["period"] and t["period"] > 1:
            confirm.append(f"サンプル行の書式が {t['period']} 行周期（縞模様）。pattern の行を順に繰り返す")
        cols_def = {}
        filled = _filled_rows(grid, t)
        for c in range(t["lo"], t["hi"] + 1):
            letter = get_column_letter(c)
            head = grid[t["header_row"]].get(c)
            samples = [grid[r][c] for r in t["body"] if c in grid[r]]
            fm = next((s for s in samples if s["formula"]), None)
            spec: dict[str, Any] = {"header": head["value"] if head else None}
            if fm:
                spec["formula"] = True
                spec["_sample"] = fm["formula"]
                for r in t["body"]:
                    cell = grid[r].get(c)
                    if cell and cell["formula"] and not _only_inside(cell["formula"], t, c):
                        confirm.append(f"{cell['ref']} の数式が表の外の行を相対参照している。固定するなら $ を付ける")
                        break
            elif not (head and head["value"] is not None) and all(_blankish(s["value"]) for s in samples):
                if not [s for s in samples if s["value"] is not None]:
                    continue   # 見出しも値も無い列は、体裁の余白。データにも定義にも入れない
                spec["keep"] = True   # 見出しが無く、× や - の既定の印だけの列も余白。印は書式としてそのまま残す
                spec["_sample"] = next(s["value"] for s in samples if s["value"] is not None)
                confirm.append(f"{letter}列は見出しが無く {spec['_sample']!r} だけなので、余白として残す（データには書かない）")
            elif head and isinstance(head["value"], str) and HUMAN_RE.search(head["value"]):
                spec["clear"] = True   # 押印・署名など、人が書き込む欄。データに入れず、空欄で出す
                confirm.append(f"{letter}列「{head['value']}」は人が書き込む欄として、データに入れず空欄で出す")
            elif _constant(samples, len(t["body"])):
                spec["keep"] = True    # どのサンプル行も同じ値（円・式など）。毎回同じなので、データに書かない
                spec["_sample"] = samples[0]["value"]
                confirm.append(f"{letter}列はどの行も {samples[0]['value']!r} の定数として残す（行ごとに変わるなら key にする）")
            elif len(samples) == len(t["body"]) and [s["value"] for s in samples] == list(range(1, len(samples) + 1)) \
                    and (not head or head["value"] is None or INDEX_HEAD_RE.match(str(head["value"]).strip())):
                spec["key"] = "$index"  # サンプルが 1, 2, 3 … の連番なら、連番の列
                spec["_sample"] = 1
            else:
                spec["key"] = _unique_key(spec["header"], letter, cols_def)
                # 例の値は、既定値（連番・× など）だけの空行を除いた行から取る
                spec["_sample"] = next((grid[r][c]["value"] for r in filled
                                        if c in grid[r] and not _blankish(grid[r][c]["value"])), None)
            if samples:
                spec["_format"] = styles.code.get(int(samples[0]["s"]), "General")
            cols_def[letter] = spec
        _readable_columns(cols_def, t, grid, _merge_origins(root), confirm, filled)
        empty = [r for r in t["body"] if r not in filled]
        if empty and len(filled) < len(t["body"]):
            confirm.append(f"{', '.join(map(str, empty))} 行目は既定値（連番・× など）だけの空行。"
                           "記入例の行数に数えるだけで、データには書かない")
        covered.update(t["body"])
        heads.add(t["header_row"])
        above = grid.get(t["header_row"] - 1, {})
        if above and all(t["lo"] <= c <= t["hi"] for c, i in above.items() if i["value"] is not None):
            heads.add(t["header_row"] - 1)   # 2 段の見出しの上の段（「判定」「対象OS」のようなまとまり）
        tables.append({
            "id": f"table{n}",
            "header_row": t["header_row"],
            "first_row": first,
            "sample_rows": len(t["body"]),
            "pattern": pattern,
            "key": "items" if n == 1 else f"items{n}",
            "columns": cols_def,
            "_total_row": t["total_row"],
            "needs_confirm": confirm,
        })
    doc = doc if doc is not None else DocValues()
    cells, keep, samples, clear = _fixed_cells(grid, covered, heads, root, doc)
    for ref, spec in cells.items():
        if isinstance(spec, dict) and "text" in spec and samples.get(ref, {}).get("label") is None:
            notes.append(f"{ref} の「{samples[ref]['value']}」は、可変の部分だけを文書の値として流し込む（{spec['text']}）。"
                         "毎回同じ文字なら keep に移す")
    header_footer = {}
    for tag, el in _header_footer(root).items():
        tpl = _header_footer_template(el.text or "", doc)
        if tpl is not None:
            header_footer[tag] = tpl
            notes.append(f"ヘッダー・フッター {tag} の「{el.text}」は、可変の部分を文書の値として流し込む")
        elif el.text and HF_CODE_RE.sub("", el.text).strip():
            notes.append(f"ヘッダー・フッター {tag} の「{el.text}」はそのまま残る。前の文書の名前・日付なら、"
                         "header_footer に文のひな形を書く")
    for r in sorted(grid):
        for i in grid[r].values():
            v = i["value"]
            if r in covered or i["ref"] in cells or not isinstance(v, str):
                continue
            hit = next((w for w in (words or {}) if w in v), None)
            if hit:
                notes.append(f"{i['ref']} の「{v}」は、前の文書の{ {'company': '会社名', 'manager': '管理者'}.get(words[hit], '作成者') }"
                             f"「{hit}」を含む。残さずに流し込むか、利用者に確かめる")
    notes += _sheet_warnings(pkg, part, root, tables)
    out = {"name": name, "tables": tables, "cells": cells, "keep": keep, "_cell_samples": samples, "_notes": notes}
    if header_footer:
        out["header_footer"] = header_footer
    if clear:
        out["clear"] = clear
    return out


def _only_inside(formula: str, t: dict, col: int) -> bool:
    lo, hi = t["body"][0], t["body"][-1]
    try:
        tok = Tokenizer(formula)
    except Exception:
        return True
    for x in tok.items:
        if x.type == Token.OPERAND and x.subtype == Token.RANGE and "!" not in x.value:
            for p in x.value.split(":"):
                m = REF_PART_RE.match(p)
                if m and not m.group(3) and not (lo <= int(m.group(4)) <= hi) and int(m.group(4)) != lo - 1:
                    return False
    return True


def _unique_key(header, letter: str, cols_def: dict) -> str:
    base = (header or "").strip() or f"col_{letter}"
    used = {c.get("key") for c in cols_def.values()}
    key, i = base, 2
    while key in used:
        key, i = f"{base}{i}", i + 1
    return key


# 連番の列の見出し。数量のサンプルがたまたま 1, 2, 3 でも、見出しが連番でなければ連番にしない
INDEX_HEAD_RE = re.compile(r"^(No\.?|NO\.?|no\.?|№|#|番号|項番|連番|通番|行番号?|順|順番|項)$")

# 人が書き込む欄（押印・署名など）の見出し・ラベル。データに入れず、空欄のまま出す
HUMAN_RE = re.compile(r"(印$|押印|捺印|検印|署名|サイン|自署|承認者?$|決裁|確認者|受付者?$|手書き|記入欄)")


def _constant(samples: list[dict], n_rows: int) -> bool:
    """どのサンプル行も同じ文字（印・仮の値を除く）なら、行ごとに変わらない定数の列。"""
    vals = [s["value"] for s in samples]
    return n_rows >= 2 and len(vals) == n_rows and isinstance(vals[0], str) and vals[0].strip() != "" \
        and len(set(vals)) == 1 and vals[0].strip() not in MARK_ON + MARK_OFF and not PLACEHOLDER_RE.search(vals[0])


def _blankish(value) -> bool:
    """空欄、または × や - の「選ばれていない」既定の印。データとしては空と同じ。"""
    return value is None or isinstance(value, str) and (not value.strip() or value.strip() in MARK_OFF)


def _filled_rows(grid, t: dict) -> list[int]:
    """表のサンプル行のうち、既定値を除いても値がある行。

    連番の列・数式の列・どの行も同じ値の列（円・式など）と、× や - の印は、どの行にも最初から入っている既定値。
    それらを除いてから空かを判定する（連番と × だけの行は、記入例ではなく空の行）。
    すべての行が空なら、すべての行を返す。
    """
    body = t["body"]
    skip = set()
    for c in range(t["lo"], t["hi"] + 1):
        samples = [grid[r][c] for r in body if c in grid[r]]
        vals = [s["value"] for s in samples]
        if any(s["formula"] for s in samples) or _constant(samples, len(body)) \
                or len(vals) == len(body) and vals == list(range(1, len(body) + 1)):
            skip.add(c)
    filled = [r for r in body
              if any(not _blankish(i["value"]) for c, i in grid[r].items() if t["lo"] <= c <= t["hi"] and c not in skip)]
    return filled or list(body)


def _merge_origins(root) -> dict[tuple[int, int], tuple[int, int]]:
    """結合範囲の各セル（列, 行）→ 左上のセル。"""
    origin: dict[tuple[int, int], tuple[int, int]] = {}
    for mc in root.iter(q("mergeCell")):
        (c1, r1), (c2, r2) = (split_ref(p) for p in (mc.get("ref").split(":") * 2)[:2])
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                origin[(c, r)] = (c1, r1)
    return origin


MARK_ON = ("○", "◯", "〇", "●", "◎", "✓", "✔", "レ", "☑", "■")
MARK_OFF = ("×", "✕", "✖", "☐", "□", "-", "－", "ー", "―")
DASHES = ("-", "－", "ー", "―")      # 「なし」の意味の横棒。印の列とは限らない（備考の - など）
MARK_PAIR = {"☐": "☑", "□": "■"}  # × や □ だけの列で、選ばれたときの印
DATE_HEADS = {"年": "year", "月": "month", "日": "day", "時": "hour", "分": "minute"}


def _readable_columns(cols_def: dict, t: dict, grid, origin, confirm: list, filled: "list | None" = None) -> None:
    """表の見た目に頼る列（○× の印・年月日に分かれた日付）を、データで意味が分かる書き方にする。

    - ○ の列が並ぶ → 見出しを値にした択一（`判定: 合格`）か、複数選択（`対象: [Windows, Linux]`）
    - ○ の列が 1 つ → true / false
    - 年・月・日の列が並ぶ → 1 つの日付（`2026-10-08`）
    """
    def text(r, c):
        v = grid.get(r, {}).get(c, {}).get("value")
        return str(v).strip() if v is not None and str(v).strip() else None

    def group_head(cols) -> "str | None":   # 見出しの 1 行上で、列のまとまりにかかる見出し（結合セル）
        above = {origin.get((c, t["header_row"] - 1), (c, t["header_row"] - 1)) for c in cols}
        if len(above) != 1:
            return None
        v = text(*reversed(next(iter(above))))
        return re.sub(r"[.\s]+", "_", v) if v else None

    def runs(cols):   # 隣り合う列のまとまり。見出しの上の段の結合セルが変わるところで切る
        out = []
        above = lambda c: origin.get((c, t["header_row"] - 1), (c, t["header_row"] - 1))
        for c in sorted(cols):
            if out and out[-1][-1] == c - 1 and (above(c) == above(c - 1) or above(c)[0] == c and not text(*reversed(above(c)))
                                                 and above(c - 1)[0] == c - 1 and not text(*reversed(above(c - 1)))):
                out[-1].append(c)
            else:
                out.append([c])
        return out

    used = {spec.get("key") for spec in cols_def.values()}
    first = t["body"].index(filled[0]) if filled else 0   # 例の値を取る行（既定値だけの空の行は飛ばす）

    def fresh(base):
        key, n = base, 2
        while key in used:
            key, n = f"{base}{n}", n + 1
        used.add(key)
        return key

    plain = {column_index_from_string(L): s for L, s in cols_def.items()
             if not s.get("formula") and s.get("key") not in (None, "$index")}
    marks, blank = {}, set()
    for c in plain:
        vals = [text(r, c) for r in t["body"]]
        hit = [v for v in vals if v]
        if hit and all(v in MARK_ON + MARK_OFF for v in hit) and (any(v in MARK_ON for v in hit)
                                                                    or not set(hit) & set(DASHES)):
            marks[c] = vals   # ○ がある列と、× や □ だけの列（どれも選ばれていない記入例）
        elif all(v in DASHES for v in hit):
            blank.add(c)   # 空か - だけの列は、空と同じ。印の列と同じまとまりの中なら選択肢にする
            plain[c]["_sample"] = None
    groups = []
    for run in runs(set(marks) | blank):
        if group_head(run):
            groups += [run] if any(c in marks for c in run) else []
        else:   # まとまりの見出しが無ければ、印の列だけが隣り合う範囲
            groups += runs([c for c in run if c in marks])
    for run in groups:
        for c in run:
            marks.setdefault(c, [None] * len(t["body"]))
        letters = [get_column_letter(c) for c in run]
        cells = [v for c in run for v in marks[c] if v]
        off = max((m for m in MARK_OFF if m in cells), key=cells.count, default=None)
        on = max((m for m in MARK_ON if m in cells), key=cells.count, default=MARK_PAIR.get(off, "○"))
        heads = [str(plain[c].get("header") or get_column_letter(c)) for c in run]
        if len(run) == 1:
            spec = plain[run[0]]
            spec["map"] = {True: on, False: off}
            spec["_sample"] = marks[run[0]][first] == on if marks[run[0]][first] else None
            confirm.append(f"{letters[0]}列の {on} を、データでは true / false で書く")
            continue
        picked = [[h for c, h in zip(run, heads) if marks[c][i] == on] for i in range(len(t["body"]))]
        multi = any(len(p) > 1 for p in picked)
        for c in run:
            used.discard(plain[c]["key"])
        key = fresh(group_head(run) or "/".join(heads))
        for c, h in zip(run, heads):
            spec = plain[c]
            spec.update({"key": key, "when": h, "mark": on})
            if off:
                spec["unmark"] = off
            spec["_sample"] = picked[first] if multi else (picked[first][0] if picked[first] else None)
        confirm.append(f"{letters[0]}〜{letters[-1]}列の {on} を、データでは {key}: "
                       + (f"[{', '.join(heads[:2])}] のような見出しの配列（複数選択）" if multi else f"{heads[0]} のような見出しの 1 つ（択一）")
                       + "で書く")
    plain.update({column_index_from_string(L): s for L, s in cols_def.items() if s.get("key") == "$index"})
    dates = [c for c in plain if str(plain[c].get("header") or "").strip() in DATE_HEADS]
    for run in runs(dates):
        if len(run) < 2:
            continue
        parts = [DATE_HEADS[str(plain[c]["header"]).strip()] for c in run]
        if len(set(parts)) < len(parts):
            continue
        for c in run:
            used.discard(plain[c]["key"])
        key = fresh(group_head(run) or "日付")
        got = dict(zip(parts, (text(t["body"][first], c) for c in run)))
        try:
            sample = dt.datetime(int(got.get("year", 2000)), int(got.get("month", 1)), int(got.get("day", 1)),
                                 int(got.get("hour", 0)), int(got.get("minute", 0)))
            sample = sample.date().isoformat() if not {"hour", "minute"} & set(parts) else sample.isoformat(sep=" ", timespec="minutes")
        except (TypeError, ValueError):
            sample = None
        for c, part in zip(run, parts):
            plain[c].update({"key": key, "part": part, "_sample": sample})
        confirm.append(f"{get_column_letter(run[0])}〜{get_column_letter(run[-1])}列の"
                       f"{'・'.join(str(plain[c]['header']).strip() for c in run)}を、データでは {key}: 2026-10-08 のような 1 つの日付で書く")


# ---------------------------------------------------------------------------
# 文書の値（表紙・タイトルの年度・期間・宛名・作成日など）。データでもラベルでもなく、文書ごとに変わる値
# ---------------------------------------------------------------------------

DOC_PREFIX = "文書"
# ラベルがこれなら、その右の値は表の外の 1 件の値ではなく、文書全体の値（どのタブでも同じ値が入る）
DOC_LABEL_RE = re.compile(r"^(作成日|作成者|作成部署|発行日|発行者|提出日|提出先|報告日|報告者|更新日|改訂日|日付|"
                          r"文書番号|文書名|資料番号|管理番号|版|版数|バージョン|Ver\.?|宛先|宛名|件名|表題|タイトル|"
                          r"プロジェクト名?|案件名|システム名|対象期間|期間|報告期間|年度)$", re.I)


def _date_fmt(text: str, m: "re.Match", codes: dict) -> str:
    """一致した日付の文字から、同じ見た目に戻す書式（`2025/07/01` → `%Y/%m/%d`、`7月` → `%-m月`）を作る。"""
    # 月・日のどれかが 0 埋め（07）なら 0 埋めの書式、どれも 1 桁か 10 以上なら 0 を付けない書式
    padded = any(m.group(g).startswith("0") for g, code in codes.items() if code != "Y")
    out, pos = "", m.start()
    for g, code in codes.items():
        a, b = m.span(g)
        out += text[pos:a].replace("%", "%%") + (f"%{code}" if code == "Y" or padded else f"%-{code}")
        pos = b
    return out + text[pos:m.end()].replace("%", "%%")


# 文の中の可変の部分。(名前, 正規表現, 種類)。上から順に、重ならないものを拾う
VAR_PATTERNS = [
    ("日付", re.compile(r"(?<!\d)(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日"), "date"),
    ("日付", re.compile(r"(?<![\d./-])(\d{4})([/.\-])(\d{1,2})\2(\d{1,2})(?![\d./-])"), "date_sep"),
    ("年月", re.compile(r"(?<!\d)(\d{4})\s*年\s*(\d{1,2})\s*月(?!\s*\d)"), "ym"),
    ("年月", re.compile(r"(?<![\d./-])(\d{4})/(\d{1,2})(?![\d./-])"), "ym_sep"),
    ("年度", re.compile(r"(?<!\d)(\d{4})\s*年度"), "int1"),
    ("和暦年度", re.compile(r"(?:令和|平成)\s*(\d{1,2}|元)\s*年度"), "wareki"),
    ("和暦年", re.compile(r"(?:令和|平成)\s*(\d{1,2}|元)\s*年(?!度)"), "wareki"),
    ("年", re.compile(r"(?<!\d)(\d{4})\s*年(?![度\d])"), "int1"),
    ("四半期", re.compile(r"第\s*([1-4])\s*四半期"), "int1"),
    ("四半期", re.compile(r"(?<![A-Za-z0-9])Q([1-4])(?![A-Za-z0-9])"), "int1"),
    ("四半期", re.compile(r"(?<![A-Za-z0-9])([1-4])Q(?![A-Za-z0-9])"), "int1"),
    ("半期", re.compile(r"(上|下)半?期"), "str1"),
    ("月", re.compile(r"(?<![\d/.\-])(\d{1,2})\s*月(?=度|分|[\s）)]|$)"), "int1"),
    ("版", re.compile(r"第\s*(\d+(?:\.\d+)*)\s*版"), "str1"),
    ("版", re.compile(r"(?<![A-Za-z])(?:Ver\.?|ver\.?|[vV])\s?(\d+(?:\.\d+)+)"), "str1"),
]
VAR_ADDRESSEE_RE = re.compile(r"^(?P<name>\S.*?)\s*(?P<tail>御中|様|殿)(?=\s|$|向け)")   # 「北斗製薬様 定例報告」の文頭の宛名も
NOT_ADDRESSEE = {"お客", "皆", "各位", "関係者各位"}
LABELLED_RE = re.compile(r"^(?P<label>[^:：\d]{1,15})\s*[:：]\s*")


def variable_parts(text: str) -> "tuple[str, list[tuple[str, Any]]] | None":
    """文の中の、文書ごとに変わりそうな部分（年度・四半期・年月・日付・版・宛名）を、欄（{名前}）にしたひな形にする。

    返すのは (ひな形, [(欄の名前, 今の値), …])。可変の部分が無ければ None。欄の名前は、まだ `文書.` を付けない仮の名前。
    注記（※）や長い文は対象にしない。
    """
    if not isinstance(text, str) or not text.strip() or len(text) > 80 or NOTE_RE.match(text) or PLACEHOLDER_RE.search(text):
        return None
    hits: list[tuple[int, int, str, str, Any]] = []   # (先頭, 末尾, 名前, 欄の中身, 値)
    taken: list[tuple[int, int]] = []
    m = VAR_ADDRESSEE_RE.match(text)
    if m and len(m.group("name").strip()) >= 2 and m.group("name").strip() not in NOT_ADDRESSEE \
            and not DOC_LABEL_RE.match(m.group("name").strip()) and not any(rx.search(m.group("name")) for _, rx, _ in VAR_PATTERNS):
        hits.append((m.start("name"), m.end("name"), "宛先", "", m.group("name").strip()))
        taken.append(m.span())
    for name, rx, kind in VAR_PATTERNS:
        for m in rx.finditer(text):
            if any(a < m.end() and m.start() < b for a, b in taken):
                continue
            if kind == "date":
                fmt = _date_fmt(text, m, {1: "Y", 2: "m", 3: "d"})
                value = f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
                span = m.span()
            elif kind == "date_sep":
                fmt = _date_fmt(text, m, {1: "Y", 3: "m", 4: "d"})
                value = f"{int(m.group(1)):04d}-{int(m.group(3)):02d}-{int(m.group(4)):02d}"
                span = m.span()
            elif kind in ("ym", "ym_sep"):
                fmt = _date_fmt(text, m, {1: "Y", 2: "m"})
                value = f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"
                span = m.span()
            else:
                fmt = ""
                g = m.group(1)
                value = 1 if g == "元" else int(g) if kind in ("int1", "wareki") else g
                span = m.span(1)
            taken.append(m.span())
            hits.append((span[0], span[1], name, fmt, value))
    if not hits:
        return None
    hits.sort()
    lab = LABELLED_RE.match(text)
    dates = [h for h in hits if h[2] == "日付"]
    names: list[str] = []
    for h in hits:
        name = h[2]
        if lab and len(hits) == len(dates) and dates:
            label = lab.group("label").strip()
            name = label if len(dates) == 1 else (f"{label}.開始" if h is dates[0] else f"{label}.終了") if len(dates) == 2 \
                else label
        names.append(name)
    # 同じ名前が 2 つ以上なら 2, 3 … を付ける（同じ文の中で、違う値を同じ欄にしない）
    seen: dict[str, int] = {}
    out, fields, pos = "", [], 0
    for (a, b, _, fmt, value), name in zip(hits, names):
        seen[name] = seen.get(name, 0) + 1
        if seen[name] > 1:
            name = f"{name}{seen[name]}"
        out += text[pos:a].replace("{", "{{").replace("}", "}}") + "{" + name + (f":{fmt}" if fmt else "") + "}"
        fields.append((name, value))
        pos = b
    out += text[pos:].replace("{", "{{").replace("}", "}}")
    return out, fields


class DocValues:
    """ブック全体の文書の値の名前。同じ名前で同じ値なら、同じキーにする（表紙とヘッダーの年度に、1 つの値が入る）。"""

    def __init__(self) -> None:
        self.values: dict[str, Any] = {}

    def key(self, name: str, value: Any) -> str:
        base = re.sub(r"[\s]+", "_", name.strip(" :：")) or "値"
        cand, n = base, 2
        while cand in self.values and self.values[cand] != value and value is not None and self.values[cand] is not None:
            cand, n = f"{base}{n}", n + 1
        if self.values.get(cand) is None:
            self.values[cand] = value
        return f"{DOC_PREFIX}.{cand}"

    def template(self, tpl: str, fields: list[tuple[str, Any]]) -> str:
        """variable_parts の仮の名前を、文書の値のキーに置き換える。"""
        for name, value in fields:
            key = self.key(name, value)
            tpl = re.sub(r"\{" + re.escape(name) + r"(?=[:}])", "{" + key, tpl, count=1)
        return tpl


HF_CODE_RE = re.compile(r'&(?:"[^"]*"|\d+|[A-Za-z&+\-])')


def _header_footer_template(text: str, doc: DocValues) -> "str | None":
    """ヘッダー・フッターの文字の、可変の部分を欄にする（& の記号は、そのまま残す）。"""
    out, pos, found = "", 0, False
    pieces = []
    for m in HF_CODE_RE.finditer(text):
        pieces += [(False, text[pos:m.start()]), (True, m.group(0))]
        pos = m.end()
    pieces.append((False, text[pos:]))
    for is_code, piece in pieces:
        got = None if is_code else variable_parts(piece)
        if got:
            found = True
            out += doc.template(*got)
        else:
            out += piece.replace("{", "{{").replace("}", "}}")
    return out if found else None


def _fixed_cells(grid, sample_rows: set[int], header_rows: set[int], root,
                 doc: "DocValues | None" = None) -> tuple[dict, list[str], dict, list[str]]:
    """表のサンプル行の外のセルを、ラベル（keep）と、流し込む欄（cells）に分ける。

    ラベルの右隣のセルは、値があってもなくても流し込む欄にする（前の文書の値を持ち越さない・空欄の記入枠も落とさない）。
    ラベルの無い数値・日付も、流し込む欄にする。それ以外の文字（タイトル・見出し・注記）がラベル（keep）。
    ラベルと同じ書式の空欄（見出しの帯の続き）は、記入枠と見なさない。
    押印・署名など人が書き込む欄は、データに入れない（前の値があれば clear で消す）。
    """
    origin = _merge_origins(root)
    cells: dict[str, Any] = {}
    keep: list[str] = []
    clear: list[str] = []
    samples: dict[str, dict] = {}
    # 2 行以上の結合の枠（所感・備考の記入欄）で、すぐ上にラベルがあるもの。空か仮の文（〜を記入）なら記入欄
    boxes: dict[tuple[int, int], str] = {}
    spans: dict[tuple[int, int], int] = {}   # 結合の左上 → 右端の列
    for (c, r), (c1, r1) in origin.items():
        spans[(c1, r1)] = max(spans.get((c1, r1), c1), c)
        if (c, r) == (c1, r1) and (c, r + 1) in origin and origin[(c, r + 1)] == (c1, r1):
            above = grid.get(r - 1, {}).get(c)
            v = grid.get(r, {}).get(c, {}).get("value")
            if above and isinstance(above["value"], str) and above["value"].strip() and not above["formula"] \
                    and (v is None or isinstance(v, str) and PLACEHOLDER_RE.search(v)):
                boxes[(c, r)] = above["value"]

    def new_key(label: "str | None", ref: str) -> str:
        base = re.sub(r"[.\s]+", "_", (label or "").strip(" :：")).strip("_") or ref
        taken = {k for v in cells.values() for k in spec_keys(v)}
        key, n = base, 2
        while key in taken:
            key, n = f"{base}{n}", n + 1
        return key

    for r in sorted(grid):
        if r in sample_rows:
            continue
        row = grid[r]
        # 「2026 年 9 月 25 日」のように、数と年・月・日の文字が交互に並ぶ欄は、1 つの日付（part）
        dates: dict[int, str] = {}
        units: set[int] = set()
        c = min(row) if row else 0
        while row and c <= max(row):
            run, k = [], c
            while k in row and (row[k]["value"] is None or isinstance(row[k]["value"], int)) and not row[k]["formula"] \
                    and str((row.get(k + 1) or {}).get("value") or "").strip() in DATE_HEADS:
                run.append(k)
                k += 2
            parts = [DATE_HEADS[str(row[k + 1]["value"]).strip()] for k in run]
            if len(run) >= 2 and len(set(parts)) == len(parts):
                dates.update(dict(zip(run, parts)))
                units.update(k + 1 for k in run)
                c = k
            else:
                c += 1
        date_key = None
        labels: dict[int, str] = {}   # この行のラベルの列 → 文字
        for c in sorted(row):
            i = row[c]
            if origin.get((c, r), (c, r)) != (c, r) or i["formula"]:
                continue   # 結合の左上以外のセル・数式は、どちらにもしない
            lc, lr = origin.get((c - 1, r), (c - 1, r))
            label = labels.get(lc) if lr == r and r not in header_rows else None
            value = i["value"]
            if c in units:
                keep.append(i["ref"])
                continue
            if c in dates:
                date_key = date_key or new_key(label or "日付", i["ref"])
                cells[i["ref"]] = {"key": date_key, "part": dates[c]}
                samples[i["ref"]] = {"label": label, "value": value}
                continue
            if (c, r) in boxes and r not in header_rows:
                cells[i["ref"]] = new_key(boxes[(c, r)], i["ref"])
                samples[i["ref"]] = {"label": boxes[(c, r)], "value": value}
                continue
            if value is None and label is not None and row[lc]["s"] == i["s"]:
                continue   # ラベルと同じ書式の空欄は、行の塗りの続き。記入枠ではない
            if r in header_rows:
                if value is not None:
                    keep.append(i["ref"])
                continue
            if label is not None and HUMAN_RE.search(label.strip(" :：")):
                if value is not None:
                    clear.append(i["ref"])   # 人が書き込む欄の前の値は消すだけ。データには入れない
                continue
            nxt = row.get(spans.get((c, r), c) + 1)
            if label is not None and isinstance(value, str) and len(value.strip()) <= 10 and nxt and not nxt["formula"] \
                    and nxt["value"] is not None and len(str(nxt["value"]).strip()) > 2:
                labels[c] = value   # 「振込先 | 銀行名 | ○○銀行」の中の見出し。右の値のラベル
                keep.append(i["ref"])
                continue
            if label is not None and doc is not None and DOC_LABEL_RE.match(label.strip(" :：")):
                # 作成日・作成者・版など、文書全体の値。どのタブでも同じキーにし、データの「文書」にまとめる
                key = doc.key(label.strip(" :："), value)
                got = variable_parts(value) if isinstance(value, str) else None
                if got and len(got[1]) == 1 and got[0].startswith("{" + got[1][0][0]) and got[0].endswith("}") \
                        and ":" in got[0] and not got[0].endswith(":%Y-%m-%d}"):   # 日付のセル（2025-04-01）はキーのまま
                    # 「2025年7月1日」のような文字の日付は、同じ見た目で入るように書式つきの文にする
                    cells[i["ref"]] = {"text": "{" + key + got[0][len(got[1][0][0]) + 1:]}
                else:
                    cells[i["ref"]] = key
                samples[i["ref"]] = {"label": label, "value": value}
            elif label is not None or (value is not None and not isinstance(value, str)):
                cells[i["ref"]] = new_key(label, i["ref"])
                samples[i["ref"]] = {"label": label, "value": value}
            elif value is not None:
                nxt_value = row.get(spans.get((c, r), c) + 1, {}).get("value")
                got = variable_parts(value) if doc is not None and nxt_value is None else None
                if got:
                    # タイトル・宛名の中の、年度・期間・日付・宛先。文の残りはそのまま、可変の部分だけを流し込む
                    cells[i["ref"]] = {"text": doc.template(*got)}
                    samples[i["ref"]] = {"label": None, "value": value}
                    continue
                labels[c] = value
                keep.append(i["ref"])
    return cells, compress_cells(keep) if keep else [], samples, compress_cells(clear) if clear else []


def _sheet_warnings(pkg, part, root, tables) -> list[str]:
    w = []
    spans = [(t["first_row"], t["first_row"] + t["sample_rows"] - 1) for t in tables]

    def inside(ref_text: str) -> bool:
        for piece in ref_text.split():
            for p in piece.split(":"):
                m = REF_PART_RE.match(p.replace("$", ""))
                if m and any(a <= int(m.group(4)) <= b for a, b in spans):
                    return True
        return False
    for mc in root.iter(q("mergeCell")):
        if inside(mc.get("ref")):
            r1, r2 = (split_ref(p)[1] for p in (mc.get("ref").split(":") * 2)[:2])
            if r1 != r2:
                w.append(f"サンプル行内の複数行にまたがる結合セル {mc.get('ref')} は、1 件が複数行（block_rows）で"
                         "その 1 件の中に収まれば各件に複製される。収まらなければ取り除かれる")
            else:
                w.append(f"サンプル行内の結合セル {mc.get('ref')} は、各行に複製される")
    for cf in root.iter(q("conditionalFormatting")):
        if inside(cf.get("sqref", "")):
            w.append(f"条件付き書式 {cf.get('sqref')} は、表の行数に合わせて範囲が伸びる")
    for dv in root.iter(q("dataValidation")):
        if inside(dv.get("sqref", "")):
            w.append(f"入力規則 {dv.get('sqref')} は、表の行数に合わせて範囲が伸びる")
    rels = pkg.rels_of(part)
    for _, (typ, target) in rels.items():
        if typ.endswith("/table"):
            w.append(f"Excel のテーブル（{target}）は、範囲を表の行数に合わせて更新する")
    return w


def summarize(definition: dict) -> str:
    """ユーザー確認用の要約テキスト。"""
    lines = []
    for s in definition["sheets"]:
        lines.append(f"■ シート「{s['name']}」")
        if not s["tables"]:
            lines.append("  可変の表は見つかりませんでした（固定セルだけのシート）")
        for t in s["tables"]:
            end = t["first_row"] + t["sample_rows"] - 1
            lines.append(f"  [{t['id']}] 見出し {t['header_row']} 行 / サンプル {t['first_row']}-{end} 行"
                         f"（{t['sample_rows']} 行）/ 繰り返し元 {t['pattern']} / データのキー: {t['key']}")
            for letter, c in t["columns"].items():
                kind = f"数式 {c['_sample']}" if c.get("formula") else f"key={c.get('key')}  例: {c.get('_sample')!r}"
                if c.get("keep"):
                    kind = f"定数として残す（データに書かない）  値: {c.get('_sample')!r}"
                elif c.get("clear"):
                    kind = "人が書き込む欄（データに書かず、空欄で出す）"
                if "when" in c:
                    kind += f"（{c['when']!r} を含むとき {c.get('mark', '○')}）"
                elif "map" in c:
                    pairs = [f"{k}→{'空欄' if v is None else v}" for k, v in c["map"].items()]
                    kind += f"（{', '.join(pairs)}）"
                elif "part" in c:
                    kind += f"（日付の {c['part']}）"
                lines.append(f"      {letter}列 「{c.get('header') or '（見出し無し）'}」 → {kind}  [{c.get('_format', '')}]")
            if t["_total_row"]:
                lines.append(f"      合計行: {t['_total_row']} 行（表の下へずらし、SUM の範囲は表に合わせて伸ばす）")
            for n in t["needs_confirm"]:
                lines.append(f"      ? {n}")
        if s["cells"]:
            lines.append("  流し込む欄（cells。テンプレートの値は残さない。データで null なら空欄になる）:")
            for ref, key in s["cells"].items():
                v = s["_cell_samples"].get(ref, {}).get("value")
                if isinstance(key, dict) and "text" in key:
                    lines.append(f"      {ref} → 文={key['text']!r}  今の値: {v!r}")
                    continue
                if isinstance(key, dict):
                    key = f"{key['key']}（日付の {key['part']}）" if "part" in key else key["key"]
                lines.append(f"      {ref} → key={key}  " + (f"今の値: {v!r}" if v is not None else "（空欄の記入枠）"))
        for tag, tpl in (s.get("header_footer") or {}).items():
            lines.append(f"  ヘッダー・フッター {tag} → 文={tpl!r}")
        if s["keep"]:
            lines.append(f"  残す（keep。見出し・ラベル・固定の文面）: {', '.join(s['keep'])}")
        if s.get("clear"):
            lines.append(f"  空にする（clear。人が書き込む欄の前の値）: {', '.join(s['clear'])}")
        for n in s["_notes"]:
            lines.append(f"  ! {n}")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# inspect（テンプレートの事実だけを、判断する側（LLM・人）が読める形で出す）
# ---------------------------------------------------------------------------

# 「仮」は 1 文字だけでは見ない（仮払金・仮受消費税のような、帳票の見出しを仮の値と取り違えないため）
PLACEHOLDER_RE = re.compile(r"(〇〇|○○|●●|◯◯|△△|□□|＊＊|\*\*|(?<![A-Za-z])x{2,}(?![A-Za-z])|サンプル|ダミー|記入例|"
                            r"^仮$|[（(]仮[）)]|仮(?:の|入力|置き|名|データ)|例[:：)）]|sample|dummy|example|yyyy|"
                            r"\b0{4}[-/]0{2}[-/]0{2}\b|\bTBD\b|を(?:記入|入力)(?:する|して|してください)?[）)]?$|^[（(].*(?:記入|入力).*[）)]$)", re.I)
NOTE_RE = re.compile(r"^\s*(※|＊|\*|注[:：）)]|備考[:：]|Note[:：])")   # 列の見出しの「備考」は注記ではない


def _style_classes(styles: Styles) -> dict[str, str]:
    """書式の説明が同じスタイル ID を、同じ種類（S1, S2, ...）にまとめる。"""
    out: dict[str, str] = {}
    labels: dict[str, str] = {}
    for sid in sorted(styles.desc):
        d = styles.desc[sid]
        labels.setdefault(d, f"S{len(labels)}")
        out[str(sid)] = labels[d]
    return out


def inspect_template(template: "str | bytes") -> dict:
    """判断の材料になる事実（値・数式・書式の種類・結合・仮値の疑い・自動検出の下書き）を集める。"""
    pkg = Package(template)
    sst = read_shared_strings(pkg)
    styles = Styles(pkg)
    classes = _style_classes(styles)
    legend = {label: desc for desc, label in {styles.desc[i]: classes[str(i)] for i in styles.desc}.items()}
    sheets = []
    for name, part in pkg.sheets():
        root = pkg.xml(part)
        expand_shared_formulas(root)
        grid = _grid(root, sst, styles)
        draft = _analyze_sheet(pkg, name, part, root, sst, styles)
        in_tables = {r for t in draft["tables"] for r in range(t["header_row"], t["first_row"] + t["sample_rows"])}
        rows = []
        for r in sorted(grid):
            cells, blanks = [], []
            for c in sorted(grid[r]):
                i = grid[r][c]
                cls = classes.get(i["s"], "S0")
                if i["value"] is None and not i["formula"]:
                    blanks.append(c)
                    continue
                item = {"ref": i["ref"], "style": cls}
                if i["formula"]:
                    item["formula"] = i["formula"]
                else:
                    item["value"] = i["value"]
                    if isinstance(i["value"], str):
                        if PLACEHOLDER_RE.search(i["value"]):
                            item["hint"] = "仮の値の疑い"
                        elif NOTE_RE.match(i["value"]):
                            item["hint"] = "注記の疑い"
                        elif r not in in_tables and variable_parts(i["value"]):
                            names = "・".join(dict.fromkeys(n for n, _ in variable_parts(i["value"])[1]))
                            item["hint"] = f"文書ごとに変わる値の疑い（{names}）"
                cells.append(item)
            rows.append({"row": r, "cells": cells,
                         "styled_blank": f"{get_column_letter(min(blanks))}-{get_column_letter(max(blanks))}" if blanks else None,
                         "shape": [(c, classes.get(grid[r][c]["s"], "S0")) for c in sorted(grid[r])]})
        sheets.append({
            "name": name,
            "merges": [m.get("ref") for m in root.iter(q("mergeCell"))],
            "conditional_formats": [c.get("sqref") for c in root.iter(q("conditionalFormatting"))],
            "validations": [d.get("sqref") for d in root.iter(q("dataValidation"))],
            "header_footer": {tag: el.text for tag, el in _header_footer(root).items() if el.text},
            "rows": rows,
            "same_shape_runs": _runs(rows),
            "auto_detected_tables": [{k: t[k] for k in ("id", "header_row", "first_row", "sample_rows", "pattern")}
                                     for t in draft["tables"]],
            "warnings": draft["_notes"],
        })
    return {"style_legend": legend, "provenance": provenance(pkg), "sheets": sheets}


def _runs(rows: list[dict]) -> list[dict]:
    """列の並びと書式の種類が同じ行が続く範囲（繰り返しの候補）。周期 2〜3 の縞模様も拾う。"""
    out = []
    i = 0
    while i < len(rows):
        j = i
        while j + 1 < len(rows) and rows[j + 1]["row"] == rows[j]["row"] + 1 and rows[j + 1]["shape"] == rows[i]["shape"]:
            j += 1
        if j > i:
            out.append({"rows": f"{rows[i]['row']}-{rows[j]['row']}", "period": 1})
            i = j + 1
            continue
        for p in (2, 3):
            k = i
            while k + p < len(rows) + 0 and rows[k + p]["row"] == rows[k]["row"] + p and rows[k + p]["shape"] == rows[k]["shape"] \
                    and all(rows[k + m + 1]["row"] == rows[k + m]["row"] + 1 for m in range(p)):
                k += 1
            if k - i >= p:
                out.append({"rows": f"{rows[i]['row']}-{rows[k + p - 1]['row']}", "period": p})
                i = k + p
                break
        else:
            i += 1
    return out


FOLD_RUN = 6   # 同じ書式の行がこれより長く続くと、テキストでは先頭 3 行と末尾 1 行だけ見せる


def _folded_rows(sh: dict) -> dict[int, int]:
    """畳む行 → その範囲で畳んだ行数（最初の 1 行にだけ数を付け、残りは 0）。"""
    out: dict[int, int] = {}
    for run in sh["same_shape_runs"]:
        a, b = (int(x) for x in run["rows"].split("-"))
        if b - a + 1 > FOLD_RUN:
            # 仮の値・注記の疑いがある行は畳まない（流用のとき、持ち越しを見落とさない）
            hidden = [r["row"] for r in sh["rows"] if a + 3 <= r["row"] < b and not any("hint" in c for c in r["cells"])]
            for k, r in enumerate(hidden):
                out[r] = len(hidden) if k == 0 else 0
    return out


def format_facts(facts: dict, fold: bool = True) -> str:
    L = ["書式の種類（S0〜: 同じ見た目は同じ番号）:"]
    L += [f"  {k} = {v}" for k, v in sorted(facts["style_legend"].items(), key=lambda kv: int(kv[0][1:]))]
    if facts.get("provenance"):
        L.append("\n来歴・持ち越しの注意（analyze の下書きの properties.scrub で消える。外部リンク・非表示・変更履歴は残るので扱いを決める）:")
        L += [f"  - {p}" for p in facts["provenance"]]
    for sh in facts["sheets"]:
        L.append(f"\n■ シート「{sh['name']}」")
        for key, label in (("merges", "結合"), ("conditional_formats", "条件付き書式"), ("validations", "入力規則")):
            if sh[key]:
                L.append(f"  {label}: {' '.join(sh[key])}")
        for tag, text in (sh.get("header_footer") or {}).items():
            L.append(f"  ヘッダー・フッター {tag}: {text!r}")
        L.append("  行（値/数式 [書式の種類]）:")
        folded = _folded_rows(sh) if fold else {}
        for row in sh["rows"]:
            if row["row"] in folded:
                if folded[row["row"]]:
                    L.append(f"         …（同じ書式の行 {folded[row['row']]} 行を省略。--all ですべて出す）")
                continue
            parts = []
            for c in row["cells"]:
                body = c["formula"] if "formula" in c else repr(c["value"])
                parts.append(f"{c['ref']}: {body} [{c['style']}]" + (f" ⚠{c['hint']}" if "hint" in c else ""))
            if row["styled_blank"]:
                parts.append(f"（値なし・書式のみ {row['styled_blank']}）")
            L.append(f"    {row['row']:>3}: " + "  ".join(parts) if parts else f"    {row['row']:>3}: （値なし）")
        if sh["same_shape_runs"]:
            L.append("  同じ書式の行が続く範囲（繰り返しの候補）: " + ", ".join(
                f"{r['rows']}（{r['period']} 行周期）" for r in sh["same_shape_runs"]))
        for t in sh["auto_detected_tables"]:
            end = t["first_row"] + t["sample_rows"] - 1
            L.append(f"  自動検出の下書き: 見出し {t['header_row']} 行 / サンプル {t['first_row']}-{end} 行 / 繰り返し元 {t['pattern']}（確定ではない）")
        for w in sh["warnings"]:
            L.append(f"  ! {w}")
    return "\n".join(L)


# ---------------------------------------------------------------------------
# render
# ---------------------------------------------------------------------------

def dig(data: Any, path: str) -> Any:
    cur = data
    for part in path.split("."):
        if isinstance(cur, dict) and part in cur:
            cur = cur[part]
        elif isinstance(cur, list) and part.isdigit() and int(part) < len(cur):
            cur = cur[int(part)]
        else:
            raise KeyError(path)
    return cur


def _excel_serial(value: str) -> float | None:
    for parser in (dt.datetime.fromisoformat, dt.date.fromisoformat):
        try:
            d = parser(value)
        except ValueError:
            continue
        if isinstance(d, dt.datetime):
            delta = d - dt.datetime(1899, 12, 30)
            return delta.days + delta.seconds / 86400
        return (d - dt.date(1899, 12, 30)).days
    return None


ILLEGAL_XML_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
MAX_CELL_TEXT = 32767  # Excel の 1 セルの文字数の上限
NUMERIC_TEXT_RE = re.compile(r"^[+-]?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?$")


def set_value(c, value: Any, styles: Styles, replace_formula: bool = False) -> None:
    """セルの値を置き換える。replace_formula なら、数式もデータの値で置き換える（流し込む欄）。"""
    for child in list(c):
        if child.tag != q("f") or replace_formula:
            c.remove(child)
    c.attrib.pop("t", None)
    if value is None:
        return
    if isinstance(value, bool):
        c.set("t", "b")
        etree.SubElement(c, q("v")).text = "1" if value else "0"
    elif isinstance(value, (int, float)) and styles.code.get(int(c.get("s") or 0)) in ("@", "builtin:49"):
        # 文字列の書式（@）の欄（版 1.0・郵便番号など）は、YAML が数値と読んでも、文字のまま入れる
        c.set("t", "inlineStr")
        t = etree.SubElement(etree.SubElement(c, q("is")), q("t"))
        t.text = str(value)
    elif isinstance(value, (int, float)):
        if isinstance(value, float) and (value != value or value in (float("inf"), float("-inf"))):
            raise TemplateError(f"{c.get('r')} に入れる数値 {value!r} は Excel に保存できません")
        etree.SubElement(c, q("v")).text = repr(value)
    else:
        text = str(value)
        if styles.is_number(c.get("s")) and NUMERIC_TEXT_RE.match(text.strip()):
            # CSV などから来た `"1,200"` を文字のまま入れると、SUM が数えず合計が黙って狂う
            number = float(text.strip().replace(",", ""))
            etree.SubElement(c, q("v")).text = repr(int(number) if number.is_integer() and "." not in text else number)
            return
        if ILLEGAL_XML_RE.search(text):
            raise TemplateError(f"{c.get('r')} に入れる文字列に、Excel に保存できない制御文字があります: {text[:40]!r}")
        if len(text) > MAX_CELL_TEXT:
            raise TemplateError(f"{c.get('r')} に入れる文字列が {len(text)} 文字あり、Excel の上限 {MAX_CELL_TEXT} 文字を超えています")
        serial = _excel_serial(text) if styles.is_date(c.get("s")) else None
        if serial is not None:
            etree.SubElement(c, q("v")).text = repr(serial)
        else:
            c.set("t", "inlineStr")
            is_ = etree.SubElement(c, q("is"))
            t = etree.SubElement(is_, q("t"))
            t.text = text
            t.set(XML_SPACE, "preserve")


def _warn_not_number(c, value: Any, styles: "Styles", where: str, warnings: list) -> None:
    """テンプレートでは数値の欄に、数値として読めない文字を入れるとき知らせる（数量「たくさん」・単価「8,200円」）。

    文字のまま入るので、その欄を使う数式（金額・小計・合計）が #VALUE! になる。
    """
    if not isinstance(value, str) or NUMERIC_TEXT_RE.match(value.strip()) or styles.is_date(c.get("s")):
        return
    v = c.find(q("v"))
    was_number = c.get("t") in (None, "n") and v is not None and v.text and c.find(q("f")) is None
    if was_number or styles.is_number(c.get("s")):
        warnings.append(f"{where} の値 {value!r} は数値ではありません（{c.get('r')} はテンプレートでは数値の欄。"
                        "文字のまま入り、この欄を使う数式は #VALUE! になる）")


def _get_or_make_cell(row, col: int, r: int):
    ref = f"{get_column_letter(col)}{r}"
    for c in row:
        cc = cell_col(c)
        if cc == col:
            return c
        if cc > col:
            new = etree.Element(q("c"), r=ref)
            c.addprevious(new)
            return new
    return etree.SubElement(row, q("c"), r=ref)


def _set_row_number(row, r: int) -> None:
    row.set("r", str(r))
    row.attrib.pop("spans", None)
    for c in row:
        col = re.match(r"[A-Za-z]+", c.get("r")).group(0)
        c.set("r", f"{col}{r}")


def _strip_cached(row) -> None:
    for c in row:
        if c.find(q("f")) is not None:
            v = c.find(q("v"))
            if v is not None:
                c.remove(v)
            c.attrib.pop("t", None) if c.get("t") in ("str", "e", "b", "n") else None


def render(template: "str | bytes", definition: dict, data: dict, output: str) -> list[str]:
    """テンプレートへ流し込んで output に書く。警告メッセージのリストを返す。"""
    if definition.get("version") != DEF_VERSION:
        raise TemplateError(f"定義ファイルの version が未対応です: {definition.get('version')!r}")
    pkg = Package(template)
    sst = read_shared_strings(pkg)
    styles = Styles(pkg)
    sheet_parts = dict(pkg.sheets())
    # scrub なら、プロパティ・プレビュー・コメントは出力から取り除く。残るもの（外部リンク・非表示など）だけを来歴として挙げる
    scrubbed = ("文書のプロパティ", "カスタムプロパティ", "プレビュー画像", "コメント", "スレッド形式のコメント")
    warnings: list[str] = [f"来歴: {p}" for p in provenance(pkg)
                           if not ((definition.get("properties") or {}).get("scrub") and p.startswith(scrubbed))]
    leftovers: dict[str, set] = {}

    # --- 1. シートごとの表の出力行数を決め、RowMap を作る -----------------------
    plans: dict[str, dict] = {}
    maps: dict[str, RowMap] = {}
    for sd in definition["sheets"]:
        name = sd["name"]
        if name not in sheet_parts:
            raise TemplateError(f"テンプレートにシート「{name}」がありません")
        tbls = []
        for t in sd.get("tables", []):
            try:
                rows = dig(data, t["key"])
            except KeyError:
                raise TemplateError(f"データに {t['key']!r} がありません（シート「{name}」表 {t['id']}）")
            if not isinstance(rows, list):
                raise TemplateError(f"{t['key']!r} は配列である必要があります")
            first, count = int(t["first_row"]), int(t["sample_rows"])
            block = int(t.get("block_rows", 1))
            if block < 1 or count % block:
                raise TemplateError(f"sample_rows ({count}) は block_rows ({block}) の倍数にしてください（表 {t['id']}）")
            pattern = [int(p) for p in t.get("pattern") or [first]]
            for p in pattern:
                if not first <= p <= first + count - block:
                    raise TemplateError(f"pattern の行 {p} がサンプル行 {first}-{first + count - block}（ブロックの先頭になれる範囲）の外です")
            warnings += _missing_key_warnings(t, block, rows, name)
            tbls.append({"def": t, "first": first, "count": count, "end": first + count - 1, "block": block,
                         "n": max(len(rows), 1) * block, "rows": rows, "pattern": pattern})
        for spec in sd.get("drop_rows") or []:  # 無視する行: 出力から取り除く
            a, b = parse_row_spec(spec)
            tbls.append({"def": {}, "first": a, "count": b - a + 1, "end": b, "block": 1, "n": 0,
                         "rows": [], "pattern": [], "drop": True})
        tbls.sort(key=lambda x: x["first"])
        for a, b in zip(tbls, tbls[1:]):
            if a["end"] >= b["first"]:
                raise TemplateError(f"シート「{name}」で表のサンプル行が重なっています")
        plans[name] = {"def": sd, "tables": tbls}
        maps[name] = RowMap(tbls)

    # --- 2. シート XML を再構成 -----------------------------------------------
    for name, part in sheet_parts.items():
        root = pkg.xml(part)
        rw = Rewriter(name, maps)
        if name in plans:
            _render_sheet(pkg, part, root, plans[name], maps[name], rw, data, styles, warnings,
                          leftovers.setdefault(name, set()))
        elif maps:
            for f in root.iter(q("f")):
                if f.text and "!" in f.text:
                    f.text = rw.formula(f.text)
        pkg.put_xml(part, root)

    # --- 3. ブック全体の参照（定義名・グラフ）と再計算の設定 --------------------
    _fix_workbook(pkg, maps)
    props = {k: fill_text(v, data, f"properties の {k}") if isinstance(v, str) else v
             for k, v in (definition.get("properties") or {}).items()}
    used = {k for sd in definition["sheets"] for k in sheet_cell_keys(sd)}
    used |= {k for v in (definition.get("properties") or {}).values() if isinstance(v, str) for k, _ in text_fields(v)}
    warnings += unused_doc_warnings(data, used)
    apply_properties(pkg, props)
    if props.get("scrub"):
        n = drop_comments(pkg)
        if n:
            warnings.append(f"コメント（メモ）{n} 件を取り除きました（properties.scrub）")
        scrub_leftover_text(pkg)
    if definition.get("strict"):
        total = sum(len(refs) for refs in leftovers.values())
        if total:
            by_sheet = "\n".join(f"  {n}: {', '.join(compress_cells(refs))}" for n, refs in leftovers.items() if refs)
            raise TemplateError(
                f"テンプレートの値がそのまま残るセルが {total} 個あります（シートごと・範囲にまとめて）:\n{by_sheet}\n"
                "残すなら keep、置き換えるなら cells / columns の key、空にするなら clear、行ごと消すなら drop_rows に入れてください")
    pkg.save(output)
    return warnings


def _missing_key_warnings(t: dict, block: int, rows: list, sheet: str) -> list[str]:
    """列の key が、データのどの行にも無い（綴りの違い）と、その列は黙って空欄になる。それを知らせる。"""
    records = [r for r in rows if isinstance(r, dict)]
    if not records:
        return []
    wanted = [spec["key"] for cols in _column_maps(t, block) for spec in cols.values()
              if spec.get("key") and spec["key"] != "$index" and not spec.get("formula")
              and not spec.get("keep") and not spec.get("clear")]
    def found(k: str) -> bool:
        return any(k in r or _row_value(r, k) is not None for r in records)
    missing = [k for k in dict.fromkeys(wanted) if not found(k)]
    heads = {k.split(".")[0] for k in wanted} | set(wanted)
    unused = [k for k in dict.fromkeys(k for r in records for k in r) if k not in heads]
    # 空欄がふつうの列（備考など）もあるので、綴り違いの手がかり（使われていないキー）があるか、どの key も当たらないときだけ
    if not missing or not unused and len(missing) < len(set(wanted)):
        return []
    return [f"シート「{sheet}」表 {t.get('id')}: 列の key {', '.join(map(repr, missing))} が、"
            f"データ {t['key']!r} のどの行にも無いため空欄になります"
            + (f"（データにあって使われていないキー: {', '.join(map(repr, unused))}）" if unused else "")]


def compress_cells(refs) -> list[str]:
    """セルの一覧を、keep などにそのまま書ける範囲（`A1:C1`・`A3:A5`）にまとめる。行ごとに横へ、同じ横幅は縦へつなぐ。"""
    rows: dict[int, list[int]] = {}
    for ref in refs:
        col, r = split_ref(ref)
        rows.setdefault(r, []).append(col)
    spans: list[tuple[int, int, int]] = []   # (行, 先頭の列, 末尾の列)
    for r in sorted(rows):
        cols = sorted(rows[r])
        start = prev = cols[0]
        for c in cols[1:] + [None]:
            if c is not None and c == prev + 1:
                prev = c
                continue
            spans.append((r, start, prev))
            if c is not None:
                start = prev = c
    blocks: list[list[int]] = []   # [先頭の行, 末尾の行, 先頭の列, 末尾の列]
    for r, a, b in spans:
        hit = next((x for x in blocks if x[1] == r - 1 and x[2] == a and x[3] == b), None)
        if hit:
            hit[1] = r
        else:
            blocks.append([r, r, a, b])
    out = []
    for r1, r2, a, b in sorted(blocks):
        first, last = f"{get_column_letter(a)}{r1}", f"{get_column_letter(b)}{r2}"
        out.append(first if first == last else f"{first}:{last}")
    return out


def _has_literal(c) -> bool:
    if c.find(q("f")) is not None:
        return False
    if c.find(q("is")) is not None:
        return True
    v = c.find(q("v"))
    return v is not None and v.text not in (None, "")


def _render_sheet(pkg, part, root, plan, rowmap, rw, data, styles, warnings, leftovers) -> None:
    sd, tables = plan["def"], plan["tables"]
    keep_set = {pos for spec in sd.get("keep") or [] for pos in parse_cell_spec(spec)}
    expand_shared_formulas(root)
    sheet_data = root.find(q("sheetData"))
    orig = sheet_rows(root)
    tbl_of: dict[int, dict] = {}
    for t in tables:
        for r in range(t["first"], t["end"] + 1):
            tbl_of[r] = t

    # 定義に残した列の見出し（analyze が書く header）と、テンプレートの見出しの行を突き合わせる。
    # 列を足した・並べ替えたテンプレートに差し替えると、黙って 1 列ずつずれて入るので止める
    sst = read_shared_strings(pkg)
    moved = []
    for t in tables:
        td, head_r = t["def"], t["def"].get("header_row")
        if t.get("drop") or not head_r or int(head_r) not in orig:
            continue
        moved += _header_mismatches(td, {cell_col(c): cell_value(c, sst) for c in orig[int(head_r)]})
    if moved:
        raise TemplateError(f"シート「{sd['name']}」の表の見出しが、定義と合いません（列を足した・並べ替えたテンプレートでは、"
                            "値が別の列に入ります）。analyze から定義を作り直すか、columns の列を直してください:\n  "
                            + "\n  ".join(moved))

    fixed_cells = {}
    cell_specs = {ref: (spec if isinstance(spec, dict) else {"key": spec}) for ref, spec in (sd.get("cells") or {}).items()}
    for ref, spec in cell_specs.items():
        col, r = split_ref(ref)
        if r in tbl_of:
            raise TemplateError(f"cells の {ref} は表のサンプル行の中です")
        if "text" in spec:   # 文の一部だけを差し替える（タイトルの年度・期間など）
            fixed_cells[(col, r)] = fill_text(spec["text"], data, f"cells の {ref}")
            continue
        try:
            value = dig(data, spec["key"])
        except KeyError:
            raise TemplateError(f"データに {spec['key']!r} がありません（cells の {ref}）")
        _check_choice(spec, value, [s for s in cell_specs.values() if s.get("key") == spec["key"]], f"cells の {ref}")
        fixed_cells[(col, r)] = convert_value(spec, value, f"cells の {ref}")

    clear_cells: set[tuple[int, int]] = set()
    for spec in sd.get("clear") or []:  # 無視する値: 書式は残して空にする
        for col, r in parse_cell_spec(spec):
            if r in tbl_of:
                raise TemplateError(f"clear の {spec} は表のサンプル行の中です（列の clear を使う）")
            clear_cells.add((col, r))

    # ヘッダー・フッター（印刷したときの上下の文字）。& は Excel の記号なので、値の中の & は && にする
    existing = _header_footer(root)
    for tag, tpl in (sd.get("header_footer") or {}).items():
        if tag not in existing:
            raise TemplateError(f"シート「{sd['name']}」の header_footer の {tag} は、テンプレートにありません"
                                f"（使えるもの: {', '.join(existing) or 'なし'}）")
        existing[tag].text = fill_text(tpl, data, f"シート「{sd['name']}」の header_footer の {tag}",
                                       escape=lambda s: s.replace("&", "&&"))

    new_rows: list[etree._Element] = []
    pattern_map: dict[int, list[int]] = {}  # 元のサンプル行 → 出力行の一覧（結合セルの複製用）

    def emit_fixed(r: int, row) -> None:
        new_r = rowmap.map(r, False)
        row = deepcopy(row)
        for f in row.iter(q("f")):
            if f.text:
                f.text = rw.formula(f.text)
            if f.get("ref"):
                f.set("ref", rw.sqref(f.get("ref")))
        _set_row_number(row, new_r)
        for (col, rr) in clear_cells:
            if rr == r:
                for c in row:
                    if cell_col(c) == col and c.find(q("f")) is None:
                        set_value(c, None, styles)
        for (col, rr), value in fixed_cells.items():
            if rr == r:
                c = _get_or_make_cell(row, col, new_r)
                _warn_not_number(c, value, styles, f"cells の {get_column_letter(col)}{r}", warnings)
                set_value(c, value, styles, replace_formula=True)
        for c in row:
            pos = (cell_col(c), r)
            if _has_literal(c) and pos not in keep_set and pos not in fixed_cells:
                leftovers.add(f"{get_column_letter(pos[0])}{r}")
        _strip_cached(row)
        new_rows.append(row)

    done_tables: set[int] = set()
    for r in sorted(set(orig) | {rr for _, rr in fixed_cells}):
        if r in tbl_of:
            t = tbl_of[r]
            if id(t) in done_tables:
                continue
            done_tables.add(id(t))
            _emit_table(t, orig, rowmap, rw, styles, new_rows, pattern_map, warnings, leftovers)
            continue
        if r in orig:
            emit_fixed(r, orig[r])
        else:  # 行そのものが無い固定セルへの書き込み
            emit_fixed(r, etree.Element(q("row"), r=str(r)))
    # 行そのものが無い表（サンプル行が未作成）にも対応
    for t in tables:
        if id(t) not in done_tables:
            _emit_table(t, orig, rowmap, rw, styles, new_rows, pattern_map, warnings, leftovers)
    new_rows.sort(key=lambda x: int(x.get("r")))

    for child in list(sheet_data):
        sheet_data.remove(child)
    sheet_data.extend(new_rows)

    filled = set(fixed_cells) | clear_cells
    for t in tables:
        if t.get("drop"):
            continue
        colmaps = _column_maps(t["def"], t["block"])
        for r in range(t["first"], t["end"] + 1):
            filled |= {(col, r) for col, spec in colmaps[(r - t["first"]) % t["block"]].items()
                       if not spec.get("formula") and not spec.get("keep")}
    _drop_stale_links(pkg, part, root, filled, warnings, sd["name"])
    dropped_rows = {r for t in tables if t.get("drop") for r in range(t["first"], t["end"] + 1)}
    _fix_comments(pkg, part, filled, set(tbl_of), rowmap, warnings, dropped_rows, sd["name"])
    _fix_sheet_parts(pkg, part, root, tables, rowmap, rw, pattern_map, warnings)


def _header_mismatches(td: dict, heads: dict) -> list[str]:
    """表の見出しの行（列番号 → 値）と、定義の columns の header を突き合わせる。

    違う見出しの列と、定義に無いのに見出しのある列（表の右・左に足された列、間に挟まった列）を挙げる。
    """
    head_r = td.get("header_row")
    text = lambda v: str(v if v is not None else "").strip()
    out = []
    defined = {}
    for letter, spec in (td.get("columns") or {}).items():
        col = column_index_from_string(letter.upper())
        want = spec.get("header") if isinstance(spec, dict) else None
        defined[col] = want
        have = heads.get(col)
        if want is not None and text(want) != text(have):
            out.append(f"{letter}{head_r}: 定義では「{want}」、テンプレートでは「{have if text(have) else '（空）'}」")
    if not defined or not any(w is not None for w in defined.values()):
        return out   # 見出しを書いていない定義（手で書いたもの）は、足された列も見ない
    lo, hi = min(defined), max(defined)
    extra = [c for c in range(lo, hi + 1) if c not in defined and text(heads.get(c))]
    c = hi + 1
    while text(heads.get(c)):
        extra.append(c)
        c += 1
    c = lo - 1
    while c >= 1 and text(heads.get(c)):
        extra.append(c)
        c -= 1
    out += [f"{get_column_letter(c)}{head_r}: テンプレートに見出し「{text(heads[c])}」の列があるが、定義の columns に無い"
            for c in sorted(extra)]
    return out


def _column_maps(td: dict, block: int) -> list[dict]:
    """行ごとの列指定 {列番号: spec} のリスト。block_rows=1 なら columns、2 以上なら block（行ごとの columns）。"""
    def conv(cols):
        return {column_index_from_string(letter.upper()): spec for letter, spec in (cols or {}).items()}
    if block == 1:
        return [conv(td.get("columns"))]
    rows = td.get("block")
    if not isinstance(rows, list) or len(rows) != block:
        raise TemplateError(f"block_rows={block} の表 {td.get('id')} には、行ごとの列指定 block を {block} 個書いてください")
    return [conv(r) for r in rows]


def _emit_table(t, orig, rowmap, rw, styles, out_rows, pattern_map, warnings, leftovers) -> None:
    if t.get("drop"):
        return
    td, k = t["def"], t["block"]
    colmaps = _column_maps(td, k)
    new_first = t["first"] + rowmap.shift_before(t["first"])
    for rec in range(t["n"] // k):
        start = t["pattern"][rec % len(t["pattern"])]
        row_data = t["rows"][rec] if rec < len(t["rows"]) else None
        for j in range(k):
            src_r = start + j
            target_r = new_first + rec * k + j
            row = deepcopy(orig[src_r]) if src_r in orig else etree.Element(q("row"), r=str(src_r))
            pattern_map.setdefault(src_r, []).append(target_r)
            offset = target_r - (src_r + rowmap.shift_before(src_r))

            def row_fn(rm, idx, last, absolute, r, _off=offset):
                if absolute:
                    return rm.map(r, last)
                if r <= t["end"]:
                    return rm.map(r, False) + _off
                return rm.map(r, last)
            for c in row:
                f = c.find(q("f"))
                if f is not None and f.text:
                    f.text = rw.formula(f.text, row_fn)
                if f is not None and f.get("ref"):  # 配列数式の範囲も、複製先の行へずらす
                    f.set("ref", re.sub(r"([A-Za-z]+)(\d+)",
                                        lambda m: f"{m.group(1)}{int(m.group(2)) - src_r + target_r}", f.get("ref")))
            _set_row_number(row, target_r)
            for col, spec in colmaps[j].items():
                if spec.get("formula") or spec.get("keep"):
                    continue
                c = _get_or_make_cell(row, col, target_r)
                if spec.get("clear") or "key" not in spec:
                    set_value(c, None, styles)
                    continue
                key = spec["key"]
                if key == "$index":
                    value = rec + 1 if row_data is not None else None
                elif row_data is None:
                    value = None
                elif isinstance(row_data, dict):
                    value = _row_value(row_data, key)
                else:
                    raise TemplateError(f"{td['key']!r} の要素はオブジェクトである必要があります")
                where = f"{td['key']}[{rec}].{key}"
                if row_data is not None:   # データの無い空の 1 行は、印も入れずに空欄
                    _check_choice(spec, value, [s for s in colmaps[j].values() if s.get("key") == key], where)
                    value = convert_value(spec, value, where)
                if isinstance(value, (dict, list)):
                    raise TemplateError(f"{where} に配列・オブジェクトは入れられません")
                _warn_not_number(c, value, styles, where, warnings)
                set_value(c, value, styles, replace_formula=True)
            for c in row:
                if cell_col(c) not in colmaps[j] and _has_literal(c):
                    leftovers.add(f"{get_column_letter(cell_col(c))}{src_r}")
            _strip_cached(row)
            out_rows.append(row)


# データを人が読める形で書き、表の見た目（○×・コード・年月日の分割）へは定義で変換する
DATE_PARTS = ("year", "month", "day", "hour", "minute")


def convert_value(spec: dict, value: Any, where: str = "") -> Any:
    """列・セルの指定（when / map / part）に従って、データの値をセルに入れる値に変える。

    - when: 値が when と同じ（配列なら when を含む）なら mark（既定 ○）、違えば unmark（既定 空欄）。
      択一（`判定: 合格`）・複数選択（`対象: [Windows, Linux]`）を、見出しの列ごとの ○ にする
    - map: 値を表の表記に置き換える（`true: ○` / `false: ×`、`高: 1`）
    - part: 日付・日時の一部（year / month / day / hour / minute）。年・月・日に分かれた欄へ 1 つの日付を入れる
    """
    if "when" in spec:
        if value is None:
            return spec.get("unmark")   # 選ばれていない（テンプレートの × の作法のまま）
        hit = spec["when"] in value if isinstance(value, list) else str(value) == str(spec["when"])
        return spec.get("mark", "○") if hit else spec.get("unmark")
    if value is None:
        return None
    if "map" in spec:
        m = spec["map"]
        for k in ((value, str(value).lower(), str(value)) if isinstance(value, bool) else (value, str(value))):
            if k in m:
                return m[k]
        raise TemplateError(f"{where} の値 {value!r} は、定義の map にありません（使える値: {', '.join(map(str, m))}）")
    if "part" in spec:
        part = spec["part"]
        if part not in DATE_PARTS:
            raise TemplateError(f"{where} の part は {', '.join(DATE_PARTS)} のどれかにしてください: {part!r}")
        d = value
        if isinstance(d, str):
            try:
                d = dt.datetime.fromisoformat(d.strip().replace("/", "-"))
            except ValueError:
                raise TemplateError(f"{where} の値 {value!r} を日付として読めません（2026-10-08 の形で書く）")
        if not isinstance(d, (dt.date, dt.datetime)):
            raise TemplateError(f"{where} の値 {value!r} は日付ではありません")
        if part in ("hour", "minute") and not isinstance(d, dt.datetime):
            return 0
        return getattr(d, part)
    return value


def _check_choice(spec: dict, value: Any, group: list, where: str) -> None:
    """択一・複数選択の値が、どの列の when にも当たらないと、黙ってどこにも ○ が付かない。それを止める。"""
    if "when" not in spec or value is None or group[0] is not spec:
        return   # 同じキーの列のうち、最初の列でだけ確かめる
    choices = [str(s["when"]) for s in group if "when" in s]
    for v in (value if isinstance(value, list) else [value]):
        if str(v) not in choices:
            raise TemplateError(f"{where} の値 {v!r} は、選べる値（{', '.join(choices)}）のどれでもありません")


def _row_value(row_data: dict, key: str) -> Any:
    """行オブジェクトから列の値を取る。キーそのもの（`No.` など）が無ければ、`test.0` のようにドット・添字でたどる。"""
    if key in row_data:
        return row_data[key]
    try:
        return dig(row_data, key)
    except KeyError:
        return None


# ---------------------------------------------------------------------------
# 文の中の一部を差し替える（`{文書.年度}年度 第{文書.四半期}四半期 売上報告書`）
# ---------------------------------------------------------------------------

TEXT_FIELD_RE = re.compile(r"\{\{|\}\}|\{([^{}:]+)(?::([^{}]*))?\}")
HEADER_FOOTER_TAGS = ("oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter")


def text_fields(tpl: str) -> list[tuple[str, str]]:
    """文のひな形の中の欄 [(データのキー, 書式), …]。`{{`・`}}` は文字の波かっこ。"""
    out, rest, pos = [], "", 0
    for m in TEXT_FIELD_RE.finditer(str(tpl)):
        rest += str(tpl)[pos:m.start()]
        pos = m.end()
        if m.group(1) is not None:
            out.append((m.group(1).strip(), m.group(2) or ""))
    rest += str(tpl)[pos:]
    if "{" in rest or "}" in rest:
        raise TemplateError(f"文のひな形 {tpl!r} の波かっこが閉じていません（文字の波かっこは {{{{ と }}}} と書く）")
    return out


def spec_keys(spec: Any) -> list[str]:
    """固定セルの指定が使う、データのキーの一覧（文のひな形なら、その中の欄のキー）。"""
    if not isinstance(spec, dict):
        return [spec]
    if "text" in spec:
        return list(dict.fromkeys(k for k, _ in text_fields(spec["text"])))
    return [spec["key"]]


def _as_date(value: Any, where: str) -> "dt.date | dt.datetime":
    if isinstance(value, (dt.date, dt.datetime)):
        return value
    s = str(value).strip().replace("/", "-")
    m = re.fullmatch(r"(\d{4})-(\d{1,2})", s)   # 年月だけ（2026-10）
    if m:
        return dt.date(int(m.group(1)), int(m.group(2)), 1)
    try:
        return dt.datetime.fromisoformat(s) if (" " in s or "T" in s) else dt.date.fromisoformat(
            "-".join(p.zfill(2) for p in s.split("-")))
    except ValueError:
        raise TemplateError(f"{where} の値 {value!r} を日付として読めません（2026-10-08 か、年月なら 2026-10 の形で書く）")


def _strftime(d: "dt.date | dt.datetime", fmt: str) -> str:
    """strftime と同じ。ただし `%-m`・`%-d`・`%-H`（0 を付けない）は、どの OS でも使える。"""
    def sub(m: "re.Match") -> str:
        code = m.group(1)
        if code.startswith("-"):
            return str(int(d.strftime("%" + code[1:])))
        return "%" if code == "%" else d.strftime("%" + code)
    return re.sub(r"%(-?[A-Za-z%])", sub, fmt)


def _field_text(value: Any, fmt: str, where: str) -> str:
    if value is None:
        return ""
    if "%" in fmt:
        return _strftime(_as_date(value, where), fmt)
    if fmt:
        try:
            return format(value, fmt)
        except (ValueError, TypeError):
            raise TemplateError(f"{where} の値 {value!r} に書式 {fmt!r} を使えません")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    return str(value)


def fill_text(tpl: str, data: Any, where: str, escape=None) -> str:
    """文のひな形の欄を、データの値で埋める。値が null の欄は空にする。

    data はデータか、キーから値を引く関数（無ければ KeyError）。
    """
    lookup = data if callable(data) else (lambda key: dig(data, key))
    text_fields(tpl)   # 閉じていない波かっこを先に止める

    def sub(m: "re.Match") -> str:
        if m.group(1) is None:
            return m.group(0)[0]
        key = m.group(1).strip()
        try:
            value = lookup(key)
        except KeyError:
            raise TemplateError(f"データに {key!r} がありません（{where}）")
        text = _field_text(value, m.group(2) or "", f"{where} の {key}")
        return escape(text) if escape else text
    return TEXT_FIELD_RE.sub(sub, str(tpl))


def read_text(tpl: str, text: Any) -> "dict | None":
    """fill_text の逆。文から欄の値を読み戻す（合わなければ None）。数字だけは数値、日付の書式は 2026-10-08 の形にする。"""
    if text is None:
        return None
    pattern, names, n = "", {}, 0
    pos = 0
    for m in TEXT_FIELD_RE.finditer(str(tpl)):
        pattern += re.escape(tpl[pos:m.start()])
        pos = m.end()
        if m.group(1) is None:
            pattern += re.escape(m.group(0)[0])
            continue
        key = m.group(1).strip()
        if key in names:
            pattern += f"(?P={names[key][0]})"
            continue
        names[key] = (f"g{n}", m.group(2) or "")
        pattern += f"(?P<g{n}>.*?)"
        n += 1
    pattern += re.escape(tpl[pos:])
    m = re.fullmatch(pattern, str(text).strip(), re.S) or re.fullmatch(pattern, str(text), re.S)
    if not m:
        return None
    out = {}
    for key, (g, fmt) in names.items():
        v = m.group(g).strip()
        if not v:
            out[key] = None
        elif "%" in fmt:
            try:
                d = dt.datetime.strptime(v, re.sub(r"%-", "%", fmt))
            except ValueError:
                out[key] = v
                continue
            if "%d" in fmt or "%-d" in fmt:
                out[key] = d.date().isoformat() if not re.search(r"%-?[HM]", fmt) else d.isoformat(sep=" ", timespec="minutes")
            else:
                out[key] = f"{d.year:04d}-{d.month:02d}"
        else:
            out[key] = int(v) if re.fullmatch(r"[+-]?\d+", v) and not (len(v) > 1 and v[0] == "0") else v
    return out


def text_notes(tpl: str) -> dict:
    """文のひな形の欄の書き方（雛形の null だけでは、文のどこに入るかが分からない）。"""
    notes = {}
    for key, fmt in text_fields(tpl):
        if "%" in fmt:
            notes[key] = f"{key}: 日付（2026-10-08）" if re.search(r"%-?d", fmt) else f"{key}: 年月（2026-10）"
        else:
            notes.setdefault(key, f"{key}: 「{tpl}」の {{{key}}} に入る文字")
    return notes


def unused_doc_warnings(data: Any, used: "set[str]") -> list[str]:
    """データの `文書:` にあって、定義のどこでも使われていない値（綴りの違いの疑い）。"""
    def leaves(obj, path):
        if isinstance(obj, dict) and obj:
            for k, v in obj.items():
                yield from leaves(v, f"{path}.{k}")
        else:
            yield path
    doc = data.get(DOC_PREFIX) if isinstance(data, dict) else None
    if not isinstance(doc, dict):
        return []
    unused = [p for p in leaves(doc, DOC_PREFIX) if not any(p == u or p.startswith(u + ".") or u.startswith(p + ".") for u in used)]
    return [f"データの {', '.join(unused)} は、定義のどこでも使われていません（綴りの違いを疑う）"] if unused else []


def _header_footer(root) -> dict:
    hf = root.find(q("headerFooter"))
    return {} if hf is None else {el.tag.split("}")[1]: el for el in hf if el.tag.split("}")[1] in HEADER_FOOTER_TAGS}


def _drop_stale_links(pkg, part, root, filled: set, warnings: list, sheet: str = "") -> None:
    """データで置き換える・空にするセルのハイパーリンクを取り除く。残すと、新しい値に元の（サンプルの）リンク先が付く。"""
    links = root.find(q("hyperlinks"))
    if links is None:
        return
    dropped = []
    for h in list(links):
        if any(pos in filled for spec in (h.get("ref") or "").split() for pos in parse_cell_spec(spec)):
            links.remove(h)
            dropped.append(h.get("ref"))
    if not dropped:
        return
    warnings.append(f"シート「{sheet}」: データを入れるセルのハイパーリンク（{', '.join(dropped)}）は、元のリンク先のままになるため取り除きました")
    if not len(links):
        root.remove(links)
    rels_name = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if rels_name in pkg.data:  # どこからも指さなくなったリンク先（元の URL）も残さない
        used = {v for el in root.iter() for k, v in el.attrib.items() if k == f"{{{NS_R}}}id"}
        rels = pkg.xml(rels_name)
        for r in list(rels):
            if r.get("Type", "").endswith("/hyperlink") and r.get("Id") not in used:
                rels.remove(r)
        pkg.put_xml(rels_name, rels)


NS_VML_X = "urn:schemas-microsoft-com:office:excel"


def _fix_comments(pkg, part, filled: set, sample_rows: set, rowmap, warnings: list, dropped_rows: set = frozenset(),
                  sheet: str = "") -> None:
    """コメント（メモ）を行のずれに追従させる。データを入れる・空にするセルのものは取り除く（新しい値に、元のメモが付く）。"""
    parts = {typ.rsplit("/", 1)[-1]: target for typ, target in pkg.rels_of(part).values() if target in pkg.data}
    if "comments" not in parts:
        return
    croot = pkg.xml(parts["comments"])
    vml = None
    if "vmlDrawing" in parts:
        try:
            vml = etree.fromstring(pkg.data[parts["vmlDrawing"]], etree.XMLParser(recover=True))
        except etree.XMLSyntaxError:
            vml = None
    shapes = {}
    for cd in (vml.iter(f"{{{NS_VML_X}}}ClientData") if vml is not None else []):
        r, c = cd.find(f"{{{NS_VML_X}}}Row"), cd.find(f"{{{NS_VML_X}}}Column")
        if r is not None and c is not None and (r.text or "").strip().isdigit() and (c.text or "").strip().isdigit():
            shapes[(int(c.text) + 1, int(r.text) + 1)] = (cd.getparent(), r)
    dropped, gone = [], []   # データを入れるセルのもの・取り除く行のもの
    for cm in list(croot.iter(q("comment"))):
        col, r = split_ref(cm.get("ref"))
        shape = shapes.get((col, r))
        if (col, r) in filled or r in dropped_rows:
            cm.getparent().remove(cm)
            if shape is not None and shape[0].getparent() is not None:
                shape[0].getparent().remove(shape[0])
            (gone if r in dropped_rows else dropped).append(cm.get("ref"))
        elif r not in sample_rows:   # 表の外の行は、行と一緒に動かす
            new_r = rowmap.map(r, False)
            cm.set("ref", f"{get_column_letter(col)}{new_r}")
            if shape is not None:
                shape[1].text = str(new_r - 1)
    if dropped:
        warnings.append(f"シート「{sheet}」: データを入れるセルのコメント（{', '.join(dropped)}）は、新しい値に元のメモが付くため取り除きました")
    if gone:
        warnings.append(f"シート「{sheet}」: 取り除く行のコメント（{', '.join(gone)}）も取り除きました")
    pkg.put_xml(parts["comments"], croot)
    if vml is not None:
        pkg.data[parts["vmlDrawing"]] = etree.tostring(vml)


def _fix_sheet_parts(pkg, part, root, tables, rowmap, rw, pattern_map, warnings) -> None:
    # 結合セル: 表の外は移動、サンプル行内は複製
    mc = root.find(q("mergeCells"))
    if mc is not None:
        keep = []
        for m in list(mc):
            a, b = (m.get("ref").split(":") * 2)[:2]
            (c1, r1), (c2, r2) = split_ref(a), split_ref(b)
            t = next((t for t in tables if t["first"] <= r1 <= t["end"] or t["first"] <= r2 <= t["end"]), None)
            mc.remove(m)
            if t and t["first"] <= r1 and r2 <= t["end"]:
                same_block = (r1 - t["first"]) // t["block"] == (r2 - t["first"]) // t["block"]
                if r1 == r2 or same_block:   # 1 件（ブロック）の中の結合は、その行を複製した先ごとに複製する
                    for target in pattern_map.get(r1, []):
                        keep.append(f"{get_column_letter(c1)}{target}:{get_column_letter(c2)}{target + r2 - r1}")
                else:
                    warnings.append(f"複数行にまたがる結合セル {m.get('ref')} は表のサンプル行内のため取り除きました")
                continue
            span = rowmap.map_span(r1, r2)
            if span is not None:
                keep.append(f"{get_column_letter(c1)}{span[0]}:{get_column_letter(c2)}{span[1]}")
        for ref in keep:
            etree.SubElement(mc, q("mergeCell"), ref=ref)
        mc.set("count", str(len(keep)))
        if not keep:
            root.remove(mc)

    # 範囲を持つ要素
    # 範囲がすべて消えた要素は取り除く（空の sqref は、Excel が壊れたファイルと見なす）
    for tag in ("conditionalFormatting", "dataValidation", "ignoredError"):
        for el in list(root.iter(q(tag))):
            if el.get("sqref"):
                el.set("sqref", rw.sqref(el.get("sqref")))
                if not el.get("sqref"):
                    el.getparent().remove(el)
    for el in list(root.iter(q("hyperlink"))):
        el.set("ref", rw.sqref(el.get("ref")))
        if not el.get("ref"):
            el.getparent().remove(el)
    for tag in ("dataValidations", "hyperlinks"):   # 子が無くなった入れ物も取り除き、件数を合わせる
        for el in list(root.iter(q(tag))):
            if len(el) == 0:
                el.getparent().remove(el)
            elif el.get("count") is not None:
                el.set("count", str(len(el)))
    for el in list(root.iter(q("ignoredErrors"))):
        if len(el) == 0:
            el.getparent().remove(el)
    af = root.find(q("autoFilter"))
    if af is not None and af.get("ref"):
        af.set("ref", rw.sqref(af.get("ref")))
    # 条件付き書式・入力規則の数式は、表の下にある参照だけを動かす（相対参照の意味を壊さない）
    for tag in ("formula", "formula1", "formula2"):
        for el in root.iter(q(tag)):
            if el.text:
                el.text = rw.formula(el.text, lambda rm, idx, last, absolute, r: rm.map(r, False))
    for brk in root.iter(q("brk")):
        if brk.get("id"):
            brk.set("id", str(rowmap.map(int(brk.get("id")), False)))

    # dimension
    dim = root.find(q("dimension"))
    if dim is not None:
        refs = [(cell_col(c), int(c.getparent().get("r"))) for c in root.iter(q("c"))]
        if refs:
            cols, rows = [x[0] for x in refs], [x[1] for x in refs]
            dim.set("ref", f"{get_column_letter(min(cols))}{min(rows)}:{get_column_letter(max(cols))}{max(rows)}")

    for _, (typ, target) in pkg.rels_of(part).items():
        if typ.endswith("/table") and target in pkg.data:  # Excel テーブル
            troot = pkg.xml(target)
            for el in (troot, troot.find(q("autoFilter"))):
                if el is not None and el.get("ref"):
                    el.set("ref", rw.sqref(el.get("ref")))
            pkg.put_xml(target, troot)
        if typ.endswith("/drawing") and target in pkg.data:  # 図・グラフの位置
            droot = pkg.xml(target)
            for tag in ("from", "to"):
                for el in droot.iter(f"{{{NS_XDR}}}{tag}"):
                    row = el.find(f"{{{NS_XDR}}}row")
                    if row is not None:
                        row.text = str(rowmap.map(int(row.text) + 1, False) - 1)
            pkg.put_xml(target, droot)


def _fix_workbook(pkg: Package, maps: dict[str, RowMap]) -> None:
    # グラフの参照範囲
    for name in list(pkg.data):
        if name.startswith("xl/charts/chart") and name.endswith(".xml"):
            croot = pkg.xml(name)
            changed = False
            for f in croot.iter(f"{{{NS_C}}}f"):
                if f.text:
                    new = Rewriter("", maps).formula(f.text)
                    if new != f.text:
                        f.text, changed = new, True
            if changed:
                pkg.put_xml(name, croot)
    wb = pkg.xml("xl/workbook.xml")
    dn = wb.find(q("definedNames"))
    if dn is not None:
        for d in dn:
            if d.text:
                d.text = Rewriter("", maps).formula(d.text)
    # 再計算: 開いた時に全再計算させ、古い計算チェーンは捨てる
    calc = wb.find(q("calcPr"))
    if calc is None:
        calc = etree.Element(q("calcPr"))
        anchor = None
        for tag in ("definedNames", "externalReferences", "functionGroups", "sheets"):
            anchor = wb.find(q(tag))
            if anchor is not None:
                break
        anchor.addnext(calc)
    calc.set("fullCalcOnLoad", "1")
    pkg.put_xml("xl/workbook.xml", wb)
    if "xl/calcChain.xml" in pkg.data:
        _drop_part(pkg, "xl/calcChain.xml")


def _drop_part(pkg: Package, part: str) -> None:
    """part を取り除き、part 自身の rels、それを指す rels、[Content_Types].xml の Override も消す。"""
    pkg.removed.add(part)
    pkg.removed.add(posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels"))
    for name in list(pkg.data):
        if not name.endswith(".rels") or name in pkg.removed:
            continue
        base = posixpath.dirname(posixpath.dirname(name))
        root = pkg.xml(name)
        changed = False
        for r in list(root):
            target = r.get("Target", "")
            if r.get("TargetMode") == "External":
                continue
            resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
            if resolved == part:
                root.remove(r)
                changed = True
        if changed:
            pkg.put_xml(name, root)
    ct = pkg.xml("[Content_Types].xml")
    for o in list(ct):
        if o.get("PartName", "").lstrip("/") == part:
            ct.remove(o)
    pkg.put_xml("[Content_Types].xml", ct)


# ---------------------------------------------------------------------------
# 来歴（他プロジェクトの成果物を流用するときの持ち越し）
# ---------------------------------------------------------------------------

NS_DC = "http://purl.org/dc/elements/1.1/"
NS_CP = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
NS_DCTERMS = "http://purl.org/dc/terms/"
NS_APP = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
CORE_FIELDS = {"creator": f"{{{NS_DC}}}creator", "lastModifiedBy": f"{{{NS_CP}}}lastModifiedBy",
               "title": f"{{{NS_DC}}}title", "subject": f"{{{NS_DC}}}subject",
               "description": f"{{{NS_DC}}}description", "keywords": f"{{{NS_CP}}}keywords",
               "category": f"{{{NS_CP}}}category"}
APP_FIELDS = {"company": f"{{{NS_APP}}}Company", "manager": f"{{{NS_APP}}}Manager"}


def read_properties(pkg: Package) -> dict[str, str]:
    """文書のプロパティ（作成者・会社名など）のうち、値があるもの。"""
    out: dict[str, str] = {}
    for part, fields in (("docProps/core.xml", CORE_FIELDS), ("docProps/app.xml", APP_FIELDS)):
        if part not in pkg.data:
            continue
        root = pkg.xml(part)
        for key, tag in fields.items():
            el = root.find(tag)
            if el is not None and (el.text or "").strip():
                out[key] = el.text.strip()
    return out


def provenance(pkg: Package) -> list[str]:
    """出力に持ち越されると困るかもしれないものを、文で返す。"""
    out = []
    props = read_properties(pkg)
    if props:
        out.append("文書のプロパティ: " + "、".join(f"{k}={v!r}" for k, v in props.items()))
    names = pkg.data.keys()
    if "docProps/custom.xml" in names:
        out.append("カスタムプロパティ（docProps/custom.xml）がある")
    if any(n.startswith("docProps/thumbnail") for n in names):
        out.append("プレビュー画像（docProps/thumbnail）がある。元の内容が写っている")
    authors = set()
    where: list[str] = []
    for sheet_name, sheet_part in pkg.sheets():  # 配置場所は Excel と openpyxl で違うので、rels の種類で探す
        for typ, target in pkg.rels_of(sheet_part).values():
            if typ.endswith("/comments") and target in pkg.data:
                root = pkg.xml(target)
                authors |= {a.text for a in root.iter(q("author")) if a.text}
                for cm in root.iter(q("comment")):
                    text = "".join(t.text or "" for t in cm.iter(q("t"))).strip().replace("\n", " ")
                    where.append(f"{sheet_name}!{cm.get('ref')}「{text[:30]}{'…' if len(text) > 30 else ''}」")
    if where:
        out.append(f"コメント（メモ）が {len(where)} 件ある。作成者: {'、'.join(sorted(authors))}。"
                   + "、".join(where[:5]) + (f" ほか {len(where) - 5} 件" if len(where) > 5 else ""))
    if any(n.startswith("xl/threadedComments") for n in names):
        out.append("スレッド形式のコメントがある")
    if any(n.startswith("xl/externalLinks/") and n.endswith(".xml") for n in names):
        out.append("他のブックへの外部リンクがある（元のファイル名・パスが残る）")
    if "xl/vbaProject.bin" in names:
        out.append("マクロ（VBA）がある。出力の拡張子は .xlsm にする")
    wb = pkg.xml("xl/workbook.xml")
    hidden = [s.get("name") for s in wb.find(q("sheets")) if s.get("state") in ("hidden", "veryHidden")]
    if hidden:
        out.append("非表示のシートがある: " + "、".join(hidden))
    n_hidden = 0
    for _, part in pkg.sheets():
        root = pkg.xml(part)
        n_hidden += sum(1 for r in root.iter(q("row")) if r.get("hidden") == "1")
        n_hidden += sum(1 for c in root.iter(q("col")) if c.get("hidden") == "1")
    if n_hidden:
        out.append(f"非表示の行・列が {n_hidden} 個ある")
    if any(n.startswith("xl/revisions/") for n in names):
        out.append("変更履歴（リビジョン）がある")
    return out


def apply_properties(pkg: Package, props: dict) -> None:
    """定義の properties を反映する。scrub: true なら、指定の無い識別情報を空にし、プレビューとカスタムプロパティを取り除く。"""
    scrub = bool(props.get("scrub"))
    for part, fields in (("docProps/core.xml", CORE_FIELDS), ("docProps/app.xml", APP_FIELDS)):
        if part not in pkg.data:
            continue
        root = pkg.xml(part)
        changed = False
        for key, tag in fields.items():
            value = props[key] if key in props else ("" if scrub else None)
            if value is None:
                continue
            el = root.find(tag)
            if el is None:
                if not value:
                    continue
                el = etree.SubElement(root, tag)
            if (el.text or "") != str(value):
                el.text = str(value) or None
                changed = True
        if changed:  # 変えないものは、元のバイト列のまま
            pkg.put_xml(part, root)
    if scrub and "docProps/core.xml" in pkg.data:
        # 作成日時・更新日時・印刷日時も、前の文書の来歴。作成と更新は今の日時にし、印刷は取り除く
        root = pkg.xml("docProps/core.xml")
        now = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        for tag in (f"{{{NS_DCTERMS}}}created", f"{{{NS_DCTERMS}}}modified"):
            el = root.find(tag)
            if el is not None:
                el.text = now
        printed = root.find(f"{{{NS_CP}}}lastPrinted")
        if printed is not None:
            root.remove(printed)
        pkg.put_xml("docProps/core.xml", root)
    if scrub:
        for name in list(pkg.data):
            if name == "docProps/custom.xml" or name.startswith("docProps/thumbnail"):
                _drop_part(pkg, name)


def drop_comments(pkg: Package) -> int:
    """コメント（メモ・スレッド形式）を、吹き出しの図形と一緒にすべて取り除く。取り除いた件数を返す。"""
    total = 0
    for _, part in pkg.sheets():
        rels = pkg.rels_of(part)
        targets = {typ.rsplit("/", 1)[-1]: (rid, target) for rid, (typ, target) in rels.items() if target in pkg.data}
        if "comments" not in targets and "threadedComment" not in targets:
            continue
        for kind in ("comments", "threadedComment"):
            if kind in targets:
                target = targets[kind][1]
                if kind == "comments":
                    total += len(list(pkg.xml(target).iter(q("comment"))))
                _drop_part(pkg, target)
        if "vmlDrawing" not in targets:
            continue
        rid, vml_part = targets["vmlDrawing"]
        try:
            vml = etree.fromstring(pkg.data[vml_part], etree.XMLParser(recover=True))
        except etree.XMLSyntaxError:
            continue
        for cd in list(vml.iter(f"{{{NS_VML_X}}}ClientData")):
            if cd.get("ObjectType") == "Note" and cd.getparent() is not None and cd.getparent().getparent() is not None:
                shape = cd.getparent()
                shape.getparent().remove(shape)
        if any(cd.get("ObjectType") != "Note" for cd in vml.iter(f"{{{NS_VML_X}}}ClientData")):
            pkg.data[vml_part] = etree.tostring(vml)   # フォームのボタンなど、メモ以外の図形は残す
            continue
        _drop_part(pkg, vml_part)
        root = pkg.xml(part)
        for ld in list(root.iter(q("legacyDrawing"))):
            if ld.get(f"{{{NS_R}}}id") == rid:
                ld.getparent().remove(ld)
        pkg.put_xml(part, root)
    if not any(n.startswith("xl/threadedComments/") for n in pkg.data if n not in pkg.removed):
        for name in list(pkg.data):
            if name.startswith("xl/persons/") and name.endswith(".xml") and name not in pkg.removed:
                _drop_part(pkg, name)   # スレッド形式のコメントの作成者の一覧
    return total


def scrub_leftover_text(pkg: Package) -> None:
    """どのセルも参照しなくなった共有文字列と、グラフの古い値のキャッシュを取り除く。"""
    name = "xl/sharedStrings.xml"
    if name in pkg.data:
        sst = pkg.xml(name)
        items = list(sst)
        roots = {part: pkg.xml(part) for _, part in pkg.sheets()}
        used = sorted({int(c.find(q("v")).text) for r in roots.values() for c in r.iter(q("c"))
                       if c.get("t") == "s" and c.find(q("v")) is not None})
        remap = {old: new for new, old in enumerate(used)}
        total = 0
        for part, root in roots.items():
            for c in root.iter(q("c")):
                if c.get("t") == "s" and c.find(q("v")) is not None:
                    c.find(q("v")).text = str(remap[int(c.find(q("v")).text)])
                    total += 1
            pkg.put_xml(part, root)
        for it in items:
            sst.remove(it)
        for old in used:
            sst.append(items[old])
        sst.set("count", str(total))
        sst.set("uniqueCount", str(len(used)))
        pkg.put_xml(name, sst)
    for part in list(pkg.data):
        if part.startswith("xl/charts/chart") and part.endswith(".xml"):
            root = pkg.xml(part)
            for tag in ("numCache", "strCache", "multiLvlStrCache"):
                for el in list(root.iter(f"{{{NS_C}}}{tag}")):
                    el.getparent().remove(el)
            pkg.put_xml(part, root)


# ---------------------------------------------------------------------------
# 入出力（JSON / YAML）
# ---------------------------------------------------------------------------

def _yaml():
    try:
        import yaml
    except ImportError:
        raise TemplateError("YAML の入出力には PyYAML が必要です（uv add pyyaml / pip install pyyaml）")
    return yaml


def parse_structured(text: str, name: str = "") -> Any:
    """拡張子が .yaml / .yml なら YAML、.json なら JSON。名前が無い（標準入力）なら JSON → YAML の順に試す。"""
    lower = name.lower()
    try:
        if lower.endswith((".yaml", ".yml")):
            return _yaml().safe_load(text)
        if lower.endswith(".json") or not lower:
            try:
                return json.loads(text)
            except json.JSONDecodeError:
                if lower:
                    raise
                return _yaml().safe_load(text)
        return json.loads(text)
    except (json.JSONDecodeError, ValueError) as e:
        raise TemplateError(f"{name or '入力'} を読めません: {e}")
    except Exception as e:  # yaml.YAMLError
        if isinstance(e, TemplateError):
            raise
        raise TemplateError(f"{name or '入力'} を読めません: {e}")


def load_structured(path: str) -> Any:
    if path == "-":
        return parse_structured(sys.stdin.read())
    with open(path, encoding="utf-8") as f:
        return parse_structured(f.read(), path)


DATA_EXTS = (".json", ".yaml", ".yml")


def data_files(paths: "str | list[str]") -> list[str]:
    """データの指定（ファイル・フォルダ・-）を、読む順のファイルの並びにする。フォルダは中のデータファイルを名前順に。"""
    out: list[str] = []
    for p in [paths] if isinstance(paths, str) else paths:
        if p != "-" and os.path.isdir(p):
            found = sorted(n for n in os.listdir(p) if n.lower().endswith(DATA_EXTS) and not n.startswith("."))
            if not found:
                raise TemplateError(f"フォルダ {p} にデータファイル（{' / '.join(DATA_EXTS)}）がありません")
            out += [os.path.join(p, n) for n in found]
        else:
            out.append(p)
    return out


def merge_data(parts: list, definition: "dict | None" = None) -> dict:
    """分けたデータ（[(ファイル名, 中身), …]）を 1 つにする。

    オブジェクトはキーごとに合わせ、配列（表の行）はファイルの順につなぐ。
    同じ欄に違う値があれば止める（どちらが正しいか分からない）。null は、ほかのファイルの値を消さない。
    表の同じ行が 2 つのファイルにあっても止める（前のデータを写したまま、ほかのファイルに残っている）。
    definition があれば、ほかのまとまり（sheets[].group）のデータも持つファイルから来た表の行も止める
    （出荷のファイルに、受注の表が前のまま写っている。行を直したあとでも、同じ行が無くても見つかる）。
    食い違いは、まとめて挙げる。
    """
    merged: dict = {}
    seen: dict[str, str] = {}   # 欄 → その値を最初に書いたファイル
    problems: list[str] = []

    def mark(v, path, src):
        seen.setdefault(path, src)
        if isinstance(v, dict):
            for k, x in v.items():
                mark(x, f"{path}.{k}", src)

    def put(a, b, path, src):
        if isinstance(a, dict) and isinstance(b, dict):
            for k, v in b.items():
                sub = f"{path}.{k}" if path else str(k)
                if k in a:
                    a[k] = put(a[k], v, sub, src)
                else:
                    a[k] = v
                    mark(v, sub, src)
            return a
        if isinstance(a, list) and isinstance(b, list):
            have = {_row_sig(x) for x in a if isinstance(x, dict)}
            dup = [x for x in b if isinstance(x, dict) and _row_sig(x) in have]
            if dup:
                problems.append(f"{path} の同じ行が、{seen.get(path, '前のファイル')} と {src} の両方にあります"
                                f"（{len(dup)} 件。前のデータを写したままなら、片方から消す）")
            return a + b
        if b is None or a == b:
            return a
        if a is None:
            seen[path] = src
            return b
        problems.append(f"{path} の値が、{seen.get(path, '前のファイル')} と {src} で違います: {a!r} / {b!r}")
        return a

    for src, obj in parts:
        if obj is None:
            continue
        if not isinstance(obj, dict):
            raise TemplateError(f"{src} の中身はオブジェクト（キーと値）にしてください")
        put(merged, deepcopy(obj), "", src)
    if definition is not None:
        problems += _stray_tables(parts, definition)
    if problems:
        raise TemplateError("分けたデータが食い違っています:\n" + "\n".join(f"  {p}" for p in problems))
    return merged


def load_data(paths: "str | list[str]", definition: "dict | None" = None) -> dict:
    """データを読む。複数のファイル・フォルダなら、1 つにまとめる（merge_data）。"""
    files = data_files(paths)
    if len(files) == 1:
        return load_structured(files[0])
    return merge_data([(f if f != "-" else "標準入力", load_structured(f)) for f in files], definition)


def _row_sig(row: dict) -> str:
    return json.dumps(row, ensure_ascii=False, sort_keys=True, default=str)


def sheet_texts(sd: dict) -> list[str]:
    """シートの、文のひな形（cells の text と header_footer）。"""
    return [s["text"] for s in (sd.get("cells") or {}).values() if isinstance(s, dict) and "text" in s] \
        + list((sd.get("header_footer") or {}).values())


def sheet_cell_keys(sd: dict) -> list[str]:
    """シートの固定セル・ヘッダー/フッターが使う、データのキー（表のキーは含まない）。"""
    keys = [k for spec in (sd.get("cells") or {}).values() for k in spec_keys(spec)]
    keys += [k for tpl in (sd.get("header_footer") or {}).values() for k, _ in text_fields(tpl)]
    return list(dict.fromkeys(keys))


def _key_groups(definition: dict) -> tuple[dict, dict]:
    """データのキー → それを使うまとまり（sheets[].group。無ければタブ）の集まり。表のキーと、固定セルのキーに分けて返す。"""
    tables: dict[str, set] = {}
    cells: dict[str, set] = {}
    for sd in definition.get("sheets", []):
        group = str(sd.get("group") or sd["name"])
        for key in sheet_cell_keys(sd):
            cells.setdefault(key, set()).add(group)
        for t in sd.get("tables", []):
            tables.setdefault(t["key"], set()).add(group)
    return tables, cells


def _stray_tables(parts: list, definition: dict) -> list[str]:
    """表の行が 2 つ以上のファイルから来て、そのうちのファイルがほかのまとまりのデータも持つなら、写したままの疑い。"""
    tables, cells = _key_groups(definition)
    owner = {k: g for k, g in {**cells, **tables}.items() if len(g) == 1}

    def groups_of(obj) -> set:
        out = set()
        for key, g in owner.items():
            try:
                if dig(obj, key) is not None:
                    out |= g
            except KeyError:
                pass
        return out
    found = []
    for key, g in tables.items():
        srcs = []
        for src, obj in parts:
            try:
                rows = dig(obj, key) if isinstance(obj, dict) else None
            except KeyError:
                rows = None
            if isinstance(rows, list) and rows:
                srcs.append((src, obj))
        if len(srcs) < 2:
            continue
        stray = [src for src, obj in srcs if groups_of(obj) - g]
        if stray:
            found.append(f"{key} の行が {'、'.join(src for src, _ in srcs)} から来ています。"
                         f"{'、'.join(stray)} はほかのまとまりのデータも持つので、前のデータを写したままに見える（不要なら消す）")
    return found


def split_data(data: dict, definition: dict) -> list:
    """データを、タブのまとまり（sheets[].group。無ければタブごと）のファイルに分ける。[(ファイル名, 中身), …]。

    いくつかのまとまりで使うキーと、定義に無いキーは 00-common に置く。
    名前順に読めば元の順に戻る（render は、フォルダを渡すと名前順に読んで 1 つにまとめる）。
    """
    owner: dict[str, set] = {}
    order: list[str] = []
    for sd in definition.get("sheets", []):
        group = str(sd.get("group") or sd["name"])
        if group not in order:
            order.append(group)
        keys = sheet_cell_keys(sd)
        keys += [t["key"] for t in sd.get("tables", [])]
        for key in keys:
            owner.setdefault(key, set()).add(group)
    rest = deepcopy(data)
    parts: dict[str, dict] = {g: {} for g in order}
    for key, groups in owner.items():
        if len(groups) != 1:
            continue
        try:
            value = dig(rest, key)
        except KeyError:
            continue
        if _drop_path(rest, key):
            _set_path(parts[next(iter(groups))], key, value)
    out = [("00-common", rest)] if _has_value(rest) else []
    for n, g in enumerate(order, start=1):
        if parts[g]:
            out.append((f"{n:02d}-" + (re.sub(r"[^\w-]+", "_", g).strip("_") or "sheet"), _as_lists(parts[g])))
    return out


def _drop_path(root: dict, path: str) -> bool:
    """ドットでたどったキーを取り除き、空になった親のオブジェクトも取り除く。

    配列の中（`test.0` など）は取り除かずに False を返す（配列ごと 00-common に残す。分けるとつなぎ直しで重なる）。
    """
    keys = path.split(".")
    chain = [root]
    for k in keys[:-1]:
        nxt = chain[-1].get(k)
        if not isinstance(nxt, dict):
            return False
        chain.append(nxt)
    if keys[-1] not in chain[-1]:
        return False
    del chain[-1][keys[-1]]
    for parent, k in zip(reversed(chain[:-1]), reversed(keys[:-1])):
        if parent[k] == {}:
            del parent[k]
    return True


def _has_value(obj) -> bool:
    return bool(obj) if isinstance(obj, (dict, list)) else obj is not None


def dump_structured(obj: Any, path: str) -> None:
    if os.path.dirname(path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        if path.lower().endswith((".yaml", ".yml")):
            _yaml().safe_dump(obj, f, allow_unicode=True, sort_keys=False, default_flow_style=False)
        else:
            json.dump(obj, f, ensure_ascii=False, indent=2)
            f.write("\n")


# ---------------------------------------------------------------------------
# データの形（定義から導く）
# ---------------------------------------------------------------------------

def _set_path(root: dict, path: str, value: Any) -> None:
    cur = root
    parts = path.split(".")
    for part in parts[:-1]:
        cur = cur.setdefault(part, {})
    cur[parts[-1]] = value


def skeleton_data(definition: dict) -> dict:
    """定義が必要とするデータの雛形（値は null）。"""
    out: dict = {}
    for sd in definition.get("sheets", []):
        for key in sheet_cell_keys(sd):
            _set_path(out, key, None)
        for t in sd.get("tables", []):
            colmaps = t["block"] if t.get("block_rows", 1) > 1 and t.get("block") else [t.get("columns")]
            row: dict = {}
            for c in (c for cols in colmaps for c in (cols or {}).values()):
                if c.get("key") and c["key"] != "$index" and not c.get("keep") and not c.get("clear"):
                    _set_path(row, c["key"], None)
            _set_path(out, t["key"], [row])
    for v in (definition.get("properties") or {}).values():
        for key, _ in (text_fields(v) if isinstance(v, str) else []):
            _set_path(out, key, None)
    return _as_lists(out)


def value_notes(definition: dict) -> list[str]:
    """値の書き方（択一・複数選択・true/false・日付）の説明。雛形の null だけでは分からないものを補う。"""
    notes: dict[str, str] = {}
    for sd in definition.get("sheets", []):
        specs = [(None, s if isinstance(s, dict) else {"key": s}) for s in (sd.get("cells") or {}).values()
                 if not (isinstance(s, dict) and "text" in s)]
        for tpl in sheet_texts(sd):
            notes.update(text_notes(tpl))
        for t in sd.get("tables", []):
            colmaps = t["block"] if t.get("block_rows", 1) > 1 and t.get("block") else [t.get("columns")]
            specs += [(t["key"], c) for cols in colmaps for c in (cols or {}).values()]
        for table, spec in specs:
            key = spec.get("key")
            if not key or key == "$index":
                continue
            name = f"{table}[].{key}" if table else key
            if "when" in spec:
                whens = [str(s["when"]) for tk, s in specs if tk == table and s.get("key") == key and "when" in s]
                notes[name] = f"{name}: {' / '.join(whens)} のどれか（複数なら配列）"
            elif "map" in spec:
                notes[name] = f"{name}: {' / '.join(str(k).lower() if isinstance(k, bool) else str(k) for k in spec['map'])} のどれか"
            elif "part" in spec:
                notes[name] = f"{name}: 日付（2026-10-08）"
    for v in (definition.get("properties") or {}).values():
        if isinstance(v, str):
            notes.update(text_notes(v))
    return list(notes.values())


def _as_lists(obj: Any) -> Any:
    """キーが 0, 1, 2 … だけのオブジェクトを配列にする（`test.0`・`test.1` の雛形を `test: [null, null]` に）。"""
    if not isinstance(obj, dict):
        return [_as_lists(x) for x in obj] if isinstance(obj, list) else obj
    obj = {k: _as_lists(v) for k, v in obj.items()}
    if obj and set(obj) == {str(i) for i in range(len(obj))}:
        return [obj[str(i)] for i in range(len(obj))]
    return obj


def _read_back(items: list) -> dict:
    """（列・セルの指定, セルの値）の組から、データの値に戻す。convert_value の逆。

    択一・複数選択は選ばれた見出しの配列で返す（1 つにするかは呼び出し側が決める）。
    × や - の「選ばれていない」印は、空（null）として読む。
    """
    out: dict = {}
    picks: dict = {}
    parts: dict = {}
    for spec, v in items:
        key = spec["key"]
        out.setdefault(key, None)   # 列の順に並べる
        if isinstance(v, str):
            v = v.strip() or None
        if "when" in spec:
            got = picks.setdefault(key, [])
            if v is not None and str(v) == str(spec.get("mark", "○")):
                got.append(spec["when"])
        elif "map" in spec:
            back = {str(m): ({"true": True, "false": False}.get(k, k) if isinstance(k, str) else k)
                    for k, m in spec["map"].items() if m is not None}
            out[key] = None if v is None else back.get(str(v), v)
        elif "part" in spec:
            parts.setdefault(key, {})[spec["part"]] = v
        else:
            out[key] = None if _blankish(v) else v
    out.update(picks)
    for key, p in parts.items():
        try:
            d = dt.datetime(int(p["year"]), int(p.get("month") or 1), int(p.get("day") or 1),
                            int(p.get("hour") or 0), int(p.get("minute") or 0))
            out[key] = d.date().isoformat() if not {"hour", "minute"} & set(p) else d.isoformat(sep=" ", timespec="minutes")
        except (KeyError, TypeError, ValueError):
            out[key] = None
    return out


def _settle_choices(records: list[dict], specs: list[dict]) -> None:
    """択一は見出し 1 つ（選ばれていなければ null）、どこかの行で 2 つ以上選ばれていれば配列にそろえる。"""
    for key in {s["key"] for s in specs if "when" in s}:
        multi = any(len(r.get(key) or []) > 1 for r in records)
        for r in records:
            got = r.get(key) or []
            r[key] = got if multi else (got[0] if got else None)


def _is_empty_record(values: dict) -> bool:
    """既定値（null・false・選ばれていない []）しか無い行。連番・定数・数式の列は、はじめから数えない。"""
    return all(v is None or v is False or v == [] for v in values.values())


def _to_data(flat: dict) -> dict:
    out: dict = {}
    for k, v in flat.items():
        _set_path(out, k, v)
    return _as_lists(out)


def extract(source: "str | bytes", definition: dict, template: "str | bytes | None" = None) -> tuple[dict, list[str]]:
    """記入済みの文書（テンプレートと同じ形）から、定義に沿ってデータを取り出す。render の逆。

    データに書くのは、定義で key を持つ欄だけ（連番・定数・数式・人が書き込む欄は書かない）。
    表の行は、既定値（連番・× など）を除いて値が無ければ空の行として書かない。
    欄（key）は、どの行も空でも記入枠として残す。
    """
    pkg = Package(source)
    sst = read_shared_strings(pkg)
    styles = Styles(pkg)
    sheet_parts = dict(pkg.sheets())
    tpl_grids: dict = {}
    if template is not None:
        tpkg = Package(template)
        tsst, tstyles = read_shared_strings(tpkg), Styles(tpkg)
        tpl_grids = {n: _grid(tpkg.xml(pt), tsst, tstyles) for n, pt in tpkg.sheets()}
    flat: dict = {}
    notes: list[str] = []
    for sd in definition.get("sheets", []):
        name = sd["name"]
        if name not in sheet_parts:
            raise TemplateError(f"文書にシート「{name}」がありません")
        grid = _grid(pkg.xml(sheet_parts[name]), sst, styles)
        tgrid = tpl_grids.get(name, {})
        shifts: list[tuple[int, int]] = []   # (テンプレートの表の最後の行, 文書で増えた行数)
        offset = 0
        for t in sorted(sd.get("tables", []), key=lambda x: int(x["first_row"])):
            first, count = int(t["first_row"]), int(t["sample_rows"])
            k = int(t.get("block_rows", 1))
            colmaps = _column_maps(t, k)
            used_cols = {c for cm in colmaps for c in cm}
            head_r = t.get("header_row")
            if head_r and int(head_r) + offset in grid:
                moved = _header_mismatches(t, {c: i["value"] for c, i in grid[int(head_r) + offset].items()})
                if moved:
                    raise TemplateError(f"シート「{name}」の表の見出しが、定義と合いません（別の形の文書からは、値が別の欄に入って"
                                        "取り出されます）。その文書の形に合う定義を使ってください:\n  " + "\n  ".join(moved))
            # 表の後ろの行（テンプレートで表のすぐ下にある見出し・ラベル）に来たら、表は終わり
            after = {c: i["value"] for c, i in tgrid.get(first + count, {}).items()
                     if isinstance(i["value"], str) and i["value"].strip()}
            records, empty = [], []
            r = first + offset
            while True:
                rows = [grid.get(r + j) for j in range(k)]
                head = rows[0]
                if head is None or _is_total_row(head) or not any(c in head for c in used_cols) \
                        or after and all(head.get(c, {}).get("value") == v for c, v in after.items()):
                    break
                items = []
                for j, cm in enumerate(colmaps):
                    for col, spec in cm.items():
                        if spec.get("formula") or spec.get("keep") or spec.get("clear") \
                                or not spec.get("key") or spec["key"] == "$index":
                            continue
                        items.append((spec, (rows[j] or {}).get(col, {}).get("value")))
                values = _read_back(items)
                (empty.append(r) if _is_empty_record(values) else records.append(values))
                r += k
            specs = [spec for cm in colmaps for spec in cm.values() if spec.get("key") and spec["key"] != "$index"]
            _settle_choices(records, specs)
            flat[t["key"]] = [_to_data(rec) for rec in records]
            if empty:
                notes.append(f"シート「{name}」表 {t.get('id')}: 既定値だけの空の行 {', '.join(map(str, empty))} は書きませんでした")
            grown = (r - first - offset) - count
            shifts.append((first + count - 1, grown))
            offset += grown
        cell_items: dict = {}
        for ref, spec in (sd.get("cells") or {}).items():
            spec = spec if isinstance(spec, dict) else {"key": spec}
            c, row = split_ref(ref)
            row += sum(g for end, g in shifts if row > end)   # 表の行数が変わった分、下の欄がずれる
            if "text" in spec:
                text = grid.get(row, {}).get(c, {}).get("value")
                got = read_text(spec["text"], text)
                if got is None:
                    notes.append(f"シート「{name}」{ref} の {text!r} は、文のひな形 {spec['text']!r} と合わないので読めませんでした")
                for key in spec_keys(spec):
                    if flat.get(key) is None:
                        flat[key] = (got or {}).get(key)
                continue
            cell_items.setdefault(spec["key"], []).append((spec, grid.get(row, {}).get(c, {}).get("value")))
        for key, items in cell_items.items():
            values = _read_back(items)
            if "when" in items[0][0]:
                got = values[key]
                values[key] = got if len(got) > 1 else (got[0] if got else None)
            flat.update(values)
        hf = _header_footer(pkg.xml(sheet_parts[name]))
        for tag, tpl in (sd.get("header_footer") or {}).items():
            got = read_text(tpl.replace("&&", "&"), (hf[tag].text or "").replace("&&", "&")) if tag in hf else None
            for key, _ in text_fields(tpl):
                if flat.get(key) is None:
                    flat[key] = (got or {}).get(key)
    return _to_data(flat), notes


def validate_definition(template: "str | bytes", definition: dict) -> None:
    """テンプレートと定義の整合（シート・行・列）を確かめる。"""
    if definition.get("version") != DEF_VERSION:
        raise TemplateError(f"定義ファイルの version が未対応です: {definition.get('version')!r}")
    sheets = dict(Package(template).sheets())
    for sd in definition.get("sheets", []):
        if sd["name"] not in sheets:
            raise TemplateError(f"テンプレートにシート「{sd['name']}」がありません")
        for ref in (sd.get("cells") or {}):
            split_ref(ref)
        for tag in (sd.get("header_footer") or {}):
            if tag not in HEADER_FOOTER_TAGS:
                raise TemplateError(f"header_footer の {tag!r} は、{', '.join(HEADER_FOOTER_TAGS)} のどれかにしてください")
        sheet_cell_keys(sd)   # 文のひな形の波かっこを確かめる
        for spec in sd.get("clear") or []:
            parse_cell_spec(spec)
        for spec in sd.get("drop_rows") or []:
            parse_row_spec(spec)
        for t in sd.get("tables", []):
            first, count = int(t["first_row"]), int(t["sample_rows"])
            block = int(t.get("block_rows", 1))
            if block < 1 or count % block:
                raise TemplateError(f"sample_rows ({count}) は block_rows ({block}) の倍数にしてください（表 {t['id']}）")
            for p in t.get("pattern") or [first]:
                if not first <= int(p) <= first + count - block:
                    raise TemplateError(f"pattern の行 {p} がサンプル行 {first}-{first + count - block} の外です（表 {t['id']}）")
            for cols in _column_maps(t, block):
                pass


def check_definition(template: "str | bytes", definition: dict) -> list[str]:
    """定義の検査。整合を確かめ、雛形データで試しに再構成する（strict なら、残る値の漏れもここで見つかる）。警告を返す。"""
    import tempfile
    validate_definition(template, definition)
    with tempfile.TemporaryDirectory() as d:
        return render(template, definition, skeleton_data(definition), os.path.join(d, "check.xlsx"))


# ---------------------------------------------------------------------------
# スタンドアローン（固有の render スクリプトを書き出す）
# ---------------------------------------------------------------------------

# PEP 723 の依存宣言。書き出したスクリプトはこのエンジンの本文も含むので、ここに `# /// script` を行頭のまま書くと、
# uv がメタデータを 2 つと数えて動かない。行頭にならないよう {pep723} で差し込む。
PEP723 = "\n".join("#" + line for line in (
    " /// script", ' requires-python = ">=3.10"', ' dependencies = ["lxml", "openpyxl", "pyyaml"]', " ///"))
STANDALONE_HEADER = '''#!/usr/bin/env python3
{pep723}
"""{title}

xlsx テンプレートへデータを流し込む、固有の render スクリプト（xlsx-report-builder の export で生成）。
スキルは不要で動く。表構造の定義はこのファイルに埋め込み済み。{template_note}

    uv run {name} --data data.yaml -o out.xlsx
    uv run {name} --data head.yaml items-1.yaml items-2.yaml -o out.xlsx   # 分けたデータをまとめて流し込む（フォルダも可）
    python {name} --data data.json -o out.xlsx        # lxml・openpyxl・pyyaml が必要
    python {name} --example-data > data.yaml          # データの雛形を出す
    python {name} --extract-def def.yaml              # 埋め込みの定義を取り出す（直したら export --from-script で再生成）

データの形:
{shape}
"""
'''


def _wrap_b64(raw: bytes) -> str:
    text = base64.b64encode(raw).decode("ascii")
    lines = [text[i:i + 100] for i in range(0, len(text), 100)] or [""]
    return "(\n" + "\n".join(f"    {line!r}" for line in lines) + "\n)"


ENGINE_VERSION = 2  # 書き出したスクリプトに入るエンジンの版。export --from-script で最新へ更新できる


def read_exported_script(path: str) -> tuple[dict, "bytes | None", str]:
    """export が書き出したスクリプトから (定義, 埋め込みテンプレート or None, テンプレートの相対パス) を取り出す。

    スクリプトは実行せず、ast で埋め込み部分の定数だけを読む。
    """
    import ast
    with open(path, encoding="utf-8") as f:
        tree = ast.parse(f.read())
    found: dict[str, Any] = {}
    for node in tree.body:
        if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name):
            name = node.targets[0].id
            if name == "DEFINITION":
                call = node.value
                if isinstance(call, ast.Call) and call.args and isinstance(call.args[0], ast.Constant):
                    found[name] = json.loads(call.args[0].value)
            elif name in ("TEMPLATE_B64", "TEMPLATE_PATH") and isinstance(node.value, ast.Constant):
                found[name] = node.value.value
    if "DEFINITION" not in found:
        raise TemplateError(f"{path} は export が書き出したスクリプトではありません（DEFINITION が無い）")
    b64 = found.get("TEMPLATE_B64") or ""
    return found["DEFINITION"], (base64.b64decode(b64) if b64 else None), found.get("TEMPLATE_PATH", "")


def export_script(template: "str | bytes", definition: dict, output: str, embed: bool = False,
                  template_name: str | None = None) -> None:
    """この文書専用の、単体で動く render スクリプトを書き出す。

    template はパス、または（埋め込み専用で）bytes。embed=False ならテンプレートは別ファイルのまま、
    スクリプトからの相対パスで参照する。
    """
    if isinstance(template, (bytes, bytearray)):
        if not embed:
            raise TemplateError("テンプレートのパスが無いため、埋め込みでしか書き出せません")
        raw, rel = bytes(template), template_name or ""
    else:
        with open(template, "rb") as f:
            raw = f.read()
        rel = os.path.relpath(os.path.abspath(template), os.path.dirname(os.path.abspath(output)))
        template_name = os.path.basename(template)
    check_definition(raw, definition)
    with open(os.path.abspath(__file__), encoding="utf-8") as f:
        engine = f.read()
    # エンジン本体（CLI の入口より前）だけを取り込む
    engine = engine.split('\nif __name__ == "__main__":')[0]
    engine = engine.split("\n", 1)[1] if engine.startswith("#!") else engine
    # docstring を 2 つ持てないため、エンジンの docstring は取り除く
    engine = re.sub(r'^"""[\s\S]*?"""\n', "", engine, count=1)
    shape = "\n".join("    " + line for line in
                       json.dumps(skeleton_data(definition), ensure_ascii=False, indent=2).splitlines()
                       + ([""] + value_notes(definition) if value_notes(definition) else []))
    title = f"{os.path.splitext(template_name or 'template')[0]} の render スクリプト"
    note = ("テンプレートも埋め込み済み（--template で差し替えられる）。" if embed
            else f"テンプレート（{rel}）は、このスクリプトからの相対パスで読む。")
    header = STANDALONE_HEADER.format(pep723=PEP723, title=title, name=os.path.basename(output), shape=shape, template_note=note)
    tpl_literal = _wrap_b64(raw) if embed else '""'
    footer = (
        "\n\n# ---------------------------------------------------------------------------\n"
        "# 埋め込み（export が書き出した部分。手で直さず、定義を直して export をやり直す）\n"
        "# ---------------------------------------------------------------------------\n"
        f"ENGINE_VERSION = {ENGINE_VERSION}\n"
        f"DEFINITION = json.loads({json.dumps(definition, ensure_ascii=False)!r})\n"
        f"TEMPLATE_B64 = {tpl_literal}\n"
        f"TEMPLATE_PATH = {rel!r}  # 埋め込まない場合の、このスクリプトからの相対パス\n"
        "\n\nif __name__ == \"__main__\":\n"
        "    raise SystemExit(standalone_main(DEFINITION, base64.b64decode(TEMPLATE_B64) or None, TEMPLATE_PATH, __doc__))\n"
    )
    if os.path.dirname(output):
        os.makedirs(os.path.dirname(output), exist_ok=True)
    with open(output, "w", encoding="utf-8") as f:
        f.write(header + engine.rstrip() + footer)
    os.chmod(output, 0o755)


def standalone_main(definition: dict, template_bytes: "bytes | None", template_path: str, doc: str | None = None) -> int:
    parser = argparse.ArgumentParser(description=(doc or "").split("\n\n")[0], formatter_class=argparse.RawDescriptionHelpFormatter,
                                     epilog="\n\n".join((doc or "").split("\n\n")[1:]))
    parser.add_argument("--data", nargs="+", action="extend",
                        help="データ（.json / .yaml / .yml。- で標準入力）。複数のファイル・フォルダを渡すと、1 つにまとめる")
    parser.add_argument("-o", "--output", help="出力 .xlsx")
    parser.add_argument("--template", help=("埋め込みのテンプレート" if template_bytes else f"既定のテンプレート（{template_path}）")
                        + "の代わりに使う .xlsx（定義と構造が同じものに限る）")
    parser.add_argument("--example-data", action="store_true", help="データの雛形（YAML）を標準出力に出す")
    if template_bytes:
        parser.add_argument("--extract-template", metavar="PATH", help="埋め込みのテンプレートを書き出す")
    parser.add_argument("--extract-def", metavar="PATH", help="埋め込みの定義を書き出す（.json / .yaml）")
    args = parser.parse_args()
    try:
        if args.extract_def:
            dump_structured(definition, args.extract_def)
            print(f"書き出しました: {args.extract_def}")
            return 0
        if args.example_data:
            print(_yaml().safe_dump(skeleton_data(definition), allow_unicode=True, sort_keys=False, default_flow_style=False), end="")
            return 0
        here = os.path.dirname(os.path.abspath(sys.argv[0]))
        if args.template:
            with open(args.template, "rb") as f:
                template = f.read()
        elif template_bytes:
            template = template_bytes
        else:
            path = os.path.join(here, template_path)
            if not os.path.exists(path):
                raise TemplateError(f"テンプレートが見つかりません: {path}（--template で指定するか、export し直す）")
            with open(path, "rb") as f:
                template = f.read()
        if getattr(args, "extract_template", None):
            if not template_bytes and not args.template:
                raise TemplateError("テンプレートは埋め込まれていません（元の .xlsx を使う）")
            with open(args.extract_template, "wb") as f:
                f.write(template)
            print(f"書き出しました: {args.extract_template}")
            return 0
        if not args.data or not args.output:
            parser.error("--data と -o が必要です")
        warnings = render(template, definition, load_data(args.data, definition), args.output)
        for w in warnings:
            print(f"警告: {w}", file=sys.stderr)
        print(f"生成しました: {args.output}")
        return 0
    except TemplateError as e:
        print(f"エラー: {e}", file=sys.stderr)
        return 1
    except FileNotFoundError as e:
        print(f"エラー: ファイルが見つかりません: {e.filename}", file=sys.stderr)
        return 1


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def cmd_analyze(args) -> int:
    definition = analyze(args.template)
    out = args.output or re.sub(r"\.xlsx?$", "", args.template, flags=re.I) + ".def.json"
    # template は、定義ファイルのある場所からの相対パスで書く（render・check・export がそこから読む）
    definition["template"] = os.path.relpath(os.path.abspath(args.template),
                                             os.path.dirname(os.path.abspath(out))).replace(os.sep, "/")
    dump_structured(definition, out)
    print(summarize(definition))
    print(f"\n定義ファイルの下書きを書きました: {out}")
    print("（? の項目と、流し込む欄・残す範囲の分け方をユーザーに確認してから、定義ファイルを直して確定する）")
    return 0


def cmd_check(args) -> int:
    definition = load_structured(args.definition)
    for w in check_definition(_template_arg(args, definition), definition):
        print(f"警告: {w}")
    print("定義は問題ありません")
    return 0


def cmd_inspect(args) -> int:
    facts = inspect_template(args.template)
    if args.json:
        json.dump(facts, sys.stdout, ensure_ascii=False, indent=2)
        print()
    else:
        print(format_facts(facts, fold=not args.all))
    return 0


def _template_arg(args, definition) -> str:
    """--template はそのまま、定義の template は定義ファイルのある場所からの相対パスとして読む。"""
    if args.template:
        return args.template
    template = definition.get("template")
    if not template:
        raise TemplateError("--template か定義ファイルの template が必要です")
    def_path = getattr(args, "definition", None)
    if os.path.isabs(template) or not def_path or def_path == "-":
        return template
    beside = os.path.join(os.path.dirname(os.path.abspath(def_path)), template)
    if os.path.exists(beside) or not os.path.exists(template):  # 以前の、作業場所からの相対パスも読めるように残す
        return beside
    return template


def cmd_render(args) -> int:
    definition = load_structured(args.definition)
    data = load_data(args.data, definition)
    warnings = render(_template_arg(args, definition), definition, data, args.output)
    for w in warnings:
        print(f"警告: {w}", file=sys.stderr)
    print(f"生成しました: {args.output}")
    return 0


def cmd_export(args) -> int:
    old_def, old_bytes, old_rel = (None, None, "")
    old_dir = ""
    if args.from_script:
        old_def, old_bytes, old_rel = read_exported_script(args.from_script)
        old_dir = os.path.dirname(os.path.abspath(args.from_script))
    if args.definition:
        definition = load_structured(args.definition)
    elif old_def is not None:
        definition = old_def
    else:
        raise TemplateError("--def か --from-script が必要です")
    embed = args.embed
    if args.template or definition.get("template") and not args.from_script:
        template: "str | bytes" = _template_arg(args, definition)
    elif old_bytes is not None:  # 以前のスクリプトに埋め込まれたテンプレートを引き継ぐ
        template, embed = old_bytes, True
    elif old_rel:
        template = os.path.normpath(os.path.join(old_dir, old_rel))
    else:
        template = _template_arg(args, definition)
    export_script(template, definition, args.output, embed=embed,
                  template_name=os.path.basename(old_rel) if isinstance(template, bytes) else None)
    print(f"書き出しました: {args.output}")
    print(f"  使い方: uv run {os.path.basename(args.output)} --data data.yaml -o out.xlsx")
    return 0


def cmd_extract(args) -> int:
    definition = load_structured(args.definition)
    try:
        template = _template_arg(args, definition)
    except TemplateError:
        template = None
    data, notes = extract(args.source, definition, template if template and os.path.exists(template) else None)
    for n in notes:
        print(n, file=sys.stderr)
    if args.split:
        os.makedirs(args.split, exist_ok=True)
        ext = args.format
        for name, part in split_data(data, definition):
            dump_structured(part, os.path.join(args.split, f"{name}.{ext}"))
            print(f"取り出しました: {os.path.join(args.split, name + '.' + ext)}")
    elif args.output:
        dump_structured(data, args.output)
        print(f"取り出しました: {args.output}")
    else:
        json.dump(data, sys.stdout, ensure_ascii=False, indent=2, default=str)
        print()
    return 0


def add_subcommands(sub) -> None:
    a = sub.add_parser("analyze", help="テンプレートを解析して表構造の定義ファイル（下書き）を作る")
    a.add_argument("template", help="テンプレート .xlsx")
    a.add_argument("-o", "--output", help="定義ファイルの出力先（.json / .yaml。省略時は <テンプレート>.def.json）")
    a.set_defaults(func=cmd_analyze)
    c = sub.add_parser("check", help="定義の検査（整合・strict の漏れ）。雛形データで試しに再構成する")
    c.add_argument("--template", help="テンプレート .xlsx（省略時は定義ファイルの template）")
    c.add_argument("--def", dest="definition", required=True, help="定義ファイル（.json / .yaml）")
    c.set_defaults(func=cmd_check)
    i = sub.add_parser("inspect", help="テンプレートの事実（値・数式・書式の種類・結合・仮値の疑い）を、判断用に出す")
    i.add_argument("template", help="テンプレート .xlsx")
    i.add_argument("--json", action="store_true", help="JSON で出す（すべての行）")
    i.add_argument("--all", action="store_true", help="同じ書式が長く続く行も、省略せずにすべて出す")
    i.set_defaults(func=cmd_inspect)
    r = sub.add_parser("render", help="テンプレート + 定義 + データから xlsx を再構成する")
    r.add_argument("--template", help="テンプレート .xlsx（省略時は定義ファイルの template）")
    r.add_argument("--def", dest="definition", required=True, help="定義ファイル（.json / .yaml）")
    r.add_argument("--data", required=True, nargs="+", action="extend",
                   help="データ（.json / .yaml。- で標準入力）。複数のファイル・フォルダを渡すと、1 つにまとめる（表の行はつなぐ）")
    r.add_argument("-o", "--output", required=True, help="出力 .xlsx")
    r.set_defaults(func=cmd_render)
    e = sub.add_parser("export", help="この文書専用の、単体で動く render スクリプトを書き出す")
    e.add_argument("--template", help="テンプレート .xlsx（省略時は定義ファイルの template）")
    e.add_argument("--def", dest="definition", help="確定した定義ファイル（.json / .yaml）。--from-script と併用すると、その定義を置き換える")
    e.add_argument("--from-script", help="以前に export したスクリプト。定義・テンプレートを引き継いで、最新のエンジンで書き出し直す")
    e.add_argument("-o", "--output", required=True, help="書き出す .py")
    e.add_argument("--embed", action="store_true", help="テンプレートもスクリプトに埋め込む（既定は別ファイルを相対パスで参照）")
    e.set_defaults(func=cmd_export)
    x = sub.add_parser("extract", help="記入済みの文書から、定義に沿ってデータ（.json / .yaml）を取り出す（render の逆）")
    x.add_argument("source", help="記入済みの .xlsx（テンプレートと同じ形の文書）")
    x.add_argument("--def", dest="definition", required=True, help="定義ファイル（.json / .yaml）")
    x.add_argument("--template", help="テンプレート .xlsx（省略時は定義ファイルの template。表の終わりを見分けるのに使う）")
    x.add_argument("-o", "--output", help="データの出力先（.json / .yaml。省略時は標準出力に JSON）")
    x.add_argument("--split", metavar="DIR",
                   help="データを、タブのまとまり（定義の sheets[].group。無ければタブ）ごとのファイルに分けて、このフォルダに書く（render --data DIR で読める）")
    x.add_argument("--format", choices=("yaml", "json"), default="yaml", help="--split で書く形式（既定 yaml）")
    x.set_defaults(func=cmd_extract)


def main() -> int:
    parser = argparse.ArgumentParser(description="xlsx テンプレートへデータを流し込む")
    add_subcommands(parser.add_subparsers(dest="command", required=True))
    args = parser.parse_args()
    try:
        return args.func(args)
    except TemplateError as e:
        print(f"エラー: {e}", file=sys.stderr)
        return 1
    except FileNotFoundError as e:
        print(f"エラー: ファイルが見つかりません: {e.filename}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
