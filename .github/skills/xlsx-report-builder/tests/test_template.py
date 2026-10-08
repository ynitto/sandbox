"""xlsx_template（analyze / render）のテスト。テンプレートは openpyxl で組み立てる。

実行: cd .github/skills/xlsx-report-builder && uv run python -m unittest discover -s tests -v
"""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import xlsx_template as xt

THIN = Side(style="thin", color="888888")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
BAND_FILL = PatternFill("solid", fgColor="DDEBF7")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")


def make_template(path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "請求書"
    ws.merge_cells("A1:E1")
    ws["A1"] = "請求書"
    ws["A1"].font = Font(name="MS Gothic", size=20, bold=True)
    ws["A3"], ws["B3"] = "請求先", "サンプル株式会社"
    ws["A4"], ws["B4"] = "請求日", "2020-01-01"
    ws["B4"].number_format = "yyyy/mm/dd"
    ws["A5"], ws["B5"] = "税率", 0.1
    ws["B5"].number_format = "0%"
    for c, h in zip("ABCDE", ["No", "品名", "数量", "単価", "金額"]):
        cell = ws[f"{c}7"]
        cell.value, cell.font, cell.fill, cell.border = h, Font(bold=True, color="FFFFFF"), HEAD_FILL, BOX
    for r, fill in ((8, None), (9, BAND_FILL)):  # 縞模様のサンプル 2 行
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = r - 7, f"サンプル{r}", 1, 100
        ws[f"E{r}"] = f"=C{r}*D{r}"
        for c in "ABCDE":
            ws[f"{c}{r}"].border = BOX
            ws[f"{c}{r}"].font = Font(name="MS Gothic", size=10)
            if fill:
                ws[f"{c}{r}"].fill = fill
        ws[f"D{r}"].number_format = ws[f"E{r}"].number_format = "#,##0"
    ws["D10"], ws["E10"] = "合計", "=SUM(E8:E9)"
    ws["D11"], ws["E11"] = "消費税", "=E10*$B$5"
    ws["D12"], ws["E12"] = "請求額", "=E10+E11"
    for r in (10, 11, 12):
        for c in "DE":
            ws[f"{c}{r}"].fill, ws[f"{c}{r}"].border = TOTAL_FILL, BOX
        ws[f"E{r}"].number_format = "#,##0"
    ws["A14"] = "お支払状況"
    for c, h in zip("ABC", ["日付", "方法", "入金額"]):
        cell = ws[f"{c}15"]
        cell.value, cell.font, cell.fill, cell.border = h, Font(bold=True, color="FFFFFF"), HEAD_FILL, BOX
    ws["A16"], ws["B16"], ws["C16"] = "2020-02-01", "振込", 500
    ws["A16"].number_format = "yyyy/mm/dd"
    for c in "ABC":
        ws[f"{c}16"].border = BOX
    ws["A18"] = "※ 振込手数料はご負担ください"
    ws["A18"].font = Font(italic=True, color="C00000")
    ws.conditional_formatting.add("E8:E9", CellIsRule(operator="greaterThan", formula=["1000"], fill=PatternFill("solid", bgColor="FFC7CE")))
    dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0")
    dv.add("C8:C9")
    ws.add_data_validation(dv)
    ws.column_dimensions["B"].width = 30
    ws.row_dimensions[8].height = 22
    wb.defined_names["請求額"] = DefinedName("請求額", attr_text="'請求書'!$E$12")
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=5, min_row=7, max_row=9), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=8, max_row=9))
    ws.add_chart(chart, "G3")
    s2 = wb.create_sheet("集計")
    s2["A1"] = "請求額"
    s2["B1"] = "='請求書'!E12"
    s2["A2"] = "合計"
    s2["B2"] = "=請求書!E10"
    wb.save(path)


DATA = {
    "customer": "株式会社テスト",
    "date": "2026-10-05",
    "rate": 0.08,
    "items": [{"name": f"商品{i}", "qty": i, "price": 1000 * i} for i in range(1, 6)],
    "payments": [{"date": "2026-10-10", "method": "振込", "amount": 1000}, {"date": "2026-10-20", "method": "現金", "amount": 2000}],
}

DEF = {
    "version": 1,
    "template": "t.xlsx",
    "sheets": [{
        "name": "請求書",
        "cells": {"B3": "customer", "B4": "date", "B5": "rate"},
        "tables": [
            {"id": "items", "header_row": 7, "first_row": 8, "sample_rows": 2, "pattern": [8, 9], "key": "items",
             "columns": {"A": {"key": "$index"}, "B": {"key": "name"}, "C": {"key": "qty"}, "D": {"key": "price"}, "E": {"formula": True}}},
            {"id": "payments", "header_row": 15, "first_row": 16, "sample_rows": 1, "pattern": [16], "key": "payments",
             "columns": {"A": {"key": "date"}, "B": {"key": "method"}, "C": {"key": "amount"}}},
        ],
    }],
}


class Base(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.dir, True)
        self.tpl = os.path.join(self.dir, "t.xlsx")
        make_template(self.tpl)

    def render(self, data=DATA, definition=DEF, name="out.xlsx"):
        out = os.path.join(self.dir, name)
        self.warnings = xt.render(self.tpl, definition, data, out)
        return out


class RenderTests(Base):
    def test_rows_expand_and_following_content_moves(self):
        ws = load_workbook(self.render())["請求書"]
        # 明細 5 行 → 8..12、合計は 13..15（+3）、支払表の見出しは 15+3+... と連動してずれる
        self.assertEqual([ws[f"B{r}"].value for r in range(8, 13)], [f"商品{i}" for i in range(1, 6)])
        self.assertEqual([ws[f"A{r}"].value for r in range(8, 13)], [1, 2, 3, 4, 5])
        self.assertEqual(ws["D13"].value, "合計")
        self.assertEqual(ws["D14"].value, "消費税")
        self.assertEqual(ws["D15"].value, "請求額")
        self.assertEqual(ws["A17"].value, "お支払状況")
        self.assertEqual(ws["A18"].value, "日付")  # 見出し行: 15 + 3 = 18
        self.assertEqual([ws[f"B{r}"].value for r in (19, 20)], ["振込", "現金"])  # 表 2 は 1 → 2 行
        self.assertEqual(ws["A22"].value, "※ 振込手数料はご負担ください")  # 18 + 3 + 1

    def test_formulas_are_translated_and_totals_extended(self):
        ws = load_workbook(self.render())["請求書"]
        self.assertEqual([ws[f"E{r}"].value for r in (8, 9, 10, 11, 12)], [f"=C{r}*D{r}" for r in (8, 9, 10, 11, 12)])
        self.assertEqual(ws["E13"].value, "=SUM(E8:E12)")
        self.assertEqual(ws["E14"].value, "=E13*$B$5")
        self.assertEqual(ws["E15"].value, "=E13+E14")

    def test_template_formatting_is_kept_and_banding_repeats(self):
        tpl = load_workbook(self.tpl)["請求書"]
        ws = load_workbook(self.render())["請求書"]
        self.assertEqual(ws["A7"].fill.fgColor.rgb, tpl["A7"].fill.fgColor.rgb)  # 見出し
        for r in range(8, 13):
            for col in "ABCDE":
                c = ws[f"{col}{r}"]
                self.assertEqual(c.border.left.style, "thin")
                self.assertEqual(c.font.name, "MS Gothic")
        self.assertEqual(ws["B9"].fill.fgColor.rgb, tpl["B9"].fill.fgColor.rgb)
        self.assertEqual(ws["B9"].fill.fgColor.rgb, ws["B11"].fill.fgColor.rgb)  # 縞は 2 行周期
        self.assertNotEqual(ws["B8"].fill.fgColor.rgb, ws["B9"].fill.fgColor.rgb)
        self.assertEqual(ws["D8"].number_format, "#,##0")
        self.assertEqual(ws["D13"].fill.fgColor.rgb, tpl["D10"].fill.fgColor.rgb)  # 合計行の色
        self.assertEqual(ws["A22"].font.italic, True)
        self.assertEqual(ws.column_dimensions["B"].width, 30)

    def test_row_height_is_copied_with_the_row(self):
        ws = load_workbook(self.render())["請求書"]
        self.assertEqual(ws.row_dimensions[8].height, 22)
        self.assertIsNone(ws.row_dimensions[9].height)

    def test_fixed_cells_and_dates(self):
        ws = load_workbook(self.render())["請求書"]
        self.assertEqual(ws["B3"].value, "株式会社テスト")
        self.assertEqual(ws["B4"].value.date().isoformat(), "2026-10-05")  # 日付書式のセルは日付型
        self.assertEqual(ws["B4"].number_format, "yyyy/mm/dd")
        self.assertEqual(ws["B5"].value, 0.08)
        self.assertEqual(ws["A19"].value.date().isoformat(), "2026-10-10")

    def test_sample_values_do_not_leak(self):
        ws = load_workbook(self.render(data={**DATA, "items": [{"name": "だけ"}]}))["請求書"]
        self.assertEqual(ws["B8"].value, "だけ")
        self.assertIsNone(ws["C8"].value)  # サンプルの 1 が残らない
        self.assertIsNone(ws["D8"].value)
        self.assertEqual(ws["E8"].value, "=C8*D8")

    def test_shrink_and_empty(self):
        ws = load_workbook(self.render(data={**DATA, "items": [{"name": "a", "qty": 1, "price": 5}] * 1}))["請求書"]
        self.assertEqual(ws["D9"].value, "合計")  # 2 行 → 1 行で -1
        self.assertEqual(ws["E9"].value, "=SUM(E8:E8)")
        ws = load_workbook(self.render(data={**DATA, "items": []}, name="e.xlsx"))["請求書"]
        self.assertIsNone(ws["B8"].value)  # 空でも 1 行（空欄）は残す
        self.assertEqual(ws["D9"].value, "合計")

    def test_merged_cf_dv_names_and_cross_sheet_refs_follow(self):
        out = self.render()
        wb = load_workbook(out)
        ws = wb["請求書"]
        self.assertIn("A1:E1", [str(m) for m in ws.merged_cells.ranges])
        cf = [str(r.sqref) for r in ws.conditional_formatting]
        self.assertEqual(cf, ["E8:E12"])
        self.assertEqual([str(d.sqref) for d in ws.data_validations.dataValidation], ["C8:C12"])
        self.assertEqual(wb.defined_names["請求額"].attr_text, "'請求書'!$E$15")
        self.assertEqual(wb["集計"]["B1"].value, "='請求書'!E15")
        self.assertEqual(wb["集計"]["B2"].value, "=請求書!E13")

    def test_chart_range_and_untouched_parts_are_byte_identical(self):
        out = self.render()
        with zipfile.ZipFile(out) as new, zipfile.ZipFile(self.tpl) as old:
            chart = [n for n in new.namelist() if n.startswith("xl/charts/chart")][0]
            self.assertIn("$E$8:$E$12", new.read(chart).decode())
            for name in ("xl/styles.xml", "xl/theme/theme1.xml", "docProps/core.xml"):
                if name in old.namelist():
                    self.assertEqual(new.read(name), old.read(name), name)
            wbxml = new.read("xl/workbook.xml").decode()
            self.assertIn('fullCalcOnLoad="1"', wbxml)

    def test_merge_in_sample_row_is_replicated(self):
        wb = load_workbook(self.tpl)
        ws = wb["請求書"]
        ws.merge_cells("B16:C16")
        wb.save(self.tpl)
        out = self.render()
        merged = [str(m) for m in load_workbook(out)["請求書"].merged_cells.ranges]
        self.assertIn("B19:C19", merged)
        self.assertIn("B20:C20", merged)

    def test_errors(self):
        with self.assertRaises(xt.TemplateError):
            self.render(data={"items": []})  # cells のキー欠落
        bad = json.loads(json.dumps(DEF))
        bad["sheets"][0]["tables"][0]["pattern"] = [20]
        with self.assertRaises(xt.TemplateError):
            self.render(definition=bad)

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice がない")
    def test_opens_and_calculates_in_libreoffice(self):
        out = self.render()
        subprocess.run(["soffice", "--headless", f"-env:UserInstallation=file://{self.dir}/lo", "--convert-to", "csv", "--outdir", self.dir, out],
                       check=True, capture_output=True, timeout=120)
        with open(os.path.join(self.dir, "out.csv"), encoding="utf-8") as fh:
            text = fh.read()
        # 商品 i: 数量 i × 単価 1000i の合計 = 1000 × (1+4+9+16+25) = 55000
        self.assertIn("合計,55000", text.replace('"', ""))
        self.assertIn("請求額,59400", text.replace('"', ""))  # 55000 × 1.08


class AnalyzeTests(Base):
    def test_detects_tables_pattern_and_columns(self):
        d = xt.analyze(self.tpl)
        sheet = d["sheets"][0]
        self.assertEqual(sheet["name"], "請求書")
        t1, t2 = sheet["tables"]
        self.assertEqual((t1["header_row"], t1["first_row"], t1["sample_rows"], t1["pattern"]), (7, 8, 2, [8, 9]))
        self.assertTrue(t1["columns"]["E"]["formula"])
        self.assertEqual(t1["columns"]["B"]["key"], "品名")
        self.assertEqual(t1["_total_row"], 10)
        self.assertEqual((t2["header_row"], t2["first_row"], t2["sample_rows"]), (15, 16, 1))
        self.assertEqual(list(t2["columns"]), ["A", "B", "C"])

    def test_fixed_cells_split_into_fill_and_keep(self):
        d = xt.analyze(self.tpl)
        sheet = d["sheets"][0]
        # ラベルの右隣は、前の値があっても流し込む欄。表の行・数式は入れない
        self.assertEqual(sheet["cells"], {"B3": "請求先", "B4": "請求日", "B5": "税率"})
        self.assertEqual(sheet["_cell_samples"]["B3"], {"label": "請求先", "value": "サンプル株式会社"})
        self.assertEqual(sheet["keep"], ["A1", "A3:A5", "A7:E7", "D10:D12", "A14", "A15:C15", "A18"])
        # 前の文書の値・来歴を持ち越さないのが既定
        self.assertTrue(d["strict"])
        self.assertEqual(d["properties"], {"scrub": True})
        xt.check_definition(self.tpl, d)   # 下書きのままで、strict の検査が通る

    def test_cli_roundtrip_with_analyzed_definition(self):
        import xlsx_builder  # noqa: F401  (サブコマンド統合の確認)
        defn = os.path.join(self.dir, "d.json")
        data = os.path.join(self.dir, "d.data.json")
        out = os.path.join(self.dir, "o.xlsx")
        script = os.path.join(os.path.dirname(__file__), "..", "scripts", "xlsx_builder.py")
        subprocess.run([sys.executable, script, "analyze", self.tpl, "-o", defn], check=True, capture_output=True)
        with open(defn, encoding="utf-8") as fh:
            d = json.load(fh)
        d["sheets"][0]["tables"][0]["key"] = "items"
        for letter, key in zip("BCD", ["name", "qty", "price"]):
            d["sheets"][0]["tables"][0]["columns"][letter]["key"] = key
        d["sheets"][0]["tables"][1]["key"] = "payments"
        for letter, key in zip("ABC", ["date", "method", "amount"]):
            d["sheets"][0]["tables"][1]["columns"][letter]["key"] = key
        d["sheets"][0]["tables"][0]["columns"]["A"]["key"] = "$index"
        d["sheets"][0]["cells"] = {"B3": "customer", "B4": "date", "B5": "rate"}
        with open(defn, "w", encoding="utf-8") as fh:
            json.dump(d, fh, ensure_ascii=False)
        with open(data, "w", encoding="utf-8") as fh:
            json.dump(DATA, fh, ensure_ascii=False)
        subprocess.run([sys.executable, script, "render", "--template", self.tpl, "--def", defn, "--data", data, "-o", out],
                       check=True, capture_output=True)
        ws = load_workbook(out)["請求書"]
        self.assertEqual(ws["B3"].value, "株式会社テスト")
        self.assertEqual(ws["E13"].value, "=SUM(E8:E12)")


def rewrite_zip(src: str, dst: str, edits: dict) -> None:
    """zip 内の part を {名前: 関数(bytes → bytes)} で書き換えて保存する。"""
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in zi.infolist():
            raw = zi.read(info.filename)
            if info.filename in edits:
                raw = edits[info.filename](raw)
            zo.writestr(info.filename, raw)


class RealWorldShapeTests(Base):
    """openpyxl が作らない形（共有数式・共有文字列・calcChain）でも壊れない。"""

    def test_shared_formulas_are_expanded_and_translated(self):
        def to_shared(raw: bytes) -> bytes:
            text = raw.decode()
            self.assertIn("<f>C8*D8</f>", text)
            text = text.replace("<f>C8*D8</f>", '<f t="shared" ref="E8:E9" si="0">C8*D8</f>')
            return text.replace("<f>C9*D9</f>", '<f t="shared" si="0"/>').encode()
        shared = os.path.join(self.dir, "shared.xlsx")
        rewrite_zip(self.tpl, shared, {"xl/worksheets/sheet1.xml": to_shared})
        self.tpl = shared
        ws = load_workbook(self.render())["請求書"]
        self.assertEqual([ws[f"E{r}"].value for r in range(8, 13)], [f"=C{r}*D{r}" for r in range(8, 13)])

    def test_calc_chain_is_dropped_consistently(self):
        chain = os.path.join(self.dir, "chain.xlsx")
        ct = lambda raw: raw.decode().replace("</Types>", '<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>').encode()
        rel = lambda raw: raw.decode().replace("</Relationships>", '<Relationship Id="rIdC" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/></Relationships>').encode()
        with zipfile.ZipFile(self.tpl) as zi, zipfile.ZipFile(chain, "w") as zo:
            for info in zi.infolist():
                raw = zi.read(info.filename)
                if info.filename == "[Content_Types].xml":
                    raw = ct(raw)
                if info.filename == "xl/_rels/workbook.xml.rels":
                    raw = rel(raw)
                zo.writestr(info.filename, raw)
            zo.writestr("xl/calcChain.xml", '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="E8" i="1"/></calcChain>')
        self.tpl = chain
        out = self.render()
        with zipfile.ZipFile(out) as z:
            self.assertNotIn("xl/calcChain.xml", z.namelist())
            self.assertNotIn("calcChain", z.read("[Content_Types].xml").decode())
            self.assertNotIn("calcChain", z.read("xl/_rels/workbook.xml.rels").decode())
        load_workbook(out)

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice がない")
    def test_template_saved_by_libreoffice_uses_shared_strings(self):
        lo_dir = os.path.join(self.dir, "lo_tpl")
        os.makedirs(lo_dir)
        subprocess.run(["soffice", "--headless", f"-env:UserInstallation=file://{self.dir}/lo", "--convert-to", "xlsx", "--outdir", lo_dir, self.tpl],
                       check=True, capture_output=True, timeout=120)
        self.tpl = os.path.join(lo_dir, "t.xlsx")
        with zipfile.ZipFile(self.tpl) as z:
            self.assertIn("xl/sharedStrings.xml", z.namelist())
        sheet = xt.analyze(self.tpl)["sheets"][0]
        self.assertEqual([(t["header_row"], t["first_row"], t["sample_rows"]) for t in sheet["tables"]], [(7, 8, 2), (15, 16, 1)])
        self.assertEqual(sheet["tables"][0]["columns"]["B"]["header"], "品名")
        ws = load_workbook(self.render())["請求書"]
        self.assertEqual(ws["B12"].value, "商品5")
        self.assertEqual(ws["A22"].value, "※ 振込手数料はご負担ください")
        self.assertEqual(ws["B9"].border.left.style, "thin")


class StandaloneAndYamlTests(Base):
    def export(self, **kw):
        import yaml
        defn = os.path.join(self.dir, "def.yaml")
        xt.dump_structured(DEF, defn)
        script = os.path.join(self.dir, "out", "render_invoice.py")
        os.makedirs(os.path.dirname(script))
        xt.export_script(self.tpl, xt.load_structured(defn), script, **kw)
        return script

    def run_script(self, script, *args, cwd=None):
        env = {k: v for k, v in os.environ.items() if k != "PYTHONPATH"}
        return subprocess.run([sys.executable, script, *args], capture_output=True, text=True, cwd=cwd or self.dir, env=env)

    def sheet_xml(self, path):
        with zipfile.ZipFile(path) as z:
            return z.read("xl/worksheets/sheet1.xml")

    def test_yaml_roundtrip_of_definition_and_data(self):
        import yaml
        d = os.path.join(self.dir, "d.yaml")
        xt.dump_structured(DEF, d)
        self.assertEqual(xt.load_structured(d), DEF)
        data = os.path.join(self.dir, "data.yml")
        with open(data, "w", encoding="utf-8") as f:
            yaml.safe_dump(DATA, f, allow_unicode=True)
        out = os.path.join(self.dir, "y.xlsx")
        xt.render(self.tpl, xt.load_structured(d), xt.load_structured(data), out)
        self.assertEqual(self.sheet_xml(out), self.sheet_xml(self.render(name="j.xlsx")))

    def test_yaml_dates_stay_usable(self):
        import yaml
        data = os.path.join(self.dir, "data.yaml")
        with open(data, "w", encoding="utf-8") as f:
            f.write("customer: テスト\ndate: 2026-10-05\nrate: 0.1\nitems: []\npayments: []\n")  # 日付は YAML では date 型
        loaded = xt.load_structured(data)
        ws = load_workbook(self.render(data=loaded))["請求書"]
        self.assertEqual(ws["B4"].value.date().isoformat(), "2026-10-05")

    def test_exported_script_runs_without_the_skill(self):
        script = self.export(embed=True)
        with open(script, encoding="utf-8") as fh:
            text = fh.read()
        self.assertNotIn("import xlsx_template", text)
        # `uv run` は PEP 723 のメタデータが 2 つあると動かない（エンジンの本文に行頭の `# /// script` を残さない）
        pep723 = re.findall(r"(?m)^# /// (?P<type>[a-zA-Z0-9-]+)$\s(?P<content>(^#(| .*)$\s)+)^# ///$", text)
        self.assertEqual([m[0] for m in pep723], ["script"])
        data = os.path.join(self.dir, "data.yaml")
        import yaml
        with open(data, "w", encoding="utf-8") as f:
            yaml.safe_dump(DATA, f, allow_unicode=True)
        out = os.path.join(self.dir, "s.xlsx")
        os.remove(self.tpl)  # テンプレートが埋め込まれていれば、元ファイルが無くても動く
        r = self.run_script(script, "--data", data, "-o", out, cwd="/")
        self.assertEqual(r.returncode, 0, r.stderr)
        ws = load_workbook(out)["請求書"]
        self.assertEqual(ws["B12"].value, "商品5")
        self.assertEqual(ws["E13"].value, "=SUM(E8:E12)")
        self.assertEqual(ws["B9"].border.left.style, "thin")

    def test_exported_script_matches_render_and_accepts_json_and_stdin(self):
        script = self.export()
        expected = self.sheet_xml(self.render())
        data = os.path.join(self.dir, "data.json")
        with open(data, "w", encoding="utf-8") as f:
            json.dump(DATA, f, ensure_ascii=False)
        out = os.path.join(self.dir, "s.xlsx")
        self.assertEqual(self.run_script(script, "--data", data, "-o", out).returncode, 0)
        self.assertEqual(self.sheet_xml(out), expected)
        r = subprocess.run([sys.executable, script, "--data", "-", "-o", out], input=json.dumps(DATA), capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)

    def test_example_data_help_and_extract(self):
        script = self.export(embed=True)
        r = self.run_script(script, "--example-data")
        import yaml
        skeleton = yaml.safe_load(r.stdout)
        self.assertEqual(sorted(skeleton), ["customer", "date", "items", "payments", "rate"])
        self.assertEqual(sorted(skeleton["items"][0]), ["name", "price", "qty"])
        h = self.run_script(script, "--help")
        self.assertIn("customer", h.stdout)
        ext = os.path.join(self.dir, "ext.xlsx")
        self.assertEqual(self.run_script(script, "--extract-template", ext).returncode, 0)
        with open(ext, "rb") as a, open(self.tpl, "rb") as b:
            self.assertEqual(a.read(), b.read())

    def test_template_stays_separate_by_default(self):
        script = self.export()
        with open(script, encoding="utf-8") as fh:
            self.assertLess(len(fh.read()), 120_000)  # 既定ではテンプレートを抱え込まない
        self.assertNotEqual(self.run_script(script, "--extract-template", os.path.join(self.dir, "x.xlsx")).returncode, 0)
        os.rename(self.tpl, self.tpl + ".bak")
        data = os.path.join(self.dir, "data.json")
        with open(data, "w", encoding="utf-8") as f:
            json.dump(DATA, f, ensure_ascii=False)
        r = self.run_script(script, "--data", data, "-o", os.path.join(self.dir, "n.xlsx"))
        self.assertEqual(r.returncode, 1)
        self.assertIn("テンプレートが見つかりません", r.stderr)
        os.rename(self.tpl + ".bak", self.tpl)

    def test_no_embed_uses_relative_path_and_template_override(self):
        script = self.export()
        # テンプレート（self.tpl）は out/ の 1 つ上にある。相対パスで引ける
        data = os.path.join(self.dir, "data.json")
        with open(data, "w", encoding="utf-8") as f:
            json.dump(DATA, f, ensure_ascii=False)
        out = os.path.join(self.dir, "n.xlsx")
        self.assertEqual(self.run_script(script, "--data", data, "-o", out, cwd="/").returncode, 0)
        r = self.run_script(script, "--data", data, "-o", out, "--template", os.path.join(self.dir, "nope.xlsx"))
        self.assertNotEqual(r.returncode, 0)

    def test_export_validates_definition(self):
        bad = json.loads(json.dumps(DEF))
        bad["sheets"][0]["name"] = "ない"
        with self.assertRaises(xt.TemplateError):
            xt.export_script(self.tpl, bad, os.path.join(self.dir, "x.py"))

    def test_bad_data_gives_message_not_traceback(self):
        script = self.export()
        data = os.path.join(self.dir, "bad.yaml")
        with open(data, "w", encoding="utf-8") as f:
            f.write("items: []\n")
        r = self.run_script(script, "--data", data, "-o", os.path.join(self.dir, "b.xlsx"))
        self.assertEqual(r.returncode, 1)
        self.assertIn("エラー", r.stderr)
        self.assertNotIn("Traceback", r.stderr)


class MaintainExportedScriptTests(Base):
    """書き出したスクリプトを、スキルで直す（定義の取り出し → 修正 → 再生成）。"""

    def setUp(self):
        super().setUp()
        self.script = os.path.join(self.dir, "render.py")
        xt.export_script(self.tpl, DEF, self.script)

    def cli(self, *args):
        entry = os.path.join(os.path.dirname(__file__), "..", "scripts", "xlsx_builder.py")
        return subprocess.run([sys.executable, entry, *args], capture_output=True, text=True)

    def run_script(self, *args):
        return subprocess.run([sys.executable, self.script, *args], capture_output=True, text=True)

    def test_definition_can_be_read_back_without_executing(self):
        definition, embedded, rel = xt.read_exported_script(self.script)
        self.assertEqual(definition, DEF)
        self.assertIsNone(embedded)
        self.assertEqual(rel, "t.xlsx")

    def test_extract_def_edit_and_regenerate(self):
        d = os.path.join(self.dir, "got.yaml")
        self.assertEqual(self.run_script("--extract-def", d).returncode, 0)
        edited = xt.load_structured(d)
        edited["sheets"][0]["cells"]["B3"] = "client.name"  # 欄のキーを変える
        xt.dump_structured(edited, d)
        new = os.path.join(self.dir, "render2.py")
        r = self.cli("export", "--from-script", self.script, "--def", d, "-o", new)
        self.assertEqual(r.returncode, 0, r.stderr)
        data = {**DATA, "client": {"name": "改修後"}}
        data_path = os.path.join(self.dir, "data.json")
        with open(data_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)
        out = os.path.join(self.dir, "o.xlsx")
        r = subprocess.run([sys.executable, new, "--data", data_path, "-o", out], capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(load_workbook(out)["請求書"]["B3"].value, "改修後")

    def test_regenerate_keeps_definition_and_relative_template(self):
        new = os.path.join(self.dir, "sub", "render3.py")
        os.makedirs(os.path.dirname(new))
        r = self.cli("export", "--from-script", self.script, "-o", new)
        self.assertEqual(r.returncode, 0, r.stderr)
        definition, embedded, rel = xt.read_exported_script(new)
        self.assertEqual(definition, DEF)
        self.assertIsNone(embedded)
        self.assertEqual(rel, os.path.join("..", "t.xlsx"))

    def test_regenerate_keeps_embedded_template(self):
        emb = os.path.join(self.dir, "emb.py")
        xt.export_script(self.tpl, DEF, emb, embed=True)
        new = os.path.join(self.dir, "emb2.py")
        self.assertEqual(self.cli("export", "--from-script", emb, "-o", new).returncode, 0)
        _, embedded, _ = xt.read_exported_script(new)
        with open(self.tpl, "rb") as f:
            self.assertEqual(embedded, f.read())

    def test_replace_template_keeps_definition(self):
        wb = load_workbook(self.tpl)
        wb["請求書"]["A1"] = "新しい請求書"
        other = os.path.join(self.dir, "t2.xlsx")
        wb.save(other)
        new = os.path.join(self.dir, "render4.py")
        r = self.cli("export", "--from-script", self.script, "--template", other, "-o", new)
        self.assertEqual(r.returncode, 0, r.stderr)
        _, _, rel = xt.read_exported_script(new)
        self.assertEqual(rel, "t2.xlsx")

    def test_rejects_a_foreign_script_and_a_broken_definition(self):
        foreign = os.path.join(self.dir, "other.py")
        with open(foreign, "w", encoding="utf-8") as f:
            f.write("print('hi')\n")
        r = self.cli("export", "--from-script", foreign, "-o", os.path.join(self.dir, "z.py"))
        self.assertEqual(r.returncode, 1)
        self.assertIn("export が書き出したスクリプトではありません", r.stderr)
        bad = json.loads(json.dumps(DEF))
        bad["sheets"][0]["tables"][0]["pattern"] = [99]
        d = os.path.join(self.dir, "bad.json")
        xt.dump_structured(bad, d)
        r = self.cli("export", "--from-script", self.script, "--def", d, "-o", os.path.join(self.dir, "z.py"))
        self.assertEqual(r.returncode, 1)
        self.assertIn("pattern", r.stderr)


def make_card_template(path: str) -> None:
    """1 件が 2 行（明細 + 備考）の「カード」形式。末尾に仮の注意書きとダミー行がある。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "案件"
    ws["A1"], ws["B1"] = "案件名", "（ここに案件名）"
    ws["A2"] = "※ 記入例: 下の 2 行は例です"  # 無視したい注意書き
    for c, h in zip("ABC", ["項目", "金額", "期限"]):
        ws[f"{c}4"].value, ws[f"{c}4"].font, ws[f"{c}4"].fill = h, Font(bold=True), HEAD_FILL
    for base, fill in ((5, None), (7, BAND_FILL)):  # 2 件 × 2 行のサンプル（縞模様はレコード単位）
        ws[f"A{base}"], ws[f"B{base}"], ws[f"C{base}"] = f"例{base}", 100, "2020-01-01"
        ws[f"A{base + 1}"] = "備考: ここに備考"
        ws.merge_cells(f"A{base + 1}:C{base + 1}")
        for r in (base, base + 1):
            for c in "ABC":
                ws[f"{c}{r}"].border = BOX
                if fill:
                    ws[f"{c}{r}"].fill = fill
        ws[f"C{base}"].number_format = "yyyy/mm/dd"
    ws["A9"], ws["B9"] = "合計", "=SUM(B5:B8)"
    ws["A11"] = "ダミー行（削除対象）"
    ws["A12"] = "ダミー行（削除対象）"
    ws["A13"] = "以上"
    wb.save(path)


CARD_DEF = {
    "version": 1,
    "sheets": [{
        "name": "案件",
        "cells": {"B1": "title"},
        "clear": ["A2"],
        "drop_rows": ["11:12"],
        "tables": [{
            "id": "cards", "header_row": 4, "first_row": 5, "sample_rows": 4, "block_rows": 2,
            "pattern": [5, 7], "key": "cards",
            "block": [
                {"A": {"key": "name"}, "B": {"key": "amount"}, "C": {"key": "due"}},
                {"A": {"key": "note"}},
            ],
        }],
    }],
}


class RolesTests(unittest.TestCase):
    """残す（keep）・流し込む（fill）・繰り返す（repeat / block）・無視する（clear / drop_rows）。"""

    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.dir, True)
        self.tpl = os.path.join(self.dir, "card.xlsx")
        make_card_template(self.tpl)
        self.data = {"title": "新規案件", "cards": [
            {"name": f"項目{i}", "amount": 100 * i, "due": f"2026-11-0{i}", "note": f"備考{i}"} for i in range(1, 4)]}

    def render(self, data=None, definition=CARD_DEF):
        out = os.path.join(self.dir, "o.xlsx")
        xt.render(self.tpl, definition, data or self.data, out)
        return load_workbook(out)["案件"]

    def test_block_repeats_two_rows_per_record(self):
        ws = self.render()
        self.assertEqual([ws[f"A{r}"].value for r in range(5, 11)],
                         ["項目1", "備考1", "項目2", "備考2", "項目3", "備考3"])
        self.assertEqual([ws[f"B{r}"].value for r in (5, 7, 9)], [100, 200, 300])
        self.assertIsNone(ws["B6"].value)  # 備考行には金額が無い
        self.assertEqual(ws["C5"].value.date().isoformat(), "2026-11-01")

    def test_block_style_pattern_and_merges_follow(self):
        ws = self.render()
        tpl = load_workbook(self.tpl)["案件"]
        self.assertEqual(ws["A7"].fill.fgColor.rgb, tpl["A7"].fill.fgColor.rgb)  # レコード単位の縞
        self.assertEqual(ws["A9"].fill.fgColor.rgb, tpl["A5"].fill.fgColor.rgb)
        merged = sorted(str(m) for m in ws.merged_cells.ranges)
        self.assertEqual(merged, ["A10:C10", "A6:C6", "A8:C8"])

    def test_total_follows_and_formula_range_extends(self):
        ws = self.render()
        self.assertEqual(ws["A11"].value, "合計")
        self.assertEqual(ws["B11"].value, "=SUM(B5:B10)")

    def test_ignore_clear_keeps_style_and_drop_rows_removes(self):
        ws = self.render()
        self.assertIsNone(ws["A2"].value)  # 注意書きは空に
        self.assertEqual(ws["A1"].value, "案件名")  # 固定は残る
        self.assertEqual(ws["B1"].value, "新規案件")
        values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertFalse(any("ダミー" in str(v) for v in values))
        self.assertEqual(ws["A13"].value, "以上")  # 11-12 行の削除と、レコード増(+2 行)で 13 のまま
        self.assertEqual(ws.max_row, 13)

    def test_block_validation(self):
        bad = json.loads(json.dumps(CARD_DEF))
        bad["sheets"][0]["tables"][0]["block"] = bad["sheets"][0]["tables"][0]["block"][:1]
        with self.assertRaises(xt.TemplateError):
            self.render(definition=bad)
        bad = json.loads(json.dumps(CARD_DEF))
        bad["sheets"][0]["tables"][0]["sample_rows"] = 3
        with self.assertRaises(xt.TemplateError):
            xt.validate_definition(self.tpl, bad)

    def test_example_data_covers_block_columns(self):
        self.assertEqual(sorted(xt.skeleton_data(CARD_DEF)["cards"][0]), ["amount", "due", "name", "note"])


class InspectTests(Base):
    def test_facts_name_style_kinds_hints_and_runs(self):
        facts = xt.inspect_template(self.tpl)
        sheet = facts["sheets"][0]
        cells = {c["ref"]: c for row in sheet["rows"] for c in row["cells"]}
        self.assertEqual(cells["B3"]["hint"], "仮の値の疑い")
        self.assertEqual(cells["A18"]["hint"], "注記の疑い")
        self.assertEqual(cells["E8"]["formula"], "=C8*D8")
        self.assertEqual(cells["A8"]["style"], cells["C8"]["style"])  # 同じ見た目は同じ種類
        self.assertNotEqual(cells["A8"]["style"], cells["A9"]["style"])  # 縞は別の種類
        legend = facts["style_legend"]
        self.assertIn("罫線", legend[cells["A8"]["style"]])
        self.assertIn("#,##0", legend[cells["D8"]["style"]])
        self.assertEqual(sheet["merges"], ["A1:E1"])
        self.assertEqual(sheet["conditional_formats"], ["E8:E9"])
        self.assertEqual([t["header_row"] for t in sheet["auto_detected_tables"]], [7, 15])
        self.assertIn("10-12", [r["rows"] for r in sheet["same_shape_runs"]])

    def test_text_and_cli(self):
        text = xt.format_facts(xt.inspect_template(self.tpl))
        self.assertIn("仮の値の疑い", text)
        self.assertIn("B8: 'サンプル8'", text)
        script = os.path.join(os.path.dirname(__file__), "..", "scripts", "xlsx_builder.py")
        r = subprocess.run([sys.executable, script, "inspect", self.tpl], capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("書式の種類", r.stdout)
        j = subprocess.run([sys.executable, script, "inspect", self.tpl, "--json"], capture_output=True, text=True)
        self.assertEqual(json.loads(j.stdout)["sheets"][0]["name"], "請求書")


class ReusedDeliverableTests(Base):
    """他のプロジェクトの成果物をテンプレートにする場合の、持ち越しの検出と除去。"""

    def setUp(self):
        super().setUp()
        from openpyxl.comments import Comment
        wb = load_workbook(self.tpl)
        wb.properties.creator = "他部署の山田"
        wb.properties.lastModifiedBy = "他部署の佐藤"
        wb.properties.title = "ProjectX 月次報告"
        wb["請求書"]["B3"].comment = Comment("旧顧客の連絡先メモ", "山田")
        wb.create_sheet("旧メモ").sheet_state = "hidden"
        wb.save(self.tpl)
        # 会社名・プレビュー画像を足す（openpyxl は書かない）
        patched = os.path.join(self.dir, "patched.xlsx")
        with zipfile.ZipFile(self.tpl) as zi, zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as zo:
            for info in zi.infolist():
                raw = zi.read(info.filename)
                if info.filename == "docProps/app.xml":
                    raw = raw.decode().replace("</Properties>", "<Company>X 社</Company></Properties>").encode()
                if info.filename == "_rels/.rels":
                    raw = raw.decode().replace("</Relationships>", '<Relationship Id="rIdT" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail" Target="docProps/thumbnail.jpeg"/></Relationships>').encode()
                if info.filename == "[Content_Types].xml":
                    text = raw.decode()
                    if 'Extension="jpeg"' not in text:
                        text = text.replace("<Override", '<Default Extension="jpeg" ContentType="image/jpeg"/><Override', 1)
                    raw = text.encode()
                zo.writestr(info.filename, raw)
            zo.writestr("docProps/thumbnail.jpeg", b"\xff\xd8preview-of-old-content")
        self.tpl = patched

    def props(self, path):
        with zipfile.ZipFile(path) as z:
            return z.read("docProps/core.xml").decode(), z.read("docProps/app.xml").decode(), z.namelist()

    def test_inspect_lists_what_would_be_carried_over(self):
        text = "\n".join(xt.inspect_template(self.tpl)["provenance"])
        for needle in ("他部署の山田", "X 社", "thumbnail", "コメント", "山田", "非表示のシート", "旧メモ"):
            self.assertIn(needle, text)
        self.assertIn("来歴・持ち越しの注意", xt.format_facts(xt.inspect_template(self.tpl)))

    def test_render_warns_but_keeps_properties_by_default(self):
        out = self.render()
        self.assertTrue(any("来歴" in w and "他部署の山田" in w for w in self.warnings))
        core, app, names = self.props(out)
        self.assertIn("他部署の山田", core)
        self.assertIn("docProps/thumbnail.jpeg", names)

    def test_scrub_clears_identity_and_applies_explicit_values(self):
        d = json.loads(json.dumps(DEF))
        d["properties"] = {"scrub": True, "title": "請求書 2026-10", "creator": "経理部"}
        core, app, names = self.props(self.render(definition=d))
        self.assertNotIn("他部署の山田", core)
        self.assertNotIn("他部署の佐藤", core)
        self.assertNotIn("ProjectX", core)
        self.assertIn("経理部", core)
        self.assertIn("請求書 2026-10", core)
        self.assertNotIn("X 社", app)
        self.assertNotIn("docProps/thumbnail.jpeg", names)
        with zipfile.ZipFile(os.path.join(self.dir, "out.xlsx")) as z:
            self.assertNotIn("thumbnail", z.read("_rels/.rels").decode())
        load_workbook(os.path.join(self.dir, "out.xlsx"))  # 壊れていない

    def test_scrub_removes_every_comment_with_its_shape(self):
        from openpyxl.comments import Comment
        wb = load_workbook(self.tpl)
        wb["請求書"]["A18"].comment = Comment("前の案件で足した注記の由来", "山田")   # 残すセルのメモ
        wb.save(self.tpl)
        d = json.loads(json.dumps(DEF))
        d["properties"] = {"scrub": True}
        out = self.render(definition=d)
        self.assertIn("コメント（メモ）1 件を取り除きました（properties.scrub）", self.warnings)
        ws = load_workbook(out)["請求書"]
        self.assertFalse([c.coordinate for row in ws.iter_rows() for c in row if c.comment])
        with zipfile.ZipFile(out) as z:
            names = z.namelist()
            sheet = z.read("xl/worksheets/sheet1.xml").decode()
            types = z.read("[Content_Types].xml").decode()
        self.assertFalse([n for n in names if "comment" in n.lower() or n.endswith(".vml")])
        self.assertNotIn("legacyDrawing", sheet)
        self.assertNotIn("comments", types)

    def test_strict_reports_template_values_that_would_leak(self):
        d = json.loads(json.dumps(DEF))
        d["strict"] = True
        with self.assertRaises(xt.TemplateError) as cm:
            self.render(definition=d)
        msg = str(cm.exception)
        # シートごとに、keep にそのまま書ける範囲で、すべてを挙げる
        self.assertIn("請求書: A1, A3:A5, A7:E7, D10:D12, A14, A15:C15, A18", msg)
        self.assertNotIn("請求書!B8", msg)  # 流し込む列は漏れではない

    def test_strict_passes_once_every_literal_is_decided(self):
        d = json.loads(json.dumps(DEF))
        sheet = d["sheets"][0]
        d["strict"] = True
        sheet["keep"] = ["A1", "A3:A5", "A7:E7", "D10:D12", "A14", "A15:C15", "A18"]
        self.render(definition=d)  # 例外なし
        xt.check_definition(self.tpl, d)

    def test_strict_catches_unlisted_sample_column_in_table(self):
        d = json.loads(json.dumps(DEF))
        d["strict"] = True
        sheet = d["sheets"][0]
        sheet["keep"] = ["A1", "A3:A5", "A7:E7", "D10:D12", "A14", "A15:C15", "A18"]
        del sheet["tables"][1]["columns"]["B"]  # 支払表の方法列を決めていない
        with self.assertRaises(xt.TemplateError) as cm:
            self.render(definition=d)
        self.assertIn("請求書: B16\n", str(cm.exception))
        sheet["tables"][1]["columns"]["B"] = {"keep": True}
        self.render(definition=d)

    def test_check_command_and_export_reject_strict_leaks(self):
        d = json.loads(json.dumps(DEF))
        d["strict"] = True
        with self.assertRaises(xt.TemplateError):
            xt.check_definition(self.tpl, d)
        with self.assertRaises(xt.TemplateError):
            xt.export_script(self.tpl, d, os.path.join(self.dir, "x.py"))
        dp = os.path.join(self.dir, "d.json")
        xt.dump_structured(DEF, dp)
        script = os.path.join(os.path.dirname(__file__), "..", "scripts", "xlsx_builder.py")
        r = subprocess.run([sys.executable, script, "check", "--template", self.tpl, "--def", dp], capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("定義は問題ありません", r.stdout)

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice がない")
    def test_scrub_removes_orphan_shared_strings_and_chart_caches(self):
        lo_dir = os.path.join(self.dir, "lo_tpl")
        os.makedirs(lo_dir)
        subprocess.run(["soffice", "--headless", f"-env:UserInstallation=file://{self.dir}/lo", "--convert-to", "xlsx", "--outdir", lo_dir, self.tpl],
                       check=True, capture_output=True, timeout=120)
        self.tpl = os.path.join(lo_dir, os.path.basename(self.tpl))
        d = json.loads(json.dumps(DEF))
        d["properties"] = {"scrub": True}
        # 置き換えるサンプル値（サンプル株式会社・サンプル8・振込 など）が、共有文字列にも残らない
        out = self.render(definition=d)
        with zipfile.ZipFile(out) as z:
            sst = z.read("xl/sharedStrings.xml").decode()
            charts = [z.read(n).decode() for n in z.namelist() if n.startswith("xl/charts/chart")]
        for old in ("サンプル株式会社", "サンプル8", "サンプル9"):
            self.assertNotIn(old, sst)
        self.assertIn("商品5", sst + "".join(self.cells_text(out)))
        self.assertTrue(charts)
        for c in charts:
            self.assertNotIn("numCache", c)
            self.assertNotIn("strCache", c)
        ws = load_workbook(out)["請求書"]
        self.assertEqual(ws["B12"].value, "商品5")
        self.assertEqual(ws["A7"].value, "No")  # 残す文字列は、再採番後も正しい
        self.assertEqual(ws["A22"].value, "※ 振込手数料はご負担ください")

    def cells_text(self, path):
        with zipfile.ZipFile(path) as z:
            return [z.read(n).decode() for n in z.namelist() if n.startswith("xl/worksheets/sheet")]


class UsabilityGuardTests(Base):
    """使う側の取り違え（綴り・置き場所・値の種類）を、黙って通さない。"""

    def cli(self, *args, cwd=None):
        entry = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "scripts", "xlsx_builder.py"))
        return subprocess.run([sys.executable, entry, *args], capture_output=True, text=True, cwd=cwd)

    def test_template_is_read_beside_the_definition_from_any_working_directory(self):
        defn = os.path.join(self.dir, "def.json")
        self.assertEqual(self.cli("analyze", self.tpl, "-o", defn).returncode, 0)
        with open(defn, encoding="utf-8") as f:
            self.assertEqual(json.load(f)["template"], "t.xlsx")
        with open(defn, "w", encoding="utf-8") as f:
            json.dump(DEF, f, ensure_ascii=False)  # template: "t.xlsx"（定義の隣）
        data = os.path.join(self.dir, "data.json")
        with open(data, "w", encoding="utf-8") as f:
            json.dump(DATA, f, ensure_ascii=False)
        elsewhere = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, elsewhere, True)
        r = self.cli("render", "--def", defn, "--data", data, "-o", os.path.join(self.dir, "o.xlsx"), cwd=elsewhere)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(self.cli("check", "--def", defn, cwd=elsewhere).returncode, 0)

    def test_missing_file_is_a_message_not_a_traceback(self):
        r = self.cli("render", "--def", os.path.join(self.dir, "none.json"), "--data", "x", "-o", "y.xlsx")
        self.assertEqual(r.returncode, 1)
        self.assertIn("見つかりません", r.stderr)
        self.assertNotIn("Traceback", r.stderr)

    def test_keys_absent_from_every_row_are_warned(self):
        data = dict(DATA, items=[{"品名": "a", "qty": 1, "price": 2}])
        self.render(data)
        msg = [w for w in self.warnings if "どの行にも無い" in w]
        self.assertEqual(len(msg), 1)
        self.assertIn("'name'", msg[0])
        self.assertIn("'品名'", msg[0])  # データ側の使われていないキーも示す
        self.render()
        self.assertFalse([w for w in self.warnings if "どの行にも無い" in w])

    def test_data_replaces_a_formula_when_the_column_or_cell_is_filled(self):
        d = json.loads(json.dumps(DEF))
        d["sheets"][0]["tables"][0]["columns"]["E"] = {"key": "amount"}
        d["sheets"][0]["cells"]["E12"] = "billed"  # 請求額（数式）を、データの値で置き換える
        data = dict(DATA, billed=12345, items=[dict(i, amount=7) for i in DATA["items"]])
        ws = load_workbook(self.render(data, d))["請求書"]
        self.assertEqual([ws[f"E{r}"].value for r in range(8, 13)], [7] * 5)
        self.assertEqual(ws["E15"].value, 12345)
        self.assertEqual(ws["E13"].value, "=SUM(E8:E12)")  # 置き換えない数式はそのまま

    def test_shrinking_table_drops_ranges_of_vanished_rows_instead_of_reversing(self):
        wb = load_workbook(self.tpl)
        ws = wb["請求書"]
        dv = DataValidation(type="list", formula1='"A,B"')
        dv.add("B9")                       # 2 行目のサンプル行だけの入力規則
        ws.add_data_validation(dv)
        ws.conditional_formatting.add("D9", CellIsRule(operator="lessThan", formula=["0"], fill=BAND_FILL))
        wb.save(self.tpl)
        out = self.render(dict(DATA, items=DATA["items"][:1]))
        with zipfile.ZipFile(out) as z:
            sheet = z.read("xl/worksheets/sheet1.xml").decode()
        refs = re.findall(r'sqref="([^"]*)"', sheet)
        for ref in refs:                   # 先頭と末尾が逆転した範囲（B9:B8）・空の範囲を作らない
            for a, b in re.findall(r"[A-Z]+(\d+):[A-Z]+(\d+)", ref):
                self.assertLessEqual(int(a), int(b), ref)
            self.assertTrue(ref.strip(), refs)
        self.assertIn("C8", refs)          # 表の全行にあった入力規則（C8:C9）は、残る 1 行に縮む
        self.assertNotIn("B9", " ".join(refs))
        load_workbook(out)

    def test_row_keys_follow_dots_and_indexes(self):
        d = json.loads(json.dumps(DEF))
        cols = d["sheets"][0]["tables"][0]["columns"]
        cols["B"]["key"], cols["C"]["key"], cols["D"]["key"] = "item.name", "nums.0", "nums.1"
        items = [{"item": {"name": f"商品{i}"}, "nums": [i, 100 * i]} for i in (1, 2)]
        ws = load_workbook(self.render(dict(DATA, items=items), d))["請求書"]
        self.assertEqual([[ws[f"{c}{r}"].value for c in "BCD"] for r in (8, 9)], [["商品1", 1, 100], ["商品2", 2, 200]])
        self.assertFalse([w for w in self.warnings if "どの行にも無い" in w])
        skel = xt.skeleton_data(d)["items"][0]
        self.assertEqual((skel["item"], skel["nums"]), ({"name": None}, [None, None]))
        cols["B"]["key"] = "No."           # ドットを含むキーそのものも引ける
        ws = load_workbook(self.render(dict(DATA, items=[dict(i, **{"No.": "X"}) for i in items]), d))["請求書"]
        self.assertEqual(ws["B8"].value, "X")

    def test_values_excel_cannot_store_are_rejected_with_the_cell(self):
        for bad, word in (("a\x0bb", "制御文字"), ("a\ufffeb", "制御文字"), (float("nan"), "保存できません"), ("x" * 40000, "上限")):
            data = dict(DATA, items=[{"name": bad, "qty": 1, "price": 1}])
            with self.assertRaises(xt.TemplateError) as cm:
                self.render(data)
            self.assertIn("B8", str(cm.exception))
            self.assertIn(word, str(cm.exception))

    def test_running_number_column_is_drafted_as_index(self):
        cols = xt.analyze(self.tpl)["sheets"][0]["tables"][0]["columns"]
        self.assertEqual(cols["A"]["key"], "$index")
        self.assertEqual(cols["C"]["key"], "数量")  # 1, 1 は連番ではない


    def test_numeric_text_becomes_a_number_in_number_formatted_cells(self):
        data = dict(DATA, items=[{"name": "1,200", "qty": 1, "price": "1,200"}, {"name": "b", "qty": 1, "price": "300"}])
        ws = load_workbook(self.render(data))["請求書"]
        self.assertEqual((ws["D8"].value, ws["D9"].value), (1200, 300))   # #,##0 のセルは数値に（SUM が数える）
        self.assertEqual(ws["B8"].value, "1,200")                          # 標準のセルは文字のまま

    def small(self, build):
        path = os.path.join(self.dir, "small.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "S"
        build(ws)
        wb.save(path)
        return path

    def test_sample_hyperlinks_do_not_point_new_values_at_the_old_target(self):
        def build(ws):
            ws.append(["名前", "URL"])
            ws.append(["a", "http://old.example"])
            ws["B2"].hyperlink = "http://old.example"
            ws["A1"].hyperlink = "http://keep.example"   # データを入れないセルのリンクは残す
        tpl = self.small(build)
        d = {"version": 1, "sheets": [{"name": "S", "tables": [{"id": "t", "header_row": 1, "first_row": 2, "sample_rows": 1,
                                                                 "key": "rows", "columns": {"A": {"key": "n"}, "B": {"key": "u"}}}]}]}
        out = os.path.join(self.dir, "o.xlsx")
        warnings = xt.render(tpl, d, {"rows": [{"n": "x", "u": "http://new1"}, {"n": "y", "u": "http://new2"}]}, out)
        ws = load_workbook(out)["S"]
        self.assertEqual([ws[c].hyperlink for c in ("B2", "B3")], [None, None])
        self.assertEqual(ws["A1"].hyperlink.target, "http://keep.example")
        self.assertTrue(any("ハイパーリンク" in w for w in warnings))
        with zipfile.ZipFile(out) as z:
            rels = z.read("xl/worksheets/_rels/sheet1.xml.rels").decode()
        self.assertNotIn("old.example", rels)    # 元のリンク先を zip に残さない

    def test_array_formula_range_follows_each_copied_row(self):
        from openpyxl.worksheet.formula import ArrayFormula

        def build(ws):
            ws.append(["a", "b", "c"])
            ws.append([1, 2, None])
            ws["C2"] = ArrayFormula("C2", "=SUM(A2:B2*1)")
        tpl = self.small(build)
        d = {"version": 1, "sheets": [{"name": "S", "tables": [{"id": "t", "header_row": 1, "first_row": 2, "sample_rows": 1,
                                                                 "key": "rows", "columns": {"A": {"key": "a"}, "B": {"key": "b"},
                                                                                            "C": {"formula": True}}}]}]}
        out = os.path.join(self.dir, "o.xlsx")
        xt.render(tpl, d, {"rows": [{"a": 1, "b": 2}, {"a": 3, "b": 4}, {"a": 5, "b": 6}]}, out)
        with zipfile.ZipFile(out) as z:
            sheet = z.read("xl/worksheets/sheet1.xml").decode()
        self.assertEqual(re.findall(r'<c r="(C\d)"><f t="array" ref="([^"]+)"', sheet), [("C2", "C2"), ("C3", "C3"), ("C4", "C4")])

    def test_placeholder_hint_skips_ledger_labels_and_catches_common_dummies(self):
        hinted = [v for v in ("仮払金", "仮受消費税", "仮説", "XXL", "山田 太郎", "Taxi") if xt.PLACEHOLDER_RE.search(v)]
        self.assertEqual(hinted, [])   # 帳票の見出し・実在しそうな値を、仮の値と取り違えない
        missed = [v for v in ("（仮）", "仮の名前", "XX株式会社", "xxx-xxxx", "Example Corp", "0000-00-00", "記入例", "〇〇様")
                  if not xt.PLACEHOLDER_RE.search(v)]
        self.assertEqual(missed, [])

    def test_inspect_text_folds_long_runs_but_keeps_hinted_rows(self):
        def build(ws):
            ws.append(["No", "顧客", "金額"])
            for c in ws[1]:
                c.font = Font(bold=True)
            for i in range(1, 201):
                ws.append([i, "サンプル株式会社" if i == 100 else f"顧客{i}", i * 10])
                for c in ws[ws.max_row]:
                    c.border = BOX
        facts = xt.inspect_template(self.small(build))
        text = xt.format_facts(facts)
        self.assertIn("同じ書式の行 195 行を省略", text)   # 2-201 の 200 行のうち、先頭 3 行・末尾 1 行・印の付いた 1 行を残す
        self.assertIn("サンプル株式会社", text)
        self.assertIn("C201", text)
        self.assertNotIn("顧客50'", text)
        self.assertEqual(len(xt.format_facts(facts, fold=False).splitlines()) - len(text.splitlines()), 194)
        self.assertEqual(len(facts["sheets"][0]["rows"]), 201)   # JSON（事実そのもの）は畳まない


class ChecklistTests(Base):
    """動作確認チェックリスト（AAA パターン・複数タブ・サマリの集計）の形。"""

    def checklist(self, build_tabs):
        path = os.path.join(self.dir, "checklist.xlsx")
        wb = Workbook()
        wb.active.title = "サマリ"
        build_tabs(wb)
        wb.save(path)
        return path

    def test_blank_input_frames_and_values_get_keys(self):
        def build(wb):
            ws = wb.create_sheet("設計書")
            ws.merge_cells("A1:D1")
            ws["A1"] = "システム設計書"
            for r, (label, value) in enumerate([("文書番号", "DOC-0815"), ("版", "1.3"), ("作成日", None),
                                                ("OS", "Windows Server 2019"), ("備考：", None)], start=3):
                ws.merge_cells(f"A{r}:B{r}")   # ラベルが結合でも、右隣は記入枠
                ws[f"A{r}"], ws[f"C{r}"] = label, value
                ws[f"C{r}"].border = BOX       # 空欄でも枠（書式）がある
            ws["A9"], ws["B9"] = 2024, "年度"
            ws["A11"] = "環境情報"
            for c in "ABCD":
                ws[f"{c}11"].fill = BAND_FILL   # 見出しの帯。右の空欄は記入枠ではない
        tpl = self.checklist(build)
        sheet = next(s for s in xt.analyze(tpl)["sheets"] if s["name"] == "設計書")
        self.assertEqual(sheet["cells"], {"C3": "文書番号", "C4": "版", "C5": "作成日", "C6": "OS", "C7": "備考", "A9": "A9"})
        self.assertEqual(sheet["keep"], ["A1", "A3:A7", "B9", "A11"])

    def test_date_code_recognizes_elapsed_time_formats(self):
        for code in ("[h]:mm", "[hh]:mm:ss", "[m]:ss", "[mm]:ss", "[s]", "[ss]"):
            with self.subTest(code=code):
                self.assertTrue(xt._is_date_code(code))

        for code in ("[Red]0.00", "[>=100]0", "[$-409]0.00", '[Blue]0 "hours"',
                     "0.0\\h", "#,##0\\ \\m\\s", "_(* #,##0_)", "0_s", "0*s", "[DBNum1]0"):
            with self.subTest(code=code):
                self.assertFalse(xt._is_date_code(code))

    def test_number_code_ignores_brackets_and_literals(self):
        styles = xt.Styles.__new__(xt.Styles)
        styles.date = {}
        for code, want in (("#,##0", True), ("[$¥-411]#,##0", True), ("[Red]0.0", True),
                           ("[$-409]@", False), ('@"件"', False), ("@\\0", False)):
            styles.code = {1: code}
            with self.subTest(code=code):
                self.assertEqual(styles.is_number("1"), want)

    def build_readable(self, wb):
        ws = wb.create_sheet("試験")
        ws.merge_cells("C1:D1"); ws["C1"] = "判定"
        ws.merge_cells("E1:G1"); ws["E1"] = "対象OS"
        ws.merge_cells("H1:J1"); ws["H1"] = "実施日"
        for i, h in enumerate(["No", "項目", "合格", "不合格", "Windows", "Linux", "macOS", "年", "月", "日", "要再試"], 1):
            cell = ws.cell(2, i, h)
            cell.font, cell.fill, cell.border = Font(bold=True, color="FFFFFF"), HEAD_FILL, BOX
        for r, row in enumerate([(1, "起動", "○", "×", "○", "○", None, 2024, 4, 1, "○"),
                                 (2, "停止", "×", "○", "○", None, None, 2024, 4, 2, None)], 3):
            for i, v in enumerate(row, 1):
                ws.cell(r, i, v).border = BOX

    def test_marks_and_split_dates_become_readable_data(self):
        tpl = self.checklist(self.build_readable)
        d = xt.analyze(tpl)
        sheet = next(s for s in d["sheets"] if s["name"] == "試験")
        cols = sheet["tables"][0]["columns"]
        # ○× の択一は見出しを値に、複数の ○ は配列に、1 列の ○ は true / false に、年月日は 1 つの日付に
        self.assertEqual({L: (cols[L]["key"], cols[L].get("when")) for L in "CDEFG"},
                         {"C": ("判定", "合格"), "D": ("判定", "不合格"), "E": ("対象OS", "Windows"),
                          "F": ("対象OS", "Linux"), "G": ("対象OS", "macOS")})
        self.assertEqual((cols["C"]["mark"], cols["C"]["unmark"]), ("○", "×"))
        self.assertEqual((cols["E"]["_sample"], cols["C"]["_sample"]), (["Windows", "Linux"], "合格"))
        self.assertEqual(cols["K"]["map"], {True: "○", False: None})
        self.assertEqual({L: (cols[L]["key"], cols[L]["part"]) for L in "HIJ"},
                         {"H": ("実施日", "year"), "I": ("実施日", "month"), "J": ("実施日", "day")})
        self.assertEqual(cols["H"]["_sample"], "2024-04-01")
        self.assertIn("C1", sheet["keep"])     # 2 段の見出しの上の段は、流し込む欄にしない
        self.assertEqual(sheet["cells"], {})
        self.assertIn("items[].判定: 合格 / 不合格 のどれか（複数なら配列）", xt.value_notes(d))

        d = json.loads(json.dumps(d))           # JSON の定義（map のキーが "true" / "false" の文字列）でも同じ
        out = os.path.join(self.dir, "o.xlsx")
        items = [{"項目": "起動", "判定": "合格", "対象OS": ["Windows", "macOS"], "実施日": "2026-10-08", "要再試": False},
                 {"項目": "停止", "判定": "不合格", "対象OS": ["Linux"], "実施日": "2026-10-09", "要再試": True},
                 {"項目": "再起動", "判定": None, "対象OS": [], "実施日": None, "要再試": None}]
        xt.render(tpl, d, {"items": items}, out)
        ws = load_workbook(out)["試験"]
        self.assertEqual([[c.value for c in row] for row in ws.iter_rows(min_row=3, max_row=5, min_col=3)],
                         [["○", "×", "○", None, "○", 2026, 10, 8, None],
                          ["×", "○", None, "○", None, 2026, 10, 9, "○"],
                          [None, None, None, None, None, None, None, None, None]])
        for bad, word in (({"判定": "保留"}, "選べる値（合格, 不合格）"), ({"要再試": "たぶん"}, "map にありません"),
                          ({"実施日": "来週"}, "日付として読めません")):
            with self.assertRaises(xt.TemplateError) as cm:
                xt.render(tpl, d, {"items": [dict(items[0], **bad)]}, out)
            self.assertIn(word, str(cm.exception))

    def test_fixed_cells_take_the_same_conversions(self):
        def build(wb):
            ws = wb.create_sheet("届")
            ws["A1"], ws["B1"], ws["C1"] = "性別", "男", "女"
            ws["A2"], ws["B2"] = "提出日", "年"
        d = {"version": 1, "sheets": [{"name": "届", "cells": {
            "B1": {"key": "性別", "when": "男", "mark": "■", "unmark": "□"},
            "C1": {"key": "性別", "when": "女", "mark": "■", "unmark": "□"},
            "C2": {"key": "提出日", "part": "year"}}}]}
        out = os.path.join(self.dir, "o.xlsx")
        xt.render(self.checklist(build), d, {"性別": "女", "提出日": "2026-10-08"}, out)
        ws = load_workbook(out)["届"]
        self.assertEqual((ws["B1"].value, ws["C1"].value, ws["C2"].value), ("□", "■", 2026))
        self.assertEqual(xt.skeleton_data(d), {"性別": None, "提出日": None})

    def test_tables_on_several_tabs_get_separate_data_keys(self):
        def build(wb):
            for name in ("ログイン", "検索"):
                ws = wb.create_sheet(name)
                ws.append(["ID", "Arrange", "Act", "Assert", "判定"])
                for c in ws[1]:
                    c.font = Font(bold=True)
                for i in (1, 2):
                    ws.append([f"X-{i}", "準備", "操作", "期待", "未実施"])
        tables = [t for sd in xt.analyze(self.checklist(build))["sheets"] for t in sd["tables"]]
        self.assertEqual([t["key"] for t in tables], ["ログイン", "検索"])   # 同じ items だと、どのタブにも同じ明細が入る

    def test_merges_inside_a_multi_row_case_are_copied_for_every_case(self):
        def build(wb):
            ws = wb.create_sheet("AAA")
            ws.append(["ID", "段階", "内容", "判定"])
            for base in (2, 5):
                for j, stage in enumerate(("Arrange", "Act", "Assert")):
                    ws[f"B{base + j}"], ws[f"C{base + j}"] = stage, "記入例"
                ws[f"A{base}"], ws[f"D{base}"] = "ID", "未実施"
                for c in "AD":
                    ws.merge_cells(f"{c}{base}:{c}{base + 2}")
            ws["A9"], ws["B9"] = "件数", "=COUNTA(A2:A7)"
        d = {"version": 1, "sheets": [{"name": "AAA", "tables": [{
            "id": "cases", "header_row": 1, "first_row": 2, "sample_rows": 6, "block_rows": 3, "pattern": [2], "key": "cases",
            "block": [{"A": {"key": "id"}, "B": {"keep": True}, "C": {"key": "arrange"}, "D": {"key": "result"}},
                      {"B": {"keep": True}, "C": {"key": "act"}}, {"B": {"keep": True}, "C": {"key": "assert"}}]}]}]}
        out = os.path.join(self.dir, "o.xlsx")
        warnings = xt.render(self.checklist(build), d, {"cases": [
            {"id": f"T-{i}", "arrange": "a", "act": "b", "assert": "c", "result": "OK"} for i in range(1, 4)]}, out)
        ws = load_workbook(out)["AAA"]
        self.assertEqual(sorted(str(m) for m in ws.merged_cells.ranges),
                         ["A2:A4", "A5:A7", "A8:A10", "D2:D4", "D5:D7", "D8:D10"])
        self.assertFalse([w for w in warnings if "結合" in w])
        self.assertEqual(ws["B12"].value, "=COUNTA(A2:A10)")

    def test_an_optional_column_left_blank_is_not_warned_but_a_misspelling_is(self):
        data = dict(DATA, items=[{"name": "a", "qty": 1}])   # price は空欄（綴り違いの手がかりなし）
        self.render(data)
        self.assertFalse([w for w in self.warnings if "どの行にも無い" in w])
        self.render(dict(DATA, items=[{"name": "a", "qty": 1, "prise": 1}]))
        self.assertTrue([w for w in self.warnings if "'prise'" in w])

    def test_remarks_column_heading_is_not_a_note(self):
        self.assertIsNone(xt.NOTE_RE.match("備考"))
        self.assertIsNotNone(xt.NOTE_RE.match("備考：振込手数料はご負担ください"))
        self.assertIsNotNone(xt.NOTE_RE.match("※ 判定は OK / NG から選ぶ"))

    def test_comments_on_filled_cells_are_dropped_and_others_follow_their_rows(self):
        from openpyxl.comments import Comment

        def build(wb):
            ws = wb.create_sheet("T")
            ws.append(["ID", "結果"])
            ws.append(["X-1", "前の結果"])
            ws["B2"].comment = Comment("前のプロジェクトのメモ", "山田")
            ws["A4"] = "※ 注記"
            ws["A4"].comment = Comment("注記の由来", "山田")
            ws["A5"] = "消す行"
            ws["A5"].comment = Comment("消す行のメモ", "山田")
        d = {"version": 1, "sheets": [{"name": "T", "drop_rows": [5], "tables": [{
            "id": "t", "header_row": 1, "first_row": 2, "sample_rows": 1, "key": "rows",
            "columns": {"A": {"key": "id"}, "B": {"key": "result"}}}]}]}
        out = os.path.join(self.dir, "o.xlsx")
        warnings = xt.render(self.checklist(build), d, {"rows": [{"id": f"N-{i}", "result": "OK"} for i in range(3)]}, out)
        ws = load_workbook(out)["T"]
        notes = {c.coordinate: c.comment.text for row in ws.iter_rows() for c in row if c.comment}
        self.assertEqual(notes, {"A6": "注記の由来"})   # 注記は行と一緒に 4 → 6、データのセルと消した行のメモは無い
        self.assertIn("データを入れるセルのコメント（B2）は、新しい値に元のメモが付くため取り除きました", warnings)
        with zipfile.ZipFile(out) as z:
            vml = next(z.read(n).decode() for n in z.namelist() if n.endswith(".vml"))
        self.assertEqual(re.findall(r"<[^>]*Row>(\d+)<", vml), ["5"])   # 図形（吹き出し）の位置も 0 始まりで 5

    def test_inspect_shows_date_cells_as_dates(self):
        import datetime as dt

        def build(wb):
            ws = wb.create_sheet("D")
            ws.append(["実施日", "時刻"])
            ws.append([dt.date(2025, 1, 2), dt.datetime(2025, 1, 2, 9, 30)])
            ws["A2"].number_format = "yyyy/mm/dd"
            ws["B2"].number_format = "yyyy/mm/dd hh:mm"
        facts = xt.inspect_template(self.checklist(build))
        cells = {c["ref"]: c["value"] for row in facts["sheets"][1]["rows"] for c in row["cells"]}
        self.assertEqual((cells["A2"], cells["B2"]), ("2025-01-02", "2025-01-02 09:30"))   # 通し番号（45659）ではなく

    def test_a_template_with_an_inserted_column_is_refused_instead_of_shifting_values(self):
        def build(heads):
            def inner(wb):
                ws = wb.create_sheet("C")
                ws.append(heads)
                for c in ws[1]:
                    c.font = Font(bold=True)
                ws.append(["1"] * len(heads))
                ws.append(["2"] * len(heads))
            return inner
        old = os.path.join(self.dir, "old.xlsx")
        shutil.move(self.checklist(build(["No", "観点", "判定"])), old)
        d = xt.analyze(old)                       # header を残した定義
        new = self.checklist(build(["No", "優先度", "観点", "判定"]))
        out = os.path.join(self.dir, "o.xlsx")
        data = {d["sheets"][1]["tables"][0]["key"]: [{"観点": "a", "判定": "OK"}]}
        xt.render(old, d, data, out)              # 同じ構造なら通る
        with self.assertRaises(xt.TemplateError) as cm:
            xt.render(new, d, data, out)
        self.assertIn("B1: 定義では「観点」、テンプレートでは「優先度」", str(cm.exception))

if __name__ == "__main__":
    unittest.main()
