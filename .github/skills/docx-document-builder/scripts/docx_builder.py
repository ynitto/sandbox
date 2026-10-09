#!/usr/bin/env python3
"""Word (.docx) の文書を作る。テンプレートへの流し込みと、テンプレート無しの新規生成。

テンプレートへの流し込み（文書のスタイル・章・段落・改行を保つ）:
    uv run python scripts/docx_builder.py inspect template.docx             # 判断用の事実を表示
    uv run python scripts/docx_builder.py analyze template.docx -o def.yaml  # 定義の下書き
    uv run python scripts/docx_builder.py check --def def.yaml               # 定義の検査
    uv run python scripts/docx_builder.py render --def def.yaml --data data.yaml -o out.docx
    uv run python scripts/docx_builder.py extract filled.docx --def def.yaml -o data.yaml
    uv run python scripts/docx_builder.py export --def def.yaml -o render_xxx.py   # 単体で動く専用スクリプト

テンプレート無し:
    uv run python scripts/docx_builder.py build --spec spec.json
    uv run python scripts/docx_builder.py example   # サンプル spec を標準出力

スペックの詳細は references/spec.md、定義ファイルは references/template.md を参照。
"""
from __future__ import annotations

import argparse
import io
import json
import os
import sys
from typing import Any

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_template as dt  # noqa: E402

import docx  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402
from lxml import etree  # noqa: E402

PAGES = {"A4": (21.0, 29.7), "A3": (29.7, 42.0), "B5": (18.2, 25.7), "Letter": (21.59, 27.94)}
DEFAULT_STYLE = {"font": "游明朝", "heading_font": "游ゴシック", "size": 10.5, "accent": "1F4E79", "text": "222222"}
HEADING_SIZES = {1: 16, 2: 13, 3: 11.5}
BULLETS = ("•", "◦", "▪")
NUMBERS = (("decimal", "%1."), ("decimalEnclosedParen", "(%2)"), ("decimal", "%3)"))


def _el(tag: str, **attrs):
    e = etree.Element(qn(tag))
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def _sub(parent, tag: str, **attrs):
    e = _el(tag, **attrs)
    parent.append(e)
    return e


def _fonts(rpr, font: str) -> None:
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = _el("w:rFonts")
        rpr.insert(0, rf)
    for k in list(rf.attrib):
        del rf.attrib[k]
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn(f"w:{a}"), font)


class Builder:
    def __init__(self, spec: dict):
        self.spec = spec
        self.style = {**DEFAULT_STYLE, **(spec.get("style") or {})}
        self.doc = docx.Document()
        self.body = self.doc.element.body
        for p in list(self.body):   # 既定の文書の空の段落を除く（用紙の設定 sectPr は残す）
            if dt.local(p) == "p":
                self.body.remove(p)
        self.numbering = self.doc.part.numbering_part.element
        self._page()
        self._styles()
        self.heading_num = self._heading_numbering() if spec.get("numbered_headings", True) else None

    # --- 文書の設定 ---
    def _page(self) -> None:
        page = self.spec.get("page") or {}
        w, h = PAGES.get(page.get("size", "A4"), PAGES["A4"])
        sec = self.doc.sections[0]
        if page.get("orientation") == "landscape":
            w, h = h, w
            sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(w), Cm(h)
        top, bottom, left, right = page.get("margins_cm") or (2.5, 2.5, 2.5, 2.5)
        sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(top), Cm(bottom), Cm(left), Cm(right)

    def _styles(self) -> None:
        st = self.doc.styles
        accent = RGBColor.from_string(self.style["accent"])
        normal = st["Normal"]
        normal.font.size = Pt(self.style["size"])
        normal.font.color.rgb = RGBColor.from_string(self.style["text"])
        _fonts(normal.element.get_or_add_rPr(), self.style["font"])
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.25
        dd = st.element.find(qn("w:docDefaults"))   # 既定の文字（テーマのフォント）も同じにする
        if dd is not None:
            rpr = dd.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}")
            if rpr is not None:
                _fonts(rpr, self.style["font"])
        for lvl in (1, 2, 3):
            s = st[f"Heading {lvl}"]
            s.font.size = Pt(HEADING_SIZES[lvl])
            s.font.bold = True
            s.font.italic = False
            s.font.color.rgb = accent
            _fonts(s.element.get_or_add_rPr(), self.style["heading_font"])
            s.paragraph_format.space_before = Pt(12 if lvl == 1 else 8)
            s.paragraph_format.space_after = Pt(4)
            s.paragraph_format.keep_with_next = True
        for name in ("Title", "Subtitle"):
            s = st[name]
            _fonts(s.element.get_or_add_rPr(), self.style["heading_font"])
            s.font.color.rgb = accent
            s.font.size = Pt(24 if name == "Title" else 13)
            s.font.italic = False
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ppr = s.element.get_or_add_pPr()
            bdr = ppr.find(qn("w:pBdr"))   # 既定の Title の下線（青い罫線）を外す
            if bdr is not None:
                ppr.remove(bdr)

    def _abstract(self, levels: list[tuple[str, str, float]], styles: "list[str] | None" = None) -> str:
        """番号の定義（abstractNum）を足し、それを使う num の id を返す。"""
        root = self.numbering
        aid = str(max([int(a.get(qn("w:abstractNumId"))) for a in root.findall(qn("w:abstractNum"))] + [0]) + 1)
        absn = _el("w:abstractNum", abstractNumId=aid)
        _sub(absn, "w:multiLevelType", val="hybridMultilevel" if not styles else "multilevel")
        for i, (fmt, text, ind_cm) in enumerate(levels):
            lvl = _sub(absn, "w:lvl", ilvl=i)
            _sub(lvl, "w:start", val=1)
            _sub(lvl, "w:numFmt", val=fmt)
            if styles and i < len(styles):
                _sub(lvl, "w:pStyle", val=styles[i])
            _sub(lvl, "w:lvlText", val=text)
            _sub(lvl, "w:lvlJc", val="left")
            ppr = _sub(lvl, "w:pPr")
            left = int(ind_cm * dt.TWIPS_PER_CM)
            hang = 0 if styles else int(0.6 * dt.TWIPS_PER_CM)
            _sub(ppr, "w:ind", left=left, hanging=hang)
        first_num = root.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(absn)
        else:
            root.append(absn)
        nid = str(max([int(n.get(qn("w:numId"))) for n in root.findall(qn("w:num"))] + [0]) + 1)
        num = _sub(root, "w:num", numId=nid)
        _sub(num, "w:abstractNumId", val=aid)
        return nid

    def _heading_numbering(self) -> str:
        ids = [self.doc.styles[f"Heading {i}"].style_id for i in (1, 2, 3)]
        nid = self._abstract([("decimal", "%1.", 0), ("decimal", "%1.%2", 0), ("decimal", "%1.%2.%3", 0)], ids)
        for i, sid in enumerate(ids):
            ppr = self.doc.styles[f"Heading {i + 1}"].element.get_or_add_pPr()
            for old in ppr.findall(qn("w:numPr")):
                ppr.remove(old)
            np_ = _el("w:numPr")
            _sub(np_, "w:ilvl", val=i)
            _sub(np_, "w:numId", val=nid)
            ppr.insert(0, np_)
        return nid

    # --- 部品 ---
    def para(self, text: str, style: "str | None" = None, align=None):
        p = self.doc.add_paragraph(style=style)
        if align is not None:
            p.alignment = align
        dt.set_para_text(p._p, text)
        return p

    def heading(self, text: str, level: int) -> None:
        if not 1 <= level <= 3:
            raise dt.TemplateError(f"見出しの level は 1〜3 です: {level}")
        self.para(text, f"Heading {level}")

    def items(self, items: Any, numbered: bool, where: str) -> None:
        if numbered:
            levels = [(fmt, text, 0.75 * (i + 1)) for i, (fmt, text) in enumerate(NUMBERS)]
        else:
            levels = [("bullet", ch, 0.75 * (i + 1)) for i, ch in enumerate(BULLETS)]
        nid = self._abstract(levels)
        for text, lvl in dt._as_items(items, where):
            if lvl > 2:
                raise dt.TemplateError(f"{where}: リストの段は 3 段（level 0〜2）までです")
            p = self.para(text, "List Paragraph")
            ppr = p._p.get_or_add_pPr()
            np_ = _el("w:numPr")
            _sub(np_, "w:ilvl", val=lvl)
            _sub(np_, "w:numId", val=nid)
            ppr.insert(1 if ppr.find(qn("w:pStyle")) is not None else 0, np_)
            p.paragraph_format.space_after = Pt(2)

    def table(self, block: dict, where: str) -> None:
        cols, rows = block.get("columns") or [], block.get("rows") or []
        if not cols:
            raise dt.TemplateError(f"{where}: columns が空です")
        t = self.doc.add_table(rows=1 + len(rows), cols=len(cols))
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        accent = self.style["accent"]
        for j, head in enumerate(cols):
            cell = t.cell(0, j)
            dt.set_cell_text(cell._tc, str(head))
            for r in cell._tc.iter(qn("w:r")):
                rpr = r.find(qn("w:rPr"))
                if rpr is None:
                    rpr = _el("w:rPr")
                    r.insert(0, rpr)
                _sub(rpr, "w:b")
                _sub(rpr, "w:color", val="FFFFFF")
            tcpr = cell._tc.get_or_add_tcPr()
            _sub(tcpr, "w:shd", val="clear", color="auto", fill=accent)
        trpr = t.rows[0]._tr.get_or_add_trPr()
        _sub(trpr, "w:tblHeader")   # ページをまたぐと、見出しの行を繰り返す
        for i, row in enumerate(rows, start=1):
            values = row if isinstance(row, list) else [row.get(str(c)) for c in cols]
            for j in range(len(cols)):
                dt.set_cell_text(t.cell(i, j)._tc, dt.to_text(values[j] if j < len(values) else ""))
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
        self.doc.add_paragraph()   # 表のすぐ後に続く段落と、表がくっつかないように

    def build(self):
        blocks = self.spec.get("blocks") or []
        if not blocks:
            raise dt.TemplateError("spec.blocks が空です。少なくとも 1 つ必要です")
        for i, b in enumerate(blocks, start=1):
            if not isinstance(b, dict):
                b = {"type": "paragraph", "text": b}
            kind = b.get("type", "paragraph")
            where = f"blocks[{i}]（{kind}）"
            if kind == "title":
                self.para(dt.to_text(b.get("text")), "Title")
                if b.get("subtitle"):
                    self.para(dt.to_text(b["subtitle"]), "Subtitle")
                for line in [b.get("date"), b.get("author")]:
                    if line:
                        self.para(dt.to_text(line), None, WD_ALIGN_PARAGRAPH.RIGHT)
            elif kind == "heading":
                self.heading(dt.to_text(b.get("text")), int(b.get("level", 1)))
            elif kind == "paragraph":
                for text in dt.split_paras(dt.to_text(b.get("text"))) or [""]:
                    self.para(text)
            elif kind in ("bullets", "numbered"):
                self.items(b.get("items"), kind == "numbered", where)
            elif kind == "table":
                self.table(b, where)
            elif kind == "page_break":
                self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            else:
                raise dt.TemplateError(f"{where}: type は title・heading・paragraph・bullets・numbered・table・page_break のどれかです")
        props = self.spec.get("properties") or {}
        cp = self.doc.core_properties
        cp.title = props.get("title", "")
        cp.author = props.get("creator", "")
        cp.last_modified_by = props.get("creator", "")
        cp.comments = ""
        return self.doc


def build(spec: dict) -> str:
    doc = Builder(spec).build()
    filename = spec.get("filename", "document.docx")
    buf = io.BytesIO()
    doc.save(buf)
    props = spec.get("properties") or {}
    given = {k: props[k] for k in ("title", "creator") if props.get(k)}
    if props.get("creator"):
        given["lastModifiedBy"] = props["creator"]
    # 既定のひな形の名残（プレビュー画像・題の一覧）を残さない
    with open(filename, "wb") as f:
        f.write(dt.finish_package(buf.getvalue(), {"scrub": True, **given}))
    return filename


EXAMPLE_SPEC = {
    "filename": "report.docx",
    "page": {"size": "A4", "margins_cm": [2.5, 2.5, 2.5, 2.5]},
    "style": {"font": "游明朝", "heading_font": "游ゴシック", "size": 10.5, "accent": "1F4E79"},
    "numbered_headings": True,
    "properties": {"title": "業務改善の報告", "creator": "docx-document-builder"},
    "blocks": [
        {"type": "title", "text": "業務改善の報告", "subtitle": "申請の確認にかかる時間を半分にする", "date": "2026 年 10 月 8 日"},
        {"type": "heading", "text": "概要", "level": 1},
        {"type": "paragraph", "text": "申請の確認に 1 件 30 分かかっている。確認者が 2 人しかいないため、月末に滞留する。\n\n"
                                      "今期は確認を自動化し、確認の時間を半分にする。"},
        {"type": "heading", "text": "進め方", "level": 1},
        {"type": "numbered", "items": ["現状を調べる（2 週間）", "確認の手順を決める", {"text": "試行する", "children": ["経理部で 1 か月"]}]},
        {"type": "heading", "text": "課題", "level": 2},
        {"type": "bullets", "items": ["確認者の偏り", "差し戻しが月 40 件"]},
        {"type": "heading", "text": "費用", "level": 1},
        {"type": "table", "columns": ["項目", "金額", "備考"],
         "rows": [["設計", "120 万円", "3 週間"], ["試行", "80 万円", "1 部署"], ["合計", "200 万円", ""]]},
        {"type": "paragraph", "text": "※ 金額は税抜き。"},
    ],
}


def main() -> int:
    parser = argparse.ArgumentParser(description="Word (.docx) を作る（テンプレートへの流し込み・新規生成）")
    sub = parser.add_subparsers(dest="command", required=True)
    b = sub.add_parser("build", help="spec から docx を新規に生成する（テンプレート無し）")
    b.add_argument("--spec", help="スペック JSON ファイル（省略時は stdin）")
    sub.add_parser("example", help="サンプル spec を標準出力に表示")
    dt.add_subcommands(sub)
    args = parser.parse_args()
    try:
        if args.command == "example":
            json.dump(EXAMPLE_SPEC, sys.stdout, ensure_ascii=False, indent=2)
            print()
            return 0
        if args.command == "build":
            spec = dt.load_structured(args.spec) if args.spec else json.load(sys.stdin)
            print(f"生成しました: {build(spec)}")
            return 0
        return args.func(args)
    except dt.TemplateError as e:
        print(f"エラー: {e}", file=sys.stderr)
        return 1
    except FileNotFoundError as e:
        print(f"エラー: ファイルが見つかりません: {e.filename}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
