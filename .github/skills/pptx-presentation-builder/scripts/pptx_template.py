#!/usr/bin/env python3
"""pptx テンプレートへ内容を流し込む（スライドのスタイル・図形はテンプレートのまま保つ）。

サブコマンド（xlsx-report-builder と同じ並び）:

    inspect  テンプレートの事実（図形・文字・大きさ・収まる字数・仮値の疑い・来歴）を出す
    analyze  定義ファイルの下書きを作る（繰り返すスライド・箇条書き・表・図の並びを機械的に拾う）
    check    定義とテンプレートの整合、決め忘れの値を検査する
    render   テンプレート + 定義 + データから pptx を再構成する
    extract  記入済みの pptx から、定義に沿ってデータを取り出す（render の逆）
    export   定義を埋め込んだ、単体で動く専用スクリプトを書き出す

図（矩形・丸・矢印線など）は、テンプレートの図形を複製・移動して使い、画像にしない。
文字は、図形の大きさと文字の大きさから収まる量を見積もり、テンプレートのサンプルと同じ粒度
（max_chars・max_items）を超えたら止める。定義ファイルとデータの書式は references/template.md。
"""
from __future__ import annotations

import argparse
import base64
import datetime as dt
import io
import json
import math
import os
import re
import sys
import unicodedata
import zipfile
from copy import deepcopy
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NSMAP = {"a": NS_A, "p": NS_P, "r": NS_R}

DEF_VERSION = 1
EMU_PER_PT = 12700
EMU_PER_CM = 360000
LINE_FACTOR = 1.2          # 行送り（文字の大きさに対する倍率）の既定
LATIN_WIDTH = 0.55         # 半角文字の幅（全角を 1 とした目安）
DEFAULT_INSETS = (91440, 45720, 91440, 45720)  # 左・上・右・下（bodyPr の既定）
GRANULARITY = 1.5          # サンプルの文字数の何倍までを、同じ粒度と見るか

SHAPE_TAGS = ("sp", "cxnSp", "grpSp", "graphicFrame", "pic")
LINE_PRSTS = {"line", "straightConnector1", "bentConnector2", "bentConnector3", "bentConnector4", "bentConnector5",
              "curvedConnector2", "curvedConnector3", "curvedConnector4", "curvedConnector5"}
ARROW_PRSTS = {"rightArrow", "leftArrow", "upArrow", "downArrow", "leftRightArrow", "upDownArrow", "notchedRightArrow",
               "stripedRightArrow"}
PLACEHOLDER_RE = re.compile(r"(〇〇|○○|●●|◯◯|△△|□□|＊＊|\*\*|(?<![A-Za-z])x{2,}(?![A-Za-z])|サンプル|ダミー|記入例|"
                            r"テキストを入力|ここに|を書く|を入力|sample|dummy|lorem|yyyy|YYYY|xxxx|XXXX|mm/dd|20\d\d[/-]0?1[/-]0?1)", re.I)
NOTE_RE = re.compile(r"^\s*(※|＊|注[:：）)]|Note[:：])")
LABEL_RE = re.compile(r"^[^:：]{1,16}[:：]$")
HUMAN_RE = re.compile(r"(印$|押印|捺印|検印|署名|サイン|自署|承認者?$|決裁|確認者|受付者?$|手書き|記入欄)")
MARK_ON = ("○", "◯", "〇", "●", "◎", "✓", "✔", "レ", "☑", "■")
MARK_OFF = ("×", "✕", "✖", "☐", "□", "-", "－", "ー", "―")
DATE_PARTS = ("year", "month", "day", "hour", "minute")


class TemplateError(Exception):
    """定義ファイル・データ・テンプレートの不整合。"""


def qa(tag: str) -> str:
    return f"{{{NS_A}}}{tag}"


def qp(tag: str) -> str:
    return f"{{{NS_P}}}{tag}"


def local(el) -> str:
    return etree.QName(el).localname if isinstance(el.tag, str) else ""


def open_prs(source: "str | bytes"):
    return Presentation(io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else source)


def read_bytes(source: "str | bytes") -> bytes:
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    with open(source, "rb") as f:
        return f.read()


# ---------------------------------------------------------------------------
# 図形の基本情報
# ---------------------------------------------------------------------------

def nv_pr(el):
    """図形の cNvPr（id・name を持つ）。"""
    for child in el:
        if local(child).startswith("nv"):
            return child.find(qp("cNvPr"))
    return None


def shape_id(el) -> int:
    c = nv_pr(el)
    return int(c.get("id")) if c is not None and c.get("id", "").isdigit() else 0


def shape_name(el) -> str:
    c = nv_pr(el)
    return c.get("name", "") if c is not None else ""


def placeholder(el) -> "tuple[str, int] | None":
    """プレースホルダーなら (type, idx)。type の省略は body（obj）。"""
    for child in el:
        if local(child).startswith("nv"):
            nv = child.find(qp("nvPr"))
            ph = nv.find(qp("ph")) if nv is not None else None
            if ph is not None:
                return ph.get("type", "body"), int(ph.get("idx", "0"))
    return None


def prst(el) -> str:
    g = el.find(f"{qp('spPr')}/{qa('prstGeom')}")
    if g is not None:
        return g.get("prst", "")
    return "custom" if el.find(f"{qp('spPr')}/{qa('custGeom')}") is not None else ""


def table_of(el):
    return el.find(f".//{qa('tbl')}") if local(el) == "graphicFrame" else None


def kind_of(el) -> str:
    tag = local(el)
    if tag == "sp":
        ph = placeholder(el)
        if ph:
            return f"placeholder:{ph[0]}"
        p = prst(el)
        if p in LINE_PRSTS:
            return f"line:{p}"
        if el.find(f"{qp('nvSpPr')}/{qp('cNvSpPr')}[@txBox='1']") is not None:
            return "textbox"
        return f"shape:{p or 'none'}"
    if tag == "cxnSp":
        return f"connector:{prst(el) or 'line'}"
    if tag == "graphicFrame":
        if table_of(el) is not None:
            return "table"
        uri = el.find(f".//{qa('graphicData')}")
        uri = uri.get("uri", "") if uri is not None else ""
        return "chart" if uri.endswith("/chart") else ("diagram" if "diagram" in uri else "object")
    if tag == "grpSp":
        return "group"
    if tag == "pic":
        return "picture"
    return tag


def is_connector(el) -> bool:
    """矢印線・線・ブロック矢印（図の要素どうしをつなぐもの）。"""
    return local(el) == "cxnSp" or prst(el) in LINE_PRSTS or prst(el) in ARROW_PRSTS


def arrow_ends(el) -> str:
    ln = el.find(f"{qp('spPr')}/{qa('ln')}")
    if ln is None:
        return ""
    ends = [n for n in ("headEnd", "tailEnd") if ln.find(qa(n)) is not None and ln.find(qa(n)).get("type", "none") != "none"]
    return "・".join({"headEnd": "始点に矢印", "tailEnd": "終点に矢印"}[n] for n in ends)


def child_shapes(container) -> list:
    return [c for c in container if local(c) in SHAPE_TAGS]


def walk(container):
    """図形を、グループの中まで順にたどる（グループ自身も返す）。"""
    for c in child_shapes(container):
        yield c
        if local(c) == "grpSp":
            yield from walk(c)


def xfrm_of(el):
    tag = local(el)
    if tag == "grpSp":
        return el.find(f"{qp('grpSpPr')}/{qa('xfrm')}")
    if tag == "graphicFrame":
        return el.find(qp("xfrm"))
    return el.find(f"{qp('spPr')}/{qa('xfrm')}")


def raw_box(el) -> "tuple[int, int, int, int] | None":
    x = xfrm_of(el)
    if x is None or x.find(qa("off")) is None or x.find(qa("ext")) is None:
        return None
    off, ext = x.find(qa("off")), x.find(qa("ext"))
    return int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))


def translate(el, dx: int, dy: int) -> None:
    x = xfrm_of(el)
    if x is None or x.find(qa("off")) is None:
        return
    off = x.find(qa("off"))
    off.set("x", str(int(off.get("x")) + int(dx)))
    off.set("y", str(int(off.get("y")) + int(dy)))


# ---------------------------------------------------------------------------
# 文字
# ---------------------------------------------------------------------------

def txbody(el):
    """図形・表のセルの txBody。"""
    if local(el) == "tc":
        return el.find(qa("txBody"))
    return el.find(qp("txBody"))


def para_level(p) -> int:
    ppr = p.find(qa("pPr"))
    return int(ppr.get("lvl", "0")) if ppr is not None else 0


def para_text(p, literal_only: bool = False) -> str:
    out = []
    for c in p:
        tag = local(c)
        if tag == "r" or (tag == "fld" and not literal_only):
            t = c.find(qa("t"))
            out.append(t.text or "" if t is not None else "")
        elif tag == "br":
            out.append("\n")
    return "".join(out)


def paragraphs(body) -> list[tuple[str, int]]:
    if body is None:
        return []
    return [(para_text(p), para_level(p)) for p in body.findall(qa("p"))]


def body_text(body) -> str:
    return "\n".join(t for t, _ in paragraphs(body)).strip("\n")


def literal_text(el) -> str:
    """図形（グループなら中のすべて）の、手で書かれた文字（スライド番号などのフィールドを除く）。"""
    bodies = [txbody(el)] if local(el) != "grpSp" else [txbody(c) for c in walk(el)]
    if table_of(el) is not None:
        bodies = list(table_of(el).iter(qa("txBody")))
    return "\n".join(para_text(p, literal_only=True) for b in bodies if b is not None for p in b.findall(qa("p"))).strip()


def text_units(s: str) -> float:
    """表示幅（全角 1・半角 LATIN_WIDTH）。"""
    return sum(1.0 if unicodedata.east_asian_width(ch) in ("W", "F", "A") else LATIN_WIDTH for ch in s if ch != "\n")


def _new_p():
    return etree.Element(qa("p"))


def _set_level(p, lvl: int) -> None:
    ppr = p.find(qa("pPr"))
    if lvl == 0:
        if ppr is not None and "lvl" in ppr.attrib:
            del ppr.attrib["lvl"]
        return
    if ppr is None:
        ppr = etree.Element(qa("pPr"))
        p.insert(0, ppr)
    ppr.set("lvl", str(lvl))


def _set_runs(p, text: str) -> None:
    runs = p.findall(qa("r"))
    rpr = None
    if runs and runs[0].find(qa("rPr")) is not None:
        rpr = deepcopy(runs[0].find(qa("rPr")))
    elif p.find(qa("endParaRPr")) is not None:
        rpr = deepcopy(p.find(qa("endParaRPr")))
        rpr.tag = qa("rPr")
    for c in list(p):
        if local(c) in ("r", "br", "fld"):
            p.remove(c)
    if not text:
        return
    r = etree.Element(qa("r"))
    if rpr is not None:
        for attr in ("dirty", "err"):
            rpr.attrib.pop(attr, None)
        r.append(rpr)
    t = etree.SubElement(r, qa("t"))
    t.text = text
    end = p.find(qa("endParaRPr"))
    if end is not None:
        end.addprevious(r)
    else:
        p.append(r)


def set_paragraphs(body, paras: list[tuple[str, int]]) -> None:
    """段落を入れ替える。段落・文字の書式は、同じ位置のサンプルの段落（段が同じとき）から取る。

    1 つの図形の中の見出しと本文のように、段落ごとに書式が違うサンプルを保つ。サンプルより多い段落は、
    同じ段（lvl）の最後のサンプルの段落に合わせる。
    """
    ps = body.findall(qa("p"))
    protos: dict[int, Any] = {}
    for p in ps:
        protos[para_level(p)] = p
    first = ps[0] if ps else None
    for p in ps:
        body.remove(p)
    for i, (text, lvl) in enumerate(paras or [("", 0)]):
        lower = [k for k in protos if k <= lvl]
        proto = ps[i] if i < len(ps) and para_level(ps[i]) == lvl else \
            protos.get(lvl, protos.get(max(lower)) if lower else first)
        p = deepcopy(proto) if proto is not None else _new_p()
        _set_level(p, lvl)
        _set_runs(p, text)
        body.append(p)


def set_text(body, text: str) -> None:
    set_paragraphs(body, [(line, 0) for line in text.split("\n")] if text else [])


def to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    if isinstance(value, (list, tuple)):
        return "\n".join(to_text(v) for v in value)
    return str(value)


# ---------------------------------------------------------------------------
# スライドの見え方（位置・文字の大きさ・収まる量）
# ---------------------------------------------------------------------------

class SlideView:
    """1 枚のスライドの図形と、継承込みの位置・文字の大きさを引く。"""

    def __init__(self, prs, slide, number: int = 0):
        self.prs, self.slide, self.number = prs, slide, number
        self.el = slide._element
        self.tree = self.el.find(f"{qp('cSld')}/{qp('spTree')}")
        self.width, self.height = int(prs.slide_width), int(prs.slide_height)
        self._objs = {s._element: s for s in slide.shapes}

    def top(self) -> list:
        return child_shapes(self.tree)

    def find(self, ref: str, where: str = ""):
        """図形を名前か `#id` で引く（グループの中も探す）。"""
        hits = [e for e in walk(self.tree)
                if (ref.startswith("#") and str(shape_id(e)) == ref[1:]) or shape_name(e) == ref]
        if not hits:
            raise TemplateError(f"{where}スライド {self.number} に図形「{ref}」がありません")
        if len(hits) > 1:
            raise TemplateError(f"{where}スライド {self.number} に「{ref}」という名前の図形が {len(hits)} つあります。"
                                f"`#id`（{', '.join('#' + str(shape_id(e)) for e in hits)}）で指定してください")
        return hits[0]

    def box(self, el) -> tuple[int, int, int, int]:
        """スライド上の位置と大きさ（EMU）。グループの中は、グループの座標を写して求める。"""
        b = raw_box(el)
        if b is None:
            obj = self._objs.get(el)
            if obj is not None and obj.width is not None:
                return int(obj.left), int(obj.top), int(obj.width), int(obj.height)
            return 0, 0, 0, 0
        x, y, w, h = b
        parent = el.getparent()
        while parent is not None and local(parent) == "grpSp":
            gx = xfrm_of(parent)
            if gx is None:
                break
            off, ext = gx.find(qa("off")), gx.find(qa("ext"))
            ch_off, ch_ext = gx.find(qa("chOff")), gx.find(qa("chExt"))
            if None in (off, ext, ch_off, ch_ext):
                break
            sx = int(ext.get("cx")) / max(int(ch_ext.get("cx")), 1)
            sy = int(ext.get("cy")) / max(int(ch_ext.get("cy")), 1)
            x = int(off.get("x")) + (x - int(ch_off.get("x"))) * sx
            y = int(off.get("y")) + (y - int(ch_off.get("y"))) * sy
            w, h = w * sx, h * sy
            parent = parent.getparent()
        return int(x), int(y), int(w), int(h)

    # --- 文字の大きさ（継承をたどる） ---
    def font_size(self, el, body, lvl: int = 0) -> float:
        def from_runs(b):
            if b is None:
                return None
            for p in b.findall(qa("p")):
                if para_level(p) != lvl:
                    continue
                for r in p.findall(qa("r")):
                    rpr = r.find(qa("rPr"))
                    if rpr is not None and rpr.get("sz"):
                        return int(rpr.get("sz")) / 100
                end = p.find(qa("endParaRPr"))
                if end is not None and end.get("sz"):
                    return int(end.get("sz")) / 100
            return None

        def from_lst(b):
            if b is None:
                return None
            d = b.find(f"{qa('lstStyle')}/{qa(f'lvl{lvl + 1}pPr')}/{qa('defRPr')}")
            return int(d.get("sz")) / 100 if d is not None and d.get("sz") else None

        scale = 1.0
        if body is not None:
            auto = body.find(f"{qa('bodyPr')}/{qa('normAutofit')}")
            if auto is not None and auto.get("fontScale"):
                scale = int(auto.get("fontScale")) / 100000
        size = from_runs(body) or from_lst(body)
        ph = placeholder(el) if local(el) != "tc" else None
        if size is None and ph is not None:
            for owner in self._inherited(ph):
                size = from_runs(owner.find(qp("txBody"))) or from_lst(owner.find(qp("txBody")))
                if size:
                    break
            if size is None:
                styles = self.slide.slide_layout.slide_master._element.find(qp("txStyles"))
                name = "titleStyle" if ph[0] in ("title", "ctrTitle") else (
                    "bodyStyle" if ph[0] in ("body", "obj", "subTitle") else "otherStyle")
                d = styles.find(f"{qp(name)}/{qa(f'lvl{lvl + 1}pPr')}/{qa('defRPr')}") if styles is not None else None
                size = int(d.get("sz")) / 100 if d is not None and d.get("sz") else None
        if size is None and local(el) == "tc":
            size = 18.0
        if size is None:
            styles = self.prs.part._element.find(qp("defaultTextStyle"))
            d = styles.find(f"{qa(f'lvl{lvl + 1}pPr')}/{qa('defRPr')}") if styles is not None else None
            size = int(d.get("sz")) / 100 if d is not None and d.get("sz") else 18.0
        return size * scale

    def _inherited(self, ph: tuple[str, int]) -> list:
        """プレースホルダーの継承元（レイアウト → マスター）の図形。"""
        out = []
        for owner in (self.slide.slide_layout, self.slide.slide_layout.slide_master):
            tree = owner._element.find(f"{qp('cSld')}/{qp('spTree')}")
            for e in walk(tree):
                p = placeholder(e)
                if p and (p == ph or (p[0] == ph[0] and ph[0] not in ("body", "obj"))
                          or (ph[0] in ("body", "obj") and p[1] == ph[1] and ph[1] != 0)):
                    out.append(e)
                    break
        return out

    def insets(self, body) -> tuple[int, int, int, int]:
        bp = body.find(qa("bodyPr")) if body is not None else None
        if bp is None:
            return DEFAULT_INSETS
        return tuple(int(bp.get(k, d)) for k, d in zip(("lIns", "tIns", "rIns", "bIns"), DEFAULT_INSETS))

    def capacity(self, el, body=None, box=None, lvl: int = 0) -> "Capacity":
        """図形に、テンプレートの文字の大きさのまま収まる量。"""
        body = body if body is not None else txbody(el)
        x, y, w, h = box or self.box(el)
        size = self.font_size(el, body, lvl)
        l, t, r, b = self.insets(body)
        line_h = size * EMU_PER_PT * line_factor(body)
        bp = body.find(qa("bodyPr")) if body is not None else None
        if bp is not None and bp.find(qa("spAutoFit")) is not None:   # 図形が文字に合わせて伸びる
            h = max(h, self.free_end(el, (x, y, w, h), "y") - y)
        cpl = max((w - l - r) / (size * EMU_PER_PT), 1.0)
        lines = max(int((h - t - b) / line_h + 0.15), 1)
        return Capacity(cpl, lines, size, body)

    def margin(self, axis: str = "x", exclude=()) -> int:
        """スライドの余白（axis の向き）。スライドとレイアウトの図形の、スライドの端からの距離の最小。

        端に付いた図形（背景・帯）と exclude（並べ直す図の見本）は数えない。図形が無ければ 5%。
        """
        size = self.width if axis == "x" else self.height
        layout = self.slide.slide_layout._element.find(f"{qp('cSld')}/{qp('spTree')}")
        ms = []
        boxes = [self.box(e) for e in self.top() if e not in exclude]
        boxes += [raw_box(e) or (0, 0, 0, 0) for e in (child_shapes(layout) if layout is not None else [])]
        for x, y, w, h in boxes:
            a, length = (x, w) if axis == "x" else (y, h)
            d = min(a, size - (a + length))
            if length > 0 and d > size * 0.01:
                ms.append(d)
        return int(min(ms)) if ms else int(size * 0.05)

    def free_end(self, el, box, axis: str, exclude=()) -> int:
        """box の先（x なら右、y なら下）に、他の図形やスライドの余白にぶつからず使える端。"""
        x, y, w, h = box
        limit = (self.width if axis == "x" else self.height) - self.margin(axis, (el, *exclude))
        for other in self.top():
            if other is el or other in exclude:
                continue
            ox, oy, ow, oh = self.box(other)
            if axis == "x":
                if oy < y + h and oy + oh > y and ox >= x + w - 1:
                    limit = min(limit, ox)
            elif ox < x + w and ox + ow > x and oy >= y + h - 1:
                limit = min(limit, oy)
        return int(limit)

    def free_start(self, el, box, axis: str, exclude=()) -> int:
        x, y, w, h = box
        limit = self.margin(axis, (el, *exclude))
        for other in self.top():
            if other is el or other in exclude:
                continue
            ox, oy, ow, oh = self.box(other)
            if axis == "x":
                if oy < y + h and oy + oh > y and ox + ow <= x + 1:
                    limit = max(limit, ox + ow)
            elif ox < x + w and ox + ow > x and oy + oh <= y + 1:
                limit = max(limit, oy + oh)
        return int(limit)


def line_factor(body) -> float:
    if body is None:
        return LINE_FACTOR
    sp = body.find(f".//{qa('lnSpc')}/{qa('spcPct')}")
    if sp is not None and sp.get("val"):
        return max(int(sp.get("val")) / 100000 * LINE_FACTOR, 0.8)
    return LINE_FACTOR


class Capacity:
    def __init__(self, cpl: float, lines: int, size: float, body=None):
        self.cpl, self.lines, self.size, self.body = cpl, lines, size, body

    def cpl_at(self, lvl: int) -> float:
        """段（lvl）の字下げを除いた、1 行の字数。"""
        if self.body is None or lvl == 0:
            return self.cpl
        for p in self.body.findall(qa("p")):
            if para_level(p) == lvl and p.find(qa("pPr")) is not None and p.find(qa("pPr")).get("marL"):
                return max(self.cpl - int(p.find(qa("pPr")).get("marL")) / (self.size * EMU_PER_PT), 1.0)
        return max(self.cpl - lvl * 2, 1.0)

    def lines_for(self, paras: list[tuple[str, int]]) -> int:
        return sum(max(1, math.ceil(text_units(t) / self.cpl_at(lvl) - 1e-9)) for t, lvl in paras) if paras else 0

    def describe(self) -> str:
        return f"1 行 約 {int(self.cpl)} 字 × {self.lines} 行（{self.size:g}pt）"


# ---------------------------------------------------------------------------
# 来歴（前の文書の値・メタデータ）
# ---------------------------------------------------------------------------

NS_DC = "http://purl.org/dc/elements/1.1/"
NS_CP = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
NS_DCT = "http://purl.org/dc/terms/"
NS_APP = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"
CORE_FIELDS = {"title": f"{{{NS_DC}}}title", "subject": f"{{{NS_DC}}}subject", "creator": f"{{{NS_DC}}}creator",
               "keywords": f"{{{NS_CP}}}keywords", "description": f"{{{NS_DC}}}description",
               "lastModifiedBy": f"{{{NS_CP}}}lastModifiedBy", "category": f"{{{NS_CP}}}category"}
APP_FIELDS = {"company": f"{{{NS_APP}}}Company", "manager": f"{{{NS_APP}}}Manager"}
APP_STALE = ("HeadingPairs", "TitlesOfParts", "Slides", "Notes", "HiddenSlides", "MMClips", "Words", "Paragraphs", "TotalTime")
NS_VT = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"


def read_properties(raw: bytes) -> dict[str, str]:
    out: dict[str, str] = {}
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        names = set(z.namelist())
        for part, fields in (("docProps/core.xml", CORE_FIELDS), ("docProps/app.xml", APP_FIELDS)):
            if part not in names:
                continue
            root = etree.fromstring(z.read(part))
            for name, tag in fields.items():
                el = root.find(tag)
                if el is not None and (el.text or "").strip():
                    out[name] = el.text.strip()
    return out


def provenance(raw: bytes) -> list[str]:
    """前の文書から持ち越しうるもの（プロパティ・ノート・コメント・非表示・リンク・埋め込み）。"""
    notes: list[str] = []
    for k, v in read_properties(raw).items():
        notes.append(f"文書のプロパティ {k}: {v}")
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        names = z.namelist()
        if "docProps/app.xml" in names:
            titles = [e.text for e in etree.fromstring(z.read("docProps/app.xml")).iter(f"{{{NS_VT}}}lpstr") if e.text]
            if titles:
                notes.append(f"文書の情報に、フォントとスライドの題の一覧が残っている（{_short(' / '.join(titles), 60)}）")
    if any(n.startswith("docProps/thumbnail") for n in names):
        notes.append("プレビュー画像（docProps/thumbnail）がある。元のスライドの見た目が残る")
    if "docProps/custom.xml" in names:
        notes.append("カスタムプロパティ（docProps/custom.xml）がある")
    if any(n.endswith("vbaProject.bin") for n in names):
        notes.append("マクロ（VBA）がある。出力の拡張子は .pptm にする")
    prs = open_prs(raw)
    seen = set()
    for master in prs.slide_masters:
        for owner, what in [(master, "スライドマスター")] + [(lay, f"レイアウト「{lay.name}」") for lay in master.slide_layouts]:
            for e in walk(owner._element.find(f"{qp('cSld')}/{qp('spTree')}")):
                text = literal_text(e) if placeholder(e) is None else ""
                if text and (what, text) not in seen:
                    seen.add((what, text))
                    notes.append(f"{what}: 「{_short(text, 30)}」（どのスライドにも出る。render は書き換えない）")
    for i, slide in enumerate(prs.slides, start=1):
        if slide._element.get("show") == "0":
            notes.append(f"スライド {i}: 非表示のスライド")
        if slide.has_notes_slide:
            text = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if text:
                notes.append(f"スライド {i}: ノート「{_short(text, 40)}」")
        for rel in slide.part.rels.values():
            if "comment" in rel.reltype.lower() and "author" not in rel.reltype.lower():
                notes.append(f"スライド {i}: コメントがある")
            elif rel.is_external and rel.reltype == RT.HYPERLINK:
                notes.append(f"スライド {i}: 外部へのリンク {rel.target_ref}")
            elif rel.reltype == RT.CHART:
                notes.append(f"スライド {i}: グラフ（値は元の文書のまま。render は書き換えない）")
            elif rel.reltype in (RT.OLE_OBJECT, RT.PACKAGE):
                notes.append(f"スライド {i}: 埋め込みのオブジェクト（元の文書の中身が入っている）")
    return notes


def _short(s: str, n: int) -> str:
    s = s.replace("\n", " / ")
    return s if len(s) <= n else s[:n - 1] + "…"


def drop_comments(prs) -> int:
    """スライドのコメント（旧形式・新形式）とコメントの作成者を取り除く。"""
    count = 0
    for slide in prs.slides:
        for rId, rel in list(slide.part.rels.items()):
            if "comment" in rel.reltype.lower():
                for ext in slide._element.iter(qp("ext")):
                    if any(v == rId for v in ext.xpath(".//@r:id", namespaces=NSMAP)):
                        ext.getparent().remove(ext)
                slide.part.rels.pop(rId)
                count += 1
    for rId, rel in list(prs.part.rels.items()):
        if "comment" in rel.reltype.lower() or "authors" in rel.reltype.lower():
            prs.part.rels.pop(rId)
    return count


def finish_package(raw: bytes, props: dict) -> bytes:
    """文書のプロパティを空にする・与える。プレビュー画像とカスタムプロパティを取り除く（scrub のとき）。"""
    scrub = bool(props.get("scrub"))
    given = {k: v for k, v in props.items() if k != "scrub" and k != "keep"}
    keep = set(props.get("keep") or [])
    if not scrub and not given:
        return raw
    src = zipfile.ZipFile(io.BytesIO(raw))
    drop: set[str] = set()
    if scrub:
        drop = {n for n in src.namelist() if n.startswith("docProps/thumbnail") or n == "docProps/custom.xml"}
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            if info.filename in drop:
                continue
            data = src.read(info.filename)
            if info.filename in ("docProps/core.xml", "docProps/app.xml"):
                root = etree.fromstring(data)
                fields = CORE_FIELDS if info.filename.endswith("core.xml") else APP_FIELDS
                for name, tag in fields.items():
                    el = root.find(tag)
                    if name in given:
                        if el is None:
                            el = etree.SubElement(root, tag)
                        el.text = str(given[name])
                    elif scrub and name not in keep and el is not None:
                        el.text = ""
                if scrub and info.filename.endswith("app.xml"):   # スライドの題の一覧と数（前の文書のまま）
                    for tag in APP_STALE:
                        for el in root.findall(f"{{{NS_APP}}}{tag}"):
                            root.remove(el)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif drop and info.filename == "_rels/.rels":
                root = etree.fromstring(data)
                for rel in list(root):
                    if rel.get("Target", "").lstrip("/") in drop:
                        root.remove(rel)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif drop and info.filename == "[Content_Types].xml":
                root = etree.fromstring(data)
                for o in list(root):
                    if o.get("PartName", "").lstrip("/") in drop:
                        root.remove(o)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dst.writestr(info, data)
    return out.getvalue()


# ---------------------------------------------------------------------------
# inspect（判断用の事実）
# ---------------------------------------------------------------------------

def _cm(emu: float) -> str:
    return f"{emu / EMU_PER_CM:.1f}"


def _shape_fact(view: SlideView, el, depth: int = 0) -> dict:
    x, y, w, h = view.box(el)
    fact: dict[str, Any] = {"id": shape_id(el), "name": shape_name(el), "kind": kind_of(el),
                            "box_cm": [round(v / EMU_PER_CM, 2) for v in (x, y, w, h)]}
    if is_connector(el) and arrow_ends(el):
        fact["arrow"] = arrow_ends(el)
    tbl = table_of(el)
    if tbl is not None:
        rows = []
        for tr in tbl.findall(qa("tr")):
            rows.append([body_text(tc.find(qa("txBody"))) for tc in tr.findall(qa("tc"))])
        fact["table"] = rows
        if any(PLACEHOLDER_RE.search(c) for r in rows for c in r):
            fact["placeholder_suspect"] = True
    body = txbody(el) if local(el) == "sp" else None
    if body is not None:
        paras = paragraphs(body)
        if any(t.strip() for t, _ in paras) or placeholder(el):
            fact["paragraphs"] = [{"text": t, "level": lvl} for t, lvl in paras]
            cap = view.capacity(el)
            fact["fits"] = cap.describe()
            fact["lines_used"] = cap.lines_for([(t, lvl) for t, lvl in paras if t])
        text = body_text(body)
        if PLACEHOLDER_RE.search(text):
            fact["placeholder_suspect"] = True
        if NOTE_RE.search(text):
            fact["note_suspect"] = True
    if local(el) == "grpSp":
        fact["children"] = [_shape_fact(view, c, depth + 1) for c in child_shapes(el)]
    return fact


def inspect_template(template: "str | bytes") -> dict:
    raw = read_bytes(template)
    prs = open_prs(raw)
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        view = SlideView(prs, slide, i)
        sl = {"number": i, "layout": slide.slide_layout.name, "hidden": slide._element.get("show") == "0",
              "shapes": [_shape_fact(view, el) for el in view.top()]}
        diagrams = find_diagrams(view)
        if diagrams:
            sl["diagram_candidates"] = [{"axis": d["axis"], "items": [[shape_name(e) for e in it] for it in d["items"]],
                                         "connectors": [shape_name(e) for e in d["connectors"]]} for d in diagrams]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            sl["notes"] = slide.notes_slide.notes_text_frame.text
        slides.append(sl)
    return {"slide_size_cm": [round(prs.slide_width / EMU_PER_CM, 2), round(prs.slide_height / EMU_PER_CM, 2)],
            "slides": slides, "provenance": provenance(raw)}


def format_facts(facts: dict) -> str:
    w, h = facts["slide_size_cm"]
    lines = [f"スライドの大きさ: {w} × {h} cm"]

    def shape_lines(f: dict, indent: str) -> None:
        x, y, bw, bh = f["box_cm"]
        flags = "".join([" ?仮値" if f.get("placeholder_suspect") else "", " ※注記" if f.get("note_suspect") else ""])
        head = f"{indent}[#{f['id']}] {f['name']}  {f['kind']}{'（' + f['arrow'] + '）' if f.get('arrow') else ''}" \
               f"  位置 {x},{y} 大きさ {bw}×{bh} cm{flags}"
        lines.append(head)
        if "fits" in f:
            lines.append(f"{indent}    収まる量: {f['fits']}（いま {f['lines_used']} 行）")
        for p in f.get("paragraphs", []):
            if p["text"]:
                lines.append(f"{indent}    {'  ' * p['level']}「{_short(p['text'], 60)}」")
        for r, row in enumerate(f.get("table", []), start=1):
            lines.append(f"{indent}    行{r}: " + " | ".join(_short(c, 16) for c in row))
        for c in f.get("children", []):
            shape_lines(c, indent + "    ")

    for sl in facts["slides"]:
        lines.append("")
        lines.append(f"== スライド {sl['number']}（レイアウト: {sl['layout']}）{' 非表示' if sl['hidden'] else ''}")
        for f in sl["shapes"]:
            shape_lines(f, "  ")
        for d in sl.get("diagram_candidates", []):
            items = " / ".join("+".join(it) for it in d["items"])
            conn = f"、つなぎ: {', '.join(d['connectors'])}" if d["connectors"] else ""
            lines.append(f"  図の並び（{ {'x': '横', 'y': '縦', 'grid': '格子'}[d['axis']] }）: {items}{conn}")
        if sl.get("notes"):
            lines.append(f"  ノート: 「{_short(sl['notes'], 60)}」")
    if facts["provenance"]:
        lines.append("")
        lines.append("来歴・持ち越しの注意:")
        lines.extend(f"  - {p}" for p in facts["provenance"])
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 図の並び（同じ図形が等間隔に並ぶもの）を見つける
# ---------------------------------------------------------------------------

def _center(b) -> tuple[float, float]:
    return b[0] + b[2] / 2, b[1] + b[3] / 2


def _union(bs: list) -> tuple[int, int, int, int]:
    x, y = min(b[0] for b in bs), min(b[1] for b in bs)
    return x, y, max(b[0] + b[2] for b in bs) - x, max(b[1] + b[3] for b in bs) - y


def _line_run(members: list, boxes: dict, view: SlideView) -> "dict | None":
    """同じ大きさの図形が、横・縦に等間隔、または格子に並ぶなら、その並び。"""
    tol_x, tol_y = view.width * 0.02, view.height * 0.02
    for axis, main, cross, tol in (("x", 0, 1, tol_y), ("y", 1, 0, tol_x)):
        ordered = sorted(members, key=lambda e: _center(boxes[e])[main])
        cs = [_center(boxes[e]) for e in ordered]
        if max(c[cross] for c in cs) - min(c[cross] for c in cs) > tol:
            continue
        steps = [b[main] - a[main] for a, b in zip(cs, cs[1:])]
        if min(steps) <= 0 or max(steps) - min(steps) > max(steps) * 0.08:
            continue
        b0, b1 = boxes[ordered[0]], boxes[ordered[1]]
        return {"axis": axis, "members": ordered, "pitch": (b1[0] - b0[0], b1[1] - b0[1])}
    # 格子
    xs = sorted({round(_center(boxes[e])[0] / tol_x) for e in members})
    ys = sorted({round(_center(boxes[e])[1] / tol_y) for e in members})
    if len(xs) >= 2 and len(ys) >= 2 and len(xs) * len(ys) == len(members):
        ordered = sorted(members, key=lambda e: (round(_center(boxes[e])[1] / tol_y), _center(boxes[e])[0]))
        cols = len(xs)
        px = boxes[ordered[1]][0] - boxes[ordered[0]][0]
        py = boxes[ordered[cols]][1] - boxes[ordered[0]][1]
        if px > 0 and py > 0:
            return {"axis": "grid", "members": ordered, "pitch": (px, py), "columns": cols}
    return None


def find_diagrams(view: SlideView) -> list[dict]:
    """スライドの図の並び。items（1 つ分の図形のリスト）を並びの順に、connectors（間をつなぐ矢印など）を間の順に。"""
    cands = [e for e in view.top() if local(e) in ("sp", "grpSp", "pic", "cxnSp") and not placeholder(e)]
    boxes = {e: view.box(e) for e in cands}
    cands = [e for e in cands if boxes[e][2] < view.width * 0.8 and boxes[e][3] < view.height * 0.8]
    items_c = [e for e in cands if not is_connector(e)]
    conns = [e for e in cands if is_connector(e)]
    tol = max(view.width, view.height) * 0.01
    clusters: list[list] = []
    for e in items_c:
        for c in clusters:
            f = c[0]
            if (local(f), prst(f)) == (local(e), prst(e)) and abs(boxes[f][2] - boxes[e][2]) <= tol \
                    and abs(boxes[f][3] - boxes[e][3]) <= tol:
                c.append(e)
                break
        else:
            clusters.append([e])
    runs = [r for r in (_line_run(c, boxes, view) for c in clusters if len(c) >= 2) if r]
    # 主になる並び: 文字枠（テキストボックス）より図形を、小さいものより大きいものを先に
    runs.sort(key=lambda r: (kind_of(r["members"][0]) == "textbox", -(boxes[r["members"][0]][2] * boxes[r["members"][0]][3])))
    used: set = set()
    found = []
    for run in runs:
        if any(m in used for m in run["members"]):
            continue
        n, items = len(run["members"]), [[m] for m in run["members"]]
        used.update(run["members"])
        for other in runs:   # 番号の丸・説明の文字など、同じ間隔で並ぶ別の図形を、1 つ分にまとめる
            if other is run or len(other["members"]) != n or other["axis"] != run["axis"] \
                    or any(m in used for m in other["members"]):
                continue
            if any(abs(a - b) > tol for a, b in zip(other["pitch"], run["pitch"])):
                continue
            offs = [(boxes[o][0] - boxes[m][0], boxes[o][1] - boxes[m][1]) for o, m in zip(other["members"], run["members"])]
            if all(abs(o[0] - offs[0][0]) <= tol and abs(o[1] - offs[0][1]) <= tol for o in offs):
                for it, o in zip(items, other["members"]):
                    it.append(o)
                used.update(other["members"])
        connectors = []
        if run["axis"] != "grid":
            main, cross = (0, 1) if run["axis"] == "x" else (1, 0)
            ibox = [_union([boxes[e] for e in it]) for it in items]
            for a, b in zip(ibox, ibox[1:]):
                lo, hi = a[main] + a[2 + main] - tol, b[main] + tol
                c_lo, c_hi = a[cross], a[cross] + a[2 + cross]
                hits = [c for c in conns if c not in used and lo <= _center(boxes[c])[main] <= hi
                        and c_lo - tol <= _center(boxes[c])[cross] <= c_hi + tol]
                if len(hits) != 1:
                    connectors = []
                    break
                connectors.append(hits[0])
            used.update(connectors)
        if not connectors and any(all(LABEL_RE.match(literal_text(it[k])) for it in items) for k in range(len(items[0]))):
            continue   # `目的：` `効果：` のラベルと値の組は、図の並びではない（ラベルは残し、値をそれぞれ流し込む）
        found.append({"axis": run["axis"], "items": items, "connectors": connectors, "pitch": run["pitch"],
                      "columns": run.get("columns")})
    return found


# ---------------------------------------------------------------------------
# analyze（定義の下書き）
# ---------------------------------------------------------------------------

GENERIC_NAME_RE = re.compile(r"^(テキスト ?ボックス|textbox|text box|正方形/長方形|長方形|rectangle|角丸四角形|"
                             r"rounded rectangle|楕円|oval|ellipse|図形|shape|コンテンツ プレースホルダー|content placeholder|"
                             r"テキスト プレースホルダー|text placeholder|title|タイトル|サブタイトル|subtitle|"
                             r"グループ化|group|表|table)?\s*\d*$", re.I)
PH_KEYS = {"title": "title", "ctrTitle": "title", "subTitle": "subtitle", "body": "body", "obj": "body", "dt": "date"}


def _unique(key: str, used: set) -> str:
    base, n = key, 2
    while key in used:
        key, n = f"{base}{n}", n + 1
    used.add(key)
    return key


def _budget(samples: list[str]) -> "int | None":
    n = max((len(s.replace("\n", "")) for s in samples if s), default=0)
    return max(math.ceil(n * GRANULARITY), n + 4) if n else None


def _line_budget(view: SlideView, el, samples: list[str]) -> "int | None":
    """1 項目の字数の上限。サンプルの 1.5 倍までだが、サンプルが使う行数を超えない（1 行の見出しは 1 行のまま）。"""
    b = _budget(samples)
    if not b:
        return None
    n = max(len(s.replace("\n", "")) for s in samples)
    cpl = int(view.capacity(el).cpl)
    if cpl < 1:
        return b
    return max(n, min(b, cpl * math.ceil(n / cpl)))


def _ref(view: SlideView, el) -> str:
    name = shape_name(el)
    same = [e for e in walk(view.tree) if shape_name(e) == name]
    return name if name and len(same) == 1 else f"#{shape_id(el)}"


def _has_bullets(body) -> bool:
    for p in body.findall(qa("p")):
        ppr = p.find(qa("pPr"))
        if para_level(p) > 0 or (ppr is not None and (ppr.find(qa("buChar")) is not None
                                                       or ppr.find(qa("buAutoNum")) is not None)):
            return True
    return False


def _index_format(samples: list[str]) -> "str | None":
    """1, 2, 3 …（STEP 1・STEP 2 … も）と並ぶなら、その書式（`{}` が番号）。"""
    if len(samples) < 2:
        return None
    fmts, nums = set(), []
    for s in samples:
        m = list(re.finditer(r"\d+", s))
        if len(m) != 1:
            return None
        fmts.add(s[:m[0].start()] + "{}" + s[m[0].end():])
        nums.append(int(m[0].group()))
    return fmts.pop() if len(fmts) == 1 and nums == list(range(nums[0], nums[0] + len(nums))) and nums[0] in (0, 1) else None


def _labels(view: SlideView, shapes: list) -> dict:
    """`目的：` のようなラベルの右か下にある図形 → ラベルの言葉（キーの名前に使う）。"""
    out = {}
    for lab in shapes:
        text = body_text(txbody(lab)) if local(lab) == "sp" else ""
        if not LABEL_RE.match(text):
            continue
        lx, ly, lw, lh = view.box(lab)
        best, dist = None, None
        for e in shapes:
            if e is lab or local(e) != "sp" or LABEL_RE.match(body_text(txbody(e))):
                continue
            x, y, w, h = view.box(e)
            right = abs((y + h / 2) - (ly + lh / 2)) < lh and x >= lx + lw * 0.9
            below = abs(x - lx) < lw and y >= ly + lh * 0.9 and y - (ly + lh) < lh * 1.5
            if right or below:
                d = (x - lx) + (y - ly)
                if dist is None or d < dist:
                    best, dist = e, d
        if best is not None:
            out[best] = text.rstrip(":：").strip()
    return out


def _title_shape(view: SlideView, shapes: list):
    """タイトルのプレースホルダーが無いスライドの、見出しの文字枠（上の方にあり、文字がいちばん大きい）。"""
    if any((placeholder(e) or ("",))[0] in ("title", "ctrTitle") for e in shapes):
        return None
    best, best_size = None, 0.0
    for e in shapes:
        if local(e) != "sp" or placeholder(e) or not literal_text(e) or LABEL_RE.match(literal_text(e)):
            continue
        if view.box(e)[1] > view.height * 0.4:
            continue
        size = view.font_size(e, txbody(e))
        if size > best_size:
            best, best_size = e, size
    return best if best_size >= 20 else None


def _key_for(el, labels: dict, list_kind: bool = False) -> str:
    if el in labels:
        return labels[el]
    ph = placeholder(el)
    if ph and ph[0] in PH_KEYS:
        return "points" if list_kind and PH_KEYS[ph[0]] == "body" else PH_KEYS[ph[0]]
    name = shape_name(el)
    if name and not GENERIC_NAME_RE.match(name):
        return re.sub(r"\s*\d+$", "", name).strip() or "text"
    return "points" if list_kind else "text"


def _analyze_table(view: SlideView, el, used: set, confirm: list) -> dict:
    tbl = table_of(el)
    trs = tbl.findall(qa("tr"))
    grid = [[body_text(tc.find(qa("txBody"))) for tc in tr.findall(qa("tc"))] for tr in trs]
    ref, n_cols = _ref(view, el), len(tbl.findall(f"{qa('tblGrid')}/{qa('gridCol')}"))
    tblpr = tbl.find(qa("tblPr"))
    first_row = tblpr is not None and tblpr.get("firstRow") == "1"
    head_runs = [r for tc in trs[0].findall(qa("tc")) for r in tc.iter(qa("r"))] if trs else []
    bold_head = bool(head_runs) and all(r.find(qa("rPr")) is not None and r.find(qa("rPr")).get("b") == "1"
                                        for r in head_runs)
    where = f"スライド {view.number}「{shape_name(el)}」"
    if n_cols == 2 and not first_row and not bold_head and len(trs) >= 2:
        cells = {}
        for r, row in enumerate(grid, start=1):
            if row and row[0].strip() and len(row) > 1:
                cells[f"{r},2"] = {"key": _unique(row[0].strip().rstrip(":："), used)}
                b = _budget([row[1]])
                if b:
                    cells[f"{r},2"]["max_chars"] = b
        confirm.append(f"{where}: 2 列の表を、左の列を見出しにした記入欄（行は増やさない）と見た。行を繰り返す表なら key と columns に直す")
        return {"shape": ref, "cells": cells}
    header = 1 if len(trs) >= 2 else 0
    footer = 1 if len(trs) - header >= 2 and grid[-1] and any(
        w in (grid[-1][0] or "").lower() for w in ("合計", "小計", "計", "total")) else 0
    body = grid[header:len(grid) - footer]
    columns: list = []
    col_keys: set = set()
    for j in range(n_cols):
        head = grid[0][j].strip() if header and j < len(grid[0]) else ""
        samples = [row[j] for row in body if j < len(row)]
        filled = [s for s in samples if s.strip()]
        if not head and not filled:
            columns.append(None)
            continue
        if HUMAN_RE.search(head):
            columns.append({"clear": True, "header": head})
            continue
        fmt = _index_format(filled) if len(filled) == len(samples) else None
        if fmt is not None:
            spec = {"key": "$index", "header": head}
            if fmt != "{}":
                spec["format"] = fmt
            columns.append(spec)
            continue
        if len(samples) >= 2 and len(set(samples)) == 1 and filled and not PLACEHOLDER_RE.search(samples[0]):
            columns.append({"keep": True, "header": head})
            continue
        if filled and all(s.strip() in MARK_ON + MARK_OFF for s in filled):
            on = next((s.strip() for s in filled if s.strip() in MARK_ON), "○")
            off = next((s.strip() for s in filled if s.strip() in MARK_OFF), "")
            spec = {"key": _unique(head or f"col{j + 1}", col_keys), "map": {True: on, False: off or None}, "header": head}
            confirm.append(f"{where}: 列「{head}」は ○ の印の列。データは true / false で書く")
            columns.append(spec)
            continue
        spec = {"key": _unique(head or f"col{j + 1}", col_keys), "header": head}
        b = _budget(filled)
        if b:
            spec["max_chars"] = b
        columns.append(spec)
    x, y, w, h = view.box(el)
    row_h = [int(tr.get("h", "0")) for tr in trs]
    sample_h = row_h[header:len(trs) - footer] or [370840]
    avail = view.free_end(el, (x, y, w, h), "y") - y - sum(row_h[:header]) - sum(row_h[len(trs) - footer:])
    cap = max(int(avail / max(sum(sample_h) / len(sample_h), 1)), len(body))
    out = {"shape": ref, "key": _unique("rows", used), "header_rows": header, "footer_rows": footer,
           "columns": columns, "max_rows": cap}
    # 合計の行: 見出し（合計）は残し、値はデータの「合計.件数」のような欄から入れる
    cells = {}
    for r in range(len(grid) - footer, len(grid)):
        base = None
        for j, text in enumerate(grid[r]):
            spec = columns[j] if j < len(columns) else None
            if j == 0 or not text.strip() or not spec or not spec.get("key") or spec.get("keep") or spec["key"] == "$index":
                continue
            base = base or _unique((grid[r][0] or "").strip() or "合計", used)
            cells[f"{r + 1},{j + 1}"] = {"key": f"{base}.{spec['key']}"}
            b = _budget([text])
            if b:
                cells[f"{r + 1},{j + 1}"]["max_chars"] = b
    if cells:
        out["cells"] = cells
        confirm.append(f"{where}: 最後の行（{grid[-1][0].strip()}）の値は、データの `{base}` から入れる"
                       "（前の値を残さない。計算はしないので、データに書く）")
    return out


def _diagram_capacity(view: SlideView, d: dict) -> int:
    items = d["items"]
    flat = [e for it in items for e in it] + d["connectors"]
    boxes = [view.box(e) for e in flat]
    gx, gy = min(b[0] for b in boxes), min(b[1] for b in boxes)
    gw, gh = max(b[0] + b[2] for b in boxes) - gx, max(b[1] + b[3] for b in boxes) - gy
    group = (gx, gy, gw, gh)
    ib = [view.box(e) for e in items[0]]
    iw = max(b[0] + b[2] for b in ib) - min(b[0] for b in ib)
    ih = max(b[1] + b[3] for b in ib) - min(b[1] for b in ib)
    n = len(items)
    px, py = d["pitch"]
    if d["axis"] == "grid":
        rows = math.ceil(n / d["columns"])
        bottom = min(view.free_end(e, group, "y", exclude=flat) for e in flat[:1])
        fit_rows = int((bottom - gy - ih) / py) + 1 if py else rows
        return max(n, d["columns"] * max(fit_rows, rows))
    axis = d["axis"]
    pitch = px if axis == "x" else py
    start, size, length = (gx, iw, gw) if axis == "x" else (gy, ih, gh)
    lo = view.free_start(flat[0], group, axis, exclude=flat)
    hi = view.free_end(flat[0], group, axis, exclude=flat)
    if d.get("align") == "center":
        center = start + length / 2
        avail = 2 * min(center - lo, hi - center)
    else:
        avail = hi - start
    return max(n, int((avail - size) / pitch) + 1 if pitch > 0 else n)


def _diagram_align(view: SlideView, d: dict) -> str:
    if d["axis"] == "grid":
        return "start"
    boxes = [view.box(e) for it in d["items"] for e in it]
    main = 0 if d["axis"] == "x" else 1
    lo = min(b[main] for b in boxes)
    hi = max(b[main] + b[2 + main] for b in boxes)
    size = view.width if main == 0 else view.height
    return "center" if abs((lo + hi) / 2 - size / 2) < size * 0.03 else "start"


def _item_paths(item: list) -> list[tuple[Any, tuple]]:
    """1 つ分の図形の中の、文字を持つ図形と、その道筋（図形の番号, グループの中の番号…）。"""
    out = []

    def rec(el, path):
        if local(el) == "grpSp":
            for k, c in enumerate(child_shapes(el)):
                rec(c, path + (k,))
        elif txbody(el) is not None:
            out.append((el, path))

    for j, el in enumerate(item):
        rec(el, (j,))
    return out


def _sorted_paths(view: SlideView, item: list) -> list[tuple[Any, tuple]]:
    """文字を持つ図形を、上から下・左から右の順に。"""
    return sorted(_item_paths(item), key=lambda ep: (round(view.box(ep[0])[1] / EMU_PER_PT / 6), view.box(ep[0])[0]))


def follow(item: list, path: tuple):
    el = item[path[0]]
    for k in path[1:]:
        el = child_shapes(el)[k]
    return el


def _analyze_diagram(view: SlideView, d: dict, n: int, used: set, confirm: list) -> dict:
    items = d["items"]
    d["align"] = _diagram_align(view, d)
    fields: dict = {}
    names = iter(("label", "text", "text3", "text4", "text5", "text6"))
    for el, path in _sorted_paths(view, items[0]):
        samples = [body_text(txbody(follow(it, path))) for it in items]
        if not any(s.strip() for s in samples):
            continue   # 文字の無い図形（飾り）
        fmt = _index_format(samples)
        if fmt is not None:
            fields[_ref(view, el)] = {"key": "$index", **({"format": fmt} if fmt != "{}" else {})}
            continue
        if len(items) >= 2 and len(set(samples)) == 1 and not PLACEHOLDER_RE.search(samples[0] or ""):
            continue   # どの図形も同じ文字（図の一部）。残す
        spec: dict = {"key": next(names)}
        b = _budget(samples)
        if b:
            spec["max_chars"] = b
        fields[_ref(view, el)] = spec
    did = _unique(f"flow{n}" if d["connectors"] else f"items{n}", used)
    out = {"id": did, "key": did, "axis": d["axis"],
           "items": [[_ref(view, e) for e in it] for it in items],
           "connectors": [_ref(view, e) for e in d["connectors"]],
           "fields": fields, "align": d["align"], "max_items": _diagram_capacity(view, d)}
    if d["axis"] == "grid":
        out["columns"] = d["columns"]
    shape = {"x": "横", "y": "縦", "grid": "格子"}[d["axis"]]
    conn = f"・つなぎ {len(d['connectors'])} 本" if d["connectors"] else ""
    confirm.append(f"スライド {view.number}: 図の並び（{shape}に {len(items)} つ{conn}）を `{did}` として繰り返す。"
                   f"最大 {out['max_items']} つ。1 つずつの図形・矢印はテンプレートのものを複製する")
    return out


def _slide_signature(view: SlideView) -> tuple:
    names = sorted(shape_name(e) for e in walk(view.tree) if literal_text(e) or table_of(e) is not None)
    return view.slide.slide_layout.name, tuple(names)


def analyze(template: "str | bytes") -> dict:
    raw = read_bytes(template)
    prs = open_prs(raw)
    confirm: list[str] = []
    slides: list[dict] = []
    views = [SlideView(prs, s, i) for i, s in enumerate(prs.slides, start=1)]
    sections = slide_sections(prs)
    sigs = [(sections.get(v.number),) + _slide_signature(v) for v in views]   # 繰り返しはセクションをまたがない
    doc = DocValues()   # 文書の値（年度・作成日・宛先など）の名前は、資料全体で 1 つ
    i = 0
    while i < len(views):
        j = i
        while j + 1 < len(views) and sigs[j + 1] == sigs[i] and sigs[i][2]:
            j += 1
        sd = _analyze_slide(views[i], confirm, doc)
        if j > i:
            sd["repeat"] = True
            for k in range(i + 1, j + 1):
                _merge_budgets(sd, _analyze_slide(views[k], []))
            _keep_constant_texts(sd, views[i:j + 1])
            confirm.append(f"スライド {i + 1}〜{j + 1} は同じ形。スライド {i + 1} を見本に、データの件数だけ繰り返す"
                           f"（`{sd['key']}` は配列）。スライド {i + 2}〜{j + 1} はサンプルとして取り除く")
        slides.append(sd)
        for k in range(i + 1, j + 1):
            slides.append({"id": f"s{k + 1}", "slide": k + 1, "drop": True})
        i = j + 1
    for v in views:
        if v.slide._element.get("show") == "0":
            confirm.append(f"スライド {v.number}: 非表示のスライド。残すか取り除く（drop）かを決める")
        if v.slide.has_notes_slide and v.slide.notes_slide.notes_text_frame is not None \
                and v.slide.notes_slide.notes_text_frame.text.strip():
            confirm.append(f"スライド {v.number}: ノートがある。properties.scrub で消える。残すなら notes にキーを書く")
    for sd in slides:
        name = sections.get(sd["slide"])
        if name and not sd.get("drop"):
            sd["group"] = name
    if sections:
        confirm.append("データを分けるときは、PowerPoint のセクションごとに 1 つのファイルにまとめる（group）。"
                       "セクションで分かれない意味のまとまりは、group を書き直す")
    return {"version": DEF_VERSION, "template": "", "strict": True, "properties": {"scrub": True},
            "slides": slides, "needs_confirm": confirm}


def slide_sections(prs) -> dict[int, str]:
    """PowerPoint のセクション（章）。スライドの番号 → セクションの名前。"""
    ids = {int(sid.get("id")): n for n, sid in enumerate(prs.slides._sldIdLst, start=1)}
    out: dict[int, str] = {}
    for sec in prs.part._element.iter(f"{{{NS_P14}}}section"):
        for sid in sec.iter(f"{{{NS_P14}}}sldId"):
            n = ids.get(int(sid.get("id", "0")))
            if n:
                out[n] = sec.get("name", "")
    return out


def _keep_constant_texts(sd: dict, views: list) -> None:
    """繰り返すスライドのどのサンプルでも同じ文字（「主な取り組み」の見出しなど）は、流し込まずに残す。"""
    for ref in list(sd.get("texts") or {}):
        try:
            texts = {literal_text(v.find(ref)) for v in views}
        except TemplateError:
            continue
        if len(texts) == 1 and texts != {""}:
            del sd["texts"][ref]
            sd.setdefault("keep", []).append(ref)
    if "texts" in sd and not sd["texts"]:
        del sd["texts"]


def _merge_budgets(base: dict, other: dict) -> None:
    """繰り返すスライドのサンプルが複数あれば、粒度（max_chars・max_items）は大きい方に合わせる。"""
    def merge(a, b):
        if isinstance(a, dict) and isinstance(b, dict):
            for k, v in b.items():
                if k in ("max_chars", "max_items") and isinstance(v, int) and isinstance(a.get(k), int):
                    a[k] = max(a[k], v)
                elif k in a:
                    merge(a[k], v)
        elif isinstance(a, list) and isinstance(b, list):
            for x, y in zip(a, b):
                merge(x, y)
    merge(base, other)


def _analyze_slide(view: SlideView, confirm: list, doc: "DocValues | None" = None) -> dict:
    n = view.number
    doc = doc if doc is not None else DocValues()
    sd: dict[str, Any] = {"id": f"s{n}", "slide": n, "key": f"s{n}"}
    texts, lists, tables, diagrams, keep, clear = {}, {}, [], [], [], []
    used: set = set()
    claimed: set = set()
    for k, d in enumerate(find_diagrams(view), start=1):
        diagrams.append(_analyze_diagram(view, d, k, used, confirm))
        claimed.update(e for it in d["items"] for e in it)
        claimed.update(d["connectors"])
    flat = [e for e in walk(view.tree) if e not in claimed and not any(a in claimed for a in e.iterancestors())]
    labels = _labels(view, [e for e in flat if local(e) == "sp"])
    title = _title_shape(view, flat)
    if title is not None and title not in labels:
        labels[title] = "title"
    for el in flat:
        if el in claimed:
            continue
        where = f"スライド {n}「{shape_name(el)}」"
        if table_of(el) is not None:
            tables.append(_analyze_table(view, el, used, confirm))
            continue
        if local(el) != "sp" or txbody(el) is None:
            continue
        ph = placeholder(el)
        text = literal_text(el)
        paras = [(t, lvl) for t, lvl in paragraphs(txbody(el))]
        if ph and ph[0] in ("sldNum", "ftr", "hdr"):
            got = variable_parts(text) if text and ph[0] != "sldNum" else None
            if got:   # フッターの表題・年度。可変の部分だけを文書の値にする
                texts[_ref(view, el)] = {"text": doc.template(*got)}
                confirm.append(f"{where}: フッターの「{_short(text, 24)}」は、可変の部分だけを文書の値として流し込む"
                               f"（{texts[_ref(view, el)]['text']}）。毎回同じ文字なら keep に移す")
            elif text:
                keep.append(_ref(view, el))
            continue
        if ph and ph[0] == "dt" and not text:
            continue
        if not text:
            if ph and ph[0] in PH_KEYS:   # 空の記入枠もキーを残す
                kind = ph[0] in ("body", "obj")
                (lists if kind else texts)[_ref(view, el)] = {"key": _unique(_key_for(el, labels, kind), used)}
            continue
        if LABEL_RE.match(text) or (el not in labels and NOTE_RE.search(text) and not PLACEHOLDER_RE.search(text)):
            keep.append(_ref(view, el))
            continue
        if NOTE_RE.search(text) and PLACEHOLDER_RE.search(text):
            clear.append(_ref(view, el))
            confirm.append(f"{where}: 書き方の注記（{_short(text, 20)}）と見て clear にした")
            continue
        filled = [t for t, _ in paras if t.strip()]
        if len(filled) >= 2 and (_has_bullets(txbody(el)) or (ph and ph[0] in ("body", "obj"))):
            spec = {"key": _unique(_key_for(el, labels, True), used), "max_items": len(filled)}
            b = _line_budget(view, el, filled)
            if b:
                spec["max_chars"] = b
            lists[_ref(view, el)] = spec
            continue
        if len(filled) >= 2 and all(re.match(r"^[・●■◆□◇○\-－*]\s*", t) for t in filled):
            confirm.append(f"{where}: 行頭の「・」を文字で打った箇条書き。1 つの文字の欄にした。"
                           "項目の数で収めるなら、段落の行頭記号に直したテンプレートで lists にする")
        label = labels.get(el)
        got = variable_parts(text) if len(filled) <= 1 else None
        if label and label != "title" and DOC_LABEL_RE.match(label):
            # 作成日・作成者・版など、文書全体の値。どのスライドでも同じキーにし、データの「文書」にまとめる
            key = doc.key(label, text)
            whole = got and len(got[1]) == 1 and got[0].startswith("{" + got[1][0][0] + ":") and got[0].endswith("}")
            spec = {"text": "{" + key + got[0][len(got[1][0][0]) + 1:]} if whole else {"key": key}
        elif got:
            # タイトル・宛名の中の、年度・期間・日付・宛先。文の残りはそのまま、可変の部分だけを流し込む
            spec = {"text": doc.template(*got)}
            confirm.append(f"{where}: 「{_short(text, 24)}」は、可変の部分だけを文書の値として流し込む（{spec['text']}）。"
                           "毎回同じ文字なら keep に移す")
        else:
            spec = {"key": _unique(_key_for(el, labels), used)}
        b = _line_budget(view, el, [text]) if len(filled) <= 1 else _budget([text])
        if b:
            spec["max_chars"] = b
        texts[_ref(view, el)] = spec
    for name, val in (("texts", texts), ("lists", lists), ("tables", tables), ("diagrams", diagrams),
                      ("keep", keep), ("clear", clear)):
        if val:
            sd[name] = val
    return sd


def summarize(definition: dict) -> str:
    lines = []
    for sd in definition.get("slides", []):
        head = f"スライド {sd['slide']}"
        if sd.get("drop"):
            lines.append(f"{head}: 取り除く（サンプル）")
            continue
        lines.append(f"{head}: データ `{sd.get('key', sd['id'])}`{'（配列。1 件 1 枚）' if sd.get('repeat') else ''}")
        for ref, spec in (sd.get("texts") or {}).items():
            spec = _spec(spec)
            lines.append(f"  文字 {ref} → " + (f"文 {spec['text']!r}" if "text" in spec else str(spec.get("key"))) + _limit(spec))
        for ref, spec in (sd.get("lists") or {}).items():
            spec = _spec(spec)
            lines.append(f"  箇条書き {ref} → {spec.get('key')}[]{_limit(spec)}")
        for t in sd.get("tables") or []:
            if t.get("key"):
                cols = ", ".join(c.get("key", "残す" if c.get("keep") else "空欄") if c else "-" for c in t["columns"])
                lines.append(f"  表 {t['shape']} → {t['key']}[]（列: {cols}。最大 {t.get('max_rows', '?')} 行）")
            else:
                lines.append(f"  表 {t['shape']} の記入欄 → {', '.join(_spec(s)['key'] for s in (t.get('cells') or {}).values())}")
        for d in sd.get("diagrams") or []:
            fields = ", ".join(_spec(s)["key"] for s in d.get("fields", {}).values())
            lines.append(f"  図 {d['id']} → {d['key']}[]（{len(d['items'])} つの見本・最大 {d.get('max_items', '?')} つ。欄: {fields}）")
        if sd.get("keep"):
            lines.append(f"  残す: {', '.join(sd['keep'])}")
        if sd.get("clear"):
            lines.append(f"  空にする: {', '.join(sd['clear'])}")
    if definition.get("needs_confirm"):
        lines.append("")
        lines.append("確認すること:")
        lines.extend(f"  ? {c}" for c in definition["needs_confirm"])
    return "\n".join(lines)



# ---------------------------------------------------------------------------
# 文書の値（表紙・タイトルの年度・期間・宛名・作成日など）。データでもラベルでもなく、文書ごとに変わる値
# ---------------------------------------------------------------------------

DOC_PREFIX = "文書"
# ラベルがこれなら、その右の値は表の外の 1 件の値ではなく、文書全体の値（どのタブでも同じ値が入る）
DOC_LABEL_RE = re.compile(r"^(作成日|作成者|作成部署|発行日|発行者|提出日|提出先|報告日|報告者|更新日|改訂日|日付|"
                          r"文書番号|文書名|資料番号|管理番号|版|版数|バージョン|Ver\.?|宛先|宛名|件名|表題|タイトル|"
                          r"プロジェクト名?|案件名|システム名|対象期間|期間|報告期間|年度)$", re.I)


def _date_fmt(text: str, m: "re.Match", codes: dict) -> str:
    """一致した日付の文字から、同じ見た目に戻す書式（`2025/07/01` → `%Y/%m/%d`、`7月` → `%-m月`）を作る。"""
    # 月・日のどれかが 0 埋め（07）なら 0 埋めの書式、どれも 1 桁か 10 以上なら 0 を付けない書式
    padded = any(m.group(g).startswith("0") for g, code in codes.items() if code != "Y")
    out, pos = "", m.start()
    for g, code in codes.items():
        a, b = m.span(g)
        out += text[pos:a].replace("%", "%%") + (f"%{code}" if code == "Y" or padded else f"%-{code}")
        pos = b
    return out + text[pos:m.end()].replace("%", "%%")


# 文の中の可変の部分。(名前, 正規表現, 種類)。上から順に、重ならないものを拾う
VAR_PATTERNS = [
    ("日付", re.compile(r"(?<!\d)(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日"), "date"),
    ("日付", re.compile(r"(?<![\d./-])(\d{4})([/.\-])(\d{1,2})\2(\d{1,2})(?![\d./-])"), "date_sep"),
    ("年月", re.compile(r"(?<!\d)(\d{4})\s*年\s*(\d{1,2})\s*月(?!\s*\d)"), "ym"),
    ("年月", re.compile(r"(?<![\d./-])(\d{4})/(\d{1,2})(?![\d./-])"), "ym_sep"),
    ("年度", re.compile(r"(?<!\d)(\d{4})\s*年度"), "int1"),
    ("和暦年度", re.compile(r"(?:令和|平成)\s*(\d{1,2}|元)\s*年度"), "wareki"),
    ("和暦年", re.compile(r"(?:令和|平成)\s*(\d{1,2}|元)\s*年(?!度)"), "wareki"),
    ("年", re.compile(r"(?<!\d)(\d{4})\s*年(?![度\d])"), "int1"),
    ("四半期", re.compile(r"第\s*([1-4])\s*四半期"), "int1"),
    ("四半期", re.compile(r"(?<![A-Za-z0-9])Q([1-4])(?![A-Za-z0-9])"), "int1"),
    ("四半期", re.compile(r"(?<![A-Za-z0-9])([1-4])Q(?![A-Za-z0-9])"), "int1"),
    ("半期", re.compile(r"(上|下)半?期"), "str1"),
    ("月", re.compile(r"(?<![\d/.\-])(\d{1,2})\s*月(?=度|分|[\s）)]|$)"), "int1"),
    ("版", re.compile(r"第\s*(\d+(?:\.\d+)*)\s*版"), "str1"),
    ("版", re.compile(r"(?<![A-Za-z])(?:Ver\.?|ver\.?|[vV])\s?(\d+(?:\.\d+)+)"), "str1"),
]
VAR_ADDRESSEE_RE = re.compile(r"^(?P<name>\S.*?)\s*(?P<tail>御中|様|殿)(?=\s|$|向け)")   # 「北斗製薬様 定例報告」の文頭の宛名も
NOT_ADDRESSEE = {"お客", "皆", "各位", "関係者各位"}
LABELLED_RE = re.compile(r"^(?P<label>[^:：\d]{1,15})\s*[:：]\s*")


def variable_parts(text: str) -> "tuple[str, list[tuple[str, Any]]] | None":
    """文の中の、文書ごとに変わりそうな部分（年度・四半期・年月・日付・版・宛名）を、欄（{名前}）にしたひな形にする。

    返すのは (ひな形, [(欄の名前, 今の値), …])。可変の部分が無ければ None。欄の名前は、まだ `文書.` を付けない仮の名前。
    注記（※）や長い文は対象にしない。
    """
    if not isinstance(text, str) or not text.strip() or len(text) > 80 or NOTE_RE.match(text) or PLACEHOLDER_RE.search(text):
        return None
    hits: list[tuple[int, int, str, str, Any]] = []   # (先頭, 末尾, 名前, 欄の中身, 値)
    taken: list[tuple[int, int]] = []
    m = VAR_ADDRESSEE_RE.match(text)
    if m and len(m.group("name").strip()) >= 2 and m.group("name").strip() not in NOT_ADDRESSEE \
            and not DOC_LABEL_RE.match(m.group("name").strip()) and not any(rx.search(m.group("name")) for _, rx, _ in VAR_PATTERNS):
        hits.append((m.start("name"), m.end("name"), "宛先", "", m.group("name").strip()))
        taken.append(m.span())
    for name, rx, kind in VAR_PATTERNS:
        for m in rx.finditer(text):
            if any(a < m.end() and m.start() < b for a, b in taken):
                continue
            if kind == "date":
                fmt = _date_fmt(text, m, {1: "Y", 2: "m", 3: "d"})
                value = f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
                span = m.span()
            elif kind == "date_sep":
                fmt = _date_fmt(text, m, {1: "Y", 3: "m", 4: "d"})
                value = f"{int(m.group(1)):04d}-{int(m.group(3)):02d}-{int(m.group(4)):02d}"
                span = m.span()
            elif kind in ("ym", "ym_sep"):
                fmt = _date_fmt(text, m, {1: "Y", 2: "m"})
                value = f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"
                span = m.span()
            else:
                fmt = ""
                g = m.group(1)
                value = 1 if g == "元" else int(g) if kind in ("int1", "wareki") else g
                span = m.span(1)
            taken.append(m.span())
            hits.append((span[0], span[1], name, fmt, value))
    if not hits:
        return None
    hits.sort()
    lab = LABELLED_RE.match(text)
    dates = [h for h in hits if h[2] == "日付"]
    names: list[str] = []
    for h in hits:
        name = h[2]
        if lab and len(hits) == len(dates) and dates:
            label = lab.group("label").strip()
            name = label if len(dates) == 1 else (f"{label}.開始" if h is dates[0] else f"{label}.終了") if len(dates) == 2 \
                else label
        names.append(name)
    # 同じ名前が 2 つ以上なら 2, 3 … を付ける（同じ文の中で、違う値を同じ欄にしない）
    seen: dict[str, int] = {}
    out, fields, pos = "", [], 0
    for (a, b, _, fmt, value), name in zip(hits, names):
        seen[name] = seen.get(name, 0) + 1
        if seen[name] > 1:
            name = f"{name}{seen[name]}"
        out += text[pos:a].replace("{", "{{").replace("}", "}}") + "{" + name + (f":{fmt}" if fmt else "") + "}"
        fields.append((name, value))
        pos = b
    out += text[pos:].replace("{", "{{").replace("}", "}}")
    return out, fields


class DocValues:
    """ブック全体の文書の値の名前。同じ名前で同じ値なら、同じキーにする（表紙とヘッダーの年度に、1 つの値が入る）。"""

    def __init__(self) -> None:
        self.values: dict[str, Any] = {}

    def key(self, name: str, value: Any) -> str:
        base = re.sub(r"[\s]+", "_", name.strip(" :：")) or "値"
        cand, n = base, 2
        while cand in self.values and self.values[cand] != value and value is not None and self.values[cand] is not None:
            cand, n = f"{base}{n}", n + 1
        if self.values.get(cand) is None:
            self.values[cand] = value
        return f"{DOC_PREFIX}.{cand}"

    def template(self, tpl: str, fields: list[tuple[str, Any]]) -> str:
        """variable_parts の仮の名前を、文書の値のキーに置き換える。"""
        for name, value in fields:
            key = self.key(name, value)
            tpl = re.sub(r"\{" + re.escape(name) + r"(?=[:}])", "{" + key, tpl, count=1)
        return tpl




# ---------------------------------------------------------------------------
# 文の中の一部を差し替える（`{文書.年度}年度 第{文書.四半期}四半期 売上報告書`）
# ---------------------------------------------------------------------------

TEXT_FIELD_RE = re.compile(r"\{\{|\}\}|\{([^{}:]+)(?::([^{}]*))?\}")


def text_fields(tpl: str) -> list[tuple[str, str]]:
    """文のひな形の中の欄 [(データのキー, 書式), …]。`{{`・`}}` は文字の波かっこ。"""
    out, rest, pos = [], "", 0
    for m in TEXT_FIELD_RE.finditer(str(tpl)):
        rest += str(tpl)[pos:m.start()]
        pos = m.end()
        if m.group(1) is not None:
            out.append((m.group(1).strip(), m.group(2) or ""))
    rest += str(tpl)[pos:]
    if "{" in rest or "}" in rest:
        raise TemplateError(f"文のひな形 {tpl!r} の波かっこが閉じていません（文字の波かっこは {{{{ と }}}} と書く）")
    return out


def spec_keys(spec: Any) -> list[str]:
    """固定セルの指定が使う、データのキーの一覧（文のひな形なら、その中の欄のキー）。"""
    if not isinstance(spec, dict):
        return [spec]
    if "text" in spec:
        return list(dict.fromkeys(k for k, _ in text_fields(spec["text"])))
    return [spec["key"]]


def _as_date(value: Any, where: str) -> "dt.date | dt.datetime":
    if isinstance(value, (dt.date, dt.datetime)):
        return value
    s = str(value).strip().replace("/", "-")
    m = re.fullmatch(r"(\d{4})-(\d{1,2})", s)   # 年月だけ（2026-10）
    if m:
        return dt.date(int(m.group(1)), int(m.group(2)), 1)
    try:
        return dt.datetime.fromisoformat(s) if (" " in s or "T" in s) else dt.date.fromisoformat(
            "-".join(p.zfill(2) for p in s.split("-")))
    except ValueError:
        raise TemplateError(f"{where} の値 {value!r} を日付として読めません（2026-10-08 か、年月なら 2026-10 の形で書く）")


def _strftime(d: "dt.date | dt.datetime", fmt: str) -> str:
    """strftime と同じ。ただし `%-m`・`%-d`・`%-H`（0 を付けない）は、どの OS でも使える。"""
    def sub(m: "re.Match") -> str:
        code = m.group(1)
        if code.startswith("-"):
            return str(int(d.strftime("%" + code[1:])))
        return "%" if code == "%" else d.strftime("%" + code)
    return re.sub(r"%(-?[A-Za-z%])", sub, fmt)


def _field_text(value: Any, fmt: str, where: str) -> str:
    if value is None:
        return ""
    if "%" in fmt:
        return _strftime(_as_date(value, where), fmt)
    if fmt:
        try:
            return format(value, fmt)
        except (ValueError, TypeError):
            raise TemplateError(f"{where} の値 {value!r} に書式 {fmt!r} を使えません")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    return str(value)


def fill_text(tpl: str, data: Any, where: str, escape=None) -> str:
    """文のひな形の欄を、データの値で埋める。値が null の欄は空にする。

    data はデータか、キーから値を引く関数（無ければ KeyError）。
    """
    lookup = data if callable(data) else (lambda key: _lookup(data, key))
    text_fields(tpl)   # 閉じていない波かっこを先に止める

    def sub(m: "re.Match") -> str:
        if m.group(1) is None:
            return m.group(0)[0]
        key = m.group(1).strip()
        try:
            value = lookup(key)
        except KeyError:
            raise TemplateError(f"データに {key!r} がありません（{where}）")
        text = _field_text(value, m.group(2) or "", f"{where} の {key}")
        return escape(text) if escape else text
    return TEXT_FIELD_RE.sub(sub, str(tpl))


def read_text(tpl: str, text: Any) -> "dict | None":
    """fill_text の逆。文から欄の値を読み戻す（合わなければ None）。数字だけは数値、日付の書式は 2026-10-08 の形にする。"""
    if text is None:
        return None
    pattern, names, n = "", {}, 0
    pos = 0
    for m in TEXT_FIELD_RE.finditer(str(tpl)):
        pattern += re.escape(tpl[pos:m.start()])
        pos = m.end()
        if m.group(1) is None:
            pattern += re.escape(m.group(0)[0])
            continue
        key = m.group(1).strip()
        if key in names:
            pattern += f"(?P={names[key][0]})"
            continue
        names[key] = (f"g{n}", m.group(2) or "")
        pattern += f"(?P<g{n}>.*?)"
        n += 1
    pattern += re.escape(tpl[pos:])
    m = re.fullmatch(pattern, str(text).strip(), re.S) or re.fullmatch(pattern, str(text), re.S)
    if not m:
        return None
    out = {}
    for key, (g, fmt) in names.items():
        v = m.group(g).strip()
        if not v:
            out[key] = None
        elif "%" in fmt:
            try:
                d = dt.datetime.strptime(v, re.sub(r"%-", "%", fmt))
            except ValueError:
                out[key] = v
                continue
            if "%d" in fmt or "%-d" in fmt:
                out[key] = d.date().isoformat() if not re.search(r"%-?[HM]", fmt) else d.isoformat(sep=" ", timespec="minutes")
            else:
                out[key] = f"{d.year:04d}-{d.month:02d}"
        else:
            out[key] = int(v) if re.fullmatch(r"[+-]?\d+", v) and not (len(v) > 1 and v[0] == "0") else v
    return out


def text_notes(tpl: str) -> dict:
    """文のひな形の欄の書き方（雛形の null だけでは、文のどこに入るかが分からない）。"""
    notes = {}
    for key, fmt in text_fields(tpl):
        if "%" in fmt:
            notes[key] = f"{key}: 日付（2026-10-08）" if re.search(r"%-?d", fmt) else f"{key}: 年月（2026-10）"
        else:
            notes.setdefault(key, f"{key}: 「{tpl}」の {{{key}}} に入る文字")
    return notes


def unused_doc_warnings(data: Any, used: "set[str]") -> list[str]:
    """データの `文書:` にあって、定義のどこでも使われていない値（綴りの違いの疑い）。"""
    def leaves(obj, path):
        if isinstance(obj, dict) and obj:
            for k, v in obj.items():
                yield from leaves(v, f"{path}.{k}")
        else:
            yield path
    doc = data.get(DOC_PREFIX) if isinstance(data, dict) else None
    if not isinstance(doc, dict):
        return []
    unused = [p for p in leaves(doc, DOC_PREFIX) if not any(p == u or p.startswith(u + ".") or u.startswith(p + ".") for u in used)]
    return [f"データの {', '.join(unused)} は、定義のどこでも使われていません（綴りの違いを疑う）"] if unused else []




_DOC_DATA: "dict | None" = None   # render・check の間だけ、データ全体（文書の値は、スライドのデータでなくここから引く）


def _lookup(obj: Any, key: str) -> Any:
    """欄の値。`文書.` で始まるキーは、スライドのデータではなく、データ全体の `文書:` から引く。"""
    if key.split(".")[0] == DOC_PREFIX and _DOC_DATA is not None:
        cur: Any = _DOC_DATA   # 書き忘れると「 御中」のような崩れた文が黙って出るので、無ければ止める（null は空）
        for part in key.split("."):
            if not isinstance(cur, dict) or part not in cur:
                raise TemplateError(f"データに {key!r} がありません（文書の値。空でよければ null と書く）")
            cur = cur[part]
        return cur
    return dig(obj, key)


def doc_keys(definition: dict) -> set:
    """定義が使う文書の値のキー（`文書.` で始まるもの）。"""
    keys = set()
    texts = [v for v in (definition.get("properties") or {}).values() if isinstance(v, str)]
    texts += list((definition.get("header_footer") or {}).values())
    for sd in definition.get("slides", []):
        for spec in (sd.get("texts") or {}).values():
            spec = _spec(spec)
            if "text" in spec:
                texts.append(spec["text"])
            elif str(spec.get("key", "")).startswith(DOC_PREFIX + "."):
                keys.add(spec["key"])
    return keys | {k for t in texts for k, _ in text_fields(t) if k.startswith(DOC_PREFIX + ".")}


def _hoist_doc(data: dict) -> dict:
    """スライドのデータの中に入った `文書:` を、データ全体の `文書:` に移す（extract・雛形）。"""
    doc: dict = {}
    for key, val in list(data.items()):
        for obj in (val if isinstance(val, list) else [val]):
            if isinstance(obj, dict) and isinstance(obj.get(DOC_PREFIX), dict):
                for k, v in obj.pop(DOC_PREFIX).items():
                    if doc.get(k) is None:
                        doc[k] = v
    if doc:
        data = {DOC_PREFIX: {**doc, **(data.get(DOC_PREFIX) or {})}, **{k: v for k, v in data.items() if k != DOC_PREFIX}}
    return data


def _spec(spec: Any) -> dict:
    return spec if isinstance(spec, dict) else {"key": spec}


def _limit(spec: dict) -> str:
    parts = []
    if spec.get("max_items"):
        parts.append(f"最大 {spec['max_items']} 項目")
    if spec.get("max_chars"):
        parts.append(f"{spec['max_chars']} 字まで")
    return f"（{'・'.join(parts)}）" if parts else ""


# ---------------------------------------------------------------------------
# 値の変換（データの意味 ⇔ スライドの見た目）
# ---------------------------------------------------------------------------

def dig(data: Any, path: str) -> Any:
    """`a.b.0` のように、ドットと添字でたどる。キーそのものがあれば、そちらを優先する。"""
    if isinstance(data, dict) and path in data:
        return data[path]
    cur = data
    for part in path.split("."):
        if isinstance(cur, dict):
            cur = cur.get(part)
        elif isinstance(cur, list) and part.isdigit() and int(part) < len(cur):
            cur = cur[int(part)]
        else:
            return None
    return cur


def _norm_key(k: Any) -> str:
    if isinstance(k, bool):
        return "true" if k else "false"
    return str(k)


def convert_value(spec: dict, value: Any, where: str = "") -> Any:
    """when（選ばれた見出しに印）・map（表記の置き換え）・part（日付の一部）。"""
    if "when" in spec:
        chosen = value if isinstance(value, list) else ([] if value is None else [value])
        hit = _norm_key(spec["when"]) in {_norm_key(v) for v in chosen}
        return spec.get("mark", "○") if hit else spec.get("unmark", "")
    if "map" in spec:
        table = {_norm_key(k): v for k, v in spec["map"].items()}
        if value is None and "null" not in table and "None" not in table:
            return table.get("false") if "false" in table and spec.get("null_as_false", True) else None
        key = _norm_key(value)
        if key not in table:
            raise TemplateError(f"{where}{spec.get('key')} の値「{value}」は map にありません（{' / '.join(table)} のどれか）")
        return table[key]
    if "part" in spec:
        if value in (None, ""):
            return None
        d = value
        if isinstance(value, str):
            try:
                d = dt.datetime.fromisoformat(value)
            except ValueError:
                raise TemplateError(f"{where}{spec.get('key')} は日付（2026-10-08）で書いてください: {value!r}")
        if spec["part"] not in DATE_PARTS:
            raise TemplateError(f"part は {', '.join(DATE_PARTS)} のどれかです: {spec['part']!r}")
        return getattr(d, spec["part"], None)
    return value


def field_text(spec: dict, obj: Any, index: int, where: str) -> str:
    if "text" in spec:   # 文の一部だけを差し替える（タイトルの年度・期間など）
        return fill_text(spec["text"], lambda k: _lookup(obj, k), where.strip() or "文")
    key = spec.get("key")
    if key == "$index":
        value: Any = index + 1
    else:
        value = convert_value(spec, _lookup(obj, key) if key else None, where)
    text = to_text(value)
    if spec.get("format") and text:
        text = spec["format"].replace("{}", text)
    return text


def _as_items(value: Any, where: str) -> list[tuple[str, int]]:
    """箇条書きのデータ（文字列の配列、{text, level}、children の入れ子）を (文字, 段) に。"""
    out: list[tuple[str, int]] = []

    def rec(v, lvl):
        if v is None:
            return
        if isinstance(v, list):
            for x in v:
                rec(x, lvl)
        elif isinstance(v, dict):
            out.append((to_text(v.get("text")), int(v.get("level", lvl))))
            rec(v.get("children"), int(v.get("level", lvl)) + 1)
        else:
            text = to_text(v)
            out.extend((line, lvl) for line in text.split("\n"))

    rec(value, 0)
    return out


class Fit:
    """収まらない値（サンプルの粒度を超える・図形からはみ出す）を集める。"""

    def __init__(self):
        self.problems: list[str] = []

    def chars(self, where: str, text: str, limit: "int | None") -> None:
        n = len(text.replace("\n", ""))
        if limit and n > limit:
            self.problems.append(f"{where}: {n} 字。サンプルの粒度は {limit} 字まで（{n - limit} 字減らす）")

    def items(self, where: str, n: int, limit: "int | None", unit: str = "項目") -> None:
        if limit and n > limit:
            self.problems.append(f"{where}: {n} {unit}。この枠に収まるのは {limit} {unit}まで（{n - limit} {unit}減らすか、まとめる）")

    def lines(self, where: str, cap: Capacity, paras: list[tuple[str, int]]) -> None:
        need = cap.lines_for([p for p in paras if p[0]] or paras)
        if need > cap.lines:
            self.problems.append(f"{where}: {need} 行になる。図形に収まるのは {cap.lines} 行"
                                 f"（1 行 約 {int(cap.cpl)} 字・{cap.size:g}pt）")


# ---------------------------------------------------------------------------
# スライドの複製・削除
# ---------------------------------------------------------------------------

def clone_slide(prs, src, after_sid):
    """src を、after_sid（sldId）の直後に複製する。ノートとコメントは写さない。"""
    new = prs.slides.add_slide(src.slide_layout)
    new_el = new._element
    for child in list(new_el):
        new_el.remove(child)
    rid_map: dict[str, str] = {}
    for rId, rel in src.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE) or "comment" in rel.reltype.lower():
            continue
        if rel.is_external:
            rid_map[rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rId] = new.part.relate_to(rel.target_part, rel.reltype)
    for k, v in src._element.attrib.items():
        new_el.set(k, v)
    for child in src._element:
        new_el.append(deepcopy(child))
    for el in new_el.iter():
        for attr, val in list(el.attrib.items()):
            if attr.startswith(f"{{{NS_R}}}"):
                if val in rid_map:
                    el.set(attr, rid_map[val])
                elif local(el) != "sldLayoutId":
                    ext = next((a for a in el.iterancestors() if local(a) == "ext"), None)
                    if ext is not None:   # コメントへの参照（新しい形式）
                        ext.getparent().remove(ext)
                        break
    for cached in ("shapes", "placeholders"):   # python-pptx が覚えている、差し替える前の図形の一覧を捨てる
        new.__dict__.pop(cached, None)
    lst = prs.slides._sldIdLst
    new_sid = lst[-1]
    lst.remove(new_sid)
    after_sid.addnext(new_sid)
    for entry in _section_entries(prs, after_sid.get("id")):   # 複製は、元のスライドと同じセクションに入れる
        entry.addnext(etree.Element(entry.tag, id=new_sid.get("id")))
    return new, new_sid


def delete_slide(prs, sid) -> None:
    rId = sid.rId
    for entry in _section_entries(prs, sid.get("id")):
        entry.getparent().remove(entry)
    prs.slides._sldIdLst.remove(sid)
    prs.part.drop_rel(rId)


def _section_entries(prs, slide_id) -> list:
    """セクション（章）の一覧の中で、そのスライドを指す要素。"""
    return [e for e in prs.part._element.iter(f"{{{NS_P14}}}sldId") if e.get("id") == str(slide_id)]


def _notes_text(slide) -> str:
    if not slide.has_notes_slide or slide.notes_slide.notes_text_frame is None:
        return ""
    return slide.notes_slide.notes_text_frame.text


def _set_notes(slide, text: str) -> None:
    if not text and not slide.has_notes_slide:
        return
    frame = slide.notes_slide.notes_text_frame
    if frame is not None:
        frame.text = text


# ---------------------------------------------------------------------------
# 定義の検査
# ---------------------------------------------------------------------------

def validate_definition(template: "str | bytes", definition: dict) -> None:
    """テンプレートと定義の整合（スライド・図形・表・図）を確かめる。"""
    if definition.get("version") != DEF_VERSION:
        raise TemplateError(f"定義ファイルの version が未対応です: {definition.get('version')!r}")
    prs = open_prs(read_bytes(template))
    slides = list(prs.slides)
    # 名前の変わった図形（テンプレートの差し替え）は、1 つずつでなく、まとめて挙げる
    missing = []
    for sd in definition.get("slides", []):
        n = int(sd.get("slide", 0))
        if sd.get("drop") or not 1 <= n <= len(slides):
            continue
        view = SlideView(prs, slides[n - 1], n)
        refs = list(sd.get("texts") or {}) + list(sd.get("lists") or {}) + (sd.get("keep") or []) + (sd.get("clear") or [])
        refs += [t["shape"] for t in sd.get("tables") or []]
        refs += [r for d in sd.get("diagrams") or [] for it in d.get("items") or [] for r in it] \
            + [r for d in sd.get("diagrams") or [] for r in d.get("connectors") or []]
        for ref in refs:
            try:
                view.find(ref)
            except TemplateError as e:
                missing.append(str(e))
    if len(missing) > 1:
        raise TemplateError("テンプレートに無い図形があります:\n" + "\n".join(f"  {m}" for m in missing))
    seen_keys: dict[str, int] = {}
    seen_slides: set[int] = set()
    for sd in definition.get("slides", []):
        n = int(sd.get("slide", 0))
        if not 1 <= n <= len(slides):
            raise TemplateError(f"スライド {n} はテンプレートにありません（{len(slides)} 枚）")
        if n in seen_slides:
            raise TemplateError(f"スライド {n} が定義に 2 回あります")
        seen_slides.add(n)
        if sd.get("drop"):
            continue
        key = sd.get("key", sd.get("id"))
        if key in seen_keys:
            raise TemplateError(f"スライド {seen_keys[key]} と {n} のデータのキー `{key}` が同じです")
        seen_keys[key] = n
        view = SlideView(prs, slides[n - 1], n)
        for section in ("texts", "lists"):
            for ref in (sd.get(section) or {}):
                el = view.find(ref)
                if local(el) != "sp" or txbody(el) is None:
                    raise TemplateError(f"スライド {n}「{ref}」は文字を入れられる図形ではありません（{kind_of(el)}）")
        for ref in (sd.get("keep") or []) + (sd.get("clear") or []):
            view.find(ref)
        for t in sd.get("tables") or []:
            el = view.find(t["shape"])
            tbl = table_of(el)
            if tbl is None:
                raise TemplateError(f"スライド {n}「{t['shape']}」は表ではありません")
            trs = tbl.findall(qa("tr"))
            if t.get("key"):
                body = len(trs) - int(t.get("header_rows", 1)) - int(t.get("footer_rows", 0))
                if body < 1:
                    raise TemplateError(f"スライド {n}「{t['shape']}」に見本の行がありません（header_rows・footer_rows を確かめる）")
                for tr in trs[int(t.get("header_rows", 1)):len(trs) - int(t.get("footer_rows", 0))]:
                    if any(tc.get("rowSpan") or tc.get("vMerge") for tc in tr.findall(qa("tc"))):
                        raise TemplateError(f"スライド {n}「{t['shape']}」の見本の行に、縦の結合があります（繰り返せない）")
                cols = t.get("columns") or []
                n_cols = len(tbl.findall(f"{qa('tblGrid')}/{qa('gridCol')}"))
                if len(cols) > n_cols:
                    raise TemplateError(f"スライド {n}「{t['shape']}」の columns が {len(cols)} 列あります（表は {n_cols} 列）")
                head_rows = int(t.get("header_rows", 1))
                for j, c in enumerate(cols):
                    if c and c.get("header") is not None and head_rows:
                        actual = body_text(trs[head_rows - 1].findall(qa("tc"))[j].find(qa("txBody"))).strip()
                        if actual != c["header"]:
                            raise TemplateError(f"スライド {n}「{t['shape']}」の {j + 1} 列目の見出しが定義と違います"
                                                f"（定義「{c['header']}」・テンプレート「{actual}」）")
            for rc in (t.get("cells") or {}):
                r, c = _rc(rc)
                if r > len(trs) or c > len(trs[r - 1].findall(qa("tc"))):
                    raise TemplateError(f"スライド {n}「{t['shape']}」に {rc} のセルがありません")
        for d in sd.get("diagrams") or []:
            items = [[view.find(r) for r in it] for it in d.get("items") or []]
            if not items:
                raise TemplateError(f"スライド {n} の図 {d.get('id')} に items がありません")
            if any(len(it) != len(items[0]) for it in items):
                raise TemplateError(f"スライド {n} の図 {d.get('id')} の items は、どれも同じ数の図形にしてください")
            conns = [view.find(r) for r in d.get("connectors") or []]
            if conns and len(items) < 2:
                raise TemplateError(f"スライド {n} の図 {d.get('id')}: つなぎ（connectors）には 2 つ以上の items が要る")
            if d.get("axis", "x") == "grid" and conns:
                raise TemplateError(f"スライド {n} の図 {d.get('id')}: 格子の並びは、つなぎを持てません")
            if d.get("axis", "x") == "grid" and not d.get("columns"):
                raise TemplateError(f"スライド {n} の図 {d.get('id')}: 格子の並びには columns が要る")
            for ref in d.get("fields") or {}:
                _field_path(view, items[0], ref, d)
        if sd.get("overflow", "error") not in ("error", "split"):
            raise TemplateError(f"スライド {n} の overflow は error か split です")


def _rc(spec: str) -> tuple[int, int]:
    m = re.match(r"^\s*(\d+)\s*,\s*(\d+)\s*$", str(spec))
    if not m:
        raise TemplateError(f"表のセルは「行,列」（1 始まり）で書いてください: {spec!r}")
    return int(m.group(1)), int(m.group(2))


def _field_path(view: SlideView, item0: list, ref: str, d: dict) -> tuple:
    target = view.find(ref)
    for el, path in _item_paths(item0):
        if el is target:
            return path
    raise TemplateError(f"スライド {view.number} の図 {d.get('id')}: fields の「{ref}」が、1 つ目の items の中にありません")


def find_leftovers(prs, definition: dict) -> list[str]:
    """strict: 定義のどこにも入らないまま、テンプレートの文字が残る図形。"""
    defs = {int(sd["slide"]): sd for sd in definition.get("slides", [])}
    out = []
    for n, slide in enumerate(prs.slides, start=1):
        sd = defs.get(n) or {}
        if sd.get("drop"):
            continue
        view = SlideView(prs, slide, n)
        covered = set()
        for section in ("texts", "lists"):
            covered.update(view.find(r) for r in (sd.get(section) or {}))
        covered.update(view.find(r) for r in (sd.get("keep") or []) + (sd.get("clear") or []))
        for d in sd.get("diagrams") or []:
            covered.update(view.find(r) for it in d.get("items") or [] for r in it)
            covered.update(view.find(r) for r in d.get("connectors") or [])
        tables = {view.find(t["shape"]): t for t in sd.get("tables") or []}
        for el in walk(view.tree):
            if el in covered or any(a in covered for a in el.iterancestors()) or local(el) == "grpSp":
                continue
            tbl = table_of(el)
            if tbl is not None and el in tables:
                t = tables[el]
                if not t.get("key"):
                    continue
                trs = tbl.findall(qa("tr"))
                lo, hi = int(t.get("header_rows", 1)), len(trs) - int(t.get("footer_rows", 0))
                cols = t.get("columns") or []
                cells = {_rc(rc) for rc in (t.get("cells") or {})}
                for r in range(hi, len(trs)):   # 合計の行: 数字の残る値の欄は、cells で入れ直す
                    for j, tc in enumerate(trs[r].findall(qa("tc"))):
                        spec = cols[j] if j < len(cols) else None
                        text = body_text(tc.find(qa("txBody")))
                        if j and spec and spec.get("key") and not spec.get("keep") and (r + 1, j + 1) not in cells \
                                and re.search(r"\d", text):
                            out.append(f"スライド {n}「{shape_name(el)}」{r + 1} 行 {j + 1} 列（最後の行）: 「{_short(text, 20)}」")
                for r in range(lo, hi):
                    for j, tc in enumerate(trs[r].findall(qa("tc"))):
                        if (j >= len(cols) or cols[j] is None) and body_text(tc.find(qa("txBody"))).strip():
                            out.append(f"スライド {n}「{shape_name(el)}」{r + 1} 行 {j + 1} 列: "
                                       f"「{_short(body_text(tc.find(qa('txBody'))), 20)}」")
                continue
            text = literal_text(el)
            if text and (tbl is not None or txbody(el) is not None):
                out.append(f"スライド {n}「{shape_name(el)}」: 「{_short(text, 24)}」")
    return out


# ---------------------------------------------------------------------------
# render
# ---------------------------------------------------------------------------

def _pages(view: SlideView, sd: dict, obj: Any) -> int:
    if sd.get("overflow") != "split":
        return 1
    pages = 1
    for ref, spec in (sd.get("lists") or {}).items():
        spec = _spec(spec)
        lim = spec.get("max_items")
        if lim:
            pages = max(pages, math.ceil(len(_as_items(dig(obj, spec["key"]), "")) / lim))
    for t in sd.get("tables") or []:
        if t.get("key") and t.get("max_rows"):
            pages = max(pages, math.ceil(len(dig(obj, t["key"]) or []) / t["max_rows"]))
    for d in sd.get("diagrams") or []:
        if d.get("max_items"):
            pages = max(pages, math.ceil(len(dig(obj, d["key"]) or []) / d["max_items"]))
    return pages


def _chunk(seq: list, limit: "int | None", page: int, pages: int) -> list:
    if pages <= 1 or not limit:
        return seq
    return seq[page * limit:(page + 1) * limit]


def render(template: "str | bytes", definition: dict, data: Any, output: str,
           allow_overflow: bool = False) -> list[str]:
    """テンプレート + 定義 + データ → pptx。警告を返す。収まらない値があれば止める（allow_overflow で警告に）。"""
    raw = read_bytes(template)
    validate_definition(raw, definition)
    if data is None:
        data = {}
    if not isinstance(data, dict):
        raise TemplateError("データは、スライドのキーを持つオブジェクトにしてください")
    prs = open_prs(raw)
    warnings: list[str] = []
    global _DOC_DATA
    _DOC_DATA = data
    if definition.get("strict"):
        left = find_leftovers(prs, definition)
        if left:
            raise TemplateError("テンプレートの値が、定義のどこにも入らないまま残ります（keep・texts・clear などに入れる）:\n  "
                                + "\n  ".join(left))
    defs = {int(sd["slide"]): sd for sd in definition.get("slides", [])}
    props = {k: fill_text(v, lambda key: _lookup(data, key), f"properties の {k}") if isinstance(v, str) else v
             for k, v in (definition.get("properties") or {}).items()}
    known = {sd.get("key", sd.get("id")) for sd in defs.values() if not sd.get("drop")} | {DOC_PREFIX}
    warnings += unused_doc_warnings(data, doc_keys(definition))
    unknown = [k for k in data if k not in known]
    if unknown:
        warnings.append(f"データの {', '.join(unknown)} は、どのスライドのキーにもありません（綴りの違いを疑う）")
    for sd in defs.values():
        key = sd.get("key", sd.get("id"))
        if sd.get("drop") or key not in data:
            continue
        objs = data[key] if isinstance(data[key], list) else [data[key]]
        fields = _top_keys(sd)
        extra = sorted({k for o in objs if isinstance(o, dict) for k in o if k not in fields})
        if extra:
            warnings.append(f"データ `{key}` の {', '.join(extra)} は、スライド {sd['slide']} のどの欄にもありません（綴りの違いを疑う）")
        # 表の行・図の項目の中の欄も確かめる
        groups = [("表", t["key"], t.get("columns") or []) for t in sd.get("tables") or [] if t.get("key")]
        groups += [("図", d["key"], [_spec(f) for f in (d.get("fields") or {}).values()]) for d in sd.get("diagrams") or []]
        for kind, gkey, specs in groups:
            names = {str(c["key"]).split(".")[0] for c in specs if c and c.get("key")}
            items = [x for o in objs if isinstance(o, dict) for x in (dig(o, gkey) or []) if isinstance(x, dict)]
            extra = sorted({k for x in items for k in x if k not in names})
            if extra and names:
                warnings.append(f"データ `{key}.{gkey}` の {', '.join(extra)} は、スライド {sd['slide']} の{kind}の"
                                "どの欄にもありません（綴りの違いを疑う）")
    slides, sids = list(prs.slides), list(prs.slides._sldIdLst)
    instances, to_delete = [], []
    for n, (slide, sid) in enumerate(zip(slides, sids), start=1):
        sd = defs.get(n)
        if sd is None:
            continue
        if sd.get("drop"):
            to_delete.append(sid)
            continue
        key = sd.get("key", sd.get("id"))
        val = data.get(key)
        if sd.get("repeat"):
            if val is not None and not isinstance(val, list):
                raise TemplateError(f"スライド {n} は繰り返すスライドです。`{key}` は配列にしてください")
            objs = list(val or [])
        else:
            if val is not None and not isinstance(val, dict):
                raise TemplateError(f"`{key}` はオブジェクト（欄: 値）にしてください")
            objs = [val or {}]
        tview = SlideView(prs, slide, n)
        plan = [(i, obj, p, k) for i, obj in enumerate(objs) for k in [_pages(tview, sd, obj)] for p in range(k)]
        if not plan:
            to_delete.append(sid)
            continue
        notes = _notes_text(slide)
        targets, anchor = [slide], sid
        for _ in plan[1:]:
            new, anchor = clone_slide(prs, slide, anchor)
            if notes and not props.get("scrub"):
                _set_notes(new, notes)
            targets.append(new)
        for (i, obj, p, k), target in zip(plan, targets):
            label = f"スライド {n}" + (f"（{key}[{i}]）" if sd.get("repeat") else "") + (f" {p + 1}/{k} 枚目" if k > 1 else "")
            instances.append((target, n, sd, obj, p, k, label))
    fit = Fit()
    for target, n, sd, obj, p, k, label in instances:
        _fill_slide(SlideView(prs, target, n), sd, obj, p, k, label, fit, warnings)
    for sid in to_delete:
        delete_slide(prs, sid)
    if fit.problems:
        msg = "収まらない値があります。テンプレートの粒度に合わせて、データの量を減らすか言い換えてください:\n  " \
              + "\n  ".join(fit.problems)
        if any("行。この枠に" in p or "項目。この枠に" in p for p in fit.problems):
            msg += "\n（全件の一覧など減らせない内容なら、利用者に確かめて、そのスライドに overflow: split を書く。続きのスライドを足す）"
        if not allow_overflow:
            raise TemplateError(msg)
        warnings.append(msg)
    if props.get("scrub"):
        n_comments = drop_comments(prs)
        if n_comments:
            warnings.append(f"コメント {n_comments} 件を取り除きました")
        mapped = {id(t[0]) for t in instances if t[2].get("notes")}
        for slide in prs.slides:
            if id(slide) not in mapped and _notes_text(slide):
                _set_notes(slide, "")
    buf = io.BytesIO()
    prs.save(buf)
    out = finish_package(buf.getvalue(), props)
    with open(output, "wb") as f:
        f.write(out)
    return warnings


def _top_keys(sd: dict) -> set:
    """スライドのデータで使う欄の名前（ドットの先頭）。"""
    keys = [_spec(s).get("key") for s in list((sd.get("texts") or {}).values()) + list((sd.get("lists") or {}).values())]
    for t in sd.get("tables") or []:
        keys += [t.get("key")] + [_spec(s).get("key") for s in (t.get("cells") or {}).values()]
    keys += [d.get("key") for d in sd.get("diagrams") or []] + [sd.get("notes")]
    return {k.split(".")[0] for k in keys if k}


def _fill_slide(view: SlideView, sd: dict, obj: dict, page: int, pages: int, label: str, fit: Fit,
                warnings: list) -> None:
    for ref, spec in (sd.get("texts") or {}).items():
        spec = _spec(spec)
        if spec.get("keep"):
            continue
        el = view.find(ref)
        where = f"{label}「{ref}」({spec.get('key')})"
        text = "" if spec.get("clear") else field_text(spec, obj, 0, f"{label} ")
        paras = [(line, 0) for line in text.split("\n")] if text else []
        fit.chars(where, text, spec.get("max_chars"))
        cap = view.capacity(el)
        if text:
            fit.lines(where, cap, paras)
        set_text(txbody(el), text)
    for ref, spec in (sd.get("lists") or {}).items():
        spec = _spec(spec)
        el = view.find(ref)
        where = f"{label}「{ref}」({spec.get('key')})"
        items = _chunk(_as_items(dig(obj, spec["key"]), where), spec.get("max_items"), page, pages)
        fit.items(where, len(items), spec.get("max_items"))
        for i, (t, _) in enumerate(items):
            fit.chars(f"{where} {i + 1} 項目め", t, spec.get("max_chars"))
        cap = view.capacity(el)
        if items:
            fit.lines(where, cap, items)
        set_paragraphs(txbody(el), items)
    for t in sd.get("tables") or []:
        _fill_table(view, t, obj, page, pages, label, fit)
    for d in sd.get("diagrams") or []:
        _fill_diagram(view, d, obj, page, pages, label, fit, warnings)
    for ref in sd.get("clear") or []:
        el = view.find(ref)
        for e in ([el] if local(el) != "grpSp" else list(walk(el))):
            if txbody(e) is not None:
                set_text(txbody(e), "")
    if sd.get("notes"):
        _set_notes(view.slide, to_text(dig(obj, sd["notes"])))


def _cell_rpr(tc):
    for e in tc.iter(qa("rPr"), qa("endParaRPr")):
        if len(e) or e.get("sz"):
            return e
    return None


def _borrow_format(tc, pool) -> None:
    """空のセル（書式を持たない）に、pool のセル（同じ列・同じ行）の文字の書式を写す。既定の 18pt にしない。"""
    if _cell_rpr(tc) is not None:
        return
    for other in pool:
        src = _cell_rpr(other) if other is not tc else None
        if src is not None:
            for p in tc.iter(qa("p")):
                end = p.find(qa("endParaRPr"))
                if end is not None:
                    p.remove(end)
                end = deepcopy(src)
                end.tag = qa("endParaRPr")
                p.append(end)
            return


def _format_pool(trs, row, col: int, head: int) -> list:
    """書式を借りるセルの順: 同じ列の本文の行 → 同じ行 → 見出しの行（太字なので最後）。"""
    def at(r):
        tcs = r.findall(qa("tc"))
        return [tcs[col]] if col < len(tcs) else []
    return [c for r in trs[head:] for c in at(r)] + row.findall(qa("tc")) + [c for r in trs[:head] for c in at(r)]


def _fill_table(view: SlideView, t: dict, obj: dict, page: int, pages: int, label: str, fit: Fit) -> None:
    el = view.find(t["shape"])
    tbl = table_of(el)
    trs = tbl.findall(qa("tr"))
    where = f"{label}「{t['shape']}」"
    for rc, spec in (t.get("cells") or {}).items():
        spec = _spec(spec)
        if spec.get("keep"):
            continue
        r, c = _rc(rc)
        tc = trs[r - 1].findall(qa("tc"))[c - 1]
        text = field_text(spec, obj, 0, f"{where} ")
        fit.chars(f"{where} {rc}({spec.get('key')})", text, spec.get("max_chars"))
        _borrow_format(tc, _format_pool(trs, trs[r - 1], c - 1, int(t.get("header_rows", 1))))
        set_text(tc.find(qa("txBody")), text)
    if not t.get("key"):
        return
    head, foot = int(t.get("header_rows", 1)), int(t.get("footer_rows", 0))
    samples = trs[head:len(trs) - foot]
    rows = dig(obj, t["key"]) or []
    if not isinstance(rows, list):
        raise TemplateError(f"{where} の `{t['key']}` は配列にしてください")
    rows = _chunk(rows, t.get("max_rows"), page, pages)
    fit.items(f"{where}({t['key']})", len(rows), t.get("max_rows"), "行")
    cols = t.get("columns") or []
    parent = samples[0].getparent()
    insert_at = list(parent).index(samples[0])
    for s in samples:
        parent.remove(s)
    new_rows = []
    for i, row in enumerate(rows or [None]):
        tr = deepcopy(samples[i % len(samples)])
        for j, tc in enumerate(tr.findall(qa("tc"))):
            spec = cols[j] if j < len(cols) else None
            if not spec or spec.get("keep"):
                continue
            body = tc.find(qa("txBody"))
            if row is None or spec.get("clear"):
                set_text(body, "")
                continue
            text = field_text(spec, row, i, f"{where} {i + 1} 行め ")
            fit.chars(f"{where} {i + 1} 行め({spec.get('key')})", text, spec.get("max_chars"))
            _borrow_format(tc, _format_pool(trs[:head] + samples, tr, j, head))
            set_text(body, text)
        new_rows.append(tr)
    for k, tr in enumerate(new_rows):
        parent.insert(insert_at + k, tr)
    # 表の高さ: 行の高さの合計。文字が増えて行が伸びる分も見積もり、下にはみ出すなら止める
    widths = [int(g.get("w")) for g in tbl.findall(f"{qa('tblGrid')}/{qa('gridCol')}")]
    total = 0
    for tr in tbl.findall(qa("tr")):
        h = int(tr.get("h", "0"))
        for j, tc in enumerate(tr.findall(qa("tc"))):
            body = tc.find(qa("txBody"))
            paras = [p for p in paragraphs(body) if p[0]]
            if not paras or j >= len(widths):
                continue
            size = view.font_size(tc, body)
            mar = int(tc.find(qa("tcPr")).get("marL", 91440)) + int(tc.find(qa("tcPr")).get("marR", 91440)) \
                if tc.find(qa("tcPr")) is not None else 182880
            span = int(tc.get("gridSpan", "1"))
            cap = Capacity(max((sum(widths[j:j + span]) - mar) / (size * EMU_PER_PT), 1.0), 1, size, body)
            need = cap.lines_for(paras) * size * EMU_PER_PT * line_factor(body) + 91440
            h = max(h, int(need))
        total += h
    x, y, w, _ = view.box(el)
    xf = el.find(qp("xfrm"))
    if xf is not None and xf.find(qa("ext")) is not None:
        xf.find(qa("ext")).set("cy", str(sum(int(tr.get("h", "0")) for tr in tbl.findall(qa("tr")))))
    bottom = view.free_end(el, (x, y, w, 1), "y")
    if y + total > bottom + EMU_PER_PT * 2:
        fit.problems.append(f"{where}({t['key']}): 表の高さが {_cm(total)} cm になり、下の余白・図形にかかる"
                            f"（使えるのは {_cm(bottom - y)} cm。行か文字を減らす）")


def _bbox(view: SlideView, els: list) -> tuple[int, int, int, int]:
    bs = [view.box(e) for e in els]
    x, y = min(b[0] for b in bs), min(b[1] for b in bs)
    return x, y, max(b[0] + b[2] for b in bs) - x, max(b[1] + b[3] for b in bs) - y


def _look(item: list) -> bytes:
    """図形の見た目（文字・id・名前・位置を除いた XML）。"""
    out = []
    for el in item:
        c = deepcopy(el)
        for sub in c.iter():
            if local(sub) == "t":
                sub.text = ""
            elif local(sub) == "cNvPr":
                sub.attrib.pop("id", None)
                sub.attrib.pop("name", None)
            elif local(sub) == "off":
                sub.set("x", "0")
                sub.set("y", "0")
        out.append(etree.tostring(c))
    return b"".join(out)


def _period(items: list) -> int:
    """見本の見た目が繰り返す周期（A・B・A なら 2）。複製は、この周期で見本を使い回す。"""
    looks = [_look(it) for it in items]
    for p in range(1, len(items) + 1):
        if all(looks[i] == looks[i % p] for i in range(len(items))):
            return p
    return len(items)


def _fill_diagram(view: SlideView, d: dict, obj: dict, page: int, pages: int, label: str, fit: Fit,
                  warnings: list) -> None:
    items = [[view.find(r) for r in it] for it in d["items"]]
    conns = [view.find(r) for r in d.get("connectors") or []]
    where = f"{label} 図 {d['id']}"
    values = dig(obj, d["key"])
    if values is None:
        values = []
    if not isinstance(values, list):
        raise TemplateError(f"{where} の `{d['key']}` は配列にしてください")
    values = _chunk(values, d.get("max_items"), page, pages)
    fit.items(f"{where}({d['key']})", len(values), d.get("max_items"), "つ")
    fields = [(_field_path(view, items[0], ref, d), _spec(spec)) for ref, spec in (d.get("fields") or {}).items()]
    caps = {path: view.capacity(follow(items[0], path)) for path, _ in fields}
    anchors = [(_bbox(view, it)[0], _bbox(view, it)[1]) for it in items]
    n_s, n = len(items), len(values)
    axis = d.get("axis", "x")
    if n_s >= 2:
        px, py = anchors[1][0] - anchors[0][0], anchors[1][1] - anchors[0][1]
        if axis == "grid":
            cols = int(d["columns"])
            px = anchors[1][0] - anchors[0][0]
            py = anchors[cols][1] - anchors[0][1] if n_s > cols else 0
    else:
        px = py = 0
        if n > 1:
            raise TemplateError(f"{where}: 見本の items が 1 つだけなので、並べる間隔が分かりません（items を 2 つ以上にする）")
    if axis == "grid":
        cols = int(d["columns"])
        targets = [(anchors[0][0] + (i % cols) * px, anchors[0][1] + (i // cols) * py) for i in range(n)]
    else:
        shift = (n_s - n) / 2 if d.get("align") == "center" else 0
        targets = [(int(anchors[0][0] + (i + shift) * px), int(anchors[0][1] + (i + shift) * py)) for i in range(n)]
    # 見本の図形を取り除き、同じ場所（重なりの順）へ複製を入れる
    tree = view.tree
    olds = [e for it in items for e in it] + conns
    positions = sorted(list(tree).index(e) for e in olds)
    first_pos = positions[0]
    conn_first = bool(conns) and list(tree).index(conns[0]) < list(tree).index(items[0][0])
    old_ids = {shape_id(e): (k, j) for k, it in enumerate(items) for j, e in enumerate(it)}
    for e in olds:
        tree.remove(e)
    next_id = max([shape_id(e) for e in walk(tree)] + list(old_ids) + [shape_id(c) for c in conns] + [1]) + 1
    new_items, new_conns = [], []
    clone_ids: list[list[int]] = []
    period = _period(items)
    for i, value in enumerate(values):
        k = i % period
        dx, dy = targets[i][0] - anchors[k][0], targets[i][1] - anchors[k][1]
        clones, ids = [], []
        for j, src in enumerate(items[k]):
            c = deepcopy(src)
            translate(c, dx, dy)
            for sub in [c] + list(walk(c)):
                pr = nv_pr(sub)
                if pr is not None:
                    pr.set("id", str(next_id))
                    if sub is c:
                        ids.append(next_id)
                    next_id += 1
            nv_pr(c).set("name", f"{d['id']}.{i + 1}.{j + 1}")
            clones.append(c)
        for path, spec in fields:
            target = follow(clones, path)
            fwhere = f"{where} {i + 1} つめ({spec.get('key')})"
            if spec.get("keep"):
                continue
            text = "" if spec.get("clear") else field_text(spec, value, i, f"{where} {i + 1} つめ ")
            fit.chars(fwhere, text, spec.get("max_chars"))
            if text:
                fit.lines(fwhere, caps[path], [(line, 0) for line in text.split("\n")])
            set_text(txbody(target), text)
        new_items.append(clones)
        clone_ids.append(ids)
    for i in range(max(n - 1, 0)):
        if not conns:
            break
        c_k = i % len(conns)
        src = conns[c_k]
        c = deepcopy(src)
        # つなぎ c_k は、見本の items[c_k] と items[c_k + 1] の間にある。複製は items[i] と items[i + 1] の間へ
        translate(c, targets[i][0] - anchors[c_k][0], targets[i][1] - anchors[c_k][1])
        nv_pr(c).set("id", str(next_id))
        nv_pr(c).set("name", f"{d['id']}.arrow.{i + 1}")
        next_id += 1
        for tag, item_i in (("stCxn", i), ("endCxn", i + 1)):
            cx = c.find(f".//{qa(tag)}")
            if cx is None:
                continue
            hit = old_ids.get(int(cx.get("id", "0")))
            if hit is None or item_i >= len(clone_ids):
                cx.getparent().remove(cx)
            else:
                cx.set("id", str(clone_ids[item_i][hit[1]]))
        new_conns.append(c)
    seq = (new_conns + [e for it in new_items for e in it]) if conn_first else \
        ([e for it in new_items for e in it] + new_conns)
    for k, e in enumerate(seq):
        tree.insert(first_pos + k, e)
    _drop_stale_timing(view, warnings, where)


def _drop_stale_timing(view: SlideView, warnings: list, where: str) -> None:
    """取り除いた図形を指すアニメーションがあれば、アニメーションごと外す（開けなくなるのを防ぐ）。"""
    timing = view.el.find(qp("timing"))
    if timing is None:
        return
    ids = {str(shape_id(e)) for e in walk(view.tree)}
    if any(t.get("spid") not in ids for t in timing.iter(qp("spTgt"))):
        view.el.remove(timing)
        warnings.append(f"{where}: アニメーションが見本の図形を指していたので、このスライドのアニメーションを外しました")


# ---------------------------------------------------------------------------
# extract（記入済みの文書からデータを取り出す。render の逆）
# ---------------------------------------------------------------------------

def _read_back(pairs: list[tuple[dict, str]]) -> dict:
    """(欄の指定, スライドの文字) の組から、データの値を起こす。when は選ばれた見出しに、part は日付にまとめる。"""
    out: dict[str, Any] = {}
    choices: dict[str, list] = {}
    parts: dict[str, dict] = {}
    for spec, text in pairs:
        if "text" in spec:
            got = read_text(spec["text"], text) or {}
            for k in spec_keys(spec):
                if out.get(k) is None:
                    out[k] = got.get(k)
            continue
        key = spec.get("key")
        if not key or key == "$index" or spec.get("keep") or spec.get("clear"):
            continue
        text = text.strip()
        if "map" in spec:
            hit = [k for k, v in spec["map"].items() if text == to_text(v).strip()]
            if hit or not text:
                k = hit[0] if hit else None
                out[key] = {"true": True, "false": False}.get(k, k) if isinstance(k, str) else k
                continue
        if text in MARK_OFF:
            text = ""
        if "when" in spec:
            choices.setdefault(key, [])
            if text and text == str(spec.get("mark", "○")).strip() or (text in MARK_ON and "mark" not in spec):
                choices[key].append(spec["when"])
            continue
        if "part" in spec:
            if text.isdigit():
                parts.setdefault(key, {})[spec["part"]] = int(text)
            else:
                parts.setdefault(key, {})
            continue
        if spec.get("format") and text:
            pat = "^" + re.escape(spec["format"]).replace(re.escape("{}"), "(.*)") + "$"
            m = re.match(pat, text, re.S)
            text = m.group(1) if m else text
        out[key] = text if text else None
    for key, chosen in choices.items():
        out[key] = chosen[0] if len(chosen) == 1 else (chosen or None)
    for key, p in parts.items():
        if {"year", "month", "day"} <= set(p):
            try:
                d = dt.datetime(p["year"], p["month"], p["day"], p.get("hour", 0), p.get("minute", 0))
                out[key] = d.date().isoformat() if not ({"hour", "minute"} & set(p)) else d.isoformat(timespec="minutes")
            except ValueError:
                out[key] = None
        else:
            out[key] = None
    nested: dict = {}
    for k, v in out.items():
        _set_path(nested, k, v)
    return nested


def _set_path(root: dict, path: str, value: Any) -> None:
    cur = root
    parts = path.split(".")
    for part in parts[:-1]:
        cur = cur.setdefault(part, {})
    cur[parts[-1]] = value


def _as_lists(obj: Any) -> Any:
    if not isinstance(obj, dict):
        return [_as_lists(x) for x in obj] if isinstance(obj, list) else obj
    obj = {k: _as_lists(v) for k, v in obj.items()}
    if obj and set(obj) == {str(i) for i in range(len(obj))}:
        return [obj[str(i)] for i in range(len(obj))]
    return obj


def _is_empty(value: Any) -> bool:
    if isinstance(value, dict):
        return all(_is_empty(v) for v in value.values())
    if isinstance(value, list):
        return all(_is_empty(v) for v in value)
    return value in (None, False, "")


def _diagram_items_in(view: SlideView, d: dict) -> list[list]:
    """render が並べた図（`<id>.<番号>.<図形>` の名前）か、なければテンプレートの見本の並び。"""
    found: dict[int, dict[int, Any]] = {}
    pat = re.compile(rf"^{re.escape(d['id'])}\.(\d+)\.(\d+)$")
    for e in view.top():
        m = pat.match(shape_name(e))
        if m:
            found.setdefault(int(m.group(1)), {})[int(m.group(2))] = e
    if found:
        return [[found[i][j] for j in sorted(found[i])] for i in sorted(found)]
    try:
        return [[view.find(r) for r in it] for it in d["items"]]
    except TemplateError:
        return []


def _read_slide(view: SlideView, sd: dict, tview: "SlideView | None", notes: list[str], label: str) -> dict:
    pairs: list[tuple[dict, str]] = []
    for ref, spec in (sd.get("texts") or {}).items():
        pairs.append((_spec(spec), body_text(txbody(view.find(ref)))))
    obj = _read_back(pairs)
    for ref, spec in (sd.get("lists") or {}).items():
        spec = _spec(spec)
        items = [(t, lvl) for t, lvl in paragraphs(txbody(view.find(ref))) if t.strip()]
        _set_path(obj, spec["key"], [t if lvl == 0 else {"text": t, "level": lvl} for t, lvl in items])
    for t in sd.get("tables") or []:
        tbl = table_of(view.find(t["shape"]))
        trs = tbl.findall(qa("tr"))
        cell_pairs = []
        tlen = len(table_of(tview.find(t["shape"])).findall(qa("tr"))) if tview and t.get("key") else len(trs)
        for rc, spec in (t.get("cells") or {}).items():
            if _spec(spec).get("keep"):
                continue
            r, c = _rc(rc)
            if r > tlen - int(t.get("footer_rows", 0)):   # 合計の行は、行が増えても表の最後から数える
                r = len(trs) - (tlen - r)
            if r <= len(trs) and c <= len(trs[r - 1].findall(qa("tc"))):
                cell_pairs.append((_spec(spec), body_text(trs[r - 1].findall(qa("tc"))[c - 1].find(qa("txBody")))))
        for k, v in _read_back(cell_pairs).items():
            obj[k] = v
        if not t.get("key"):
            continue
        head, foot = int(t.get("header_rows", 1)), int(t.get("footer_rows", 0))
        rows = []
        cols = t.get("columns") or []
        for r, tr in enumerate(trs[head:len(trs) - foot], start=head + 1):
            tcs = tr.findall(qa("tc"))
            rec = _read_back([(c, body_text(tcs[j].find(qa("txBody")))) for j, c in enumerate(cols) if c and j < len(tcs)])
            if _is_empty(rec):
                notes.append(f"{label}「{t['shape']}」{r} 行め: 空の行なので書かない")
                continue
            rows.append(rec)
        keys = [c["key"] for c in cols if c and c.get("key") and c["key"] != "$index" and not c.get("keep") and not c.get("clear")]
        for rec in rows:
            for k in keys:
                if dig(rec, k) is None and k not in rec:
                    _set_path(rec, k, None)
        _set_path(obj, t["key"], rows)
    for d in sd.get("diagrams") or []:
        titems = [[tview.find(r) for r in it] for it in d["items"]] if tview else None
        items = _diagram_items_in(view, d)
        fields = []
        for ref, spec in (d.get("fields") or {}).items():
            base = titems[0] if titems else [view.find(r) for r in d["items"][0]]
            fields.append(((_field_path(tview or view, base, ref, d)), _spec(spec)))
        values = []
        for i, it in enumerate(items):
            try:
                rec = _read_back([(spec, body_text(txbody(follow(it, path)))) for path, spec in fields])
            except (IndexError, TypeError):
                continue
            if _is_empty(rec):
                notes.append(f"{label} 図 {d['id']} の {i + 1} つめ: 空なので書かない")
                continue
            values.append(rec)
        _set_path(obj, d["key"], values)
    if sd.get("notes"):
        _set_path(obj, sd["notes"], _notes_text(view.slide) or None)
    return _as_lists(obj)


def _slide_matches(view: SlideView, sd: dict, tview: "SlideView | None") -> bool:
    if tview is not None and view.slide.slide_layout.name != tview.slide.slide_layout.name:
        return False
    refs = list(sd.get("texts") or {}) + list(sd.get("lists") or {}) + [t["shape"] for t in sd.get("tables") or []]
    try:
        for r in refs:
            view.find(r)
    except TemplateError:
        return False
    return True


def _scalars(sd: dict, obj: dict) -> dict:
    """split の続きのページを見分けるための、繰り返さない欄の値。"""
    many = {_spec(s)["key"] for s in (sd.get("lists") or {}).values()}
    many |= {t["key"] for t in sd.get("tables") or [] if t.get("key")}
    many |= {d["key"] for d in sd.get("diagrams") or []}
    return {k: v for k, v in obj.items() if k not in many}


def _merge_page(sd: dict, base: dict, page: dict) -> None:
    for spec in (sd.get("lists") or {}).values():
        k = _spec(spec)["key"]
        base[k] = (base.get(k) or []) + (page.get(k) or [])
    for t in sd.get("tables") or []:
        if t.get("key"):
            base[t["key"]] = (base.get(t["key"]) or []) + (page.get(t["key"]) or [])
    for d in sd.get("diagrams") or []:
        base[d["key"]] = (base.get(d["key"]) or []) + (page.get(d["key"]) or [])


def _extract_in_place(prs, tprs, defs: dict, notes: list) -> "dict | None":
    """テンプレートと同じ枚数の文書を、1 枚ずつ対応させて読む。取り除く見本（drop）のスライドが、
    直前の繰り返すスライドと同じ形なら、その 1 件として読む。形が合わないスライドがあれば None。"""
    data: dict[str, Any] = {}
    last = None
    for n, (slide, tslide) in enumerate(zip(prs.slides, tprs.slides), start=1):
        if slide.slide_layout.name != tslide.slide_layout.name:
            return None
        sd = defs.get(n)
        view = SlideView(prs, slide, n)
        if sd is None:
            last = None
            continue
        if sd.get("drop"):
            if last is not None and _slide_matches(view, last[0], last[1]):
                data[last[0].get("key", last[0]["id"])].append(_read_slide(view, last[0], last[1], notes, f"スライド {n}"))
            continue
        tview = SlideView(tprs, tslide, n)
        if not _slide_matches(view, sd, tview):
            return None
        obj = _read_slide(view, sd, tview, notes, f"スライド {n}")
        key = sd.get("key", sd["id"])
        data[key] = [obj] if sd.get("repeat") else obj
        last = (sd, tview) if sd.get("repeat") else None
    return data


def extract(source: "str | bytes", definition: dict, template: "str | bytes | None" = None) -> tuple[dict, list[str]]:
    """記入済みの pptx から、定義の key の欄だけをデータにする。"""
    prs = open_prs(read_bytes(source))
    tprs = open_prs(read_bytes(template)) if template else None
    defs = {int(sd["slide"]): sd for sd in definition.get("slides", [])}
    total = len(tprs.slides) if tprs else max(defs or {0: None})
    if not tprs and set(defs) != set(range(1, total + 1)):
        raise TemplateError("定義に無いスライドがあるので、並びを合わせるのにテンプレートが要ります（--template）")
    tslides = list(tprs.slides) if tprs else []
    out_slides = list(prs.slides)
    notes: list[str] = []
    if tprs and len(out_slides) == len(tslides):
        found = _extract_in_place(prs, tprs, defs, notes)
        if found is not None:
            return _hoist_doc(found), notes
        notes.clear()
    order = [n for n in range(1, total + 1) if not (defs.get(n) or {}).get("drop")]
    data: dict[str, Any] = {}
    pos = 0
    for idx, n in enumerate(order):
        sd = defs.get(n)
        rest = sum(1 for m in order[idx + 1:] if not (defs.get(m) or {}).get("repeat"))
        if sd is None:
            pos += 1
            continue
        tview = SlideView(tprs, tslides[n - 1], n) if tprs else None
        key = sd.get("key", sd.get("id"))
        records: list[dict] = []
        last_scalars = None
        while pos < len(out_slides) and len(out_slides) - pos > rest:
            view = SlideView(prs, out_slides[pos], pos + 1)
            if not _slide_matches(view, sd, tview):
                break
            label = f"スライド {pos + 1}"
            obj = _read_slide(view, sd, tview, notes, label)
            sc = _scalars(sd, obj)
            if records and sd.get("overflow") == "split" and sc == last_scalars:
                _merge_page(sd, records[-1], obj)
            elif records and not sd.get("repeat"):
                break
            else:
                records.append(obj)
            last_scalars = sc
            pos += 1
        if sd.get("repeat"):
            data[key] = records
        else:
            if not records:
                notes.append(f"テンプレートのスライド {n} に当たるスライドが見つかりません")
            data[key] = records[0] if records else {}
    if pos < len(out_slides):
        notes.append(f"スライド {pos + 1}〜{len(out_slides)} は、定義のどのスライドにも当たらないので読まない")
    return _hoist_doc(data), notes


# ---------------------------------------------------------------------------
# データの形（定義から導く）
# ---------------------------------------------------------------------------

def _field_keys(specs) -> list[str]:
    return [s["key"] for s in specs if s and s.get("key") and s["key"] != "$index" and not s.get("keep") and not s.get("clear")]


def skeleton_data(definition: dict) -> dict:
    """定義が必要とするデータの雛形（値は null。配列は 1 件）。"""
    out: dict = {}
    for sd in definition.get("slides", []):
        if sd.get("drop"):
            continue
        obj: dict = {}
        for k in _field_keys(_spec(s) for s in (sd.get("texts") or {}).values()):
            _set_path(obj, k, None)
        for k in (k for s in (sd.get("texts") or {}).values() if "text" in _spec(s) for k in spec_keys(s)):
            _set_path(obj, k, None)
        for s in (sd.get("lists") or {}).values():
            _set_path(obj, _spec(s)["key"], [None])
        for t in sd.get("tables") or []:
            for k in _field_keys(_spec(s) for s in (t.get("cells") or {}).values()):
                _set_path(obj, k, None)
            if t.get("key"):
                row: dict = {}
                for k in _field_keys(t.get("columns") or []):
                    _set_path(row, k, None)
                _set_path(obj, t["key"], [row])
        for d in sd.get("diagrams") or []:
            item: dict = {}
            for k in _field_keys(_spec(s) for s in (d.get("fields") or {}).values()):
                _set_path(item, k, None)
            _set_path(obj, d["key"], [item])
        if sd.get("notes"):
            _set_path(obj, sd["notes"], None)
        out[sd.get("key", sd["id"])] = [_as_lists(obj)] if sd.get("repeat") else _as_lists(obj)
    return _hoist_doc(out)


def value_notes(definition: dict) -> list[str]:
    """値の書き方と、収まる量（サンプルの粒度）。雛形の null だけでは分からないものを補う。"""
    notes: dict[str, str] = {}

    def note(name: str, spec: dict, items_word: str = "") -> None:
        bits = []
        if "when" in spec:
            return
        if "map" in spec:
            bits.append(" / ".join(_norm_key(k) for k in spec["map"]) + " のどれか")
        if "part" in spec:
            bits.append("日付（2026-10-08）")
        if spec.get("max_items"):
            bits.append(f"最大 {spec['max_items']} {items_word or '項目'}")
        if spec.get("max_chars"):
            bits.append(f"{'1 項目 ' if items_word else ''}{spec['max_chars']} 字まで")
        if bits:
            notes[name] = f"{name}: {'・'.join(bits)}"

    for sd in definition.get("slides", []):
        if sd.get("drop"):
            continue
        base = sd.get("key", sd["id"]) + ("[]" if sd.get("repeat") else "")
        for s in (sd.get("texts") or {}).values():
            s = _spec(s)
            if s.get("key"):
                note(s["key"] if s["key"].startswith(DOC_PREFIX + ".") else f"{base}.{s['key']}", s)
            if "text" in s:
                notes.update(text_notes(s["text"]))
        for s in (sd.get("lists") or {}).values():
            s = _spec(s)
            note(f"{base}.{s['key']}[]", s, "項目")
        for t in sd.get("tables") or []:
            for s in (t.get("cells") or {}).values():
                note(f"{base}.{_spec(s)['key']}", _spec(s))
            if t.get("key"):
                if t.get("max_rows"):
                    notes[f"{base}.{t['key']}"] = f"{base}.{t['key']}: 最大 {t['max_rows']} 行"
                for c in t.get("columns") or []:
                    if c and c.get("key") and c["key"] != "$index":
                        note(f"{base}.{t['key']}[].{c['key']}", c)
        for d in sd.get("diagrams") or []:
            if d.get("max_items"):
                notes[f"{base}.{d['key']}"] = f"{base}.{d['key']}: 最大 {d['max_items']} つ（図形 1 つに 1 件）"
            for s in (d.get("fields") or {}).values():
                s = _spec(s)
                if s.get("key") and s["key"] != "$index":
                    note(f"{base}.{d['key']}[].{s['key']}", s)
        if sd.get("overflow") == "split":
            notes[f"{base}.*"] = f"{base}: 項目・行が上限を超えたら、同じ形のスライドを続けて足す（split）"
    for v in list((definition.get("properties") or {}).values()) + list((definition.get("header_footer") or {}).values()):
        if isinstance(v, str):
            notes.update(text_notes(v))
    return list(notes.values())


def check_definition(template: "str | bytes", definition: dict) -> list[str]:
    """定義の検査。整合を確かめ、雛形データで試しに再構成する（strict なら、残る値の漏れもここで見つかる）。"""
    import tempfile
    raw = read_bytes(template)
    validate_definition(raw, definition)
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "check.pptx")
        warnings = render(raw, definition, skeleton_data(definition), out, allow_overflow=True)
        left = provenance(read_bytes(out))   # 出力に残るものだけ（scrub・drop で消えるものは挙げない）
    warnings = [w for w in warnings if not w.startswith("収まらない値")]
    return warnings + [f"来歴（出力に残る）: {p}" for p in left]


# ---------------------------------------------------------------------------
# 入出力（JSON / YAML）
# ---------------------------------------------------------------------------

def _yaml():
    try:
        import yaml
    except ImportError:
        raise TemplateError("YAML の入出力には PyYAML が必要です（uv add pyyaml / pip install pyyaml）")
    return yaml


def parse_structured(text: str, name: str = "") -> Any:
    lower = name.lower()
    try:
        if lower.endswith((".yaml", ".yml")):
            return _yaml().safe_load(text)
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            if lower.endswith(".json"):
                raise
            return _yaml().safe_load(text)
    except TemplateError:
        raise
    except Exception as e:  # json.JSONDecodeError・yaml.YAMLError
        raise TemplateError(f"{name or '入力'} を読めません: {e}")


def load_structured(path: str) -> Any:
    if path == "-":
        return parse_structured(sys.stdin.read())
    with open(path, encoding="utf-8") as f:
        return parse_structured(f.read(), path)


def dump_structured(obj: Any, path: str) -> None:
    with open(path, "w", encoding="utf-8") as f:
        if path.lower().endswith((".yaml", ".yml")):
            _yaml().safe_dump(obj, f, allow_unicode=True, sort_keys=False, default_flow_style=False)
        else:
            json.dump(obj, f, ensure_ascii=False, indent=2, default=str)
            f.write("\n")


DATA_EXTS = (".json", ".yaml", ".yml")


def data_files(paths: "str | list[str]") -> list[str]:
    """データの指定（ファイル・フォルダ・-）を、読む順のファイルの並びにする。フォルダは中のデータファイルを名前順に。"""
    out: list[str] = []
    for p in [paths] if isinstance(paths, str) else paths:
        if p != "-" and os.path.isdir(p):
            found = sorted(n for n in os.listdir(p) if n.lower().endswith(DATA_EXTS) and not n.startswith("."))
            if not found:
                raise TemplateError(f"フォルダ {p} にデータファイル（{' / '.join(DATA_EXTS)}）がありません")
            out += [os.path.join(p, n) for n in found]
        else:
            out.append(p)
    return out


def merge_data(parts: list, groups: "dict | None" = None) -> dict:
    """分けたデータ（[(ファイル名, 中身), …]）を 1 つにする。

    オブジェクトはキーごとに合わせ、配列（繰り返すスライド・箇条書き・表の行・図）はファイルの順につなぐ。
    同じ欄に違う値があれば止める（どちらが正しいか分からない）。null は、ほかのファイルの値を消さない。
    配列の同じ項目が 2 つのファイルにあっても止める（前のデータを写したまま、ほかのファイルに残っている）。
    groups（スライドのキー → まとまり）があれば、ほかのまとまりのファイルにも書かれた繰り返しのスライドも止める
    （地域別のファイルに、事業別の配列が前のまま写っている）。食い違いは、まとめて挙げる。
    """
    merged: dict = {}
    seen: dict[str, str] = {}   # 欄 → その値を最初に書いたファイル
    problems: list[str] = []

    def mark(v, path, src):
        seen.setdefault(path, src)
        if isinstance(v, dict):
            for k, x in v.items():
                mark(x, f"{path}.{k}", src)

    def put(a, b, path, src):
        if isinstance(a, dict) and isinstance(b, dict):
            for k, v in b.items():
                sub = f"{path}.{k}" if path else str(k)
                if k in a:
                    a[k] = put(a[k], v, sub, src)
                else:
                    a[k] = v
                    mark(v, sub, src)
            return a
        if isinstance(a, list) and isinstance(b, list):
            dup = [x for x in b if isinstance(x, dict) and x in a]
            if dup:
                problems.append(f"{path} の同じ項目が、{seen.get(path, '前のファイル')} と {src} の両方にあります"
                                f"（{len(dup)} 件。前のデータを写したままなら、片方から消す）")
            return a + b
        if b is None or a == b:
            return a
        if a is None:
            seen[path] = src
            return b
        problems.append(f"{path} の値が、{seen.get(path, '前のファイル')} と {src} で違います: {a!r} / {b!r}")
        return a

    for src, obj in parts:
        if obj is None:
            continue
        if not isinstance(obj, dict):
            raise TemplateError(f"{src} の中身はオブジェクト（キーと値）にしてください")
        put(merged, deepcopy(obj), "", src)
    if groups:
        where: dict[str, list] = {}
        for src, obj in parts:
            for k, v in (obj or {}).items():
                if isinstance(v, list) and k in groups:
                    where.setdefault(k, []).append(src)
        for k, srcs in where.items():
            if len(srcs) < 2:
                continue
            mixed = [src for src, obj in parts if src in srcs and {groups.get(x) for x in obj if x in groups} != {groups[k]}]
            if mixed:
                problems.append(f"{k}（{groups[k]}）が {', '.join(srcs)} にあります。{', '.join(mixed)} は別のまとまりのファイルなので、"
                                "前のデータを写したままなら消す")
    if problems:
        raise TemplateError("分けたデータが食い違います:\n" + "\n".join(f"  {p}" for p in problems))
    return merged


def load_data(paths: "str | list[str]", definition: "dict | None" = None) -> dict:
    """データを読む。複数のファイル・フォルダなら、1 つにまとめる（merge_data）。定義があれば、まとまり（group）も確かめる。"""
    files = data_files(paths)
    if len(files) == 1:
        return load_structured(files[0])
    groups = {sd.get("key", sd["id"]): sd["group"] for sd in (definition or {}).get("slides", []) if sd.get("group")}
    return merge_data([(f if f != "-" else "標準入力", load_structured(f)) for f in files], groups)


def split_data(data: dict, definition: dict) -> list:
    """データを、スライドのまとまり（slides[].group。無ければスライドごと）のファイルに分ける。[(ファイル名, 中身), …]。

    定義に無いキーは 00-common に置く。名前順に読めば元の順に戻る（render は、フォルダを渡すと名前順に読んで 1 つにまとめる）。
    """
    order: list[str] = []
    owner: dict[str, str] = {}
    for sd in definition.get("slides", []):
        if sd.get("drop"):
            continue
        group = str(sd.get("group") or sd.get("key", sd["id"]))
        if group not in order:
            order.append(group)
        owner[sd.get("key", sd["id"])] = group
    parts: dict[str, dict] = {g: {} for g in order}
    rest = {}
    for key, value in data.items():
        if value == {} and key in owner:   # 流し込む欄の無いスライド（残すだけ）は、ファイルに書かない
            continue
        (parts[owner[key]] if key in owner else rest)[key] = value
    out = [("00-common", rest)] if rest else []
    for n, g in enumerate(order, start=1):
        if parts[g]:
            out.append((f"{n:02d}-" + (re.sub(r"[^\w-]+", "_", g).strip("_") or "slide"), parts[g]))
    return out


# ---------------------------------------------------------------------------
# スタンドアローン（この文書専用の render スクリプトを書き出す）
# ---------------------------------------------------------------------------

PEP723 = "\n".join("#" + line for line in (
    " /// script", ' requires-python = ">=3.10"', ' dependencies = ["lxml", "python-pptx", "pyyaml"]', " ///"))
STANDALONE_HEADER = '''#!/usr/bin/env python3
{pep723}
"""{title}

pptx テンプレートへ内容を流し込む、専用の render スクリプト（pptx-presentation-builder の export で生成）。
スキルは不要で動く。定義はこのファイルに埋め込み済み。{template_note}

    uv run {name} --data data.yaml -o out.pptx
    uv run {name} --data data/ -o out.pptx            # 分けたデータ（フォルダの中を名前順に）をまとめて流し込む
    python {name} --data data.json -o out.pptx        # lxml・python-pptx・pyyaml が必要
    python {name} --example-data > data.yaml          # データの雛形を出す
    python {name} --extract-def def.yaml              # 埋め込みの定義を取り出す（直したら export --from-script で再生成）

データの形:
{shape}
"""
'''
ENGINE_VERSION = 1


def _wrap_b64(raw: bytes) -> str:
    text = base64.b64encode(raw).decode("ascii")
    lines = [text[i:i + 100] for i in range(0, len(text), 100)] or [""]
    return "(\n" + "\n".join(f"    {line!r}" for line in lines) + "\n)"


def read_exported_script(path: str) -> tuple[dict, "bytes | None", str]:
    """export が書き出したスクリプトから (定義, 埋め込みのテンプレート or None, テンプレートの相対パス) を読む（実行しない）。"""
    import ast
    with open(path, encoding="utf-8") as f:
        tree = ast.parse(f.read())
    found: dict[str, Any] = {}
    for node in tree.body:
        if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name):
            name = node.targets[0].id
            if name == "DEFINITION":
                call = node.value
                if isinstance(call, ast.Call) and call.args and isinstance(call.args[0], ast.Constant):
                    found[name] = json.loads(call.args[0].value)
            elif name in ("TEMPLATE_B64", "TEMPLATE_PATH") and isinstance(node.value, ast.Constant):
                found[name] = node.value.value
    if "DEFINITION" not in found:
        raise TemplateError(f"{path} は export が書き出したスクリプトではありません（DEFINITION が無い）")
    b64 = found.get("TEMPLATE_B64") or ""
    return found["DEFINITION"], (base64.b64decode(b64) if b64 else None), found.get("TEMPLATE_PATH", "")


def export_script(template: "str | bytes", definition: dict, output: str, embed: bool = False,
                  template_name: str | None = None) -> None:
    if isinstance(template, (bytes, bytearray)):
        if not embed:
            raise TemplateError("テンプレートのパスが無いため、埋め込みでしか書き出せません")
        raw, rel = bytes(template), template_name or ""
    else:
        raw = read_bytes(template)
        rel = os.path.relpath(os.path.abspath(template), os.path.dirname(os.path.abspath(output))).replace(os.sep, "/")
        template_name = os.path.basename(template)
    check_definition(raw, definition)
    with open(os.path.abspath(__file__), encoding="utf-8") as f:
        engine = f.read()
    engine = engine.split('\nif __name__ == "__main__":')[0]
    engine = engine.split("\n", 1)[1] if engine.startswith("#!") else engine
    engine = re.sub(r'^"""[\s\S]*?"""\n', "", engine, count=1)
    shape_lines = json.dumps(skeleton_data(definition), ensure_ascii=False, indent=2).splitlines()
    notes = value_notes(definition)
    shape = "\n".join("    " + line for line in shape_lines + ([""] + notes if notes else []))
    title = f"{os.path.splitext(template_name or 'template')[0]} の render スクリプト"
    note = ("テンプレートも埋め込み済み（--template で差し替えられる）。" if embed
            else f"テンプレート（{rel}）は、このスクリプトからの相対パスで読む。")
    header = STANDALONE_HEADER.format(pep723=PEP723, title=title, name=os.path.basename(output), shape=shape.replace('"""', "'''"),
                                      template_note=note)
    footer = (
        "\n\n# ---------------------------------------------------------------------------\n"
        "# 埋め込み（export が書き出した部分。手で直さず、定義を直して export をやり直す）\n"
        "# ---------------------------------------------------------------------------\n"
        f"ENGINE_VERSION = {ENGINE_VERSION}\n"
        f"DEFINITION = json.loads({json.dumps(definition, ensure_ascii=False, default=str)!r})\n"
        f"TEMPLATE_B64 = {_wrap_b64(raw) if embed else chr(34) * 2}\n"
        f"TEMPLATE_PATH = {rel!r}  # 埋め込まない場合の、このスクリプトからの相対パス\n"
        "\n\nif __name__ == \"__main__\":\n"
        "    raise SystemExit(standalone_main(DEFINITION, base64.b64decode(TEMPLATE_B64) or None, TEMPLATE_PATH, __doc__))\n"
    )
    with open(output, "w", encoding="utf-8") as f:
        f.write(header + engine.rstrip() + footer)
    os.chmod(output, 0o755)


def standalone_main(definition: dict, template_bytes: "bytes | None", template_path: str, doc: str | None = None) -> int:
    parser = argparse.ArgumentParser(description=(doc or "").split("\n\n")[0], formatter_class=argparse.RawDescriptionHelpFormatter,
                                     epilog="\n\n".join((doc or "").split("\n\n")[1:]))
    parser.add_argument("--data", nargs="+", action="extend",
                        help="データ（.json / .yaml / .yml。- で標準入力）。複数のファイルかフォルダを渡すと 1 つにまとめる")
    parser.add_argument("-o", "--output", help="出力 .pptx")
    parser.add_argument("--template", help="埋め込みのテンプレートの代わりに使う .pptx（定義と構造が同じものに限る）")
    parser.add_argument("--allow-overflow", action="store_true", help="収まらない値があっても止めず、警告にする")
    parser.add_argument("--example-data", action="store_true", help="データの雛形（YAML）を標準出力に出す")
    parser.add_argument("--extract-template", metavar="PATH", help="埋め込みのテンプレートを書き出す")
    parser.add_argument("--extract-def", metavar="PATH", help="埋め込みの定義を書き出す（.json / .yaml）")
    args = parser.parse_args()
    try:
        if args.extract_def:
            dump_structured(definition, args.extract_def)
            print(f"書き出しました: {args.extract_def}")
            return 0
        if args.example_data:
            print(_yaml().safe_dump(skeleton_data(definition), allow_unicode=True, sort_keys=False, default_flow_style=False), end="")
            for n in value_notes(definition):
                print(f"# {n}")
            return 0
        here = os.path.dirname(os.path.abspath(sys.argv[0]))
        if args.template:
            template = read_bytes(args.template)
        elif template_bytes:
            template = template_bytes
        else:
            path = os.path.join(here, template_path)
            if not os.path.exists(path):
                raise TemplateError(f"テンプレートが見つかりません: {path}（--template で指定するか、export し直す）")
            template = read_bytes(path)
        if args.extract_template:
            if not template_bytes and not args.template:
                raise TemplateError("テンプレートは埋め込まれていません（元の .pptx を使う）")
            with open(args.extract_template, "wb") as f:
                f.write(template)
            print(f"書き出しました: {args.extract_template}")
            return 0
        if not args.data or not args.output:
            parser.error("--data と -o が必要です")
        for w in render(template, definition, load_data(args.data, definition), args.output, args.allow_overflow):
            print(f"警告: {w}", file=sys.stderr)
        print(f"生成しました: {args.output}")
        return 0
    except TemplateError as e:
        print(f"エラー: {e}", file=sys.stderr)
        return 1
    except FileNotFoundError as e:
        print(f"エラー: ファイルが見つかりません: {e.filename}", file=sys.stderr)
        return 1


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def _template_arg(args, definition) -> str:
    """--template はそのまま、定義の template は定義ファイルのある場所からの相対パスとして読む。"""
    if getattr(args, "template", None):
        return args.template
    template = definition.get("template")
    if not template:
        raise TemplateError("--template か定義ファイルの template が必要です")
    def_path = getattr(args, "definition", None)
    if os.path.isabs(template) or not def_path or def_path == "-":
        return template
    return os.path.join(os.path.dirname(os.path.abspath(def_path)), template)


def cmd_analyze(args) -> int:
    definition = analyze(args.template)
    out = args.output or re.sub(r"\.pptx?$", "", args.template, flags=re.I) + ".def.json"
    definition["template"] = os.path.relpath(os.path.abspath(args.template),
                                             os.path.dirname(os.path.abspath(out))).replace(os.sep, "/")
    dump_structured(definition, out)
    print(summarize(definition))
    print(f"\n定義ファイルの下書きを書きました: {out}")
    print("（? の項目と、流し込む欄・残す図形・繰り返す範囲の分け方をユーザーに確認してから、定義ファイルを直して確定する）")
    return 0


def cmd_inspect(args) -> int:
    facts = inspect_template(args.template)
    if args.json:
        json.dump(facts, sys.stdout, ensure_ascii=False, indent=2)
        print()
    else:
        print(format_facts(facts))
    return 0


def cmd_check(args) -> int:
    definition = load_structured(args.definition)
    for w in check_definition(_template_arg(args, definition), definition):
        print(f"警告: {w}")
    print("定義は問題ありません")
    notes = value_notes(definition)
    if notes:
        print("\nデータの書き方と、収まる量:")
        for n in notes:
            print(f"  {n}")
    return 0


def cmd_render(args) -> int:
    definition = load_structured(args.definition)
    data = load_data(args.data, definition)
    for w in render(_template_arg(args, definition), definition, data, args.output, args.allow_overflow):
        print(f"警告: {w}", file=sys.stderr)
    print(f"生成しました: {args.output}")
    return 0


def cmd_export(args) -> int:
    old_def, old_bytes, old_rel, old_dir = None, None, "", ""
    if args.from_script:
        old_def, old_bytes, old_rel = read_exported_script(args.from_script)
        old_dir = os.path.dirname(os.path.abspath(args.from_script))
    if args.definition:
        definition = load_structured(args.definition)
    elif old_def is not None:
        definition = old_def
    else:
        raise TemplateError("--def か --from-script が必要です")
    embed = args.embed
    # --def の template が指すファイルがあれば、それを使う（テンプレートを差し替えたとき）。無ければスクリプトのものを使う
    given = bool(args.definition and definition.get("template") and os.path.exists(_template_arg(args, definition)))
    if args.template or given or (definition.get("template") and not args.from_script):
        template: "str | bytes" = _template_arg(args, definition)
    elif old_bytes is not None:
        template, embed = old_bytes, True
    elif old_rel:
        template = os.path.normpath(os.path.join(old_dir, old_rel))
    else:
        template = _template_arg(args, definition)
    export_script(template, definition, args.output, embed=embed,
                  template_name=os.path.basename(old_rel) if isinstance(template, bytes) else None)
    print(f"書き出しました: {args.output}")
    print(f"  使い方: uv run {os.path.basename(args.output)} --data data.yaml -o out.pptx")
    return 0


def cmd_extract(args) -> int:
    definition = load_structured(args.definition)
    try:
        template = _template_arg(args, definition)
    except TemplateError:
        template = None
    data, notes = extract(args.source, definition, template if template and os.path.exists(template) else None)
    for n in notes:
        print(n, file=sys.stderr)
    if args.split:
        os.makedirs(args.split, exist_ok=True)
        for name, part in split_data(data, definition):
            path = os.path.join(args.split, f"{name}.{args.format}")
            dump_structured(part, path)
            print(f"取り出しました: {path}")
    elif args.output:
        dump_structured(data, args.output)
        print(f"取り出しました: {args.output}")
    else:
        json.dump(data, sys.stdout, ensure_ascii=False, indent=2, default=str)
        print()
    return 0


def add_subcommands(sub) -> None:
    i = sub.add_parser("inspect", help="テンプレートの事実（図形・文字・大きさ・収まる量・仮値の疑い・来歴）を、判断用に出す")
    i.add_argument("template", help="テンプレート .pptx")
    i.add_argument("--json", action="store_true", help="JSON で出す")
    i.set_defaults(func=cmd_inspect)
    a = sub.add_parser("analyze", help="テンプレートを解析して定義ファイル（下書き）を作る")
    a.add_argument("template", help="テンプレート .pptx")
    a.add_argument("-o", "--output", help="定義ファイルの出力先（.json / .yaml。省略時は <テンプレート>.def.json）")
    a.set_defaults(func=cmd_analyze)
    c = sub.add_parser("check", help="定義の検査（整合・strict の漏れ）。雛形データで試しに再構成する")
    c.add_argument("--template", help="テンプレート .pptx（省略時は定義ファイルの template）")
    c.add_argument("--def", dest="definition", required=True, help="定義ファイル（.json / .yaml）")
    c.set_defaults(func=cmd_check)
    r = sub.add_parser("render", help="テンプレート + 定義 + データから pptx を再構成する")
    r.add_argument("--template", help="テンプレート .pptx（省略時は定義ファイルの template）")
    r.add_argument("--def", dest="definition", required=True, help="定義ファイル（.json / .yaml）")
    r.add_argument("--data", required=True, nargs="+", action="extend",
                   help="データ（.json / .yaml。- で標準入力）。複数のファイルかフォルダ（中を名前順に）を渡すと 1 つにまとめる")
    r.add_argument("-o", "--output", required=True, help="出力 .pptx")
    r.add_argument("--allow-overflow", action="store_true", help="収まらない値があっても止めず、警告にする")
    r.set_defaults(func=cmd_render)
    x = sub.add_parser("extract", help="記入済みの pptx から、定義に沿ってデータを取り出す（render の逆）")
    x.add_argument("source", help="記入済みの .pptx（テンプレートと同じ形の文書）")
    x.add_argument("--def", dest="definition", required=True, help="定義ファイル（.json / .yaml）")
    x.add_argument("--template", help="テンプレート .pptx（省略時は定義ファイルの template。スライドの並びを合わせるのに使う）")
    x.add_argument("-o", "--output", help="データの出力先（.json / .yaml。省略時は標準出力に JSON）")
    x.add_argument("--split", metavar="DIR", help="データを、スライドのまとまり（group。無ければスライド）ごとのファイルに分けて DIR に書く")
    x.add_argument("--format", choices=("yaml", "json"), default="yaml", help="--split で書く形式（既定 yaml）")
    x.set_defaults(func=cmd_extract)
    e = sub.add_parser("export", help="この文書専用の、単体で動く render スクリプトを書き出す")
    e.add_argument("--template", help="テンプレート .pptx（省略時は定義ファイルの template）")
    e.add_argument("--def", dest="definition", help="確定した定義ファイル。--from-script と併用すると、その定義を置き換える")
    e.add_argument("--from-script", help="以前に export したスクリプト。定義・テンプレートを引き継いで、最新のエンジンで書き出し直す")
    e.add_argument("-o", "--output", required=True, help="書き出す .py")
    e.add_argument("--embed", action="store_true", help="テンプレートもスクリプトに埋め込む（既定は別ファイルを相対パスで参照）")
    e.set_defaults(func=cmd_export)


def main() -> int:
    parser = argparse.ArgumentParser(description="pptx テンプレートへ内容を流し込む")
    add_subcommands(parser.add_subparsers(dest="command", required=True))
    args = parser.parse_args()
    try:
        return args.func(args)
    except TemplateError as e:
        print(f"エラー: {e}", file=sys.stderr)
        return 1
    except FileNotFoundError as e:
        print(f"エラー: ファイルが見つかりません: {e.filename}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
