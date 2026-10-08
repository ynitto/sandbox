#!/usr/bin/env python3
"""PowerPoint (.pptx) を作る。テンプレートへの流し込みと、テンプレート無しの新規生成。

テンプレートへの流し込み（スライドのスタイル・図形を保つ）:
    uv run python scripts/pptx_builder.py inspect template.pptx             # 判断用の事実を表示
    uv run python scripts/pptx_builder.py analyze template.pptx -o def.yaml  # 定義の下書き
    uv run python scripts/pptx_builder.py check --def def.yaml               # 定義の検査
    uv run python scripts/pptx_builder.py render --def def.yaml --data data.yaml -o out.pptx
    uv run python scripts/pptx_builder.py extract filled.pptx --def def.yaml -o data.yaml
    uv run python scripts/pptx_builder.py export --def def.yaml -o render_xxx.py   # 単体で動く専用スクリプト

テンプレート無し:
    uv run python scripts/pptx_builder.py build --spec spec.json
    uv run python scripts/pptx_builder.py example   # サンプル spec を標準出力

スペックの詳細は references/spec.md、定義ファイルは references/template.md を参照。
"""
from __future__ import annotations

import argparse
import json
import math
import os
import sys
from typing import Any

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pptx_template as pt  # noqa: E402

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.text.text import Font  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402

SIZES = {"16:9": (12192000, 6858000), "4:3": (9144000, 6858000)}
DEFAULT_STYLE = {"font": "Meiryo", "accent": "1F4E79", "text": "222222", "light": "DDEBF7", "line": "7F7F7F"}
SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "rounded": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL,
          "chevron": MSO_SHAPE.CHEVRON}
MARGIN = 457200          # 左右・下の余白（0.5 インチ）
TITLE_TOP, TITLE_H = 304800, 762000
BODY_TOP = 1219200
FONT = {"title": 28, "cover": 40, "subtitle": 20, "body": 18, "table": 14, "label": 16, "note": 12, "card_title": 16,
        "card_body": 12}
MIN_BOX = 1188720        # 図の要素 1 つの最小の幅（1.3 インチ）。これより狭くなる数は、粒度が崩れるので止める
ARROW_GAP = 548640


class Builder:
    def __init__(self, spec: dict):
        self.spec = spec
        self.style = {**DEFAULT_STYLE, **(spec.get("style") or {})}
        self.prs = Presentation()
        w, h = SIZES.get(spec.get("size", "16:9"), SIZES["16:9"])
        self.prs.slide_width, self.prs.slide_height = Emu(w), Emu(h)
        self.W, self.H = w, h
        self.blank = self.prs.slide_layouts[6]
        self.problems: list[str] = []

    # --- 部品 ---
    def color(self, name: str) -> RGBColor:
        return RGBColor.from_string(self.style[name])

    def text(self, slide, x, y, w, h, value: Any, size: int, bold=False, color="text", align=None, anchor=None,
             where: str = "", paras: "list[tuple[str, int]] | None" = None, bullets: bool = False):
        box = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
        tf = box.text_frame
        tf.word_wrap = True
        if anchor:
            tf.vertical_anchor = anchor
        paras = paras if paras is not None else [(line, 0) for line in pt.to_text(value).split("\n")]
        self.fill(tf, paras, size, bold, color, align, bullets)
        self.check(box._element, (x, y, w, h), paras, size, where)
        return box

    def fill(self, tf, paras, size, bold=False, color="text", align=None, bullets=False) -> None:
        for i, (line, lvl) in enumerate(paras or [("", 0)]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.level = lvl
            if align:
                p.alignment = align
            if bullets:   # 行頭の記号は文字に混ぜず、段落の書式（箇条書き）で付ける。流し込みでも記号が残る
                ppr = p._p.get_or_add_pPr()
                ppr.set("marL", str(342900 + lvl * 342900))
                ppr.set("indent", "-342900")
                etree.SubElement(ppr, pt.qa("buFont"), typeface="Arial")
                etree.SubElement(ppr, pt.qa("buChar"), char="•" if lvl == 0 else "–")
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size - 2 * lvl if lvl else size)
            r.font.bold = bold
            r.font.name = self.style["font"]
            r.font.color.rgb = self.color(color)

    def check(self, el, box, paras, size, where) -> None:
        x, y, w, h = box
        l, t, r, b = pt.DEFAULT_INSETS
        cap = pt.Capacity(max((w - l - r) / (size * pt.EMU_PER_PT), 1), max(int((h - t - b) / (size * pt.EMU_PER_PT * pt.LINE_FACTOR) + 0.15), 1), size)
        need = cap.lines_for([p for p in paras if p[0]])
        if need > cap.lines:
            self.problems.append(f"{where}: {need} 行になる。収まるのは {cap.lines} 行（1 行 約 {int(cap.cpl)} 字）")

    def title(self, slide, text: str, where: str) -> None:
        self.text(slide, MARGIN, TITLE_TOP, self.W - 2 * MARGIN, TITLE_H, text, FONT["title"], bold=True,
                  color="accent", anchor=MSO_ANCHOR.MIDDLE, where=f"{where} title")
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(MARGIN), Emu(TITLE_TOP + TITLE_H), Emu(self.W - 2 * MARGIN), Emu(27432))
        bar.fill.solid()
        bar.fill.fore_color.rgb = self.color("accent")
        bar.line.fill.background()

    def shape(self, slide, kind: str, x, y, w, h, fill="accent", line=None):
        s = slide.shapes.add_shape(SHAPES.get(kind, MSO_SHAPE.ROUNDED_RECTANGLE), Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
        s.fill.solid()
        s.fill.fore_color.rgb = self.color(fill)
        if line:
            s.line.color.rgb = self.color(line)
        else:
            s.line.fill.background()
        return s

    # --- スライドの種類 ---
    def slide_title(self, sd: dict, where: str) -> None:
        s = self.prs.slides.add_slide(self.blank)
        self.text(s, MARGIN, self.H * 0.32, self.W - 2 * MARGIN, 1371600, sd.get("title"), FONT["cover"], bold=True,
                  color="accent", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM, where=f"{where} title")
        if sd.get("subtitle"):
            self.text(s, MARGIN, self.H * 0.32 + 1463040, self.W - 2 * MARGIN, 914400, sd["subtitle"], FONT["subtitle"],
                      align=PP_ALIGN.CENTER, where=f"{where} subtitle")

    def slide_bullets(self, sd: dict, where: str) -> None:
        s = self.prs.slides.add_slide(self.blank)
        self.title(s, sd.get("title", ""), where)
        paras = pt._as_items(sd.get("bullets"), where)
        self.text(s, MARGIN, BODY_TOP, self.W - 2 * MARGIN, self.H - BODY_TOP - MARGIN, None, FONT["body"],
                  where=f"{where} bullets", paras=paras, bullets=True)

    def slide_table(self, sd: dict, where: str) -> None:
        s = self.prs.slides.add_slide(self.blank)
        self.title(s, sd.get("title", ""), where)
        cols, rows = sd.get("columns") or [], sd.get("rows") or []
        if not cols:
            raise pt.TemplateError(f"{where}: columns が空です")
        row_h = 411480
        avail = self.H - BODY_TOP - MARGIN
        max_rows = int(avail / row_h) - 1
        if len(rows) > max_rows:
            self.problems.append(f"{where}: {len(rows)} 行。1 枚に収まるのは {max_rows} 行まで（行を減らすか、スライドを分ける）")
        n = min(len(rows), max_rows) if rows else 0
        frame = s.shapes.add_table(n + 1, len(cols), Emu(MARGIN), Emu(BODY_TOP), Emu(self.W - 2 * MARGIN), Emu(row_h * (n + 1)))
        table = frame.table
        col_w = (self.W - 2 * MARGIN) / len(cols)
        for j, head in enumerate(cols):
            cell = table.cell(0, j)
            cell.text = str(head)
            cell.fill.solid()
            cell.fill.fore_color.rgb = self.color("accent")
            self._cell_font(cell, bold=True, color_name=None)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for i, row in enumerate(rows[:n], start=1):
            values = row if isinstance(row, list) else [row.get(str(c)) for c in cols]
            for j in range(len(cols)):
                cell = table.cell(i, j)
                cell.text = pt.to_text(values[j] if j < len(values) else "")
                self._cell_font(cell)
                need = pt.Capacity(max((col_w - 182880) / (FONT["table"] * pt.EMU_PER_PT), 1), 1, FONT["table"]).lines_for(
                    [(cell.text, 0)]) if cell.text else 1
                if need > 2:
                    self.problems.append(f"{where} {i} 行 {j + 1} 列: {need} 行に折り返す（1 セル 2 行まで）")

    def _cell_font(self, cell, bold=False, color_name="text") -> None:
        for p in cell.text_frame.paragraphs:
            # 空のセルにも書式を持たせる（テンプレートにしたとき、あとで入れた文字が既定の 18pt にならない）
            for f in [r.font for r in p.runs] + [Font(p._p.get_or_add_endParaRPr())]:
                f.size = Pt(FONT["table"])
                f.bold = bold
                f.name = self.style["font"]
                if color_name:
                    f.color.rgb = self.color(color_name)

    def slide_flow(self, sd: dict, where: str) -> None:
        """手順・流れ。図形（矩形・角丸・丸・山形）と矢印のコネクタで描く（画像にしない）。"""
        s = self.prs.slides.add_slide(self.blank)
        self.title(s, sd.get("title", ""), where)
        steps = sd.get("steps") or []
        kind = sd.get("shape", "rounded")
        n = len(steps)
        if not n:
            return
        gap = 0 if kind == "chevron" else ARROW_GAP
        width = self.W - 2 * MARGIN
        max_n = int((width + gap) / (MIN_BOX + gap))
        if n > max_n:
            self.problems.append(f"{where}: {n} 段。横に収まるのは {max_n} 段まで（段をまとめる）")
            steps, n = steps[:max_n], max_n
        box_w = (width - (n - 1) * gap) / n
        box_h = 1188720 if kind != "ellipse" else min(box_w, 1645920)
        y = BODY_TOP + 640080
        boxes = []
        for i, step in enumerate(steps):
            step = step if isinstance(step, dict) else {"label": step}
            x = MARGIN + i * (box_w + gap)
            b = self.shape(s, kind, x, y, box_w, box_h)
            b.name = f"step.{i + 1}"
            tf = b.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            self.fill(tf, [(pt.to_text(step.get("label")), 0)], FONT["label"], bold=True, align=PP_ALIGN.CENTER)
            for r in tf.paragraphs[0].runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)   # 濃い塗りの上は白い文字
            self.check(b._element, (x, y, box_w, box_h), [(pt.to_text(step.get("label")), 0)], FONT["label"], f"{where} {i + 1} 段め label")
            if step.get("note"):
                self.text(s, x, y + box_h + 91440, box_w, 1371600, step["note"], FONT["note"],
                          align=PP_ALIGN.CENTER, where=f"{where} {i + 1} 段め note")
            boxes.append(b)
        if kind == "chevron":
            return
        for i, (a, b) in enumerate(zip(boxes, boxes[1:]), start=1):
            x1 = int(a.left + a.width)
            x2 = int(b.left)
            cy = int(a.top + a.height / 2)
            c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x1), Emu(cy), Emu(x2), Emu(cy))
            c.begin_connect(a, 3)
            c.end_connect(b, 1)
            c.line.color.rgb = self.color("line")
            c.line.width = Pt(2)
            ln = c.line._get_or_add_ln()
            etree.SubElement(ln, pt.qa("tailEnd"), type="triangle")
            c.name = f"arrow.{i}"

    def slide_cards(self, sd: dict, where: str) -> None:
        """並列の項目。角丸の図形を格子に並べる。"""
        s = self.prs.slides.add_slide(self.blank)
        self.title(s, sd.get("title", ""), where)
        cards = sd.get("cards") or []
        cols = int(sd.get("columns") or min(max(len(cards), 1), 3))
        rows = math.ceil(len(cards) / cols) if cards else 0
        gap = 228600
        width, height = self.W - 2 * MARGIN, self.H - BODY_TOP - MARGIN
        if rows > 2:
            self.problems.append(f"{where}: {len(cards)} 枚。{cols} 列なら {cols * 2} 枚まで（まとめるか、スライドを分ける）")
            cards, rows = cards[:cols * 2], 2
        cw = (width - (cols - 1) * gap) / cols
        ch = min((height - (rows - 1) * gap) / max(rows, 1), 2743200)
        for i, card in enumerate(cards):
            card = card if isinstance(card, dict) else {"title": card}
            x = MARGIN + (i % cols) * (cw + gap)
            y = BODY_TOP + (i // cols) * (ch + gap)
            b = self.shape(s, "rounded", x, y, cw, ch, fill="light")
            b.name = f"card.{i + 1}"
            self.text(s, x + 91440, y + 91440, cw - 182880, 548640, card.get("title"), FONT["card_title"], bold=True,
                      color="accent", where=f"{where} {i + 1} 枚め title")
            if card.get("body"):
                self.text(s, x + 91440, y + 640080, cw - 182880, ch - 731520, card["body"], FONT["card_body"],
                          where=f"{where} {i + 1} 枚め body")

    def build(self) -> Presentation:
        kinds = {"title": self.slide_title, "bullets": self.slide_bullets, "table": self.slide_table,
                 "flow": self.slide_flow, "cards": self.slide_cards}
        slides = self.spec.get("slides") or []
        if not slides:
            raise pt.TemplateError("spec.slides が空です。少なくとも 1 枚必要です")
        for i, sd in enumerate(slides, start=1):
            kind = sd.get("type", "bullets")
            if kind not in kinds:
                raise pt.TemplateError(f"スライド {i}: type は {', '.join(kinds)} のどれかです: {kind!r}")
            kinds[kind](sd, f"スライド {i}")
            if sd.get("notes"):
                self.prs.slides[-1].notes_slide.notes_text_frame.text = str(sd["notes"])
        props = self.spec.get("properties") or {}
        cp = self.prs.core_properties
        cp.title = props.get("title", "")
        cp.author = props.get("creator", "")
        cp.last_modified_by = props.get("creator", "")
        return self.prs


def build(spec: dict, allow_overflow: bool = False) -> tuple[str, list[str]]:
    b = Builder(spec)
    prs = b.build()
    if b.problems and not allow_overflow:
        raise pt.TemplateError("収まらない内容があります。量を減らすか言い換えてください:\n  " + "\n  ".join(b.problems))
    filename = spec.get("filename", "presentation.pptx")
    prs.save(filename)
    return filename, b.problems


EXAMPLE_SPEC = {
    "filename": "proposal.pptx",
    "size": "16:9",
    "style": {"font": "Meiryo", "accent": "1F4E79"},
    "properties": {"title": "業務改善のご提案", "creator": "pptx-presentation-builder"},
    "slides": [
        {"type": "title", "title": "業務改善のご提案", "subtitle": "2026 年 10 月"},
        {"type": "bullets", "title": "現状の課題",
         "bullets": ["申請の確認に 1 件 30 分かかる", {"text": "確認者が 2 人しかいない", "level": 1}, "差し戻しが月 40 件ある"]},
        {"type": "flow", "title": "進め方", "shape": "rounded",
         "steps": [{"label": "現状調査", "note": "2 週間"}, {"label": "設計", "note": "3 週間"},
                   {"label": "試行", "note": "1 か月"}, {"label": "展開", "note": "全部署"}]},
        {"type": "cards", "title": "効果", "columns": 3,
         "cards": [{"title": "時間", "body": "確認の時間を半分に"}, {"title": "品質", "body": "差し戻しを 1/4 に"},
                   {"title": "負荷", "body": "確認者の偏りを解消"}]},
        {"type": "table", "title": "費用", "columns": ["項目", "金額", "備考"],
         "rows": [["設計", "120 万円", "3 週間"], ["試行", "80 万円", "1 部署"], ["展開", "200 万円", "全部署"]]},
    ],
}


def main() -> int:
    parser = argparse.ArgumentParser(description="PowerPoint (.pptx) を作る（テンプレートへの流し込み・新規生成）")
    sub = parser.add_subparsers(dest="command", required=True)
    b = sub.add_parser("build", help="spec から pptx を新規に生成する（テンプレート無し）")
    b.add_argument("--spec", help="スペック JSON ファイル（省略時は stdin）")
    b.add_argument("--allow-overflow", action="store_true", help="収まらない内容があっても止めず、警告にする")
    sub.add_parser("example", help="サンプル spec を標準出力に表示")
    pt.add_subcommands(sub)
    args = parser.parse_args()
    try:
        if args.command == "example":
            json.dump(EXAMPLE_SPEC, sys.stdout, ensure_ascii=False, indent=2)
            print()
            return 0
        if args.command == "build":
            spec = pt.load_structured(args.spec) if args.spec else json.load(sys.stdin)
            filename, problems = build(spec, args.allow_overflow)
            for p in problems:
                print(f"警告: {p}", file=sys.stderr)
            print(f"生成しました: {filename}")
            return 0
        return args.func(args)
    except pt.TemplateError as e:
        print(f"エラー: {e}", file=sys.stderr)
        return 1
    except FileNotFoundError as e:
        print(f"エラー: ファイルが見つかりません: {e.filename}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
