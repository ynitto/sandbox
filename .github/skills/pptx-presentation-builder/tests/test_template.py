"""pptx_template（inspect / analyze / check / render / extract / export）のテスト。

テンプレートは python-pptx で組み立てる（既定のテンプレートのレイアウト・プレースホルダーを使う）。
実行: cd .github/skills/pptx-presentation-builder && uv run python -m unittest discover -s tests -v
"""
from __future__ import annotations

import json
import os
import re
import subprocess
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Emu, Pt
from lxml import etree

import pptx_template as pt

CM = 360000
LEFT = 12192000 / CM / 2 - 11   # 3 つの四角（幅 6・間隔 8 cm）を、スライドの中央に置く
MID = 12192000 / 2
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
BAND = RGBColor(0x2E, 0x75, 0xB6)


def _font(shape, size, bold=False):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold


def _box(slide, name, x, y, w, h, text, size=14):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tb.name = name
    tb.text_frame.word_wrap = True
    tb.text_frame.text = text
    _font(tb, size)
    return tb


def make_template(path: str) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    prs.core_properties.author = "前任者"
    prs.core_properties.title = "前の案件の報告"
    # 1. 表紙
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "2025年度 第1四半期 報告"
    s.placeholders[1].text = "サンプル株式会社 御中"
    # 2. 箇条書き（ノートつき）
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "背景"
    body = s.placeholders[1].text_frame
    body.text = "申請の確認に時間がかかる"
    for t in ("確認者が少ない", "差し戻しが多い"):
        body.add_paragraph().text = t
    s.notes_slide.notes_text_frame.text = "前回の発表メモ: 顧客名は伏せる"
    # 3・4. 同じ形のスライド（施策ごとに 1 枚）
    for name, goal in (("施策 A", "確認の時間を半分にする"), ("施策 B", "差し戻しを減らす")):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = name
        _box(s, "目的ラベル", 2 * CM, 5 * CM, 3 * CM, 1 * CM, "目的：")
        _box(s, "目的", 5.2 * CM, 5 * CM, 20 * CM, 1 * CM, goal)
    # 5. 流れの図（番号の丸 + 角丸の四角 + 矢印のコネクタ）
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "進め方"
    steps = []
    for i, label in enumerate(("調査", "設計", "試行")):
        x = int((LEFT + i * 8) * CM)
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(7 * CM), Emu(6 * CM), Emu(3 * CM))
        b.name = f"ステップ {i + 1}"
        b.fill.solid()
        b.fill.fore_color.rgb = ACCENT if i % 2 == 0 else BAND
        b.text_frame.text = label
        _font(b, 18, bold=True)
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(x + int(2.25 * CM)), Emu(5 * CM), Emu(int(1.5 * CM)), Emu(int(1.5 * CM)))
        c.name = f"番号 {i + 1}"
        c.text_frame.text = str(i + 1)
        _font(c, 14)
        steps.append(b)
    for i, (a, b) in enumerate(zip(steps, steps[1:]), start=1):
        con = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, a.left + a.width, a.top + a.height // 2, b.left, b.top + b.height // 2)
        con.begin_connect(a, 3)
        con.end_connect(b, 1)
        etree.SubElement(con.line._get_or_add_ln(), pt.qa("tailEnd"), type="triangle")
        con.name = f"矢印 {i}"
    # 6. 表（行を繰り返す）と記入例の注記
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "作業の状況"
    gf = s.shapes.add_table(3, 4, Emu(2 * CM), Emu(4.5 * CM), Emu(28 * CM), Emu(3 * CM))
    gf.name = "状況表"
    for j, h in enumerate(("作業", "担当", "完了", "確認印")):
        gf.table.cell(0, j).text = h
    for r, row in enumerate((("設計書の作成", "山田", "○", ""), ("試験の実施", "佐藤", "×", "")), start=1):
        for j, v in enumerate(row):
            gf.table.cell(r, j).text = v
    _box(s, "注記", 2 * CM, 16 * CM, 20 * CM, 1 * CM, "※ 記入例です。実際の作業に置き換える", 10)
    # 7. 2 列の記入欄の表（左が見出し）
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "開催の案内"
    gf = s.shapes.add_table(2, 2, Emu(2 * CM), Emu(5 * CM), Emu(20 * CM), Emu(2 * CM))
    gf.name = "案内表"
    tbl = gf.table
    tbl.first_row = False
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "日時", "2020/01/01 10:00"
    tbl.cell(1, 0).text, tbl.cell(1, 1).text = "場所", "会議室 A"
    prs.save(path)


def full_data() -> dict:
    return {
        "s1": {"title": "2026年度 第3四半期 報告", "subtitle": "株式会社テスト 御中"},
        "s2": {"title": "課題", "points": ["確認に時間がかかる", {"text": "担当が 2 人", "level": 1}, "様式がばらばら"]},
        "s3": [{"title": "施策 1", "目的": "確認を自動化する"}, {"title": "施策 2", "目的": "様式を統一する"},
               {"title": "施策 3", "目的": "差し戻し理由を記録する"}],
        "s5": {"title": "進め方", "flow1": [{"label": "調査"}, {"label": "設計"}, {"label": "試作"}, {"label": "展開"}]},
        "s6": {"title": "作業の状況", "rows": [{"作業": "要件の整理", "担当": "田中", "完了": True},
                                          {"作業": "設計書", "担当": "山田", "完了": False},
                                          {"作業": "試験", "担当": "佐藤", "完了": True}]},
        "s7": {"title": "開催の案内", "日時": "2026/11/05 14:00", "場所": "本社 3F"},
    }


def slide_texts(path: str) -> list[list[str]]:
    prs = Presentation(path)
    out = []
    for slide in prs.slides:
        texts = []
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text:
                texts.append(sh.text_frame.text)
            if sh.has_table:
                texts.extend(c.text for r in sh.table.rows for c in r.cells)
        out.append(texts)
    return out


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = self.tmp.name
        self.template = os.path.join(self.dir, "report.pptx")
        make_template(self.template)
        self.definition = pt.analyze(self.template)
        self.definition["template"] = "report.pptx"
        self.out = os.path.join(self.dir, "out.pptx")

    def tearDown(self):
        self.tmp.cleanup()

    def sd(self, n: int) -> dict:
        return next(s for s in self.definition["slides"] if s["slide"] == n)

    def render(self, data=None, **kw):
        return pt.render(self.template, self.definition, full_data() if data is None else data, self.out, **kw)


class InspectTest(Base):
    def test_facts_show_shapes_capacity_and_provenance(self):
        facts = pt.inspect_template(self.template)
        self.assertEqual(len(facts["slides"]), 7)
        s5 = facts["slides"][4]
        cand = s5["diagram_candidates"][0]
        self.assertEqual(cand["axis"], "x")
        self.assertEqual(len(cand["items"]), 3)
        self.assertEqual(cand["connectors"], ["矢印 1", "矢印 2"])
        arrow = next(f for f in s5["shapes"] if f["name"] == "矢印 1")
        self.assertIn("終点に矢印", arrow["arrow"])
        title = next(f for f in facts["slides"][0]["shapes"] if f["kind"] == "placeholder:ctrTitle")
        self.assertIn("収まる", pt.format_facts(facts))
        self.assertIn("字", title["fits"])
        sub = next(f for f in facts["slides"][0]["shapes"] if f["kind"] == "placeholder:subTitle")
        self.assertTrue(sub["placeholder_suspect"])
        prov = "\n".join(facts["provenance"])
        self.assertIn("creator: 前任者", prov)
        self.assertIn("ノート", prov)


class AnalyzeTest(Base):
    def test_slide_roles(self):
        d = self.definition
        self.assertTrue(d["strict"])
        self.assertTrue(d["properties"]["scrub"])
        s1 = self.sd(1)
        self.assertEqual(sorted(pt._spec(v)["key"] for v in s1["texts"].values()), ["subtitle", "title"])
        s2 = self.sd(2)
        spec = next(iter(s2["lists"].values()))
        self.assertEqual((spec["key"], spec["max_items"]), ("points", 3))
        self.assertTrue(self.sd(3)["repeat"])
        self.assertTrue(self.sd(4)["drop"])
        self.assertIn("目的ラベル", self.sd(3)["keep"])
        self.assertEqual(self.sd(3)["texts"]["目的"]["key"], "目的")   # ラベルの言葉をキーにする
        # 繰り返すスライドの粒度は、サンプルのうち長い方に合わせる
        self.assertGreaterEqual(self.sd(3)["texts"]["目的"]["max_chars"], len("確認の時間を半分にする"))

    def test_diagram(self):
        d = self.sd(5)["diagrams"][0]
        self.assertEqual(d["axis"], "x")
        self.assertEqual(d["items"], [["ステップ 1", "番号 1"], ["ステップ 2", "番号 2"], ["ステップ 3", "番号 3"]])
        self.assertEqual(d["connectors"], ["矢印 1", "矢印 2"])
        self.assertEqual(d["fields"]["番号 1"], {"key": "$index"})
        self.assertEqual(d["fields"]["ステップ 1"]["key"], "label")
        self.assertEqual(d["align"], "center")
        self.assertEqual(d["max_items"], 4)   # スライドの余白（タイトルの左端）までなら 4 つ並ぶ

    def test_table_columns(self):
        t = self.sd(6)["tables"][0]
        cols = t["columns"]
        self.assertEqual(cols[0]["key"], "作業")
        self.assertEqual(cols[2]["map"], {True: "○", False: "×"})
        self.assertEqual(cols[3], {"clear": True, "header": "確認印"})
        self.assertIn("注記", self.sd(6)["clear"])
        kv = self.sd(7)["tables"][0]
        self.assertNotIn("key", kv)
        self.assertEqual(kv["cells"]["1,2"]["key"], "日時")


class RenderTest(Base):
    def test_full_render(self):
        warnings = self.render()
        texts = slide_texts(self.out)
        self.assertEqual(len(texts), 8)   # 7 枚 - 見本の 1 枚 + 施策 3 枚 - 1 枚
        self.assertIn("株式会社テスト 御中", texts[0])
        self.assertEqual([t for t in texts[2] if t.startswith("施策")], ["施策 1"])
        self.assertIn("目的：", texts[3])   # 残すラベル
        self.assertIn("差し戻し理由を記録する", texts[4])
        self.assertFalse(any("サンプル" in t for ts in texts for t in ts))
        self.assertFalse(any("記入例" in t for ts in texts for t in ts))
        self.assertIn("2026/11/05 14:00", texts[7])
        self.assertIn("日時", texts[7])
        self.assertEqual(warnings, [])

    def test_list_keeps_paragraph_style_and_levels(self):
        self.render()
        prs = Presentation(self.out)
        body = prs.slides[1].placeholders[1].text_frame
        self.assertEqual([(p.text, p.level) for p in body.paragraphs],
                         [("確認に時間がかかる", 0), ("担当が 2 人", 1), ("様式がばらばら", 0)])

    def test_diagram_uses_shapes_and_connectors(self):
        self.render()
        prs = Presentation(self.out)
        slide = prs.slides[5]
        boxes = [sh for sh in slide.shapes if re.match(r"flow1\.\d+\.1$", sh.name)]
        circles = [sh for sh in slide.shapes if re.match(r"flow1\.\d+\.2$", sh.name)]
        arrows = [sh for sh in slide.shapes if ".arrow." in sh.name]
        self.assertEqual([b.text_frame.text for b in boxes], ["調査", "設計", "試作", "展開"])
        self.assertEqual([c.text_frame.text for c in circles], ["1", "2", "3", "4"])
        self.assertEqual(len(arrows), 3)
        # 図形は画像にせず、テンプレートの図形（角丸の四角・丸）のまま。大きさ・塗り・文字の大きさも同じ
        self.assertTrue(all(b.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE for b in boxes))
        self.assertTrue(all(c.auto_shape_type == MSO_SHAPE.OVAL for c in circles))
        self.assertEqual({b.width for b in boxes}, {Emu(6 * CM)})
        self.assertEqual([b.fill.fore_color.rgb for b in boxes], [ACCENT, BAND, ACCENT, BAND])   # 見本の色の順を繰り返す
        self.assertEqual({r.font.size for b in boxes for r in b.text_frame.paragraphs[0].runs}, {Pt(18)})
        self.assertFalse(any(sh.shape_type == 13 for sh in slide.shapes))   # PICTURE
        # 同じ間隔で並び、中央に寄せる
        lefts = [b.left for b in boxes]
        self.assertEqual({b - a for a, b in zip(lefts, lefts[1:])}, {Emu(8 * CM)})
        center = (lefts[0] + boxes[-1].left + boxes[-1].width) / 2
        self.assertAlmostEqual(center, MID, delta=CM * 0.01)
        # 矢印は、隣り合う四角をつなぐ
        ids = [b.shape_id for b in boxes]
        for i, a in enumerate(sorted(arrows, key=lambda s: s.name)):
            st = a._element.find(f".//{pt.qa('stCxn')}")
            end = a._element.find(f".//{pt.qa('endCxn')}")
            self.assertEqual((int(st.get("id")), int(end.get("id"))), (ids[i], ids[i + 1]))
            self.assertIsNotNone(a._element.find(f".//{pt.qa('tailEnd')}"))
        # 図形の id は重ならない
        all_ids = [sh.shape_id for sh in slide.shapes]
        self.assertEqual(len(all_ids), len(set(all_ids)))

    def test_fewer_items_center_and_drop_extra_arrows(self):
        data = full_data()
        data["s5"]["flow1"] = [{"label": "調査"}, {"label": "展開"}]
        self.render(data)
        slide = Presentation(self.out).slides[5]
        boxes = [sh for sh in slide.shapes if re.match(r"flow1\.\d+\.1$", sh.name)]
        self.assertEqual(len(boxes), 2)
        self.assertEqual(len([sh for sh in slide.shapes if ".arrow." in sh.name]), 1)
        center = (boxes[0].left + boxes[-1].left + boxes[-1].width) / 2
        self.assertAlmostEqual(center, MID, delta=CM * 0.01)
        data["s5"]["flow1"] = []
        self.render(data)
        slide = Presentation(self.out).slides[5]
        self.assertFalse([sh for sh in slide.shapes if sh.name.startswith("flow1.") or sh.name.startswith("ステップ")])

    def test_table_rows_marks_and_cleared_column(self):
        self.render()
        prs = Presentation(self.out)
        table = next(sh for sh in prs.slides[6].shapes if sh.has_table).table
        rows = [[c.text for c in r.cells] for r in table.rows]
        self.assertEqual(rows, [["作業", "担当", "完了", "確認印"], ["要件の整理", "田中", "○", ""],
                                ["設計書", "山田", "×", ""], ["試験", "佐藤", "○", ""]])
        frame = next(sh for sh in prs.slides[6].shapes if sh.has_table)
        self.assertEqual(frame.height, sum(r.height for r in table.rows))

    def test_scrub_properties_and_notes(self):
        self.render()
        prs = Presentation(self.out)
        self.assertEqual(prs.core_properties.author, "")
        self.assertEqual(prs.core_properties.title, "")
        self.assertEqual(prs.slides[1].notes_slide.notes_text_frame.text, "")
        self.definition["properties"] = {"scrub": True, "title": "第3四半期 報告"}
        self.render()
        self.assertEqual(Presentation(self.out).core_properties.title, "第3四半期 報告")

    def test_style_is_kept(self):
        self.render()
        tpl, out = Presentation(self.template), Presentation(self.out)
        self.assertEqual(out.slide_width, tpl.slide_width)
        self.assertEqual(out.slides[0].slide_layout.name, tpl.slides[0].slide_layout.name)
        self.assertEqual(out.slides[3].slide_layout.name, tpl.slides[2].slide_layout.name)
        purpose = next(sh for sh in out.slides[3].shapes if sh.name == "目的")
        self.assertEqual(purpose.text_frame.paragraphs[0].runs[0].font.size, Pt(14))

    def test_missing_slide_data_removes_repeat_slides(self):
        data = full_data()
        data["s3"] = []
        self.render(data)
        self.assertEqual(len(Presentation(self.out).slides), 5)

    def test_unknown_keys_warn(self):
        data = full_data()
        data["s1"]["subtitel"] = "x"
        data["s9"] = {}
        warnings = self.render(data)
        self.assertTrue(any("subtitel" in w for w in warnings))
        self.assertTrue(any("s9" in w for w in warnings))


class FitTest(Base):
    def test_over_granularity_stops_with_all_problems(self):
        data = full_data()
        data["s2"]["points"] = ["a", "b", "c", "d"]
        data["s5"]["flow1"][0]["label"] = "現状の業務フローを細かく調査する"
        data["s3"][0]["目的"] = "確認の作業を、申請の受付から承認まで、すべて自動化して人の手を減らす"
        with self.assertRaises(pt.TemplateError) as cm:
            self.render(data)
        msg = str(cm.exception)
        self.assertIn("4 項目", msg)
        self.assertIn("flow1", msg)
        self.assertIn("サンプルの粒度", msg)
        self.assertIn("スライド 3（s3[0]）", msg)
        warnings = self.render(data, allow_overflow=True)
        self.assertTrue(any("収まらない" in w for w in warnings))

    def test_geometric_overflow_without_budget(self):
        spec = next(iter(self.sd(1)["texts"].values()))
        spec.pop("max_chars")
        data = full_data()
        data["s1"]["title"] = "とても長い表題" * 12
        with self.assertRaises(pt.TemplateError) as cm:
            self.render(data)
        self.assertIn("行になる", str(cm.exception))

    def test_too_many_diagram_items(self):
        data = full_data()
        cap = self.sd(5)["diagrams"][0]["max_items"]
        data["s5"]["flow1"] = [{"label": f"段{i}"} for i in range(cap + 1)]
        with self.assertRaises(pt.TemplateError) as cm:
            self.render(data)
        self.assertIn(f"{cap} つまで", str(cm.exception))

    def test_table_height(self):
        data = full_data()
        data["s6"]["rows"] = [{"作業": "長い作業の名前" * 3, "担当": "田中", "完了": True}] * 9
        self.sd(6)["tables"][0]["max_rows"] = 20
        for c in self.sd(6)["tables"][0]["columns"]:
            if c:
                c.pop("max_chars", None)
        with self.assertRaises(pt.TemplateError) as cm:
            self.render(data)
        self.assertIn("表の高さ", str(cm.exception))

    def test_split_adds_continuation_slides(self):
        sd = self.sd(2)
        sd["overflow"] = "split"
        data = full_data()
        data["s2"]["points"] = [f"項目 {i}" for i in range(1, 8)]
        self.render(data)
        texts = slide_texts(self.out)
        self.assertEqual(len(texts), 10)
        self.assertEqual(texts[1], ["課題", "項目 1\n項目 2\n項目 3"])
        self.assertEqual(texts[3], ["課題", "項目 7"])
        back, _ = pt.extract(self.out, self.definition, self.template)
        self.assertEqual(back["s2"]["points"], data["s2"]["points"])


class StrictTest(Base):
    def test_leftover_value_stops(self):
        sd = self.sd(6)
        sd["clear"].remove("注記")
        with self.assertRaises(pt.TemplateError) as cm:
            self.render()
        self.assertIn("注記", str(cm.exception))
        self.assertIn("記入例", str(cm.exception))

    def test_table_column_without_spec(self):
        self.sd(6)["tables"][0]["columns"][1] = None
        with self.assertRaises(pt.TemplateError) as cm:
            self.render()
        self.assertIn("山田", str(cm.exception))

    def test_undefined_slide_text(self):
        self.definition["slides"] = [s for s in self.definition["slides"] if s["slide"] != 7]
        with self.assertRaises(pt.TemplateError) as cm:
            self.render()
        self.assertIn("スライド 7", str(cm.exception))
        self.definition["strict"] = False
        self.render()


class ValidateTest(Base):
    def test_unknown_shape(self):
        self.sd(1)["texts"]["無い図形"] = "x"
        with self.assertRaises(pt.TemplateError) as cm:
            pt.check_definition(self.template, self.definition)
        self.assertIn("無い図形", str(cm.exception))

    def test_header_mismatch(self):
        self.sd(6)["tables"][0]["columns"][1]["header"] = "担当者"
        with self.assertRaises(pt.TemplateError) as cm:
            pt.check_definition(self.template, self.definition)
        self.assertIn("見出しが定義と違います", str(cm.exception))

    def test_check_passes_and_reports_provenance(self):
        warnings = pt.check_definition(self.template, self.definition)
        self.assertTrue(any("前任者" in w for w in warnings))

    def test_skeleton_and_notes(self):
        sk = pt.skeleton_data(self.definition)
        self.assertEqual(sk["s3"], [{"title": None, "目的": None}])
        self.assertEqual(sk["s5"]["flow1"], [{"label": None}])
        self.assertEqual(sk["s6"]["rows"], [{"作業": None, "担当": None, "完了": None}])
        notes = "\n".join(pt.value_notes(self.definition))
        self.assertIn("s2.points[]: 最大 3 項目", notes)
        self.assertIn("s6.rows[].完了: true / false のどれか", notes)


class ExtractTest(Base):
    def test_round_trip(self):
        data = full_data()
        self.render(data)
        back, notes = pt.extract(self.out, self.definition, self.template)
        self.assertEqual(back, data)

    def test_from_template_itself(self):
        back, _ = pt.extract(self.template, self.definition, self.template)
        # テンプレートの形のまま（見本のスライドを取り除く前）の文書。見本の 2 枚目も 1 件として読む
        self.assertEqual(back["s3"], [{"title": "施策 A", "目的": "確認の時間を半分にする"},
                                      {"title": "施策 B", "目的": "差し戻しを減らす"}])
        self.assertEqual(back["s5"]["flow1"], [{"label": "調査"}, {"label": "設計"}, {"label": "試行"}])
        self.assertEqual(back["s6"]["rows"][0], {"作業": "設計書の作成", "担当": "山田", "完了": True})


class ExportTest(Base):
    def run_script(self, *args):
        return subprocess.run([sys.executable, *args], capture_output=True, text=True, cwd=self.dir)

    def test_standalone_script(self):
        script = os.path.join(self.dir, "render_report.py")
        pt.export_script(self.template, self.definition, script)
        data = os.path.join(self.dir, "data.json")
        with open(data, "w", encoding="utf-8") as f:
            json.dump(full_data(), f, ensure_ascii=False)
        r = self.run_script(script, "--data", data, "-o", "x.pptx")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(slide_texts(os.path.join(self.dir, "x.pptx")), (self.render() or True) and slide_texts(self.out))
        r = self.run_script(script, "--example-data")
        self.assertIn("flow1", r.stdout)
        self.assertIn("最大 3 項目", r.stdout)
        definition, embedded, rel = pt.read_exported_script(script)
        self.assertEqual(definition["slides"], json.loads(json.dumps(self.definition["slides"], default=str)))
        self.assertIsNone(embedded)
        self.assertEqual(rel, "report.pptx")

    def test_embedded_and_regenerate(self):
        script = os.path.join(self.dir, "render_report.py")
        pt.export_script(self.template, self.definition, script, embed=True)
        moved = os.path.join(self.dir, "sub")
        os.makedirs(moved)
        os.rename(script, os.path.join(moved, "r.py"))
        os.remove(self.template)
        data = os.path.join(self.dir, "data.json")
        with open(data, "w", encoding="utf-8") as f:
            json.dump(full_data(), f, ensure_ascii=False)
        r = self.run_script(os.path.join(moved, "r.py"), "--data", data, "-o", "y.pptx")
        self.assertEqual(r.returncode, 0, r.stderr)
        _, embedded, _ = pt.read_exported_script(os.path.join(moved, "r.py"))
        self.assertTrue(embedded.startswith(b"PK"))


class CliTest(Base):
    def test_cli_flow(self):
        cli = os.path.join(os.path.dirname(__file__), "..", "scripts", "pptx_builder.py")

        def run(*args):
            return subprocess.run([sys.executable, cli, *args], capture_output=True, text=True, cwd=self.dir)

        r = run("inspect", "report.pptx")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("図の並び（横）", r.stdout)
        r = run("analyze", "report.pptx", "-o", "def.yaml")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("確認すること", r.stdout)
        r = run("check", "--def", "def.yaml")
        self.assertEqual(r.returncode, 0, r.stderr)
        with open(os.path.join(self.dir, "data.json"), "w", encoding="utf-8") as f:
            json.dump(full_data(), f, ensure_ascii=False)
        r = run("render", "--def", "def.yaml", "--data", "data.json", "-o", "o.pptx")
        self.assertEqual(r.returncode, 0, r.stderr)
        r = run("extract", "o.pptx", "--def", "def.yaml", "-o", "back.yaml")
        self.assertEqual(r.returncode, 0, r.stderr)
        r = run("render", "--def", "def.yaml", "--data", "back.yaml", "-o", "o2.pptx")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(slide_texts(os.path.join(self.dir, "o.pptx")), slide_texts(os.path.join(self.dir, "o2.pptx")))
        bad = full_data()
        bad["s2"]["points"] = ["a", "b", "c", "d"]
        with open(os.path.join(self.dir, "bad.json"), "w", encoding="utf-8") as f:
            json.dump(bad, f, ensure_ascii=False)
        r = run("render", "--def", "def.yaml", "--data", "bad.json", "-o", "o3.pptx")
        self.assertEqual(r.returncode, 1)
        self.assertIn("4 項目", r.stderr)


class OpensCleanlyTest(Base):
    def test_package_is_consistent(self):
        """書き出した pptx の部品が、参照の切れなく揃っていること（PowerPoint の修復を招かない）。"""
        self.render()
        with zipfile.ZipFile(self.out) as z:
            names = set(z.namelist())
            for name in names:
                if not name.endswith(".rels"):
                    continue
                base = os.path.dirname(os.path.dirname(name))
                for rel in etree.fromstring(z.read(name)):
                    if rel.get("TargetMode") == "External":
                        continue
                    target = os.path.normpath(os.path.join(base, rel.get("Target"))).lstrip("/")
                    if rel.get("Target").startswith("/"):
                        target = rel.get("Target").lstrip("/")
                    self.assertIn(target, names, f"{name} → {rel.get('Target')}")
            self.assertFalse(any(n.startswith("docProps/thumbnail") for n in names))
        prs = Presentation(self.out)
        for slide in prs.slides:
            rids = set(slide.part.rels.keys())
            for el in slide._element.iter():
                for k, v in el.attrib.items():
                    if k.startswith(f"{{{pt.NS_R}}}"):
                        self.assertIn(v, rids)


if __name__ == "__main__":
    unittest.main()


class VerticalAndCommentsTest(unittest.TestCase):
    def test_vertical_flow_and_comment_removed(self):
        with tempfile.TemporaryDirectory() as d:
            path, out = os.path.join(d, "v.pptx"), os.path.join(d, "o.pptx")
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[6])
            boxes = []
            for i, t in enumerate(("受付", "確認", "承認")):
                b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(8 * CM), Emu(int((1 + i * 5) * CM)), Emu(6 * CM), Emu(3 * CM))
                b.name, b.text_frame.text = f"箱 {i + 1}", t
                boxes.append(b)
            for i, (a, b) in enumerate(zip(boxes, boxes[1:]), start=1):
                c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, a.left + a.width // 2, a.top + a.height, b.left + b.width // 2, b.top)
                c.begin_connect(a, 2)
                c.end_connect(b, 0)
                c.name = f"線 {i}"
            # コメント（旧形式）を 1 つ付ける
            from pptx.opc.package import Part
            from pptx.opc.packuri import PackURI
            part = Part(PackURI("/ppt/comments/comment1.xml"), "application/vnd.openxmlformats-officedocument.presentationml.comments+xml",
                        prs.part.package, b'<p:cmLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>')
            s.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments")
            prs.save(path)
            self.assertTrue(any("コメント" in p for p in pt.provenance(pt.read_bytes(path))))
            definition = pt.analyze(path)
            d0 = definition["slides"][0]["diagrams"][0]
            self.assertEqual((d0["axis"], len(d0["connectors"])), ("y", 2))
            self.assertEqual(d0["max_items"], 3)   # 下の余白までに 3 つ
            pt.render(path, definition, {"s1": {d0["key"]: [{"label": "受付"}, {"label": "承認"}]}}, out)
            slide = Presentation(out).slides[0]
            got = [sh for sh in slide.shapes if re.match(r"flow1\.\d+\.1$", sh.name)]
            self.assertEqual([g.text_frame.text for g in got], ["受付", "承認"])
            self.assertEqual(got[1].top - got[0].top, Emu(5 * CM))
            self.assertEqual(got[0].left, Emu(8 * CM))
            with zipfile.ZipFile(out) as z:
                self.assertFalse(any("comments/" in n for n in z.namelist()))
