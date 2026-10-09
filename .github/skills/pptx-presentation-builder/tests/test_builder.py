"""pptx_builder（build / example）と、build で作った資料をテンプレートにして流し込むテスト。

実行: cd .github/skills/pptx-presentation-builder && uv run python -m unittest discover -s tests -v
"""
from __future__ import annotations

import copy
import os
import re
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

import pptx_builder as pb
import pptx_template as pt


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = self.tmp.name

    def tearDown(self):
        self.tmp.cleanup()

    def build(self, spec: dict, **kw) -> str:
        spec = copy.deepcopy(spec)
        spec["filename"] = os.path.join(self.dir, spec.get("filename", "deck.pptx"))
        return pb.build(spec, **kw)[0]


class BuildTest(Base):
    def test_example(self):
        path = self.build(pb.EXAMPLE_SPEC)
        prs = Presentation(path)
        self.assertEqual(len(prs.slides), 5)
        self.assertEqual(prs.core_properties.title, "業務改善のご提案")
        self.assertAlmostEqual(prs.slide_width / prs.slide_height, 16 / 9, places=2)

    def test_flow_is_drawn_with_shapes_not_pictures(self):
        prs = Presentation(self.build(pb.EXAMPLE_SPEC))
        flow = prs.slides[2]
        boxes = [s for s in flow.shapes if s.name.startswith("step.")]
        arrows = [s for s in flow.shapes if s.name.startswith("arrow.")]
        self.assertEqual(len(boxes), 4)
        self.assertTrue(all(b.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE for b in boxes))
        self.assertEqual(len(arrows), 3)
        self.assertTrue(all(a.shape_type == MSO_SHAPE_TYPE.LINE for a in arrows))
        for a, (left, right) in zip(arrows, zip(boxes, boxes[1:])):
            self.assertEqual(int(a._element.find(f".//{pt.qa('stCxn')}").get("id")), left.shape_id)
            self.assertEqual(int(a._element.find(f".//{pt.qa('endCxn')}").get("id")), right.shape_id)
            self.assertIsNotNone(a._element.find(f".//{pt.qa('tailEnd')}"))
        self.assertFalse(any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for sl in prs.slides for s in sl.shapes))

    def test_empty_table_cell_keeps_font_for_later_filling(self):
        spec = {"slides": [{"type": "table", "title": "費用", "columns": ["プラン", "特徴"], "rows": [["ライト", ""]]}]}
        prs = Presentation(self.build(spec))
        table = next(s for s in prs.slides[0].shapes if s.has_table).table
        end = table.cell(1, 1).text_frame.paragraphs[0]._p.find(pt.qa("endParaRPr"))
        self.assertEqual(end.get("sz"), str(pb.FONT["table"] * 100))   # あとで入れた文字が既定の 18pt にならない
        self.assertEqual(str(table.cell(0, 0).fill.fore_color.rgb), pb.DEFAULT_STYLE["accent"])

    def test_bullets_use_paragraph_bullets(self):
        prs = Presentation(self.build(pb.EXAMPLE_SPEC))
        body = next(s for s in prs.slides[1].shapes if s.has_text_frame and len(s.text_frame.paragraphs) > 1)
        self.assertEqual([p.level for p in body.text_frame.paragraphs], [0, 1, 0])
        self.assertFalse(body.text_frame.paragraphs[0].text.startswith("・"))   # 記号は文字に混ぜない
        self.assertIsNotNone(body.text_frame.paragraphs[0]._p.find(f".//{pt.qa('buChar')}"))

    def test_too_much_stops(self):
        spec = copy.deepcopy(pb.EXAMPLE_SPEC)
        spec["slides"][2]["steps"] = [{"label": f"段 {i}"} for i in range(12)]
        spec["slides"][1]["bullets"] = ["とても長い説明の文章です。" * 8] * 10
        with self.assertRaises(pt.TemplateError) as cm:
            self.build(spec)
        self.assertIn("段まで", str(cm.exception))
        self.assertIn("行になる", str(cm.exception))
        _, problems = pb.build(dict(spec, filename=os.path.join(self.dir, "x.pptx")), allow_overflow=True)
        self.assertTrue(problems)

    def test_unknown_type(self):
        with self.assertRaises(pt.TemplateError):
            self.build({"slides": [{"type": "chart"}]})

    def test_malformed_spec_is_reported_not_crashed(self):
        for spec in ([1], {"slides": "x"}, {"slides": ["x"]}):
            with self.subTest(spec=spec), self.assertRaises(pt.TemplateError):
                pb.build(spec)


class BuiltDeckAsTemplateTest(Base):
    """build で作った資料を、そのままテンプレートにして流し込む（図の並び・格子の並び）。"""

    def test_flow_and_grid(self):
        spec = copy.deepcopy(pb.EXAMPLE_SPEC)
        spec["slides"][3]["cards"] = [{"title": t, "body": b} for t, b in (
            ("時間", "確認を半分に"), ("品質", "差し戻しを減らす"), ("負荷", "偏りをなくす"),
            ("費用", "外注を減らす"), ("安全", "誤送信を防ぐ"), ("教育", "新人が早く慣れる"))]
        path = self.build(spec)
        definition = pt.analyze(path)
        flow = next(s for s in definition["slides"] if s["slide"] == 3)["diagrams"][0]
        self.assertEqual(len(flow["connectors"]), 3)
        grid = next(s for s in definition["slides"] if s["slide"] == 4)["diagrams"][0]
        self.assertEqual((grid["axis"], grid["columns"]), ("grid", 3))
        data, _ = pt.extract(path, definition, path)
        self.assertEqual([x["label"] for x in data["s3"]["flow1"]], ["現状調査", "設計", "試行", "展開"])
        data["s3"]["flow1"] = data["s3"]["flow1"][:2]
        data["s4"]["items1"] = data["s4"]["items1"][:4]
        out = os.path.join(self.dir, "out.pptx")
        pt.render(path, definition, data, out)
        prs = Presentation(out)
        cards = [s for s in prs.slides[3].shapes if re.match(r"items1\.\d+\.1$", s.name)]
        self.assertEqual(len(cards), 4)
        self.assertEqual(len({c.top for c in cards}), 2)   # 3 列の格子なので 2 段目に折り返す
        self.assertEqual(cards[3].left, cards[0].left)
        self.assertEqual(len([s for s in prs.slides[2].shapes if ".arrow." in s.name]), 1)
        back, _ = pt.extract(out, definition, path)
        self.assertEqual(back, data)


if __name__ == "__main__":
    unittest.main()
